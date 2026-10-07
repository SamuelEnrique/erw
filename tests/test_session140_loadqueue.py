"""Session 140: NYISO's load interconnection requests (warehouse/connectors/nyiso_load_queue.py).

No request is made. The parser reads a cut of the real workbook NYISO served on 7 October 2026
(tests/fixtures/session140/loadqueue/nyiso_load_sheets_cut.xlsx): the header row of the sheet "Load Projects", twenty
of its 74 request rows as NYISO typed them (values and cell formats the workbook's own; the twenty are requests whose
developer is an organization), its End-Use Key and NOTES, and the whole sheet "Load Project Tracking" as the values
the workbook holds for its formula cells. The tracking sheet's megawatts are therefore NYISO's sums over all 74
requests, not over the twenty rows of the cut."""
import csv
import datetime as dt
import decimal
import hashlib
import json
import os
import shutil
import sys
import tempfile
import unittest
from unittest import mock

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(ROOT, "warehouse", "connectors"))
sys.path.insert(0, os.path.join(ROOT, "warehouse", "validate"))
import iso_prices as ip  # noqa: E402
import nyiso_load_queue as q  # noqa: E402
import erw_validate  # noqa: E402

FIX = os.path.join(ROOT, "tests", "fixtures", "session140", "loadqueue", "nyiso_load_sheets_cut.xlsx")
SITE = os.path.join(ROOT, "site", "data", "nyiso_load_queue.json")
REC = {"file": "cut.xlsx", "url": q.URL, "answered_at": q.URL, "status": "200", "bytes": "11907", "sha256": "0" * 64,
       "last_modified": "Fri, 11 Sep 2026 17:26:02 GMT", "retrieved_at": "2026-10-07T09:09:14Z", "origin": "a test"}
IP_DIRS = ("OUT_DIR", "LOG_DIR", "RAW_DIR", "METADATA_DIR", "STATUS_DIR")


def quiet(_msg):
    pass


def parsed():
    return q.parse_workbook(FIX, quiet)


def site_doc(rows=None):
    got = parsed()
    rows = got[0] if rows is None else rows
    return q.build_site(rows, got[1], got[2], got[3], got[4], REC, "2026-10-07T00:00:00Z")


def total(values):
    return sum((decimal.Decimal(str(v)) for v in values), decimal.Decimal(0))


class Parser(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rows, cls.key, cls.end_use_key, cls.notes, cls.tracking = parsed()
        cls.by_id = {r["queue_id"]: r for r in cls.rows}

    def test_every_request_row_of_the_cut_and_nothing_else(self):
        self.assertEqual(len(self.rows), 20)
        self.assertEqual(len(self.by_id), 20)
        self.assertTrue(all(r["sheet"] == "Load Projects" for r in self.rows))
        self.assertNotIn("End-Use Key", self.by_id)  # the key and the notes under the list are not requests
        self.assertNotIn("NOTES:", self.by_id)
        self.assertEqual(self.end_use_key[0], "End-Use Key")
        self.assertEqual(self.notes[0], "NOTES:")

    def test_queue_number_and_megawatts_as_printed(self):
        self.assertIn("0290A", self.by_id)  # text, its leading zero kept
        self.assertIn("205", self.by_id)    # a number cell, no ".0" added
        self.assertEqual(self.by_id["1465"]["capacity_mw"], "50.2")
        self.assertEqual(self.by_id["1713"]["capacity_mw"], "22.3")
        self.assertEqual(self.by_id["1743"]["capacity_mw"], "1935")
        self.assertEqual(self.by_id["1743"]["name"], "St. Lawrence Infrastructure 2")
        self.assertEqual(self.by_id["1743"]["operator"], "St. Lawrence Infrastructure, LLC")

    def test_dates_as_printed_and_iso_only_where_the_cell_holds_a_date(self):
        r = self.by_id["1465"]
        self.assertEqual((r["queue_date_printed"], r["queue_date"]), ("11/14/22", "2022-11-14"))
        # NYISO's cell holds 14 January 1900 here: kept as it is, not corrected and not dropped
        self.assertEqual((r["last_updated_printed"], r["last_updated_date"]), ("01-14-00", "1900-01-14"))
        self.assertEqual((r["proposed_initial_backfeed_printed"], r["proposed_initial_backfeed_date"], r["proposed_initial_backfeed_month"]),
                         ("I/S", "", ""))
        r = self.by_id["0580"]
        self.assertEqual((r["queue_date_printed"], r["queue_date"]), ("09-27-16", "2016-09-27"))
        self.assertEqual((r["proposed_initial_backfeed_printed"], r["proposed_initial_backfeed_date"], r["proposed_initial_backfeed_month"]),
                         ("11-2026", "", "2026-11"))  # a month and a year: no day is added

    def test_a_blank_cell_stays_blank(self):
        r = self.by_id["1446"]  # withdrawn: no end use, no bundle, no studies, no backfeed date printed
        for col in ("end_use", "end_use_words", "sis_bundle", "availability_of_studies", "affected_transmission_owner",
                    "ia_tender_printed", "ia_tender_date", "fs_completion_printed", "fs_completion_date",
                    "proposed_initial_backfeed_printed", "proposed_initial_backfeed_date", "proposed_initial_backfeed_month"):
            self.assertEqual(r[col], "", col)
        self.assertTrue(all(r["ia_tender_date"] == "" and r["fs_completion_date"] == "" for r in self.rows))

    def test_status_in_nyisos_own_words_from_the_sheets_key(self):
        self.assertEqual(len(self.key), 19)
        self.assertEqual(self.key["0"], "Withdrawn")
        self.assertEqual(self.key["14"], "In Service Commercial")
        self.assertEqual(self.key["5P"], "SRIS Commenced, Stopped and Pending Adoption of IP")
        self.assertEqual(self.key["P"], "Pending Adoption of IP Compliance with Order 2023")
        self.assertEqual(self.by_id["1315"]["iso_status"], "8")
        self.assertEqual(self.by_id["1315"]["iso_status_words"], "Rejected Cost Allocation/Next FS Pending")
        self.assertEqual(self.by_id["0580"]["iso_status_words"], "Under Construction")
        self.assertEqual(self.by_id["205"]["iso_status_words"], "In Service Commercial")

    def test_end_use_label_is_the_tracking_sheets(self):
        self.assertEqual(self.by_id["1730"]["end_use"], "DAT-AI")
        self.assertEqual(self.by_id["1730"]["end_use_words"], "AI Datacenter")
        self.assertEqual(self.by_id["1536"]["end_use_words"], "Microchip fabrication")
        self.assertEqual(self.by_id["0776"]["end_use_words"], "Cryptocurrency mining")

    def test_tracking_sheet_as_printed(self):
        t = self.tracking
        self.assertEqual(t["title"], "Load Project Tracking Summary (by NYISO Zone and Requested MW)")
        self.assertEqual([b["block"] for b in t["blocks"]], ["MW By Status", "MW By End-Use", "MW By IR Submission Date"])
        status = t["blocks"][0]
        self.assertEqual(status["columns"], list("ABCDEFGHIJK") + ["NYCA"])
        tot = status["rows"][-1]
        self.assertTrue(tot["total_row"])
        self.assertEqual(tot["mw"]["NYCA"], 14473.1)  # the figure session 138 saw
        self.assertEqual(tot["mw"]["G"], 252.3)
        withdrawn = status["rows"][0]
        self.assertEqual((withdrawn["code"], withdrawn["label"], withdrawn["mw"]["NYCA"]), ("-", "Withdrawn (not included in NYCA total)", 3078.88))
        self.assertEqual(t["blocks"][1]["rows"][-1]["mw"]["NYCA"], 13611.9)  # stored as 13611.900000000001; printed 13611.9
        self.assertEqual(t["blocks"][2]["rows"][-1]["mw"]["NYCA"], 14232.9)
        self.assertEqual(len(t["notes"]), 2)
        self.assertIn("excludes withdrawn and in-service", t["notes"][0]["text"])

    def test_a_new_column_a_missing_column_and_an_unknown_status_fail(self):
        import openpyxl
        for change in ("new", "missing", "status", "zone", "no queue number"):
            wb = openpyxl.load_workbook(FIX, data_only=True)
            ws = wb["Load Projects"]
            if change == "new":
                ws.cell(row=1, column=22, value="A Column NYISO Added")
            elif change == "missing":
                ws.cell(row=1, column=5).value = None
            elif change == "status":
                ws.cell(row=8, column=15, value="77")
            elif change == "zone":
                ws.cell(row=8, column=11, value="Z")
            else:
                ws.cell(row=8, column=1).value = None
            with self.assertRaises(RuntimeError, msg=change):
                q.read_projects(ws, quiet)

    def test_date_cell(self):
        d = dt.datetime(2005, 11, 2)
        self.assertEqual(q.date_cell(d, "m/d/yy;@"), ("11/2/05", "2005-11-02", "2005-11", True))
        self.assertEqual(q.date_cell(d, "mm-dd-yy"), ("11-02-05", "2005-11-02", "2005-11", True))
        self.assertEqual(q.date_cell(d, "a format nobody knows"), ("2005-11-02", "2005-11-02", "2005-11", False))
        self.assertEqual(q.date_cell("07-2037", "General"), ("07-2037", "", "2037-07", True))
        self.assertEqual(q.date_cell("13-2037", "General"), ("13-2037", "", "", True))  # not a month: kept as typed, read as nothing
        self.assertEqual(q.date_cell("I/S", "mmm-yy"), ("I/S", "", "", True))
        self.assertEqual(q.date_cell(None, "General"), ("", "", "", True))


class ZoneMap(unittest.TestCase):
    def test_letters_a_to_k_onto_the_price_tables_zone_names(self):
        self.assertEqual(list(q.ZONES), list("ABCDEFGHIJK"))
        self.assertEqual(sorted(q.ZONES.values()), sorted(ip.NYISO_ZONES))  # the names nyiso_dam_zone_prices carries
        self.assertEqual(len(set(q.ZONES.values())), 11)
        self.assertEqual((q.ZONES["A"], q.ZONES["B"], q.ZONES["C"], q.ZONES["D"], q.ZONES["E"], q.ZONES["F"]),
                         ("WEST", "GENESE", "CENTRL", "NORTH", "MHK VL", "CAPITL"))
        self.assertEqual((q.ZONES["G"], q.ZONES["H"], q.ZONES["I"], q.ZONES["J"], q.ZONES["K"]),
                         ("HUD VL", "MILLWD", "DUNWOD", "N.Y.C.", "LONGIL"))

    def test_table_and_site_file_carry_both(self):
        rows, key = parsed()[:2]
        df = q.build_table(rows, key, REC, quiet)
        row = df[df["queue_id"] == "1738"].iloc[0]
        self.assertEqual((row["zone"], row["zone_entity"]), ("H", "nyiso:MILLWD"))
        doc = site_doc()
        self.assertEqual([(z["zone"], z["zone_entity"]) for z in doc["zones"]], [(k, "nyiso:" + v) for k, v in q.ZONES.items()])
        r = next(r for r in doc["rows"] if r["queue_position"] == "1721")
        self.assertEqual((r["zone"], r["zone_name"], r["zone_entity"]), ("K", "LONGIL", "nyiso:LONGIL"))


class SiteFile(unittest.TestCase):
    def check_sums(self, doc):
        line = [r for r in doc["rows"] if r["in_line"]]
        for z in doc["zones"]:
            mine = [r for r in line if r["zone"] == z["zone"]]
            self.assertEqual(z["requests"], len(mine), z["zone"])
            self.assertEqual(decimal.Decimal(str(z["mw"])), total(r["mw"] for r in mine if r["mw"] is not None), z["zone"])
            self.assertEqual(sum(s["requests"] for s in z["by_status"]), z["requests"], z["zone"])
            self.assertEqual(total(s["mw"] for s in z["by_status"]), decimal.Decimal(str(z["mw"])), z["zone"])
            for group, t in z["not_in_line"].items():
                apart = [r for r in doc["rows"] if r["zone"] == z["zone"] and r["group"] == group]
                self.assertEqual((t["requests"], decimal.Decimal(str(t["mw"]))), (len(apart), total(r["mw"] for r in apart if r["mw"] is not None)))
        self.assertEqual(doc["total"]["requests"], len(line))
        self.assertEqual(doc["total"]["requests"], sum(z["requests"] for z in doc["zones"]) + doc["total"]["requests_without_zone"])
        self.assertEqual(decimal.Decimal(str(doc["total"]["mw"])), total(r["mw"] for r in line if r["mw"] is not None))
        self.assertEqual(doc["all_requests_on_sheet"], len(doc["rows"]))
        self.assertEqual(len(line) + sum(t["requests"] for t in doc["not_in_line"].values()), len(doc["rows"]))
        for r in doc["rows"]:
            self.assertEqual(r["in_line"], r["group"] == "in line")
            self.assertEqual(r["sheet"], "Load Projects")

    def test_zone_sums_equal_the_rows_sums_in_the_cut(self):
        doc = site_doc()
        self.check_sums(doc)
        zones = {z["zone"]: z for z in doc["zones"]}
        self.assertEqual((doc["total"]["requests"], doc["total"]["mw"]), (14, 6520.9))  # by hand from the twenty rows
        self.assertEqual((zones["C"]["requests"], zones["C"]["mw"]), (4, 1190))
        self.assertEqual((zones["D"]["requests"], zones["D"]["mw"]), (3, 2490))
        self.assertEqual((zones["E"]["requests"], zones["E"]["mw"]), (3, 1342))
        self.assertEqual((zones["J"]["requests"], zones["J"]["mw"], zones["J"]["by_status"]), (0, 0, []))

    def test_zone_sums_equal_the_rows_sums_in_the_file_the_page_reads(self):
        self.assertTrue(os.path.exists(SITE), "site/data/nyiso_load_queue.json is not there")
        with open(SITE, encoding="utf-8") as f:
            text = f.read()
        doc = json.loads(text)
        self.check_sums(doc)
        self.assertEqual(doc["source"]["sheet_names"], ["Load Projects", "Load Project Tracking"])
        self.assertTrue(doc["source"]["url"].startswith("https://www.nyiso.com/"))
        self.assertEqual(doc["rule"]["sums_cover_sheet"], "Load Projects")
        self.assertIn("Withdrawn", doc["rule"]["in_line"])
        self.assertIn("In Service", doc["rule"]["in_line"])
        self.assertEqual({s["status_code"] for s in doc["rule"]["statuses_not_in_line"]}, {"0", "13", "14", "15"})
        self.assertTrue(all(z["zone_entity"] == "nyiso:" + q.ZONES[z["zone"]] for z in doc["zones"]))
        self.assertNotIn(chr(0x2014), text)
        self.assertTrue(all("developer" not in r and "operator" not in r for r in doc["rows"]))  # the page's rows carry no developer's name

    def test_a_withdrawn_request_is_not_counted_in_line(self):
        doc = site_doc()
        withdrawn = [r for r in doc["rows"] if r["status_code"] == "0"]
        self.assertEqual(len(withdrawn), 2)
        self.assertTrue(all(not r["in_line"] and r["group"] == "withdrawn" for r in withdrawn))
        self.assertEqual(doc["not_in_line"]["withdrawn"], {"requests": 2, "mw": 158})
        # the same request, in line and then withdrawn: the count and the megawatts leave "in line" and nothing else moves
        rows = parsed()[0]
        for r in rows:
            if r["queue_id"] == "1738":
                r["iso_status"], r["iso_status_words"] = "0", "Withdrawn"
        after = site_doc(rows)
        self.assertEqual(after["total"]["requests"], doc["total"]["requests"] - 1)
        self.assertEqual(decimal.Decimal(str(after["total"]["mw"])), decimal.Decimal(str(doc["total"]["mw"])) - 1000)
        self.assertEqual(after["not_in_line"]["withdrawn"], {"requests": 3, "mw": 1158})
        self.assertEqual(next(z for z in after["zones"] if z["zone"] == "H")["requests"], 0)
        self.check_sums(after)

    def test_a_request_in_service_is_not_counted_in_line(self):
        doc = site_doc()
        self.assertEqual(doc["not_in_line"]["in service"], {"requests": 4, "mw": 360.2})
        self.assertTrue(all(not r["in_line"] for r in doc["rows"] if r["status_code"] == "14"))
        key = parsed()[1]
        self.assertEqual([q.group_of(c, key) for c in ("0", "13", "14", "15")], ["withdrawn", "in service", "in service", "in service"])
        self.assertEqual([q.group_of(c, key) for c in ("1", "8", "12", "P")], ["in line"] * 4)
        self.assertEqual(q.group_of("", key), "no status printed")
        self.assertEqual([q.harmonized(c, key) for c in ("0", "14", "12", "13", "")], ["withdrawn", "completed", "active", "active", ""])

    def test_differences_from_the_tracking_sheet_are_listed_not_hidden(self):
        doc = site_doc()
        whats = " | ".join(c["what"] for c in doc["checks"])
        self.assertIn('status 8 ("Rejected Cost Allocation/Next FS Pending")', whats)  # the sheet's status block has no line for 8
        self.assertEqual(doc["nyiso_summary"]["sheet"], "Load Project Tracking")
        self.assertIn("holds no requests", doc["nyiso_summary"]["relation"])


class Stages(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="erw_s140_")
        self.saved = {k: getattr(ip, k) for k in IP_DIRS}
        self.raw, self.manifest = q.RAW, q.MANIFEST
        q.RAW = os.path.join(self.tmp, "saved")
        q.MANIFEST = os.path.join(q.RAW, "manifest.csv")
        os.makedirs(q.RAW)
        self.out = os.path.join(self.tmp, "out")

    def tearDown(self):
        for k, v in self.saved.items():
            setattr(ip, k, v)
        q.RAW, q.MANIFEST = self.raw, self.manifest
        shutil.rmtree(self.tmp, ignore_errors=True)

    def save_cut(self):
        shutil.copy(FIX, os.path.join(q.RAW, "cut.xlsx"))
        with open(FIX, "rb") as f:
            content = f.read()
        q.append_manifest(dict(REC, bytes=str(len(content)), sha256=hashlib.sha256(content).hexdigest()))
        return content

    def run_main(self, *args):
        with mock.patch("sys.stdout"), mock.patch("sys.stderr"):
            return q.main(list(args) + ["--out-dir", self.out])

    def test_write_builds_a_table_the_validator_passes(self):
        self.save_cut()
        self.assertEqual(self.run_main("--write"), 0)
        path = os.path.join(self.out, "nyiso_load_queue.csv")
        report = erw_validate.validate(path)
        self.assertEqual(report["errors"], [])
        with open(path, encoding="utf-8") as f:
            text = f.read()
        lines = [x for x in text.splitlines() if not x.startswith("#")]
        self.assertEqual(lines[0].split(","), q.COLS)
        self.assertEqual(len(lines) - 1, 20)
        self.assertIn("# Source: nyiso:load_queue", text)
        self.assertIn("https://www.nyiso.com/legal-notice", text)
        with open(os.path.join(self.out, "site", "nyiso_load_queue.json"), encoding="utf-8") as f:
            self.assertEqual(json.load(f)["total"]["requests"], 14)
        with open(os.path.join(self.out, "metadata", "sources.csv"), encoding="utf-8", newline="") as f:
            reg = list(csv.DictReader(f))
        self.assertEqual([(r["source"], r["license"], r["tables"]) for r in reg], [("nyiso:load_queue", "public", "nyiso_load_queue")])
        with open(os.path.join(self.out, "status", "nyiso_load_queue.json"), encoding="utf-8") as f:
            self.assertEqual(json.load(f)["results"][0]["status"], "ok")

    def test_ceiling(self):
        self.assertEqual(q.CEILING, 50000)
        self.save_cut()
        self.assertEqual(self.run_main("--write", "--ceiling", "19"), 1)  # twenty rows against nineteen: nothing written
        self.assertFalse(os.path.exists(os.path.join(self.out, "nyiso_load_queue.csv")))
        self.assertFalse(os.path.exists(os.path.join(self.out, "site", "nyiso_load_queue.json")))
        with open(os.path.join(self.out, "status", "nyiso_load_queue.json"), encoding="utf-8") as f:
            res = json.load(f)["results"][0]
        self.assertEqual(res["status"], "failed")
        self.assertIn("ceiling", res["detail"])
        self.assertEqual(self.run_main("--write", "--ceiling", "20"), 0)

    def test_paused_publisher_is_never_requested(self):
        self.save_cut()
        row = {"scope": "nyiso", "paused_on": "2026-10-07", "reason": "a test", "terms_url": "https://example.org", "until": "a person"}
        with mock.patch.object(ip, "paused", return_value=row) as asked, \
                mock.patch.object(ip, "pause_line", return_value="PAUSED since 2026-10-07: a test"), \
                mock.patch.object(q.requests, "get", side_effect=AssertionError("a request was made")) as get:
            self.assertEqual(self.run_main("--pull", "--terms", "--write"), 0)
        asked.assert_called_with("nyiso")
        get.assert_not_called()
        self.assertFalse(os.path.exists(os.path.join(self.out, "nyiso_load_queue.csv")))
        with open(os.path.join(self.out, "status", "nyiso_load_queue.json"), encoding="utf-8") as f:
            self.assertEqual(json.load(f)["results"][0]["status"], "skipped")

    def test_the_connector_asks_before_any_stage_and_names_no_other_operator(self):
        with open(os.path.join(ROOT, "warehouse", "connectors", "nyiso_load_queue.py"), encoding="utf-8") as f:
            src = f.read()
        body = src[src.index("def main("):]
        self.assertLess(body.index("ip.paused(ISO)"), body.index("terms(log)"))
        self.assertLess(body.index("ip.paused(ISO)"), body.index("pull(log)"))
        self.assertNotIn("misoenergy", src)
        self.assertNotIn("pjm.com", src)
        self.assertNotIn(chr(0x2014), src)
        with open(__file__, encoding="utf-8") as f:
            self.assertNotIn(chr(0x2014), f.read())

    def fake_get(self, content):
        resp = mock.Mock(content=content, status_code=200, url=q.URL, headers={"Last-Modified": "Fri, 11 Sep 2026 17:26:02 GMT", "Content-Type": "x"})
        return mock.patch.object(q.requests, "get", return_value=resp)

    def test_pull_makes_one_request_and_keeps_a_manifest(self):
        with open(FIX, "rb") as f:
            content = f.read()
        raw_root = os.path.join(self.tmp, "raw")
        with self.fake_get(content) as get, mock.patch("sys.stdout"):
            q.pull(quiet, raw_root)
        self.assertEqual(get.call_count, 1)
        self.assertEqual(get.call_args[0][0], q.URL)
        (m,) = q.read_manifest()
        self.assertEqual((m["url"], m["status"], m["bytes"], m["sha256"]), (q.URL, "200", str(len(content)), hashlib.sha256(content).hexdigest()))
        self.assertTrue(m["retrieved_at"].endswith("Z"))
        path, rec = q.saved_workbook()
        with open(path, "rb") as f:
            self.assertEqual(f.read(), content)

    def test_pull_refuses_an_answer_that_is_not_the_workbook(self):
        with self.fake_get(b"<html>Access denied</html>"), mock.patch("sys.stdout"):
            with self.assertRaises(RuntimeError):
                q.pull(quiet, os.path.join(self.tmp, "raw"))
        with self.assertRaises(RuntimeError):
            q.saved_workbook()  # the refused answer is kept on disk and never read as the workbook

    def queue_run_copy(self, raw_root, age_days):
        with open(FIX, "rb") as f:
            content = f.read()
        got = dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=age_days)
        d = os.path.join(raw_root, "iso_queues", got.strftime("%Y%m%dT%H%M%SZ"))
        os.makedirs(d)
        with open(os.path.join(d, "00001_NYISO-Interconnection-Queue.xlsx"), "wb") as f:
            f.write(content)
        with open(os.path.join(d, "manifest.csv"), "w", encoding="utf-8", newline="\n") as f:  # the layout iso_prices.RawStore writes
            f.write("retrieved_at,status,bytes,sha256,last_modified,file,url\n")
            f.write(",".join('"' + v + '"' for v in [got.strftime("%Y-%m-%dT%H:%M:%SZ"), "200", str(len(content)), hashlib.sha256(content).hexdigest(),
                                                     "", "00001_NYISO-Interconnection-Queue.xlsx", q.URL]) + "\n")
        return content

    def test_pull_uses_the_weekly_queue_runs_copy_when_it_is_seven_days_old_or_less(self):
        raw_root = os.path.join(self.tmp, "raw")
        content = self.queue_run_copy(raw_root, 2)
        with mock.patch.object(q.requests, "get", side_effect=AssertionError("a request was made")) as get, mock.patch("sys.stdout"):
            q.pull(quiet, raw_root)
        get.assert_not_called()
        (m,) = q.read_manifest()
        self.assertTrue(m["origin"].startswith("copied from "))
        self.assertEqual(m["sha256"], hashlib.sha256(content).hexdigest())

    def test_a_copy_older_than_seven_days_is_not_used(self):
        raw_root = os.path.join(self.tmp, "raw")
        self.queue_run_copy(raw_root, 8)
        self.assertIsNone(q.daily_copy(dt.datetime.now(dt.timezone.utc), raw_root))


if __name__ == "__main__":
    unittest.main()
