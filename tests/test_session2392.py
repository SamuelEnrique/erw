"""Session 2392: ILR's electricity sector statistics (the data.public.lu workbook), Luxembourg.

Energy Research Warehouse (ERW). Tests warehouse/connectors/ilr_luxembourg_electricity.py on a workbook built here in the
download's form (made up for the test, not ILR's): the monthly withdrawals by class in MWh, the current year's months
marked provisional, the customers by class a year, a source's capacity in MW and its production in MWh, the trade by
neighbour a year, a blank not written; a block missing, a value that is not a number and a row given twice stopping the
run; the record's xlsx resource and licence; the sector rule and the holds. No network.

    python -m unittest tests.test_session2392 -v
"""

import io
import json
import os
import sys
import unittest

import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(ROOT, "warehouse", "connectors"))
sys.path.insert(0, os.path.join(ROOT, "warehouse", "metadata"))
sys.path.insert(0, os.path.join(ROOT, "warehouse", "redivis"))
import ilr_luxembourg_electricity as S  # noqa: E402
import build_coverage as B  # noqa: E402
import release_v1 as R  # noqa: E402

BASE = {"source": "s", "source_url": "u", "retrieved_at": "r", "vintage": "v"}
MONTHS = ["Janvier", "Février ", "Mars", "Avril", "Mai", "Juin", "Juillet", "Août ", "Septembre", "Octobre", "Novembre", "Décembre"]
YEARS3 = [None, "2023", "2024", 2025]


def build(withdrawn="1000", drop=None, twice=False):
    wb = openpyxl.Workbook()
    wb.active.title = "Sommaire"
    c = wb.create_sheet("1. Consommation")
    rows = [["Nombre de clients au 31/12"], [], [None, "Clients industriels \n(>2GWh/an)", "Clients résidentiels", "Total"],
            ["2024", 200, 285662, 285862], [2025, 188, None, 288885], [], [],
            ["Electricité prélevée en kWh", "dernière mise à jour le 01.09.2026 - Les données mensuelles publiées d’électricité prélevée de l’année en cours sont provisoires."],
            [], [None] + MONTHS, ["2025"], ["Clients résidentiels"] + [2000.0] * 12, [2026],
            ["Clients résidentiels", 3000.0, float(withdrawn) if withdrawn.replace(".", "").isdigit() else withdrawn] + [None] * 10]
    if twice:
        rows.append(["Clients résidentiels", 1.0] + [None] * 11)
    for r in rows:
        c.append(r)
    p = wb.create_sheet("2. Production")
    for title in ("Puissance installée par source d'énergie en kW", "Volume injecté en kWh", "Nombre de raccordement d'installations de production par type",
                  "Production nationale en kWh"):
        if title.startswith("Production") and drop == "generation":
            continue
        for r in ([title], [], YEARS3, ["Eolienne", 208000, 213735.5, 230715.77], ["Total", 816000, None, 1177320.583], [], []):
            p.append(r)
    a = wb.create_sheet("3. Autoconsommation et gdp")
    for title in ("Nombre d'autoconsommateurs ou partageurs d'énergie", "Puissance installée des autoconsommateurs et groupes de partages en kW",
                  "Volume autoconsommée ou partagé en KWh", "Volume injecté des autoconsommateurs et partageurs en KWh"):
        for r in ([title], [], YEARS3, ["Photovoltaïque", 4249, 12744, 23459], [], []):
            a.append(r)
    x = wb.create_sheet("7. Import_Export")
    for r in (["Données annuelles: Import - Export en kWh"], [], ["Données annuelles"], [], [None, "2024", 2025], ["Export"], ["Belgique", 86174091, 66345818],
              ["Import"], ["Allemagne", 2813862707, 2627257804], [], ["Données mensuelles: Import - Export en kWh"], [None] + MONTHS):
        x.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


class Parse(unittest.TestCase):
    def test_rows(self):
        counts = {}
        t = S.parse(build(), BASE, counts).set_index(["entity", "variable", "ts_utc"])
        r = t.loc[("ilr:lu:clients_residentiels", "withdrawal", "2026-02-01T00:00:00Z")]
        self.assertEqual((r.value, r.unit, r.freq, r.geo, r.x_status, r.node), (1.0, "MWh", "P1M", "LU", "provisional", "Clients résidentiels"))
        self.assertEqual(t.loc[("ilr:lu:clients_residentiels", "withdrawal", "2025-12-01T00:00:00Z")].x_status, "")
        self.assertNotIn(("ilr:lu:clients_residentiels", "withdrawal", "2026-03-01T00:00:00Z"), t.index)
        self.assertEqual(t.loc[("ilr:lu:clients_industriels_2gwh_an", "customers", "2025-01-01T00:00:00Z")].value, 188.0)
        self.assertNotIn(("ilr:lu:clients_residentiels", "customers", "2025-01-01T00:00:00Z"), t.index)
        cap = t.loc[("ilr:lu:eolienne", "capacity", "2024-01-01T00:00:00Z")]
        self.assertEqual((cap.value, cap.unit, cap.freq), (213.7355, "MW", "P1Y"))
        self.assertEqual(t.loc[("ilr:lu:total", "generation", "2025-01-01T00:00:00Z")].value, 1177.320583)
        self.assertNotIn(("ilr:lu:total", "generation", "2024-01-01T00:00:00Z"), t.index)
        self.assertEqual(t.loc[("ilr:lu:photovoltaique", "self_consumers", "2023-01-01T00:00:00Z")].unit, "count")
        self.assertEqual(t.loc[("ilr:lu:allemagne", "import", "2025-01-01T00:00:00Z")].value, 2627257.804)
        self.assertEqual(t.loc[("ilr:lu:belgique", "export", "2024-01-01T00:00:00Z")].value, 86174.091)
        self.assertEqual(counts, {"sheet 7's monthly import and export not taken": 1})

    def test_stops(self):
        for content in (build(drop="generation"), build(withdrawn="about"), build(twice=True)):
            with self.assertRaises(RuntimeError):
                S.parse(content, BASE, {})

    def test_record(self):
        rec = {"license": "cc-by", "resources": [{"format": "pdf", "url": "https://x/a.pdf"}, {"format": "xlsx", "url": "https://x/b.xlsx"}]}
        self.assertEqual(S.xlsx_of(json.dumps(rec).encode()), ("https://x/b.xlsx", "cc-by"))
        rec["resources"].append({"format": "xlsx", "url": "https://x/c.xlsx"})
        with self.assertRaises(RuntimeError):
            S.xlsx_of(json.dumps(rec))


class Registered(unittest.TestCase):
    def test_sector_and_holds(self):
        self.assertEqual(B.sector_of(S.NAME), "power")
        self.assertIn(S.NAME, R.STORAGE_HELD)
        with open(os.path.join(ROOT, "warehouse", "supabase", "live_set.yaml"), encoding="utf-8") as f:
            self.assertIn(f"  - {S.NAME}\n", f.read().replace("\r\n", "\n"))


if __name__ == "__main__":
    unittest.main()
