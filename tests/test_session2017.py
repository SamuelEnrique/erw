"""Session 2017: EIA's buildings surveys (CBECS 2018, RECS 2020).

Energy Research Warehouse (ERW). Tests warehouse/connectors/eia_buildings.py on a sheet built here in the form of EIA's
Table C2 (figures made up for the test, not EIA's): the All buildings row, a bold heading, a level-1 part named with its
parent, million dollars to USD, a code counted with no row; a header that does not read as expected, a row before any
heading, text where a figure belongs, a category twice and an indent under nothing stopping the run; the sector rules
and the holds. No network.

    python -m unittest tests.test_session2017 -v
"""

import io
import os
import sys
import unittest

import openpyxl
from openpyxl.styles import Alignment, Font

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(ROOT, "warehouse", "connectors"))
sys.path.insert(0, os.path.join(ROOT, "warehouse", "metadata"))
sys.path.insert(0, os.path.join(ROOT, "warehouse", "redivis"))
import eia_buildings as E  # noqa: E402
import build_coverage as B  # noqa: E402
import release_v1 as R  # noqa: E402

BASE = {"source": "s", "source_url": "u", "retrieved_at": "r", "vintage": "v"}
HEAD = ["Number of buildings (thousand)", "Total floorspace", "Sum of major fuels", "Electricity", "Natural gas", "Fuel oil", "District heat"]


def book(body, head=HEAD):
    """body: (label, [figures], indent, bold) rows after the All buildings row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(["Release date: December 2022"])
    ws.append(["Table C2. Total energy expenditures by major fuel, 2018"])
    ws.append([None, "All buildings", None, "Total energy expenditures (million dollars)"])
    ws.append([None] + head)
    ws.append(["All buildings", 10, 100, 50, 40, 8, 1, 1])
    for label, figs, indent, bold in body:
        ws.append([label] + figs)
        c = ws.cell(row=ws.max_row, column=1)
        c.alignment = Alignment(indent=indent)
        c.font = Font(b=bold)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


ACT = ("Principal building activity", [], 0, True)


class Parse(unittest.TestCase):
    def test_rows(self):
        counts = {}
        rows = E.parse(book([ACT, ("Health care", [2, 20, 10, 8, 2, "Q", "N"], 0, False),
                             ("Inpatient", [1, 10, 6, 5, 1, "Q", "N"], 1, False), ("Office", [8, 80, 40, 32, 6, 1, 1], 0, False)]),
                       "C2", BASE, counts)
        got = {(r["entity"], r["variable"]): r for r in rows}
        self.assertEqual(counts, {"Q": 2, "N": 2})
        a = got[("cbecs:all", "expenditure_sum_major_fuels")]
        self.assertEqual((a["value"], a["unit"], a["x_level"], a["x_heading"]), (50e6, "USD", 0, ""))
        p = got[("cbecs:principal_building_activity_health_care_inpatient", "expenditure_electricity")]
        self.assertEqual((p["value"], p["x_level"], p["x_parent"], p["node"]), (5e6, 1, "Health care", "Inpatient"))
        self.assertNotIn(("cbecs:principal_building_activity_health_care", "expenditure_fuel_oil"), got)
        self.assertEqual(got[("cbecs:principal_building_activity_office", "expenditure_district_heat")]["geo"], "US")

    def test_stops(self):
        office = ("Office", [8, 80, 40, 32, 6, 1, 1], 0, False)
        for body in ([office],                                                   # a row before any heading
                     [ACT, ("Office", [8, 80, "n/a", 32, 6, 1, 1], 0, False)],    # text where a figure belongs
                     [ACT, office, office],                                       # a category twice
                     [ACT, ("Inpatient", [1, 10, 6, 5, 1, 0, 0], 1, False)],     # indented under nothing
                     [ACT, ("Climate", [], 0, False)]):                           # no figure, not bold
            with self.assertRaises(RuntimeError):
                E.parse(book(body), "C2", BASE, {})
        with self.assertRaises(RuntimeError):
            E.parse(book([ACT, office], head=HEAD[:3] + ["Coal"] + HEAD[4:]), "C2", BASE, {})


class Footnotes(unittest.TestCase):
    """Session 2023: a footnote letter two tables print differently is set aside, so a category joins across them."""

    def row(self, table, heading, *path):
        return {"x_table": table, "x_heading": heading, "x_parent": path[-2] if len(path) > 1 else "", "_path": tuple(path),
                "entity": "x"}

    def test_letters_set_aside_across_tables_only(self):
        rows = [self.row("C1", "Climate zonec", "Cool"), self.row("C2", "Climate Zonea", "Cool"),
                self.row("C1", "Primary source", "Other sourcesd"), self.row("C2", "Primary source", "Other sourcesb"),
                self.row("C1", "Heated", "Gas"), self.row("C1", "Heated", "Gal"),        # one table: not footnotes, kept
                self.row("C1", "Activity", "Health care", "Inpatient")]
        out, canon = E.rename(rows, "cbecs")
        ids = [r["entity"] for r in out]
        self.assertEqual(ids[0], ids[1])
        self.assertEqual(ids[0], "cbecs:climate_zone_cool")
        self.assertEqual(ids[2], ids[3])
        self.assertEqual(ids[2], "cbecs:primary_source_other_sources")
        self.assertNotEqual(ids[4], ids[5])
        self.assertEqual(ids[6], "cbecs:activity_health_care_inpatient")
        self.assertEqual(out[1]["x_heading"], "Climate zone")
        self.assertNotIn("_path", out[0])


class Registered(unittest.TestCase):
    def test_sectors_and_holds(self):
        with open(os.path.join(ROOT, "warehouse", "supabase", "live_set.yaml"), encoding="utf-8") as f:
            live = f.read().replace("\r\n", "\n")
        for sv in E.SURVEYS.values():
            self.assertEqual(B.sector_of(sv["table"]), "power;gas;oil")
            self.assertIn(sv["table"], R.STORAGE_HELD)
            self.assertIn(f"  - {sv['table']}\n", live)


if __name__ == "__main__":
    unittest.main()
