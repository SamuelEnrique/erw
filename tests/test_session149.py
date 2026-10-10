"""Session 149: the owner's rulings applied, and the data loose ends.

  1. No tracked code carries an e-mail address inside a User-Agent, a From header or a contact default. The only
     contact string the ERW sends is "ERW research project, github.com/SamuelEnrique/erw" (the owner's ruling of
     7 October 2026). The digest's own sender and recipient settings are read from the environment
     (DIGEST_FROM, DIGEST_RECIPIENTS) and are allowed, as are addresses at the reserved example domains.

The later sections of this file are named where they begin. No network. Nothing here changes the environment at
import, and every test that needs a table skips on a machine without it.
"""
import os
import re
import subprocess
import sys
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

CONTACT = "ERW research project, github.com/SamuelEnrique/erw"

CODE_SUFFIXES = (".py", ".mjs", ".ts", ".tsx", ".js", ".sh", ".ps1", ".yml", ".yaml")
ADDRESS = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}")
# What makes a line a request's contact: a user agent, a From header, curl's -A, a contact default, a mailto.
CONTEXT = re.compile(
    r"user[-_ ]?agent|\bUA\b|\bcurl\b[^\n]*\s-A\s|[\"']from[\"']\s*:|\bFrom:|contact|mailto:",
    re.IGNORECASE,
)
# The digest's sender and recipients: settings read from the environment, each with its name on the line.
ALLOWED_SETTING = re.compile(r"DIGEST_FROM|DIGEST_RECIPIENTS")
# Domains that reach nobody (RFC 2606 and RFC 6761), and Resend's own test sender.
ALLOWED_DOMAIN = re.compile(r"(?:^|\.)(?:example\.(?:com|org|net)|invalid|test|example|localhost)$|^resend\.dev$",
                            re.IGNORECASE)


def tracked_code(root=ROOT):
    """The tracked code files, by git. None when git cannot say (not a checkout)."""
    try:
        out = subprocess.run(["git", "-C", root, "ls-files", "-z"], capture_output=True, timeout=120)
    except (OSError, subprocess.SubprocessError):
        return None
    if out.returncode != 0:
        return None
    names = [n for n in out.stdout.decode("utf-8", "replace").split("\0") if n]
    return [n for n in names if n.lower().endswith(CODE_SUFFIXES)]


def contact_addresses(text):
    """(line number, domain) of every e-mail address that sits in a request's contact, in one file's text.

    An address counts when its own line, or one of the two lines above it, names a user agent, a From header, a
    contact default or a mailto, unless the line names one of the digest's settings or the address is at a
    reserved domain. Only the domain is returned: a finding never repeats the address itself.
    """
    found = []
    lines = text.splitlines()
    for i, line in enumerate(lines):
        hits = ADDRESS.findall(line)
        if not hits:
            continue
        if ALLOWED_SETTING.search(line):
            continue
        window = "\n".join(lines[max(0, i - 2): i + 1])
        if not CONTEXT.search(window):
            continue
        for address in hits:
            domain = address.rsplit("@", 1)[1]
            if ALLOWED_DOMAIN.search(domain):
                continue
            found.append((i + 1, domain))
    return found


class NoPersonalAddressInAContact(unittest.TestCase):
    """Item 1: the personal contact string is gone, and cannot come back into tracked code unnoticed."""

    def test_the_scanner_finds_an_address_in_a_user_agent(self):
        # built in pieces, so that this file holds no address in a contact of its own
        someone = "some.one" + "@" + "mailhost" + ".com"
        for line in (
            'curl -sS -A "ERW research pilot %s" "$url"' % someone,
            'code=$(curl -sS -L -A "${UA:-ERW research pilot %s}" -o "$out" "$url")' % someone,
            'headers = {"User-Agent": "erw (%s)"}' % someone,
            'UA = "ERW research %s"' % someone,
            'req.add_header("From", "%s")  # From: header' % someone,
            'CONTACT = "%s"' % someone,
        ):
            self.assertEqual(contact_addresses(line), [(1, "mailhost.com")], line)

    def test_the_scanner_reads_a_user_agent_split_over_lines(self):
        someone = "some.one" + "@" + "mailhost" + ".com"
        text = 'headers = {\n    "User-Agent":\n        "erw research (%s)",\n}\n' % someone
        self.assertEqual(contact_addresses(text), [(3, "mailhost.com")])

    def test_the_scanner_allows_the_digest_settings_and_reserved_domains(self):
        sender = "onboarding" + "@" + "resend.dev"
        person = "reader" + "@" + "mailhost.com"
        for line in (
            'sender = env("DIGEST_FROM") or "ERW Energy Digest <%s>"' % sender,
            'KEYS = dict(DIGEST_RECIPIENTS="%s", DIGEST_FROM="ERW <%s>")' % (person, person),
            'self.assertEqual(call["from"], "ERW <%s>")' % ("erw" + "@" + "example.org"),
            'headers = {"User-Agent": "%s"}' % CONTACT,
            'UA = "x (%s)"' % ("a" + "@" + "b.example.invalid"),
        ):
            self.assertEqual(contact_addresses(line), [], line)

    def test_an_address_outside_a_contact_is_not_this_tests_business(self):
        self.assertEqual(contact_addresses('subscriber = "%s"' % ("reader" + "@" + "mailhost.com")), [])

    def test_the_ruled_contact_string_holds_no_address(self):
        self.assertEqual(CONTACT, "ERW research project, github.com/SamuelEnrique/erw")
        self.assertIsNone(ADDRESS.search(CONTACT))

    def test_no_tracked_code_sends_an_address_as_its_contact(self):
        names = tracked_code()
        if names is None:
            self.skipTest("git cannot list the tracked files here")
        self.assertGreater(len(names), 200, "the list of tracked code is too short to be the repository's")
        bad = []
        for name in names:
            path = os.path.join(ROOT, name)
            try:
                with open(path, encoding="utf-8", errors="replace") as f:
                    text = f.read()
            except OSError:
                continue
            for line_no, domain in contact_addresses(text):
                bad.append("%s:%d (an address at %s)" % (name, line_no, domain))
        self.assertEqual(bad, [], "an e-mail address inside a User-Agent, a From header or a contact default: "
                                  "send exactly %r" % CONTACT)


# ---------------------------------------------------------------------------------------------------------------
# 2. The daily run's failed steps (run 37633030030 of 7 October 2026, and every run from 2 or 4 October).
#    grid_network: the raw folder of the interchange connector holds the data lock's own check, whose whole answer
#    is the JSON value true; read as an EIA page it raised AttributeError. build_status: the interchange connector's
#    gap row names no day (its market is "pairs") and its table has two columns of its own, which read_series refuses.
# ---------------------------------------------------------------------------------------------------------------
import json
import tempfile
from unittest import mock


def src(*parts):
    with open(os.path.join(ROOT, *parts), encoding="utf-8") as f:
        return f.read()


class TheNetworkBuildPassesOverAnAnswerThatIsNotAPage(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        for d in ("connectors", "derived"):
            p = os.path.join(ROOT, "warehouse", d)
            if p not in sys.path:
                sys.path.insert(0, p)
        try:
            import grid_network
        except ImportError as exc:  # a machine without pandas
            raise unittest.SkipTest("grid_network cannot be imported here: %s" % exc)
        cls.gn = grid_network

    def raw(self, files):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        run = os.path.join(tmp.name, "eia930_interchange", "20261007T142035Z")
        os.makedirs(run)
        for name, text in files.items():
            with open(os.path.join(run, name), "w", encoding="utf-8") as f:
                f.write(text)
        return tmp.name

    def test_the_lock_check_saved_beside_the_pages_does_not_stop_the_build(self):
        # the file the runner holds: 20261007T142059.123456Z_00009_erw_lock_check, four bytes
        page = {"response": {"total": 1, "data": [
            {"period": "2026-10-05T01", "fromba": "ERCO", "fromba-name": "Electric Reliability Council of Texas, Inc.",
             "toba": "SWPP", "toba-name": "Southwest Power Pool", "value": "-55"}]}}
        root = self.raw({
            "20261007T142040.000001Z_00001_data_frequency_hourly": json.dumps(page),
            "20261007T142059.000001Z_00009_erw_lock_check": "true",
            "20261007T142059.000002Z_00010_erw_lock_renew": "false",
            "20261007T142059.000003Z_00011_a_list": "[1, 2]",
            "20261007T142059.000004Z_00012_a_number": "3",
            "20261007T142059.000005Z_00013_response_true": json.dumps({"response": True}),
            "20261007T142059.000006Z_00014_data_null": json.dumps({"response": {"data": None}}),
            "20261007T142059.000007Z_00015_rows_not_rows": json.dumps({"response": {"data": [True, "x"]}}),
            "20261007T142059.000008Z_00016_not_json": "<html>busy</html>",
            "manifest.csv": "retrieved_at,status,bytes,sha256,last_modified,file,url\n",
        })
        with mock.patch.object(self.gn.ip, "RAW_DIR", root):
            names = self.gn.ba_names()
        self.assertEqual(names["SWPP"], "Southwest Power Pool")
        self.assertEqual(names["ERCO"], "Electric Reliability Council of Texas, Inc.")

    def test_the_fault_was_the_bool(self):
        # what the builder did until session 149, on the same four bytes
        with self.assertRaises(AttributeError) as c:
            json.loads("true").get("response", {})
        self.assertIn("'bool' object has no attribute 'get'", str(c.exception))


class StatusListsAGapThatNamesNoDay(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        for d in ("connectors", "metadata"):
            p = os.path.join(ROOT, "warehouse", d)
            if p not in sys.path:
                sys.path.insert(0, p)
        try:
            import build_status
        except ImportError as exc:
            raise unittest.SkipTest("build_status cannot be imported here: %s" % exc)
        cls.bs = build_status

    def test_a_day_is_a_calendar_day(self):
        for good in ("2026-10-06", "2024-02-29"):
            self.assertTrue(self.bs.is_day(good), good)
        for bad in ("pairs", "", "2026-10-06..2026-10-07", "2026-13-01", "2026-10-06T00", "20261006", None):
            self.assertFalse(self.bs.is_day(bad), bad)

    def test_the_interchange_gap_is_not_read_as_a_day(self):
        # the runner's file: the series columns and the table's own two, which read_series refuses
        ip = self.bs.ip
        with tempfile.TemporaryDirectory() as tmp:
            with open(os.path.join(tmp, "eia930_all_interchange.csv"), "w", encoding="utf-8", newline="\n") as f:
                f.write("# Energy Research Warehouse (ERW): a header line\n" + ",".join(ip.SERIES_COLS + ["ba", "x_to_ba"]) + "\n")
            with mock.patch.object(ip, "OUT_DIR", tmp):
                with self.assertRaises(RuntimeError):  # the refusal that failed the step, still the reader's rule
                    ip.read_series(os.path.join(tmp, "eia930_all_interchange.csv"))
                self.assertIsNone(self.bs.day_complete("eia930_all_interchange", "pairs"))
                self.assertIsNone(self.bs.day_complete("eia930_all_interchange", "2026-10-06"))

    def test_the_builder_lists_such_a_gap(self):
        text = src("warehouse", "metadata", "build_status.py")
        self.assertIn("if not is_day(g.day):", text)
        self.assertIn("listed, not re-checked", text)



# ---------------------------------------------------------------------------------------------------------------
# 3. ISO-NE's monthly file (Aggregate Monthly DDG Undelivered Energy): pulled under its ceilings, held internal.
#    No ISO-NE workbook is in the repository (the table is internal). The layout is tested on a workbook made here
#    that holds the labels and the one line session 144's report already printed (wind, January 2026: delivered
#    385,416.50 MWh, undelivered 5,879.80 MWh); the whole pull is tested on the workbooks themselves where they are
#    (the data machine's raw store) and skipped elsewhere.
# ---------------------------------------------------------------------------------------------------------------
import io

RAW_ISONE = os.path.join(os.path.dirname(ROOT), "erw", "warehouse", "raw", "isone_ddg_undelivered")
LABELS = [None, "DE [MWH]", "UE [MWH]", "UE-NonBind [MWH]", "UE-Bind [MWH]", "percent UE", "percent UE-NonBind", "percent UE-Bind"]
MONTH_NAMES = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def isone_module():
    p = os.path.join(ROOT, "warehouse", "connectors")
    if p not in sys.path:
        sys.path.insert(0, p)
    try:
        import isone_ddg_undelivered
        import openpyxl  # noqa: F401  (the workbook made here, and pandas' reader)
    except ImportError as exc:
        raise unittest.SkipTest("the ISO-NE connector cannot be imported here: %s" % exc)
    return isone_ddg_undelivered


def workbook(sheet="2026 System", labels=LABELS, months=MONTH_NAMES, lines=None):
    """A workbook in ISO-NE's layout: the sheet's mark, the labels, twelve month lines. lines: {month number: cells}."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Nomenclature"
    wb.active.append(["ISO-NE Public"])
    ws = wb.create_sheet(sheet)
    ws.append([None, None, None, "ISO-NE Public"])
    ws.append(labels)
    for i, m in enumerate(months):
        ws.append([m] + list((lines or {}).get(i + 1, [" ", " ", " ", " ", None, None, None])))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class IsoNeUndelivered(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.c = isone_module()

    def test_the_ceilings_and_the_contact_string(self):
        c = self.c
        self.assertEqual((c.MAX_ROWS, c.MAX_REQUESTS, c.MAX_BYTES), (2000, 20, 50 * 1024 * 1024))
        self.assertEqual(c.UA, {"User-Agent": CONTACT})
        self.assertIsNone(ADDRESS.search(json.dumps(c.UA)))
        self.assertLessEqual(c.ROWS_EST, 12 * len(c.COLUMNS))

    def test_a_month_line_as_printed_and_blank_months(self):
        c = self.c
        rows = c.parse(workbook(lines={1: [385416.5, 5879.8, " ", " ", None, None, None]}), "2026 Undelivered Wind Energy Aggregate Report")
        self.assertEqual([(r["variable"], r["year"], r["month"], r["value"], r["unit"]) for r in rows],
                         [("wind_delivered_mwh", 2026, 1, 385416.5, "MWh"), ("wind_undelivered_mwh", 2026, 1, 5879.8, "MWh")])
        self.assertEqual(c.parse(workbook(), "2026 Undelivered Solar Energy Aggregate Report"), [])   # nothing published: no row

    def test_the_fuel_and_the_year_come_from_iso_nes_own_description(self):
        c = self.c
        line = {1: [385416.5, 5879.8, " ", " ", None, None, None]}
        self.assertEqual({r["fuel"] for r in c.parse(workbook(sheet="2025 System", lines=line), "2025 Undelivered Solar Energy Aggregate Report")}, {"solar"})
        for bad in ("", "2026 Something Else", "Undelivered Wind Energy Aggregate Report"):
            with self.assertRaises(RuntimeError):
                c.parse(workbook(lines=line), bad)
        with self.assertRaises(RuntimeError):   # the description's year has no sheet
            c.parse(workbook(sheet="2025 System", lines=line), "2026 Undelivered Wind Energy Aggregate Report")

    def test_a_layout_that_was_not_read_stops_the_read(self):
        c = self.c
        name = "2026 Undelivered Wind Energy Aggregate Report"
        with self.assertRaises(RuntimeError):   # a column session 149 did not read
            c.parse(workbook(labels=LABELS + ["UE-Other [MWH]"]), name)
        with self.assertRaises(RuntimeError):   # the months out of order
            c.parse(workbook(months=["Feb", "Jan"] + MONTH_NAMES[2:]), name)
        with self.assertRaises(RuntimeError):   # eleven month lines
            c.parse(workbook(months=MONTH_NAMES[:11]), name)
        with self.assertRaises(RuntimeError):   # words where a number should be
            c.parse(workbook(lines={1: [385416.5, "n/a", " ", " ", None, None, None]}), name)

    def test_the_stop_sits_before_the_request(self):
        c = self.c
        asked = []
        with tempfile.TemporaryDirectory() as tmp, mock.patch.object(c, "RAW", tmp), mock.patch.object(c, "PAUSE", 0), \
                mock.patch.object(c, "http_get", side_effect=lambda url: asked.append(url)):
            for i in range(c.MAX_REQUESTS):
                c.record(retrieved_at="2026-10-07T00:00:00Z", status=200, bytes=10, sha256="", kind="list", file="", url="u%d" % i)
            with self.assertRaises(c.CeilingStop):
                c.fetch("https://www.iso-ne.com/x", "list", "x.json", lambda m: None)
        self.assertEqual(asked, [])   # the twenty-first request is never made
        with tempfile.TemporaryDirectory() as tmp, mock.patch.object(c, "RAW", tmp), mock.patch.object(c, "http_get", side_effect=lambda url: asked.append(url)):
            with mock.patch.object(c, "counted", return_value=(c.MAX_ROWS - 83, 3, 0)):
                with self.assertRaises(c.CeilingStop):   # 84 more rows would pass 2,000
                    c.fetch("https://www.iso-ne.com/y.xlsx", "workbook", "workbooks/y.xlsx", lambda m: None, rows_est=c.ROWS_EST)
            with mock.patch.object(c, "counted", return_value=(0, 3, c.MAX_BYTES - 10)):
                with self.assertRaises(c.CeilingStop):   # the bytes
                    c.fetch("https://www.iso-ne.com/y.xlsx", "workbook", "workbooks/y.xlsx", lambda m: None, bytes_est=c.BYTES_EST)
        self.assertEqual(asked, [])

    def test_an_access_control_is_recorded_and_left(self):
        c = self.c

        class Answer:
            def __init__(self, code, body):
                self.status_code, self.content = code, body
        for answer, kind in ((Answer(403, b"denied"), "workbook"), (Answer(200, b"<html>are you a person?</html>"), "workbook")):
            with tempfile.TemporaryDirectory() as tmp, mock.patch.object(c, "RAW", tmp), mock.patch.object(c, "PAUSE", 0), \
                    mock.patch.object(c, "http_get", return_value=answer):
                with self.assertRaises(c.AccessStop):
                    c.fetch("https://www.iso-ne.com/z.xlsx", kind, "workbooks/z.xlsx", lambda m: None)
                self.assertEqual([r["kind"] for r in c.manifest()], ["refused"])
                self.assertEqual(c.workbooks(), {})

    def test_it_is_internal_everywhere_it_is_named(self):
        c = self.c
        text = src("warehouse", "connectors", "isone_ddg_undelivered.py")
        self.assertIn('f"License: internal. The owner\'s ruling of 7 October 2026: pulled and held internal.', text)
        self.assertIn('license="internal"', text)
        live = src("warehouse", "supabase", "live_set.yaml")
        self.assertIn("- " + c.NAME, live.split("catalogue_hold:", 1)[1].split("review_hold:", 1)[0])
        self.assertIn("- " + c.SOURCE, live.split("sources_hold:", 1)[1])
        self.assertIn('(r"^isone_ddg_undelivered_monthly$", "power")', src("warehouse", "metadata", "build_coverage.py"))
        import csv
        with open(os.path.join(ROOT, "warehouse", "metadata", "sources.csv"), encoding="utf-8", newline="") as f:
            row = {r["source"]: r for r in csv.DictReader(f)}[c.SOURCE]
        self.assertEqual((row["license"], row["tables"]), ("internal", c.NAME))
        self.assertIn("held internal", row["report"])
        self.assertNotIn("isone_ddg_undelivered", src("warehouse", "run_daily.sh"))   # pulled once: on no schedule

    def test_no_page_reads_it_and_the_placeholder_says_held(self):
        hits = []
        for folder in ("app", "lib", "components", "data", "public"):
            for base, dirs, files in os.walk(os.path.join(ROOT, "site", folder)):
                dirs[:] = [d for d in dirs if d not in ("node_modules", ".next")]
                for name in files:
                    if not name.endswith((".ts", ".tsx", ".json", ".mjs", ".js", ".csv")):
                        continue
                    try:
                        with open(os.path.join(base, name), encoding="utf-8", errors="replace") as f:
                            if "isone_ddg_undelivered" in f.read():
                                hits.append(os.path.join(base, name))
                    except OSError:
                        continue
        # session 153: Ask ERCOT's tool names the table in one file, in order to refuse it by its name ("held, not
        # shown"); that file reads no file and no row of it (tests/test_session153.py holds what its lines may say)
        refuser = os.path.normcase(os.path.join(ROOT, "site", "lib", "chat", "pagefiles.ts"))
        self.assertEqual([h for h in hits if os.path.normcase(h) != refuser], [])
        face = src("site", "lib", "freeenergy.ts")
        self.assertIn("the undelivered energy of its dispatchable wind and solar plants. Held, not shown.", face)
        self.assertNotIn("Not yet in the ERW", face)
        note = src("docs", "methods", "curtailment.md")
        self.assertIn("pulled and held internal, not shown", note)
        self.assertIn('"You are also hereby put on notice that the Content is protected by copyright under United States laws. '
                      'Any duplication of the Content or non-personal use may violate copyright, trademark, and other laws."', note)


@unittest.skipUnless(os.path.exists(os.path.join(RAW_ISONE, "manifest.csv")), "ISO-NE's workbooks are not on this machine")
class IsoNeUndeliveredOnTheWorkbooks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.c = isone_module()
        cls.patch = mock.patch.object(cls.c, "RAW", RAW_ISONE)
        cls.patch.start()
        cls.table, cls.lines = cls.c.build()

    @classmethod
    def tearDownClass(cls):
        cls.patch.stop()

    def test_the_pull_stayed_under_its_ceilings(self):
        c = self.c
        rows, requests_made, received = c.counted()
        self.assertEqual(rows, len(self.table))
        self.assertLessEqual(rows, c.MAX_ROWS)
        self.assertLessEqual(requests_made, c.MAX_REQUESTS)
        self.assertLessEqual(received, c.MAX_BYTES)
        self.assertEqual(sorted({r["kind"] for r in c.manifest()}), ["list", "terms", "workbook"])   # nothing refused, nothing unread
        self.assertTrue(all(r["url"].startswith("https://www.iso-ne.com/") for r in c.manifest()))

    def test_the_line_session_144_printed(self):
        t = self.table.set_index(["variable", "ts_utc"])["value"]
        self.assertAlmostEqual(t[("wind_delivered_mwh", "2026-01-01T00:00:00Z")], 385416.50, places=2)
        self.assertAlmostEqual(t[("wind_undelivered_mwh", "2026-01-01T00:00:00Z")], 5879.80, places=2)

    def test_every_month_once_and_iso_nes_figures_agree_with_each_other(self):
        t = self.table
        self.assertFalse(t.duplicated(["entity", "variable", "ts_utc"]).any())
        self.assertEqual(set(t["entity"]), {"isone:system"})
        self.assertEqual(set(t["freq"]), {"P1M"})
        self.assertTrue(t["ts_utc"].str.endswith("-01T00:00:00Z").all())
        w = t.pivot(index="ts_utc", columns="variable", values="value")
        for fuel in ("wind", "solar"):
            g = w[[c for c in w.columns if c.startswith(fuel + "_")]].dropna()
            self.assertGreater(len(g), 0)
            ue, de = g[fuel + "_undelivered_mwh"], g[fuel + "_delivered_mwh"]
            self.assertLess((ue - g[fuel + "_undelivered_nonbinding_mwh"] - g[fuel + "_undelivered_binding_mwh"]).abs().max(), 0.01)
            self.assertLess((g[fuel + "_undelivered_pct"] - 100 * ue / (ue + de)).abs().max(), 1e-6)   # a percent, not a fraction
            self.assertTrue((g[fuel + "_undelivered_pct"] <= 100).all() and (g >= 0).all().all())
        self.assertEqual(w["wind_delivered_mwh"].dropna().index.min()[:7], "2018-01")

    def test_the_notice_saved_is_the_notice_quoted(self):
        import hashlib
        import html
        with open(os.path.join(RAW_ISONE, "legal_notice.html"), "rb") as f:
            body = f.read()
        row = [r for r in self.c.manifest() if r["kind"] == "terms"][0]
        self.assertEqual(hashlib.sha256(body).hexdigest(), row["sha256"])
        words = " ".join(html.unescape(re.sub(r"(?s)<[^>]+>", " ", body.decode("utf-8", "replace"))).split())
        self.assertIn("You are also hereby put on notice that the Content is protected by copyright under United States laws. "
                      "Any duplication of the Content or non-personal use may violate copyright, trademark, and other laws.", words)


# ---------------------------------------------------------------------------------------------------------------
# 7. NYISO's terms and the load queue rows. The owner's ruling of 7 October 2026: the terms are quoted, and the rows
#    are shown only if the words allow it. They confer no license and reserve every right, so the rows are held
#    internal: the table, the registry, the page's file (no request, zone or megawatt) and the page's face.
# ---------------------------------------------------------------------------------------------------------------
RAW_NYISO = os.path.join(os.path.dirname(ROOT), "erw", "warehouse", "raw", "nyiso_load_queue")


def nyiso_module():
    p = os.path.join(ROOT, "warehouse", "connectors")
    if p not in sys.path:
        sys.path.insert(0, p)
    try:
        import nyiso_load_queue
    except ImportError as exc:
        raise unittest.SkipTest("the NYISO load queue connector cannot be imported here: %s" % exc)
    return nyiso_load_queue


class NyisoLoadQueueIsHeld(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.q = nyiso_module()

    def test_the_switch_is_off_and_the_license_follows_it(self):
        q = self.q
        self.assertIs(q.SHOWN, False)
        self.assertEqual(q.LICENSE, "internal")
        text = src("warehouse", "connectors", "nyiso_load_queue.py")
        self.assertIn('LICENSE = "public" if SHOWN else "internal"', text)
        self.assertIn("license=LICENSE, tables=[NAME]", text)
        self.assertIn("write_site(doc if SHOWN else held_site(doc), site_file, log)", text)
        self.assertIn('("License: internal. Session 149', text)

    def test_every_sentence_that_bears_on_copying_is_quoted(self):
        q = self.q
        self.assertEqual(len(q.TERMS_QUOTES), 6)
        joined = " ".join(q.TERMS_QUOTES)
        for words in ("does not confer any license or ownership interest", "expressly reserves such rights and property in its entirety",
                      "republishing, retransmitting, reproducing", "prior written permission", "All Rights Reserved"):
            self.assertIn(words, joined)
        note = src("docs", "methods", "datacenter_cost.md")
        for sentence in q.TERMS_QUOTES:
            self.assertIn(" ".join(sentence.split()), " ".join(note.split()), sentence[:60])   # word for word in the Method note
        self.assertIn("They do not allow it.", note)
        self.assertIn("NYISO's terms do not allow it", note)
        self.assertIn("does not allow showing the rows", q.TERMS_READING)

    def test_the_file_the_page_reads_holds_no_request(self):
        q = self.q
        doc = json.loads(src("site", "data", "nyiso_load_queue.json"))
        self.assertEqual(sorted(doc), ["built", "license", "shown", "source", "why", "words"])
        self.assertEqual((doc["shown"], doc["license"], doc["words"], doc["why"]), (False, "internal", q.HELD_WORDS, q.HELD_WHY))
        self.assertEqual(sorted(doc["source"]), ["name", "publisher", "report_page", "sheet_names", "terms", "url"])
        self.assertEqual(doc["source"]["terms"]["quotes"], q.TERMS_QUOTES)
        self.assertEqual(doc["source"]["terms"]["license"], "internal")
        self.assertLess(len(json.dumps(doc)), 4000)   # it held 74 requests in 74 kB

    def test_the_face_reads_the_placeholder(self):
        page = src("site", "app", "cost-of-power", "page.tsx")
        self.assertIn("const NY_SHOWN = NYLOAD.shown !== false && Array.isArray(NYLOAD.rows) && Array.isArray(NYLOAD.zones);", page)
        self.assertIn('<span key="ny" data-nyload-held="1"><Missing words={NY_HELD_WORDS} why={NY_HELD_WHY} /></span>', page)
        self.assertIn('{x.grid === "nyiso" && NY_SHOWN ? (', page)
        self.assertEqual(page.count("<NyLoadCell"), 1)          # only behind the switch
        self.assertIn("NY_SHOWN ? <NyLoadCell", page)
        self.assertIn('"/cost-of-power": "review"', src("site", "lib", "release.ts").replace("\n", " ").replace("  ", " "))

    def test_the_registry_says_internal_and_the_live_set_holds_it(self):
        import csv
        with open(os.path.join(ROOT, "warehouse", "metadata", "sources.csv"), encoding="utf-8", newline="") as f:
            row = {r["source"]: r for r in csv.DictReader(f)}["nyiso:load_queue"]
        self.assertEqual((row["license"], row["tables"]), ("internal", "nyiso_load_queue"))
        live = src("warehouse", "supabase", "live_set.yaml")
        self.assertIn("- nyiso_load_queue", live.split("catalogue_hold:", 1)[1].split("review_hold:", 1)[0])
        self.assertIn("- nyiso:load_queue", live.split("sources_hold:", 1)[1])

    @unittest.skipUnless(os.path.exists(os.path.join(RAW_NYISO, "manifest.csv")), "NYISO's saved legal notice is not on this machine")
    def test_the_quotes_are_in_the_notice_as_saved(self):
        import csv
        import hashlib
        q = self.q
        with open(os.path.join(RAW_NYISO, "manifest.csv"), encoding="utf-8", newline="") as f:
            row = [r for r in csv.DictReader(f) if r["url"] == q.TERMS_URL and r["status"] == "200"][-1]
        with open(os.path.join(RAW_NYISO, row["file"]), "rb") as f:
            body = f.read()
        self.assertEqual(hashlib.sha256(body).hexdigest(), row["sha256"])
        text = q.page_text(body)
        for sentence in q.TERMS_QUOTES:
            self.assertIn(sentence, text, sentence[:60])
        for grant in ("may be used", "may be reproduced", "you may copy", "permission is granted", "is hereby granted"):
            self.assertNotIn(grant, text.lower(), grant)   # no sentence of the notice grants leave


# ---------------------------------------------------------------------------------------------------------------
# 4. ERCOT's load zones, daily, at the ceiling asked. Session 140 asked "a new ceiling, or count only new rows" and
#    named no number; the owner approved "at the ceiling asked"; the narrower reading is taken: the ceiling stays
#    3,000,000 and a row counts once. On the saved real days of session 140's fixtures (tests/fixtures/session140).
# ---------------------------------------------------------------------------------------------------------------
import hashlib
import shutil

FIX140 = os.path.join(ROOT, "tests", "fixtures", "session140")


def zone_module():
    p = os.path.join(ROOT, "warehouse", "connectors")
    if p not in sys.path:
        sys.path.insert(0, p)
    try:
        import ercot_zone_prices
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise unittest.SkipTest("the load zone connector cannot be imported here: %s" % exc)
    return ercot_zone_prices


def first_day_only(content):
    """ERCOT's two-day fixture workbook with its second sheet (the November day) taken out: an earlier copy of a
    year's workbook, as ERCOT posts it before the later days exist. Real rows; none is changed."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content))
    for ws in wb.worksheets[1:]:
        wb.remove(ws)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class TheLoadZoneCeilingCountsARowOnce(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.z = zone_module()
        with open(os.path.join(FIX140, "dam_2015_two_days.xlsx"), "rb") as f:
            cls.whole = f.read()
        cls.early = first_day_only(cls.whole)

    def setUp(self):
        self.raw = tempfile.mkdtemp(prefix="erw_s149_zone_raw_")
        self.addCleanup(shutil.rmtree, self.raw, ignore_errors=True)
        self.lines = []

    def log(self, m):
        self.lines.append(str(m))

    def save(self, content, doc, published, year=2015):
        z = self.z
        name = "dam_%d_%s.zip" % (year, doc)
        with open(os.path.join(self.raw, name), "wb") as f:
            f.write(content)
        z._append_csv(os.path.join(self.raw, "manifest.csv"), z.MANIFEST_COLS,
                      {"market": "dam", "year": year, "file": name, "doc_id": doc, "url": z.FILE_URL.format(doc=doc), "bytes": len(content),
                       "sha256": hashlib.sha256(content).hexdigest(), "published": published, "friendly_name": "",
                       "retrieved_at": "2026-10-07T09:00:00Z", "how": "requested"})

    def read_all(self):
        for m in self.z.manifest(self.raw):
            self.z.parsed(m, self.log, self.raw)

    def test_the_ceiling_is_the_one_already_set(self):
        self.assertEqual(self.z.CEILING, 3_000_000)
        text = src("warehouse", "connectors", "ercot_zone_prices.py")
        self.assertIn("A ROW COUNTS", text)
        self.assertIn('"at the ceiling\nasked"'.replace("\n", " "), " ".join(text.split()))

    def test_a_later_copy_counts_only_the_rows_it_adds(self):
        z = self.z
        self.save(self.early, "1", "2015-03-15T08:01:37-05:00")       # the March day: 23 hours, eight zones
        self.read_all()
        self.assertEqual(z.counted(self.raw), 8 * 23)
        self.save(self.whole, "2", "2015-11-08T08:01:37-05:00")       # ERCOT posts the year again: the March day and a November day
        self.assertEqual(z.counted(self.raw), z.bound("dam", 2015))    # not read yet: the year as if full, never an undercount
        self.read_all()
        self.assertEqual(z.counted(self.raw), 8 * 48)                  # each row once: 184 held and 200 new
        every_copy = sum(int(z.counts(self.raw)[m["sha256"]]["rows_kept"]) for m in z.manifest(self.raw))
        self.assertEqual(every_copy, 8 * 23 + 8 * 48)                  # what the count was until session 149
        self.assertEqual(z.counts(self.raw)["distinct:dam:2015"]["rows"], 8 * 48)
        self.save(self.whole, "3", "2015-11-15T08:01:37-05:00")       # the same workbook downloaded again
        self.assertEqual(z.counted(self.raw), 8 * 48)                  # counts nothing

    def test_the_request_that_would_pass_the_ceiling_is_refused_before_it_is_made(self):
        z = self.z
        import test_session140_zones as t140
        self.save(self.early, "1", "2015-03-15T08:01:37-05:00")
        self.read_all()
        room = z.room_needed("dam", 2015, self.raw)
        self.assertEqual(room, z.bound("dam", 2015) - 8 * 23)          # the most a new copy can add: the year less what is held
        get = t140.Getter([2015], "dam")
        a = t140.args(years=[2015], markets=["dam"], ceiling=8 * 23 + room - 1)
        self.assertEqual(z.pull(a, self.log, get=get, sleep=lambda s: None, raw=self.raw), 3)
        self.assertEqual([u for u in get.asked if "mirDownload" in u], [])      # the list was read; no workbook was asked for
        self.assertTrue(any(line.startswith("STOPPED before dam 2015") for line in self.lines))
        get = t140.Getter([2015], "dam")
        a = t140.args(years=[2015], markets=["dam"], ceiling=8 * 23 + room)     # exactly enough room: asked
        self.assertEqual(z.pull(a, self.log, get=get, sleep=lambda s: None, raw=self.raw), 0)
        self.assertEqual(len([u for u in get.asked if "mirDownload" in u]), 1)
        self.assertTrue(any("counted against the ceiling" in line for line in self.lines))   # the count is in every run's log

    def test_the_refresh_asks_only_about_the_years_still_growing(self):
        z = self.z
        self.assertEqual(z.growing_years(self.raw, now_year=2026), [2025, 2026])      # no final copy of 2025 held
        self.save(self.whole, "9", "2026-01-01T08:44:09-05:00", year=2025)            # published after its year ended
        self.assertEqual(z.growing_years(self.raw, now_year=2026), [2025, 2026])      # real time's final copy is not held yet
        text = src("warehouse", "connectors", "ercot_zone_prices.py")
        self.assertIn("a.pull = a.write = True", text)
        self.assertIn("a.years = growing_years(raw)", text)


sys.path.insert(0, os.path.join(ROOT, "tests"))


# ---------------------------------------------------------------------------------------------------------------
# 5 and 6. The page files rebuilt with SPP South's years, and the data machine's daily run that rebuilds them.
# ---------------------------------------------------------------------------------------------------------------
class SppSouthHasItsYears(unittest.TestCase):
    def test_the_datacenter_page_holds_spp_south_from_2019(self):
        index = json.loads(src("site", "data", "datacenter", "index.json"))
        spp = index["grids"]["spp"]
        self.assertEqual(spp["years"], list(range(2019, 2027)))
        south = [r for r in spp["regions"] if "SOUTH" in json.dumps(r).upper()][0]
        self.assertEqual(south["da"]["first"], "2019-01-01T06:00:00Z")
        self.assertGreater(south["da"]["hours"], 68000)
        for year in range(2019, 2027):
            self.assertTrue(os.path.exists(os.path.join(ROOT, "site", "data", "datacenter", "spp_%d.json" % year)), year)

    def test_the_capture_file_holds_spp_south_from_2019(self):
        cap = json.loads(src("site", "data", "seller", "capture.json"))
        south = [h for h in cap["grids"]["spp"]["hubs"] if "SOUTH" in json.dumps({k: v for k, v in h.items() if not isinstance(v, dict)}).upper()][0]
        self.assertEqual(south["da"]["first"], "2019-01-01T06:00:00Z")
        self.assertGreater(south["da"]["hours"], 68000)
        self.assertIn("2019-01", south["da"]["wind"])


class TheDataMachinesDailyRun(unittest.TestCase):
    SCRIPT = os.path.join("warehouse", "run_data_machine.sh")

    def run_script(self, *args, **env):
        bash = shutil.which("bash")
        if not bash:
            self.skipTest("bash is not on this machine")
        full = dict(os.environ)
        full.update(env)
        return subprocess.run([bash, self.SCRIPT.replace(os.sep, "/")] + list(args), cwd=ROOT, env=full, capture_output=True, text=True, timeout=120)

    def test_it_refuses_on_the_runner_before_anything_else(self):
        r = self.run_script(GITHUB_ACTIONS="true", ERW_ROLE="data")
        self.assertEqual(r.returncode, 0)
        self.assertTrue(r.stdout.startswith("run_data_machine REFUSED: this is the GitHub runner"), r.stdout[:200])
        text = src("warehouse", "run_data_machine.sh").split("set -uo pipefail", 1)[1]   # the commands, past the header's words
        self.assertLess(text.index('"${GITHUB_ACTIONS:-}" = "true"'), text.index("warehouse/lock.py role"))
        self.assertLess(text.index("REFUSED: this is the GitHub runner"), text.index("soft_step dm_sync"))

    def test_it_refuses_on_a_machine_that_is_not_a_data_machine_or_holds_no_history(self):
        r = self.run_script("--check", GITHUB_ACTIONS="", ERW_ROLE="code")
        self.assertEqual(r.returncode, 0)
        self.assertIn("REFUSED: this machine's role is 'code', not data", r.stdout)
        if not os.path.exists(os.path.join(ROOT, "warehouse", "output", "ercot_all_hub_prices_history.csv")):
            r = self.run_script("--check", GITHUB_ACTIONS="", ERW_ROLE="data")
            self.assertEqual(r.returncode, 0)
            self.assertIn("REFUSED: this machine does not hold the price histories", r.stdout)

    def test_every_builder_is_a_recorded_step_and_no_workflow_calls_it(self):
        text = src("warehouse", "run_data_machine.sh")
        self.assertIn(". warehouse/soft_step.sh", text)            # under warehouse/health.py: tried once more, recorded, never stops the run
        for step, builder in (("dm_datacenter_page", "warehouse/derived/datacenter_page.py"), ("dm_capture_price", "warehouse/derived/capture_price.py"),
                              ("dm_curtailment_shares", "warehouse/derived/curtailment_shares.py"), ("dm_free_energy", "warehouse/derived/free_energy.py"),
                              ("dm_curtailment_worth", "warehouse/derived/curtailment_worth.py"),
                              ("dm_ercot_zone_prices", "warehouse/connectors/ercot_zone_prices.py --refresh")):
            self.assertIn('soft_step %s "$PYTHON" %s' % (step, builder), text)
        self.assertIn('"$PYTHON" warehouse/health.py run --step "$name" -- "$@"', src("warehouse", "soft_step.sh"))
        self.assertNotIn("soft_step dm_ercot_estimate", text)       # the runner builds the Texas file every day, from a newer table
        self.assertNotIn("ercot_estimate_page.py", text.split("set -uo pipefail", 1)[1])
        self.assertNotIn(chr(0x2014), text)
        for name in os.listdir(os.path.join(ROOT, ".github", "workflows")):
            self.assertNotIn("run_data_machine", src(".github", "workflows", name), name)
        self.assertNotIn("run_data_machine.sh", src("warehouse", "run_daily.sh").split("set -", 1)[-1])
        self.assertIn("The data machine's daily run (session 149)", src("docs", "machines.md"))

    def test_the_gate_is_not_piped_and_the_commit_names_its_paths(self):
        text = src("warehouse", "run_data_machine.sh")
        gate = [ln for ln in text.splitlines() if ln.lstrip().startswith('"$PYTHON" warehouse/validate/erw_validate.py')]
        self.assertEqual(len(gate), 1)
        self.assertNotIn("|", gate[0])
        self.assertIn("rc=$?", text.split(gate[0], 1)[1].splitlines()[1])
        self.assertIn('if [ "$branch" != "main" ]; then', text)
        self.assertIn("site/data/datacenter site/data/seller/capture.json site/data/curtailment/shares.json", text)
        self.assertNotIn("git add -A", text)
        self.assertNotIn("git add .", text)
        self.assertNotIn("--force", text)


if __name__ == "__main__":
    unittest.main()
