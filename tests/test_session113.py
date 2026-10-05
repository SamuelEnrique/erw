"""Session 113: ERCOT's yearly reserve price file in both forms, and the deals tracker's second version.

Energy Research Warehouse (ERW). No request leaves the machine. Part one reads two real documents saved in
tests/fixtures/session113/ (public: ERCOT's terms, item 5), each as ERCOT served it on 5 October 2026:

    ercot_np4_181_er_2026_workbook_20261004.zip   the 2026 file published 4 October 2026: one Excel workbook, the
                                                  form that stopped the daily run that day
    ercot_np4_181_er_2025_csv.zip                 the 2025 file: one CSV, the form the connector was written for

Part two reads the deals table and the news table, both in git, and runs site/lib/deals2.ts in Node on the rows as
the live set serves them: every count and sum the page shows is compared with the same figure computed here from the
CSV. The figures are computed, not written down, so the daily extraction's new deals do not break the tests.

    python -m unittest tests.test_session113 -v
"""

import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree

import openpyxl
import pandas as pd

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
for p in ("warehouse", "warehouse/connectors", "warehouse/deals"):
    sys.path.insert(0, os.path.join(ROOT, p))

import iso_prices as ip  # noqa: E402
import ercot_as_prices as ercot  # noqa: E402
import duplicates as dup  # noqa: E402

FIX = os.path.join(ROOT, "tests", "fixtures", "session113")
WORKBOOK = "ercot_np4_181_er_2026_workbook_20261004.zip"
CSV = "ercot_np4_181_er_2025_csv.zip"
HEAD = ["Delivery Date", "Hour Ending", "Repeated Hour Flag", "REGDN", "REGUP", "RRS", "NSPIN", "ECRS"]
ORDER = ("REGDN", "REGUP", "RRS", "NSPIN", "ECRS")


def quiet(msg):
    pass


def fixture(name):
    with open(os.path.join(FIX, name), "rb") as f:
        return f.read()


def zipped(members):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as o:
        for name, body in members.items():
            o.writestr(name, body)
    return out.getvalue()


def workbook(rows, lead=7, sheets=1):
    """A made-up workbook laid out as ERCOT's: `lead` rows above the header (its logo and title), then the rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(sheets - 1):
        wb.create_sheet()
    for i in range(lead):
        ws.append(["Historical DAM Clearing Prices for Capacity"] if i == 5 else [])
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def hour(rows, service, day, he, flag="N"):
    ts = ercot.to_utc(pd.Series([day]), pd.Series([he]), pd.Series([flag]))[0]
    got = rows[(rows["service"] == service) & (rows["ts"] == ts)]["value"].tolist()
    return got[0] if len(got) == 1 else got


class TheWorkbook(unittest.TestCase):
    """The 2026 file as ERCOT posts it since 4 October 2026."""

    @classmethod
    def setUpClass(cls):
        cls.rows = ercot.parse_year(fixture(WORKBOOK))

    def test_it_is_the_form_that_failed(self):
        names = zipfile.ZipFile(io.BytesIO(fixture(WORKBOOK))).namelist()
        # the name run_status.csv quotes for the failed run of 4 October
        self.assertEqual(names, ["rpt.00013091.0000000000000000.20261004.080005.DAMASMCPC_2026.xlsx"])

    def test_every_hour_of_every_service(self):
        # 1 January to 3 October 2026 is 276 Central days, one of them 23 hours long (8 March)
        self.assertEqual(self.rows.groupby("service").size().to_dict(), {s: 276 * 24 - 1 for s in ercot.SERVICES})
        self.assertEqual((ip.utc_iso(self.rows["ts"].min()), ip.utc_iso(self.rows["ts"].max())),
                         ("2026-01-01T06:00:00Z", "2026-10-04T04:00:00Z"))
        kept, gaps = ercot.complete_days(self.rows, quiet)
        self.assertEqual((len(kept), gaps), (len(self.rows), []))
        self.assertEqual(int((self.rows["day"] == "03/08/2026").sum()), 23 * 5)

    def test_the_first_and_the_last_hour_by_hand(self):
        # read in the workbook's rows 9 and 6631: 01/01/2026 01:00 and 10/03/2026 24:00
        self.assertEqual([hour(self.rows, s, "01/01/2026", "01:00") for s in ORDER], [0.92, 0.87, 0.5, 1.97, 0.64])
        self.assertEqual([hour(self.rows, s, "10/03/2026", "24:00") for s in ORDER], [1.25, 7.42, 5.95, 20.3, 9.5])

    def test_no_price_is_rounded(self):
        """Every price against the text the workbook's own sheet stores, read here without openpyxl."""
        outer = zipfile.ZipFile(io.BytesIO(fixture(WORKBOOK)))
        book = zipfile.ZipFile(io.BytesIO(outer.read(outer.namelist()[0])))
        m = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        strings = ["".join(t.text or "" for t in si.iter(m + "t"))
                   for si in ElementTree.fromstring(book.read("xl/sharedStrings.xml")).findall(m + "si")]
        sheet = [n for n in book.namelist() if re.fullmatch(r"xl/worksheets/sheet1\.xml", n, re.I)]
        self.assertEqual(len(sheet), 1)
        stored, text = {}, {}
        for row in ElementTree.fromstring(book.read(sheet[0])).iter(m + "row"):
            for c in row.findall(m + "c"):
                v = c.find(m + "v")
                if v is None:
                    continue
                at = (int(row.get("r")), re.match(r"[A-Z]+", c.get("r")).group(0))
                if c.get("t") == "s":
                    text[at] = strings[int(v.text)]
                else:
                    stored[at] = v.text
        head = {col: text[(8, col)].strip() for col in "DEFGH"}
        self.assertEqual(head, dict(zip("DEFGH", HEAD[3:])))
        by = {(s, t): v for s, t, v in zip(self.rows["service"], self.rows["ts"], self.rows["value"])}
        days = {}
        for (r, col), v in stored.items():
            self.assertRegex(v, r"^\d+(\.\d{1,2})?$")  # dollars and cents: the sheet stores nothing longer
            key = (text[(r, "A")], text[(r, "B")], text[(r, "C")])
            if key not in days:
                days[key] = ercot.to_utc(pd.Series([key[0]]), pd.Series([key[1]]), pd.Series([key[2]]))[0]
            self.assertEqual(by[(head[col], days[key])], float(v), f"row {r} column {col}")
        self.assertEqual(len(stored), len(self.rows))


class TheCsv(unittest.TestCase):
    """The 2025 file: the form of every year before, still read as it was."""

    @classmethod
    def setUpClass(cls):
        cls.rows = ercot.parse_year(fixture(CSV))

    def test_a_whole_year(self):
        self.assertEqual(zipfile.ZipFile(io.BytesIO(fixture(CSV))).namelist(),
                         ["rpt.00013091.0000000000000000.20260101081728.DAMASMCPC_2025.csv"])
        self.assertEqual(self.rows.groupby("service").size().to_dict(), {s: 8760 for s in ercot.SERVICES})
        kept, gaps = ercot.complete_days(self.rows, quiet)
        self.assertEqual((len(kept), gaps), (43800, []))

    def test_by_hand(self):
        # the file's first and last lines, and the two 01:00 to 02:00 hours of 2 November
        self.assertEqual([hour(self.rows, s, "01/01/2025", "01:00") for s in ORDER], [0.64, 1.09, 0.63, 0.63, 0.03])
        self.assertEqual([hour(self.rows, s, "12/31/2025", "24:00") for s in ORDER], [0.85, 1.34, 0.99, 2.59, 0.99])
        self.assertEqual((hour(self.rows, "REGDN", "11/02/2025", "02:00", "N"),
                          hour(self.rows, "REGDN", "11/02/2025", "02:00", "Y")), (1.5, 1.75))
        self.assertEqual((int((self.rows["day"] == "03/09/2025").sum()), int((self.rows["day"] == "11/02/2025").sum())),
                         (23 * 5, 25 * 5))


class BothForms(unittest.TestCase):
    def test_one_year_reads_the_same_in_either_form(self):
        """The real 2025 CSV, written out as a workbook laid out like ERCOT's, reads to the same rows."""
        z = zipfile.ZipFile(io.BytesIO(fixture(CSV)))
        lines = [ln.split(",") for ln in z.read(z.namelist()[0]).decode("utf-8").splitlines()]
        body = [ln[:3] + [float(c) if c.strip() else None for c in ln[3:]] for ln in lines[1:]]
        book = ercot.parse_year(zipped({"year.xlsx": workbook([[c.strip() for c in lines[0]]] + body)}))
        pd.testing.assert_frame_equal(book, ercot.parse_year(fixture(CSV)))

    def test_an_empty_cell_is_no_row_in_a_workbook_too(self):
        rows = ercot.parse_year(zipped({"x.xlsx": workbook([HEAD, ["06/09/2023", "01:00", "N", 1, 2, 3, 4, None],
                                                            ["06/10/2023", "01:00", "N", 1, 2, 3, 4, 5.25]])}))
        self.assertEqual(rows[rows["service"] == "ECRS"]["value"].tolist(), [5.25])
        self.assertEqual(len(rows), 9)

    def test_a_zip_with_neither_or_both_is_refused(self):
        good = workbook([HEAD, ["06/10/2023", "01:00", "N", 1, 2, 3, 4, 5]])
        text = ",".join(HEAD) + "\n06/10/2023,01:00,N,1,2,3,4,5\n"
        for members in ({"a.txt": "x"}, {"a.csv": text, "b.xlsx": good}, {"a.xlsx": good, "b.xlsx": good},
                        {"a.csv": text, "b.csv": text}):
            with self.assertRaisesRegex(RuntimeError, "expected one CSV or one Excel workbook"):
                ercot.parse_year(zipped(members))

    def test_a_workbook_that_changed_again_is_refused(self):
        row = ["06/10/2023", "01:00", "N", 1, 2, 3, 4, 5]
        bad = {
            "a price that is not a number": workbook([HEAD, row[:7] + ["n/a"]]),
            "0 header rows": workbook([["Date", "Hour"] + HEAD[2:], row]),
            "2 header rows": workbook([HEAD, row, HEAD, row]),
            "expected one$": workbook([HEAD, row], sheets=2),
            "is not text": workbook([HEAD, [20230610] + row[1:]]),
            "layout changed": workbook([HEAD[:7] + ["FFR"], row]),
            "holds no hour": workbook([HEAD]),
        }
        for message, body in bad.items():
            with self.assertRaisesRegex(RuntimeError, message):
                ercot.parse_year(zipped({"x.xlsx": body}))


# ---------------------------------------------------------------------------
# Part two: the deals tracker, version 2
# ---------------------------------------------------------------------------

EVENT_COLS = ["event_id", "event_date", "status", "mw", "price", "currency", "parties", "source", "source_url"]
STANDARD = EVENT_COLS + ["event_type", "entity_ids"]
SELECTIONS = [{}, {"view": "ai"}, {"view": "storage"}, {"type": "ppa"}, {"type": "offtake"}, {"type": "project_finance"},
              {"type": "acquisition"}, {"type": "tolling"}, {"type": "other"}, {"grid": "ERCOT"}, {"grid": "PJM"},
              {"year": "2025"}, {"year": "2026"}, {"party": "amazon"}, {"tech": "nuclear"}, {"tech": "not stated"},
              {"view": "ai", "type": "ppa", "year": "2026"}, {"type": "nonsense", "grid": "Mars", "year": "1999"}]
NAMED = {"ppa": "ppa", "offtake": "offtake", "project_finance": "project_finance", "m_and_a": "acquisition", "tolling": "tolling"}
PJM = {"VA": "PJM", "MD": "PJM", "OH": "PJM", "PA": "PJM", "WV": "PJM", "NJ": "PJM", "DE": "PJM", "DC": "PJM"}


def src(*parts):
    with open(os.path.join(ROOT, *parts), encoding="utf-8") as f:
        return f.read()


def table(name):
    return pd.read_csv(os.path.join(ROOT, "warehouse", "output", name + ".csv"), comment="#", dtype=str, keep_default_na=False)


def live_rows(d):
    """The table as the live set serves it to the site: the events columns, and every other stated field in extra."""
    out = []
    for r in d.to_dict("records"):
        row = {c: (r[c] if r[c] != "" else None) for c in EVENT_COLS}
        row["mw"] = float(r["mw"]) if r["mw"] != "" else None
        row["price"] = float(r["price"]) if r["price"] != "" else None
        row["extra"] = {k: v for k, v in r.items() if k not in STANDARD and v != ""}
        out.append(row)
    return out


def node(js):
    exe = shutil.which("node")
    if not exe:
        raise unittest.SkipTest("node is not on this machine")
    r = subprocess.run([exe, "--input-type=module", "-e", js], cwd=os.path.join(ROOT, "site"), capture_output=True, text=True, timeout=120)
    if r.returncode != 0:
        if "ERR_UNKNOWN_FILE_EXTENSION" in r.stderr or "Unknown file extension" in r.stderr:
            raise unittest.SkipTest("this node does not read TypeScript files")
        raise AssertionError(r.stderr[-2000:])
    return json.loads(r.stdout.strip().splitlines()[-1])


def words(r):
    return {w for c in ("buyer", "seller", "other_parties") for w in re.findall(r"[a-z]{4,}", r[c].lower())}


def number(v):
    return float(v) if v != "" else None


class DealsRegister(unittest.TestCase):
    def test_the_register_is_sound_and_the_site_holds_it(self):
        reg, d = dup.register(), table("energy_deals")
        self.assertEqual(dup.check(reg, d), [])
        held = json.loads(src("site", "data", "deal_duplicates.json"))
        self.assertEqual(held["pairs"], dup.snapshot(reg)["pairs"])
        self.assertEqual(set(reg["ruling"]), {"fold", "doubtful"})

    def test_a_folded_pair_is_close_in_time_and_shares_a_party(self):
        reg, d = dup.register(), table("energy_deals").set_index("event_id")
        for p in reg[reg["ruling"] == "fold"].to_dict("records"):
            a, b = d.loc[p["duplicate_event_id"]], d.loc[p["kept_event_id"]]
            days = abs((pd.Timestamp(a["event_date"][:10]) - pd.Timestamp(b["event_date"][:10])).days)
            self.assertLessEqual(days, 14, p["duplicate_event_id"])
            self.assertTrue(words(a) & words(b), p["duplicate_event_id"])

    def test_a_bad_register_is_refused(self):
        reg, d = dup.register(), table("energy_deals")
        one = reg.iloc[:1].copy()
        self.assertTrue(dup.check(one.assign(kept_event_id="deal:none"), d))
        self.assertTrue(dup.check(one.assign(ruling="merge"), d))
        self.assertTrue(dup.check(one.assign(kept_event_id=one["duplicate_event_id"]), d))
        chain = pd.concat([one, one.assign(duplicate_event_id=one["kept_event_id"], kept_event_id=reg.iloc[1]["kept_event_id"])])
        self.assertTrue(dup.check(chain, d))

    def test_every_deal_has_a_source_link(self):
        d = table("energy_deals")
        for c in ("source_url", "story_urls"):
            self.assertTrue(d[c].str.startswith("http").all(), c)
        self.assertTrue((d["story_urls"].str.split(";").map(len) == d["n_stories"].astype(int)).all())


class DealsLibrary(unittest.TestCase):
    """site/lib/deals2.ts on the table's rows, against the same figures computed here."""

    @classmethod
    def setUpClass(cls):
        d = table("energy_deals")
        folds = dup.register()
        folds = folds[folds["ruling"] == "fold"]
        cls.raw = d
        cls.d = d[~d["event_id"].isin(folds["duplicate_event_id"])].copy()
        cls.folds = folds
        tmp = tempfile.mkdtemp(prefix="erw113_")
        cls.addClassCleanup(shutil.rmtree, tmp, True)
        path = os.path.join(tmp, "rows.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump(live_rows(d), f)
        cls.got = node(
            "const fs = (await import('node:fs')).default; const m = await import('./lib/deals2.ts');"
            f"const rows = JSON.parse(fs.readFileSync({json.dumps(path)}, 'utf-8'));"
            "const pairs = JSON.parse(fs.readFileSync('data/deal_duplicates.json', 'utf-8')).pairs;"
            "const all = m.fold(rows.map(m.toDeal), pairs);"
            f"const out = {json.dumps(SELECTIONS)}.map((sp) => {{ const x = m.inputsOf(sp, all); const shown = m.select(all, x);"
            " return { x, href: m.hrefOf(x), ids: shown.map((d) => d.id), counts: m.counts(shown), sentence: m.summary(all, shown, x),"
            " months: m.months(all, shown) }; });"
            "console.log(JSON.stringify({ n: all.length, deals: all.map((d) => ({ id: d.id, mw: d.mw, price: d.price, other: d.priceOther,"
            " dollars: d.dollars, grid: d.grid, links: d.links, folded: d.folded })), out, empty: m.summary([], [], m.inputsOf({}, [])) }));")
        cls.by = {x["id"]: x for x in cls.got["deals"]}

    def sel(self, sp):
        return self.got["out"][SELECTIONS.index(sp)]

    def test_folded_pairs_are_one_deal_with_both_rows_links(self):
        self.assertEqual(self.got["n"], len(self.raw) - len(self.folds))
        raw = self.raw.set_index("event_id")
        for p in self.folds.to_dict("records"):
            self.assertNotIn(p["duplicate_event_id"], self.by)
            kept = self.by[p["kept_event_id"]]
            self.assertEqual(kept["folded"], [p["duplicate_event_id"]])
            want = raw.loc[p["kept_event_id"], "story_urls"].split(";") + raw.loc[p["duplicate_event_id"], "story_urls"].split(";")
            self.assertEqual(kept["links"], list(dict.fromkeys(want)))
            self.assertEqual(kept["dollars"], number(raw.loc[p["kept_event_id"], "dollars"]))  # the kept row's own figure

    def test_a_figure_not_stated_is_never_a_number(self):
        raw = self.d.set_index("event_id")
        for i, x in self.by.items():
            r = raw.loc[i]
            self.assertEqual((x["mw"], x["dollars"], x["price"]), (number(r["mw"]), number(r["dollars"]), number(r["price"])), i)
            self.assertEqual(x["other"] is None, r["price_value"] == "", i)
            self.assertTrue(x["links"] and all(u.startswith("http") for u in x["links"]), i)

    def test_every_selection_against_the_table(self):
        d = self.d
        group = d["deal_type"].map(lambda t: NAMED.get(t, "other"))
        ai = d["ai_power"].str.lower() == "true"
        storage = (d["technology"].str.contains(r"batter|storage|\bbess\b", case=False)
                   | d["asset"].str.contains(r"batter|energy storage|\bbess\b", case=False))
        grid = d["state"].map({"TX": "ERCOT", **PJM}).fillna("")
        party = (d["buyer"] + ";" + d["seller"] + ";" + d["other_parties"]).str.lower()
        year = d["event_date"].str[:4]
        masks = [None, ai, storage, group == "ppa", group == "offtake", group == "project_finance", group == "acquisition",
                 group == "tolling", group == "other", grid == "ERCOT", grid == "PJM", year == "2025", year == "2026",
                 party.str.contains("amazon"), d["technology"] == "nuclear", d["technology"] == "",
                 ai & (group == "ppa") & (year == "2026"), None]
        self.assertEqual(len(masks), len(SELECTIONS))
        for sp, m, got in zip(SELECTIONS, masks, self.got["out"]):
            want = d if m is None else d[m]
            self.assertEqual(sorted(got["ids"]), sorted(want["event_id"]), sp)
            c = got["counts"]
            mw = pd.to_numeric(want.loc[want["mw"] != "", "mw"])
            usd = pd.to_numeric(want.loc[want["dollars"] != "", "dollars"])
            self.assertEqual((c["deals"], c["withMw"], c["mw"], c["withDollars"], c["withPrice"]),
                             (len(want), len(mw), float(mw.sum()), len(usd), int((want["price"] != "").sum())), sp)
            self.assertAlmostEqual(c["dollars"], float(usd.sum()), places=2)
            self.assertEqual({x["month"]: x["n"] for x in got["months"] if x["n"]}, want["event_date"].str[:7].value_counts().to_dict(), sp)
        self.assertEqual(sum(self.got["out"][0]["counts"]["byGroup"].values()), len(d))

    def test_the_months_run_without_a_hole(self):
        ms = [x["month"] for x in self.got["out"][0]["months"]]
        first, last = self.d["event_date"].str[:7].min(), self.d["event_date"].str[:7].max()
        self.assertEqual(ms, [p.strftime("%Y-%m") for p in pd.period_range(first, last, freq="M")])

    def test_an_address_the_page_does_not_offer_shows_every_deal(self):
        got = self.sel({"type": "nonsense", "grid": "Mars", "year": "1999"})
        self.assertEqual((got["href"], got["counts"]["deals"]), ("/deals/v2", len(self.d)))
        self.assertEqual(self.sel({"view": "ai", "type": "ppa", "year": "2026"})["href"], "/deals/v2?view=ai&type=ppa&year=2026")

    def test_the_sentence_is_the_counts(self):
        for sp in SELECTIONS:
            got = self.sel(sp)
            c, s = got["counts"], got["sentence"]
            if c["deals"] == 0:
                self.assertTrue(s.startswith("No deal held matches this selection"), s)
                continue
            self.assertIn(f"{len(self.d):,} deals", s)
            if c["deals"] > 1:
                self.assertIn(f"{c['mw']:,.0f} MW in all" if c["withMw"] else "None of them states a size", s)
                self.assertIn("none states a price" if c["withPrice"] == 0 else f"{c['withPrice']:,}", s)
            if got["href"] != "/deals/v2":
                self.assertTrue(s.startswith(f"{c['deals']:,} of the {len(self.d):,} deals held"), s)
        self.assertEqual(self.got["empty"], "No deal is held.")

    def test_tolling_is_not_a_type_the_extraction_has(self):
        body = re.search(r"^DEAL_TYPES = \[(.*?)\]", src("warehouse", "deals", "extract.py"), re.S | re.M).group(1)
        types = re.findall(r'"([a-z_]+)"', body)
        self.assertEqual(len(types), 14)  # the page says fourteen
        self.assertNotIn("tolling", types)
        self.assertEqual(self.sel({"type": "tolling"})["counts"]["deals"], int((self.d["deal_type"] == "tolling").sum()))
        self.assertTrue(set(self.raw["deal_type"]) <= set(types))

    def test_the_grid_is_only_a_state_one_operator_mostly_serves(self):
        raw = self.d.set_index("event_id")
        for i, x in self.by.items():
            if x["grid"]:
                self.assertTrue(raw.loc[i, "state"], i)
        states = re.search(r"STATE_GRID: Record<string, string> = \{(.*?)\};", src("site", "lib", "deals2.ts"), re.S).group(1)
        for split in ("IL", "IN", "MO", "KY", "NC", "MT", "NM", "AR", "MS", "ND", "SD"):  # two grids, or one and none
            self.assertNotRegex(states, rf"\b{split}:")


class DealsPage(unittest.TestCase):
    def test_the_page_says_what_it_is(self):
        page = src("site", "app", "deals", "v2", "page.tsx")
        for said in ("not a complete record of the market", "state a size in MW", "state a price in US dollars per MWh", "Nothing is estimated",
                     "Storage only", "AI power", "Deals by month", 'method="get"', "data-deals-summary"):
            self.assertIn(said, page)
        for head in ("Date", "Parties", "Technology", "Size", "Price", "Source"):
            self.assertIn(f'"{head}"', page)
        self.assertIn('export const NOT_STATED = "not stated"', src("site", "lib", "deals2.ts"))
        self.assertNotRegex(page, r"anthropic|\bfetch\(")  # no model call, no request of its own

    def test_it_is_in_review_and_the_older_page_is_as_it_was(self):
        release = src("site", "lib", "release.ts")
        self.assertRegex(release, r'"/deals/v2": "review"')
        self.assertRegex(release, r'"/deals": "review"')
        r = subprocess.run(["git", "status", "--porcelain", "--", "site/app/deals/page.tsx", "site/app/deals/DealsTable.tsx"],
                           cwd=ROOT, capture_output=True, text=True)
        self.assertEqual(r.stdout.strip(), "")

    def test_the_empty_months_are_the_unscored_ones(self):
        """The chart's note: the stories of June to September 2025 are held but not scored, and no deal is dated there."""
        n, d = table("news_stories"), table("energy_deals")
        month = n["event_date"].str[:7]
        for m in ("2025-06", "2025-07", "2025-08", "2025-09"):
            self.assertGreater(int((month == m).sum()), 0, m)
            self.assertEqual(int(((month == m) & (n["significance"] != "")).sum()), 0, m)
            self.assertEqual(int((d["event_date"].str[:7] == m).sum()), 0, m)

    def test_no_em_dash(self):
        for parts in (("site", "app", "deals", "v2", "page.tsx"), ("site", "lib", "deals2.ts"), ("warehouse", "deals", "duplicates.py"),
                      ("warehouse", "deals", "duplicates.csv"), ("site", "data", "deal_duplicates.json"), ("tests", "test_session113.py"),
                      ("warehouse", "connectors", "ercot_as_prices.py")):
            self.assertNotIn(chr(0x2014), src(*parts), parts)


if __name__ == "__main__":
    unittest.main()
