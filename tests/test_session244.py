"""Session 244: ERCOT's real-time reserve price adders by hour (warehouse/connectors/ercot_rt_adders.py).

Energy Research Warehouse (ERW). docs/methods/ercot_rt_price_adders.md. What is held here: how SCED intervals become
hours (the time weights, the 15 minutes a value may hold, an hour written only when all of it is covered), ERCOT's
stamps at the two clock changes, the three forms of the yearly workbook and which document of 2015 is read, and the
wiring (the sector rule, the live set's holds, the data machine's step, the method).

The intervals of the first test are real: ERCOT's rows for 12:00 to 13:00 Central on 15 February 2021, copied from the
workbook RTM_ORDC_REL_DPLY_PRC_ADDR_RSRV_2021 as session 244 read it. No request: workbooks are built in memory.
"""

import datetime as dt
import io
import os
import sys
import unittest
import zipfile

import openpyxl
import pandas as pd
import yaml

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(ROOT, "warehouse", "connectors"))
sys.path.insert(0, os.path.join(ROOT, "warehouse", "metadata"))
import ercot_rt_adders as c  # noqa: E402

EM = chr(0x2014)
# (Central stamp, RTORPA) as ERCOT published them, 15 February 2021
URI = [("11:57:55", 194.5017), ("12:00:33", 150.6694), ("12:05:22", 199.3594), ("12:10:24", 297.7572), ("12:15:20", 490.7322),
       ("12:20:22", 318.3860), ("12:25:21", 131.9920), ("12:30:22", 88.1409), ("12:35:20", 108.3254), ("12:40:21", 77.2204),
       ("12:45:21", 79.8974), ("12:50:21", 29.3315), ("12:55:20", 30.3973), ("12:59:15", 114.8161), ("13:04:00", 100.0)]


def src(*parts):
    with open(os.path.join(ROOT, *parts), encoding="utf-8") as f:
        return f.read()


def utc(text):
    return pd.Timestamp(text, tz="UTC")


def stamps(start, minutes):
    return [utc(start) + pd.Timedelta(minutes=m) for m in minutes]


def workbook(title, head, months, info=False, first_row=8):
    """A yearly zip as ERCOT posts it: one workbook, a sheet a month, the header under a title."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if info:
        ws = wb.create_sheet("Report Info")
        ws.append(["Historical Real-Time Price Adders by SCED Interval"])
        ws.append(["SCED Timestamp", "SCED Timestamp"])   # the note names the column too: it is not a month's sheet
    for name, rows in months.items():
        ws = wb.create_sheet(name)
        for _ in range(first_row - 3):
            ws.append([None])
        ws.append([title])
        ws.append([None])
        ws.append(head)
        for r in rows:
            ws.append(r)
    book = io.BytesIO()
    wb.save(book)
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w") as f:
        f.writestr("workbook.xlsx", book.getvalue())
    return z.getvalue()


OLD = ["Batch ID", "      SCED Timestamp      ", "Repeated Hour Flag", "System Lamda", "PRC", "RTOLCAP", "RTOFFCAP", "RTORPA", "RTOFFPA", "RTOLHSL"]
MID = OLD + ["RTORDPA", "RTRRUC"]
NEW = ["      SCED Timestamp      ", "Repeated Hour Flag", "System Lamda", "RTRDPA", "RTRDPARUS", "RTOLHSL"]


class FromIntervalsToHours(unittest.TestCase):
    def test_the_hour_of_15_february_2021_is_the_time_weighted_mean_of_ercots_own_rows(self):
        local = pd.Series([pd.Timestamp(f"2021-02-15 {t}") for t, _ in URI])
        ts = c.to_utc(local, pd.Series(["N"] * len(URI)))
        h = c.hourly(ts, [v for _, v in URI])
        hour = h.loc[utc("2021-02-15T18:00:00Z")]
        self.assertEqual(hour["covered"], 3600)
        self.assertEqual(round(hour["mean"], 4), 169.0334)          # the figure of docs/methods/ercot_rt_price_adders.md
        self.assertEqual(hour["max"], 490.7322)
        # by hand: the value at 11:57:55 holds the hour's first 33 seconds, the one at 12:59:15 its last 45
        held = [33, 289, 302, 296, 302, 299, 301, 298, 301, 300, 300, 299, 235, 45]
        self.assertEqual(sum(held), 3600)
        self.assertAlmostEqual(hour["mean"], sum(s * v for s, (_, v) in zip(held, URI)) / 3600, places=9)
        # the hours on either side are open: not all of their seconds are covered
        self.assertLess(h.loc[utc("2021-02-15T17:00:00Z"), "covered"], 3600)
        self.assertLess(h.loc[utc("2021-02-15T19:00:00Z"), "covered"], 3600)

    def test_a_value_holds_to_the_next_stamp_and_for_no_longer_than_15_minutes(self):
        # five-minute stamps through 10:00 to 12:00, but for a silence from 10:40 to 10:58
        minutes = [m for m in range(-5, 125, 5) if not 40 < m < 58] + [58]
        ts = stamps("2024-07-18T10:00:00Z", sorted(minutes))
        h = c.hourly(ts, [1.0] * len(ts))
        self.assertEqual(h.loc[utc("2024-07-18T10:00:00Z"), "covered"], 3600 - 3 * 60)   # 10:40 holds to 10:55, not to 10:58
        self.assertEqual(h.loc[utc("2024-07-18T11:00:00Z"), "covered"], 3600)
        self.assertEqual(c.MAX_GAP_MIN, 15)
        # a silence of exactly 15 minutes is covered: the value holds to the next stamp
        ts = stamps("2024-07-18T10:00:00Z", [-5, 0, 15, 30, 45, 60, 65])
        self.assertEqual(c.hourly(ts, [2.0] * len(ts)).loc[utc("2024-07-18T10:00:00Z"), "covered"], 3600)

    def test_the_mean_weights_by_seconds_and_the_max_is_the_largest_value_that_held(self):
        ts = stamps("2022-07-13T19:00:00Z", [-1, 0, 6, 9, 50, 60, 61])
        ts[4] = utc("2022-07-13T19:20:00Z")                       # 19:09 holds 11 minutes; then 19:20, 19:30, ... every ten
        ts = sorted(set(ts + stamps("2022-07-13T19:00:00Z", [30, 40, 50])))
        v = {0: 0.0, 6: 600.0, 9: 0.0}
        values = [v.get(int((t - utc("2022-07-13T19:00:00Z")).total_seconds() // 60), 0.0) for t in ts]
        h = c.hourly(ts, values)
        hour = h.loc[utc("2022-07-13T19:00:00Z")]
        self.assertEqual(hour["covered"], 3600)
        self.assertAlmostEqual(hour["mean"], 600.0 * 180 / 3600)     # three minutes at 600 in an hour of zeros: 30
        self.assertEqual(hour["max"], 600.0)

    def test_an_empty_cell_covers_nothing_and_the_last_stamp_holds_nothing(self):
        ts = stamps("2015-06-25T05:00:00Z", range(0, 125, 5))
        values = [float("nan")] * 12 + [0.5] * 13                   # no value before 06:00: ERCOT's column is empty until then
        h = c.hourly(ts, values)
        self.assertNotIn(utc("2015-06-25T05:00:00Z"), h.index)
        self.assertEqual(h.loc[utc("2015-06-25T06:00:00Z"), "covered"], 3600)
        self.assertNotIn(utc("2015-06-25T07:00:00Z"), h.index)      # the stamp at 07:00 is the last: nothing says how long it held

    def test_stamps_out_of_order_or_repeated_are_an_error_not_a_guess(self):
        ts = stamps("2024-01-01T06:00:00Z", [0, 10, 5])
        with self.assertRaises(RuntimeError):
            c.hourly(ts, [0.0, 0.0, 0.0])
        with self.assertRaises(RuntimeError):
            c.hourly(stamps("2024-01-01T06:00:00Z", [0, 5, 5]), [0.0, 0.0, 0.0])


class ErcotsStamps(unittest.TestCase):
    def test_central_time_to_utc_in_winter_and_in_summer(self):
        local = pd.Series([pd.Timestamp("2021-02-15 12:00:33"), pd.Timestamp("2022-07-13 14:00:10")])
        got = c.to_utc(local, pd.Series(["N", "N"]))
        self.assertEqual(list(got), [utc("2021-02-15T18:00:33Z"), utc("2022-07-13T19:00:10Z")])

    def test_the_repeated_hour_of_november_is_told_apart_by_ercots_flag(self):
        local = pd.Series([pd.Timestamp("2024-11-03 01:55:10"), pd.Timestamp("2024-11-03 01:00:12"), pd.Timestamp("2024-11-03 02:00:11")])
        got = c.to_utc(local, pd.Series(["N", "Y", "N"]))
        self.assertEqual(list(got), [utc("2024-11-03T06:55:10Z"), utc("2024-11-03T07:00:12Z"), utc("2024-11-03T08:00:11Z")])
        self.assertTrue(got.is_monotonic_increasing)

    def test_a_stamp_in_the_hour_march_skips_is_an_error(self):
        with self.assertRaises(Exception):
            c.to_utc(pd.Series([pd.Timestamp("2024-03-10 02:30:00")]), pd.Series(["N"]))


class TheYearlyWorkbooks(unittest.TestCase):
    def row(self, head, stamp, **adders):
        r = [None] * len(head)
        names = [h.strip() for h in head]
        r[names.index("SCED Timestamp")] = stamp
        r[names.index("Repeated Hour Flag")] = "N"
        for k, v in adders.items():
            r[names.index(k)] = v
        return r

    def test_each_form_is_read_by_its_own_header_row(self):
        t = dt.datetime(2014, 6, 1, 0, 0, 27)
        raw = workbook("Historical Real-Time Reserve ORDC Price Adder - June 2014", OLD,
                       {"Jun_1": [self.row(OLD, t, RTORPA=0.0003, RTOFFPA=0.0001), self.row(OLD, dt.datetime(2014, 6, 1, 0, 1, 5, 999000), RTORPA=0.0033, RTOFFPA=0.0005)],
                        "Jul_2": [self.row(OLD, dt.datetime(2014, 7, 1, 0, 0, 12), RTORPA=0, RTOFFPA=0)]})
        x = c.parse_doc(raw, "ordc", 2014, "a test")
        self.assertEqual(list(x.columns), ["local", "flag", "RTORPA", "RTOFFPA"])
        self.assertEqual(len(x), 3)
        self.assertEqual(x["local"][1], pd.Timestamp("2014-06-01 00:01:06"))       # read to the second
        self.assertEqual(list(x["RTORPA"]), [0.0003, 0.0033, 0.0])
        raw = workbook("Historical Real-Time ORDC and Reliability Deployment Price Adders and Reserves - January 2015", MID,
                       {"Jan": [self.row(MID, dt.datetime(2015, 1, 1, 0, 0, 16), RTORPA=0, RTOFFPA=0)],
                        "Dec": [self.row(MID, dt.datetime(2015, 12, 1, 0, 0, 30), RTORPA=0, RTOFFPA=0, RTORDPA=1.5)]}, first_row=9)
        x = c.parse_doc(raw, "ordc_rdpa", 2015, "a test")
        self.assertTrue(pd.isna(x["RTORDPA"][0]))                                   # an empty cell is not a zero
        self.assertEqual(x["RTORDPA"][1], 1.5)
        self.assertEqual(len(c.parse_doc(raw, "ordc_rdpa", 2015, "a test", sheets=[12])), 1)   # the last sheet alone
        raw = workbook("Historical Real-Time Price Adders by SCED Interval - December 2025", NEW,
                       {"Nov": [], "Dec": [self.row(NEW, dt.datetime(2025, 12, 5, 0, 0, 18, 999000), RTRDPA=0)]}, info=True, first_row=9)
        x = c.parse_doc(raw, "rdpa", 2025, "a test")
        self.assertEqual(list(x.columns), ["local", "flag", "RTRDPA"])
        self.assertEqual(len(x), 1)

    def test_a_layout_that_changed_is_an_error(self):
        t = dt.datetime(2020, 1, 1, 0, 0, 20)
        for what, raw in {
            "a column gone": workbook("x", [h for h in MID if h != "RTORDPA"], {"Jan": [[1, t, "N", 20.0, 1, 1, 1, 0, 0, 1, 0]]}),
            "a price that is text": workbook("x", MID, {"Jan": [self.row(MID, t, RTORPA="n/a", RTOFFPA=0, RTORDPA=0)]}),
            "a stamp that is text": workbook("x", MID, {"Jan": [self.row(MID, "01/01/2020 00:00:20", RTORPA=0, RTOFFPA=0, RTORDPA=0)]}),
            "a stamp of another month": workbook("x", MID, {"Jan": [self.row(MID, dt.datetime(2020, 3, 1, 0, 0, 20), RTORPA=0, RTOFFPA=0, RTORDPA=0)]}),
            "a sheet that is no month": workbook("x", MID, {"Summary": [self.row(MID, t, RTORPA=0, RTOFFPA=0, RTORDPA=0)]}),
            "a flag that is neither": workbook("x", MID, {"Jan": [[1, t, "X", 20.0, 1, 1, 1, 0, 0, 1, 0, 0]]}),
        }.items():
            with self.assertRaises(RuntimeError, msg=what):
                c.parse_doc(raw, "ordc_rdpa", 2020, "a test")

    def test_the_whole_year_of_2015_is_read_and_the_half_year_document_is_not(self):
        def doc(name, published, doc_id):
            return {"FriendlyName": name, "PublishDate": published, "DocID": doc_id}
        docs = [doc("HIST_RT_SCED_PRC_ADDR_2026", "2026-10-11T10:13:35-05:00", 9), doc("HIST_RT_SCED_PRC_ADDR_2025", "2026-01-08T18:58:14-05:00", 8),
                doc("RTM_ORDC_REL_DPLY_PRC_ADDR_RSRV_2025", "2025-12-08T12:21:55-05:00", 7), doc("RTM_ORDC_REL_DPLY_PRC_ADDR_RSRV_2015", "2016-01-07T16:44:38-05:00", 3),
                doc("Historical Real-Time Reserve ORDC Price Adder_2015", "2015-06-21T10:05:30-05:00", 2), doc("Historical Real-Time Reserve ORDC Price Adder_2014", "2015-01-02T11:00:28-05:00", 1)]
        got = [(d["DocID"], form, year) for d, form, year in c.chosen(docs)]
        self.assertEqual(got, [(1, "ordc", 2014), (3, "ordc_rdpa", 2015), (7, "ordc_rdpa", 2025), (8, "rdpa", 2025), (9, "rdpa", 2026)])
        # a year posted again: the one published last
        again = docs + [doc("HIST_RT_SCED_PRC_ADDR_2026", "2026-10-18T10:13:35-05:00", 10)]
        self.assertEqual([d["DocID"] for d, _, _ in c.chosen(again)][-1], 10)
        # a name this connector does not know is a new form: an error, not a file passed over
        with self.assertRaises(RuntimeError):
            c.chosen(docs + [doc("RT_PRICE_ADDERS_V3_2027", "2027-01-08T00:00:00-05:00", 11)])

    def test_the_month_a_sheet_stands_for(self):
        self.assertEqual([c.month_of(s) for s in ("Jun_1", "Dec_7", "Jan", "Oct_10", "Report Info", "Sheet1")], [6, 12, 1, 10, None, None])


class TheTableAndItsWiring(unittest.TestCase):
    def test_a_row_per_adder_statistic_and_hour_in_the_series_shape(self):
        rows = pd.DataFrame({"adder": ["RTORPA", "RTRDPA"], "hour": [utc("2021-02-15T18:00:00Z"), utc("2026-01-01T07:00:00Z")], "mean": [169.03341234, 0.0],
                             "max": [490.7322, 0.0], "source_url": ["https://www.ercot.com/misdownload/servlets/mirDownload?doclookupId=1"] * 2,
                             "retrieved_at": ["2026-10-13T06:37:00Z"] * 2, "vintage": ["2022-01-01T15:26:31Z"] * 2})
        s = c.to_series(rows)
        self.assertEqual(list(s.columns), c.ip.SERIES_COLS)
        self.assertEqual(sorted(s["variable"]), ["rtorpa_max", "rtorpa_mean", "rtrdpa_max", "rtrdpa_mean"])
        r = s[s["variable"] == "rtorpa_mean"].iloc[0]
        self.assertEqual((r["entity"], r["ts_utc"], r["value"], r["unit"], r["freq"], r["geo"], r["market"], r["node"], r["source"]),
                         ("ercot:system", "2021-02-15T18:00:00Z", 169.0334, "USD/MWh", "PT1H", "US-TX", "ercot_rtm", "RTORPA", "ercot:NP6-792-ER"))

    def test_the_name_the_ceiling_and_the_report(self):
        self.assertEqual((c.NAME, c.CEILING, c.RTID, c.SOURCE), ("ercot_rtm_price_adders", 1_000_000, 13231, "ercot:NP6-792-ER"))
        self.assertLessEqual(len(c.NAME), 40)
        text = src("warehouse", "connectors", "ercot_rt_adders.py")
        self.assertIn('if ip.paused("ercot"):', text)                       # a paused publisher is never requested
        self.assertIn("ip.fetch_raw(CONNECTOR,", text)                      # a document held is not asked for again
        self.assertNotIn("fillna(0)", text.replace('h["covered"].fillna(0)', ""))   # seconds covered may be zero; a price is never filled

    def test_coverage_the_live_set_and_the_daily_run_know_the_table(self):
        import build_coverage
        self.assertTrue(any(__import__("re").match(p, c.NAME) and s == "power" for p, s in build_coverage.SECTOR_RULES))
        live = yaml.safe_load(src("warehouse", "supabase", "live_set.yaml"))
        self.assertIn("^ercot_rtm_price_adders$", live["full"])
        self.assertIn(c.NAME, live["review_hold"])                          # loaded, and in no count a visitor sees
        self.assertIn(c.SOURCE, live["sources_hold"])
        self.assertNotIn(c.NAME, live["catalogue_hold"])
        sh = src("warehouse", "run_data_machine.sh")
        step = 'soft_step dm_ercot_rt_adders "$PYTHON" warehouse/connectors/ercot_rt_adders.py'
        self.assertIn(step, sh)
        self.assertLess(sh.index("table_stores dm_orpuc_held_stores"), sh.index(step))
        self.assertLess(sh.index(step), sh.index('table_stores dm_rt_adders_stores "$ADDERS" "$was"'))
        self.assertLess(sh.index('table_stores dm_rt_adders_stores "$ADDERS" "$was"'), sh.index("# ---- 2. ERCOT's load zones"))

    def test_the_method_says_what_the_connector_does(self):
        doc = src("docs", "methods", "ercot_rt_price_adders.md")
        for words in ("RTORPA", "RTOFFPA", "RTORDPA", "RTRDPA", "4 December 2025", "for at most 15 minutes", "all 3,600 of its seconds", "169.0334",
                      "Hours a year with an adder above zero", "NP6-792-ER", "What it cannot say"):
            self.assertIn(words, doc, words)
        for f in ("docs/methods/ercot_rt_price_adders.md", "warehouse/connectors/ercot_rt_adders.py", "tests/test_session244.py"):
            self.assertNotIn(EM, src(*f.split("/")), f)


if __name__ == "__main__":
    unittest.main()
