#!/usr/bin/env python3
"""Samoa's electricity a year: production by source (diesel, hydro, coconut oil, solar, wind, biomass), 2008 to 2025, and
the Electric Power Corporation's (EPC) sales by customer category, 2008 to 2024, as the Samoa Bureau of Statistics (SBS)
gives them in tables E1 and E3 of its Electricity Statistics 2025 workbook, from EPC (session 2574, chain 127; the ERW held
Samoa only in EIA's international table).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/sbs_samoa_electricity_annual.csv, from one
workbook (no key):

    https://www.sbs.gov.ws/documents/social/Environment/2025/Electricity-Statistics-2025.xlsx

    entity      sbs:ws:<item>: total, diesel, hydro, coconut, solar, wind, biomass (production, table E1); commercial,
                domestic, government, hotels, industrial, religious, schools, local_consumption, street_lights, total
                (sales, table E3)
    variable    production, sales
    unit        MWh, as SBS gives them; ts_utc the year's first day; freq P1Y; geo WS
    node        SBS's label; x_table the table

Table E1 gives each source's MWh and its share; the shares are not taken. Each year's sources must add to its total within
1 MWh. Table E3's categories are taken as given; its TOTAL row is taken only for a year where it equals the categories' sum
within 1 MWh: for 2008 to 2015 the workbook prints one figure, 153,987.5 MWh, in every year, which its categories do not
add to, so those eight totals are not written (2016 to 2024 are). Table E2 (diesel against renewable) repeats E1 and is not
taken. Nothing is filled.

    python warehouse/connectors/sbs_samoa_electricity.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/sbs_ws/ and reused unless --fresh; --offline asks nothing.
www.sbs.gov.ws's robots file closes /wp-admin/ only (session 2574). License: internal. SBS's site footer reads
"(c) 2022 | Samoa Bureau of Statistics, All Right Reserved."; no terms of reuse found; internal until a ruling.
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "sbs_samoa_electricity_annual"
CONNECTOR = "sbs_ws"
SOURCE = "sbs_ws:electricity_statistics_2025"
HOST = "www.sbs.gov.ws"
URL = "https://www.sbs.gov.ws/documents/social/Environment/2025/Electricity-Statistics-2025.xlsx"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = ("Samoa Bureau of Statistics (SBS), Environment Statistics: Electricity Statistics 2025, tables E1 and E3 (source: Electric "
            "Power Corporation, EPC)")
SOURCES = [("Total", "total"), ("Diesel", "diesel"), ("Hydro", "hydro"), ("Coconut", "coconut"), ("Solar", "solar"), ("Wind", "wind"),
           ("Biomass", "biomass")]
CATEGORIES = [("Commercial", "commercial"), ("Domestic", "domestic"), ("Gov't Dept.", "government"), ("Hotels", "hotels"),
              ("Industrial", "industrial"), ("Religions", "religious"), ("Schools", "schools"), ("Local Consumption", "local_consumption"),
              ("Street lightsa", "street_lights")]
X_COLS = ["x_table"]


def lab(v):
    return " ".join(str(v).split()) if v is not None else ""


def num(v, where):
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        raise RuntimeError(f"{where}: not a number: {v!r}")
    if v < 0:
        raise RuntimeError(f"{where}: negative: {v!r}")
    return float(v)


def row(item, var, year, v, node, table, base):
    return dict(entity=f"sbs:ws:{item}", variable=var, ts_utc=f"{year}-01-01T00:00:00Z", value=round(v, 3), unit="MWh", freq="P1Y",
                geo="WS", market="", node=node, **base, x_table=table)


def parse_e1(rows, base):
    title = [lab(r[1]) for r in rows if lab(r[1]).startswith("Table E1")]
    if len(title) != 1 or "(MWh)" not in title[0]:
        raise RuntimeError(f"E1: the title not found or not in MWh ({title})")
    hd = [i for i, r in enumerate(rows) if lab(r[1]) == "Year"]
    if len(hd) != 1:
        raise RuntimeError("E1: the header row not found")
    h = rows[hd[0]]
    cols = {}
    for name, item in SOURCES:
        at = [j for j, v in enumerate(h) if lab(v) == name]
        if len(at) != 1:
            raise RuntimeError(f"E1: the column {name!r} found {len(at)} times")
        cols[item] = at[0]
    if [lab(rows[hd[0] + 1][j]) for j in cols.values()] != ["MWh"] * len(cols):
        raise RuntimeError("E1: the unit row is not MWh under every source")
    out, years = [], []
    for r in rows[hd[0] + 2:]:
        y = r[1]
        if isinstance(y, str) and y.startswith("Source"):
            break
        if y is None:
            continue
        if not isinstance(y, int) or not 1990 <= y <= 2100:
            raise RuntimeError(f"E1: not a year: {y!r}")
        vals = {item: num(r[j], f"E1 {y} {item}") for item, j in cols.items()}
        parts = sum(v for k, v in vals.items() if k != "total")
        if abs(parts - vals["total"]) > 1:
            raise RuntimeError(f"E1 {y}: the sources make {parts}, the total {vals['total']}")
        years.append(y)
        out += [row(item, "production", y, v, name, "E1", base) for (name, item), v in zip(SOURCES, vals.values())]
    if not years or years != list(range(years[0], years[-1] + 1)):
        raise RuntimeError(f"E1: the years are not a run: {years}")
    return out


def parse_e3(rows, base):
    title = [lab(r[1]) for r in rows if lab(r[1]).startswith("Table E3")]
    if len(title) != 1 or "(MWh)" not in title[0]:
        raise RuntimeError(f"E3: the title not found or not in MWh ({title})")
    hd = [i for i, r in enumerate(rows) if sum(isinstance(v, int) and 1990 <= v <= 2100 for v in r) >= 5]
    if len(hd) != 1:
        raise RuntimeError("E3: the years' row not found")
    yc = [(j, v) for j, v in enumerate(rows[hd[0]]) if isinstance(v, int) and 1990 <= v <= 2100]
    if [v for _, v in yc] != list(range(yc[0][1], yc[-1][1] + 1)):
        raise RuntimeError("E3: the years are not a run")
    found, out, skipped = {}, [], []
    for r in rows[hd[0] + 1:]:
        name = lab(r[1])
        if name.startswith("Source"):
            break
        if not name:
            continue
        found[name] = {y: num(r[j], f"E3 {name} {y}") for j, y in yc}
    for name, _ in CATEGORIES + [("TOTAL", "total")]:
        if name not in found:
            raise RuntimeError(f"E3: the row {name!r} not found")
    extra = set(found) - {c[0] for c in CATEGORIES} - {"TOTAL"}
    if extra:
        raise RuntimeError(f"E3: rows not in the map: {sorted(extra)}")
    for name, item in CATEGORIES:
        out += [row(item, "sales", y, v, name, "E3", base) for y, v in found[name].items()]
    for _, y in yc:
        s = sum(found[c][y] for c, _ in CATEGORIES)
        if abs(s - found["TOTAL"][y]) <= 1:
            out.append(row("total", "sales", y, found["TOTAL"][y], "TOTAL", "E3", base))
        else:
            skipped.append(y)
    return out, skipped


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    for s in ("E1", "E3"):
        if s not in wb.sheetnames:
            raise RuntimeError(f"no sheet {s} ({wb.sheetnames})")
    e1 = [list(r) + [None] * 3 for r in wb["E1"].iter_rows(values_only=True)]
    e3 = [list(r) + [None] * 3 for r in wb["E3"].iter_rows(values_only=True)]
    p = parse_e1(e1, base)
    s, skipped = parse_e3(e3, base)
    return pd.DataFrame(p + s), skipped


def main(argv=None):
    ap = argparse.ArgumentParser(description="SBS: Samoa's electricity production by source and EPC's sales by category")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t, skipped = parse(content, base)
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["variable", "entity", "ts_utc"]).reset_index(drop=True)
        v = {(x, e.split(":")[-1], ts[:4]): val for x, e, ts, val in zip(t.variable, t.entity, t.ts_utc, t.value)}
        lp = max(t[t.variable == "production"].ts_utc.str[:4])
        head = [
            "Energy Research Warehouse (ERW): Samoa Bureau of Statistics (SBS), Electricity Statistics 2025, tables E1 (production by "
            "source) and E3 (EPC's sales by category), from the Electric Power Corporation (session 2574)",
            "Shape: series (docs/datastandard.md v0). entity sbs:ws:<item>; variable production or sales; unit MWh as SBS gives them; "
            "ts_utc the year's first day; freq P1Y; geo WS; node SBS's label; x_table the table.",
            f"{len(t)} rows; in {lp} production {v[('production', 'total', lp)] / 1e3:,.1f} GWh (diesel {v[('production', 'diesel', lp)] / 1e3:,.1f}, "
            f"hydro {v[('production', 'hydro', lp)] / 1e3:,.1f}, solar {v[('production', 'solar', lp)] / 1e3:,.1f}). E3's total is written only "
            f"where it equals its categories (not for {', '.join(str(y) for y in skipped) or 'no year'}, where the workbook prints one figure "
            "every year); the shares and table E2 are not taken; nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/sbs_samoa_electricity.py; the workbook retrieved {base['retrieved_at']}; last "
            f"modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL}",
            f"License: internal. SBS's site states only \"All Right Reserved\"; internal until a ruling. Cite: {CITATION}.",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities; E3 totals not written for {skipped}")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="Samoa Bureau of Statistics (SBS)",
                                    report="Environment Statistics: Electricity Statistics 2025, tables E1 and E3 (source EPC)", report_url=URL,
                                    document_list="https://www.sbs.gov.ws/environment/", license="internal", tables=[NAME])])
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
