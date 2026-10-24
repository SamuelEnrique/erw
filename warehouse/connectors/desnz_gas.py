#!/usr/bin/env python3
"""The UK's natural gas supply and consumption, a quarter from 1998, from the Department for Energy Security and Net
Zero's Energy Trends table 4.1 (session 2226, chain 099: the UK's gas balance beside its electricity, which
desnz_uk_electricity_quarterly holds).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/desnz_uk_gas_quarterly.csv, from the
Energy Trends section 4 page's table ET 4.1 (an xlsx on assets.publishing.service.gov.uk), its sheet "Quarter (GWh)":

    https://www.gov.uk/government/statistics/gas-section-4-energy-trends     (the page, for the file's address)

The sheet is one table: a header row ("Column1", then "<year>" and "<n>th quarter" in one cell a column, the latest
marked "[provisional]"), and one row a line of the balance (indigenous production, imports, of which LNG, exports, stock
change, transfers, total supply, statistical difference, total demand, transformation and its electricity and heat
generation, energy industry use, losses, final consumption and its sectors).

    entity      desnz_et41:uk
    variable    the line in snake case, its notes dropped (indigenous_production, imports, imports_of_liquefied_natural_gas,
                exports, stock_change, transfers, total_supply, statistical_difference, total_demand, transformation,
                electricity_generation, heat_generation, energy_industry_use, losses, final_consumption, iron_steel,
                other_industries, transport, domestic, other_final_users, non_energy_use); node the line as DESNZ writes it.
                Totals are DESNZ's lines, never sums made here
    unit        MWh (DESNZ's GWh times 1,000)
    ts_utc      the quarter's first day at 00:00:00Z; freq P3M; geo GB (the United Kingdom)
    x_provisional  yes for a quarter DESNZ marks "[provisional]" (revised in a later issue)

DESNZ's notes are in the header: imports and exports on a physical flow basis from 2000; LNG imports include cargoes from
2016; stock change positive is a draw; the series breaks they name. The million cubic metre sheets and the annual sheets
are not taken. A blank cell has no row; nothing is filled. A sheet missing, a header not of this form, a value that is
not a number or a line given twice stops the run.

    python warehouse/connectors/desnz_gas.py [--offline] [--out-dir DIR]

Ceiling: two requests (the page and the file it links), reused from warehouse/raw/desnz_gas/ with --offline.
www.gov.uk's robots file disallows print and search pages only, and assets.publishing.service.gov.uk has no robots file
(sessions 2168 and 2226). License: the Open Government Licence v3.0 (GOV.UK's licence for its content, DESNZ's verdict);
attribution: "Contains public sector information licensed under the Open Government Licence v3.0".
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "desnz_uk_gas_quarterly"
CONNECTOR = "desnz_gas"
SOURCE = "desnz:energy_trends_4_1"
PAGE = "https://www.gov.uk/government/statistics/gas-section-4-energy-trends"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
SHEET = "Quarter (GWh)"
LINK = re.compile(r'href="(https://assets\.publishing\.service\.gov\.uk/media/[^"]+/ET_4\.1_[A-Z]{3}_\d\d\.xlsx)"')
QUARTER = re.compile(r"^(\d{4}) ([1-4])(?:st|nd|rd|th) quarter( \[provisional\])?$")
COLS = ["entity", "variable", "ts_utc", "value", "unit", "freq", "geo", "market", "node", "source", "source_url",
        "retrieved_at", "vintage", "x_provisional"]


def text(v):
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def slug(s):
    return re.sub(r"[^a-z0-9]+", "_", re.sub(r"\[.*?\]", "", s.lower())).strip("_")


def link_of(page_html):
    m = sorted(set(LINK.findall(page_html)))
    if len(m) != 1:
        raise RuntimeError(f"the page links {len(m)} ET 4.1 workbooks: {m}")
    return m[0]


def build(workbook, url, retrieved, vintage):
    wb = openpyxl.load_workbook(io.BytesIO(workbook), read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise RuntimeError(f"no sheet {SHEET!r}: {wb.sheetnames}")
    rows = [list(r) for r in wb[SHEET].iter_rows(values_only=True)]
    heads = [i for i, r in enumerate(rows) if r and text(r[0]) == "Column1"]
    if len(heads) != 1:
        raise RuntimeError(f"{len(heads)} header rows ('Column1')")
    h = heads[0]
    cols = []
    for j, c in enumerate(rows[h][1:], start=1):
        c = text(c)
        if not c:
            continue
        m = QUARTER.fullmatch(c)
        if not m:
            raise RuntimeError(f"a column header not '<year> <n>th quarter': {c!r}")
        y, q = int(m.group(1)), int(m.group(2))
        cols.append((j, f"{y}-{3 * q - 2:02d}-01T00:00:00Z", "yes" if m.group(3) else "no"))
    out = []
    for r in rows[h + 1:]:
        line = text(r[0])
        if not line:
            continue
        var = slug(line)
        for j, ts, prov in cols:
            v = r[j] if j < len(r) else None
            if v is None or text(v) == "":
                continue
            if not isinstance(v, (int, float)):
                raise RuntimeError(f"{line}, {ts[:7]}: not a number: {v!r}")
            out.append({"entity": "desnz_et41:uk", "variable": var, "ts_utc": ts, "value": round(float(v) * 1000, 3), "unit": "MWh",
                        "freq": "P3M", "geo": "GB", "market": "", "node": line, "source": SOURCE, "source_url": url,
                        "retrieved_at": retrieved, "vintage": vintage, "x_provisional": prov})
    t = pd.DataFrame(out, columns=COLS)
    if t.empty:
        raise RuntimeError("no line in the sheet")
    twice = t[t.duplicated(["entity", "variable", "ts_utc"], keep=False)]
    if len(twice):
        raise RuntimeError(f"a line given twice: {twice.variable.iloc[0]} {twice.ts_utc.iloc[0]}")
    notes = []
    if "Notes" in wb.sheetnames:
        for r in wb["Notes"].iter_rows(values_only=True):
            if r and re.fullmatch(r"Note \d+", text(r[0])) and len(r) > 1 and r[1]:
                notes.append(f"{text(r[0])}: {text(r[1])}")
    return t.sort_values(["variable", "ts_utc"]).reset_index(drop=True), notes


def main(argv=None):
    ap = argparse.ArgumentParser(description="DESNZ Energy Trends 4.1: the UK's natural gas supply and consumption, a quarter")
    ap.add_argument("--offline", action="store_true", help="the saved files only, no request")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host("www.gov.uk") or ip.paused_host("assets.publishing.service.gov.uk"):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        page, _ = ip.fetch_raw(CONNECTOR, PAGE, log, offline=args.offline, fresh=not args.offline, headers=UA, timeout=60, pause=1)
        url = link_of(page.decode("utf-8", errors="replace"))
        book, rec = ip.fetch_raw(CONNECTOR, url, log, offline=args.offline, headers=UA, timeout=120, pause=1)
        t, notes = build(book, url, rec.get("retrieved_at") or "", ip.vintage_of(rec))
        last = t[t.variable == "total_demand"].iloc[-1]
        prov = sorted(set(t[t.x_provisional == "yes"].ts_utc.str[:7]))
        head = [
            f"Energy Research Warehouse (ERW): UK Department for Energy Security and Net Zero (DESNZ), Energy Trends table 4.1, natural gas supply and "
            f"consumption, a quarter, {t.ts_utc.min()[:7]} to {t.ts_utc.max()[:7]} (session 2226)",
            "Shape: series (docs/datastandard.md v0). entity desnz_et41:uk; variable the balance's line in snake case (indigenous_production, imports, "
            "imports_of_liquefied_natural_gas, exports, stock_change, transfers, total_supply, statistical_difference, total_demand, transformation, "
            "electricity_generation, heat_generation, energy_industry_use, losses, final_consumption and its sectors); unit MWh (DESNZ's GWh times "
            "1,000); ts_utc the quarter's first day; freq P3M; geo GB; node the line as DESNZ writes it; x_provisional.",
            f"{len(t)} rows, {t.variable.nunique()} lines; total demand {last.value / 1e6:,.1f} TWh in the quarter of {last.ts_utc[:7]}; quarters DESNZ "
            f"marks provisional: {', '.join(prov) or 'none'}. Totals are DESNZ's lines, never sums made here. A blank cell has no row; nothing is filled.",
            "DESNZ's notes: " + " | ".join(notes),
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/desnz_gas.py; the workbook retrieved {rec.get('retrieved_at') or ''}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} DESNZ, Energy Trends section 4 (gas), {PAGE}",
            f"  access: {url}",
            "License: public (the Open Government Licence v3.0). Contains public sector information licensed under the Open Government Licence v3.0.",
        ]
        ip.write_csv(t, NAME, head, log, cols=COLS)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.variable.nunique()} lines, provisional {prov}")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="UK Department for Energy Security and Net Zero (DESNZ)",
                                    report="Energy Trends section 4: gas (table ET 4.1, natural gas supply and consumption)", report_url=PAGE,
                                    document_list=PAGE, license="public", tables=[NAME])])
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
