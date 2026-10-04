#!/usr/bin/env python3
"""CAISO wind and solar curtailment at the finest interval CAISO publishes, by reason, from 2019 (session 98).

Energy Research Warehouse (ERW). One table, caiso_curtailment_intervals. An approved pull (Samuel, 4 October 2026):
from 2019, ceiling 500,000 rows, USD 0. Method: docs/methods/caiso_curtailment_intervals.md.

    python warehouse/connectors/caiso_curtailment_intervals.py --from-dir runs/session98/raw   # the workbooks already saved there
    python warehouse/connectors/caiso_curtailment_intervals.py                                # request them (and the daily reports)
    python warehouse/connectors/caiso_curtailment_intervals.py --out-dir DIR                  # a trial: nothing in warehouse/output
    --reports-dir DIR reads the daily reports an earlier run saved (drr_<day>.html) instead of requesting them again.

Sources (both CAISO's, both public; credit the California ISO):
- to 2025-12-31: "Production and curtailments data" (www.caiso.com/library/production-curtailments-data), one workbook
  a year, sheet Curtailments: Date, Hour (hour ending, 1 to 24), Interval (1 to 12), Wind Curtailment and Solar
  Curtailment in MW, and from 2022 a Reason (Local or System; blank on 9,241 rows of 2022, and absent before). Only
  the 5-minute intervals with a curtailment are listed. The two 2025 workbooks overlap; a row is read once.
- from 2026-01-01: the Daily Renewable Report (www.caiso.com/library/daily-renewable-reports), one page a day. By fuel
  it gives the hour: curt_hr_tot_<fuel>_<category>_<local|system>_mwh, 24 values (23 or 25 when the clocks change).
  Its 5-minute curtailment series is wind and solar together, not by fuel, and is not taken: by fuel the hour is the
  finest interval offered, and the 5-minute series would pass the pull's ceiling.

Rows (entity caiso:ISO; ts_utc the interval's start, UTC):
    freq PT5M, MW    curtailed_<fuel>_mw (no reason published), curtailed_<fuel>_local_mw, curtailed_<fuel>_system_mw
    freq PT1H, MWh   curtailed_<fuel>_<econ|ss|oi>_<local|system>_mwh   (economic; self-schedule cut; operator instruction)
    freq P1D,  MWh   curtailed_<fuel>_day_mwh: one row for every day the source covers, zero included, so that a day
                     with no interval row is known to be a day with no curtailment and not a day that is not held.
                     ts_utc is the local (Pacific) day at 00:00:00Z (Decision 11).
Only intervals with a curtailment above zero are written, as in CAISO's workbooks. A 5-minute MW is the interval's
average: its MWh is MW x 5/60.

The clock: Date is the Pacific day; a 5-minute interval starts at (Hour - 1) hours and (Interval - 1) x 5 minutes of
that day's clock. On the autumn day the clock hour 01:00 to 02:00 happens twice and the workbook does not say which;
such a row is dated the first (daylight time), and the log counts them.

The ceiling: the run stops before writing if the table would hold more than 500,000 rows.
"""

import argparse
import datetime as dt
import glob
import io
import os
import re
import sys
import time

import pandas as pd
import requests

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "caiso_curtailment_intervals"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-curtailment/0.1; +https://github.com/SamuelEnrique/erw)"}
CAISO = "https://www.caiso.com"
PAGE = CAISO + "/library/production-curtailments-data"
DRR_PAGE = CAISO + "/library/daily-renewable-reports"
DRR_FROM = "2026-01-01"
FIRST_YEAR = 2019
CEILING = 500_000
TZ = "America/Los_Angeles"
CATS = ["econ_local", "econ_system", "ss_local", "ss_system", "oi_local", "oi_system"]
TERMS = ("CAISO's Privacy and Terms of Use (https://www.caiso.com/privacy-terms-of-use, read 2026-10-04): the materials and information on its "
         "website \"may be used by you provided that you keep intact all copyright, trademark and other proprietary notices and that you credit the "
         "California ISO when using such materials and/or information.\"")
SOURCES = {"caiso:production_curtailments": ("Production and curtailments data (5-minute)", PAGE),
           "caiso:daily_renewable_report": ("Daily Renewable Report (curtailment and output)", DRR_PAGE)}


def get(url, log, what):
    def call():
        r = requests.get(url, headers=UA, timeout=600)
        if r.status_code != 200:
            raise RuntimeError(f"{what}: HTTP {r.status_code} for {url}")
        return r
    return ip.with_retries(what, call, log)


def workbooks(from_dir, raw_dir, log):
    """[(file name, url, bytes, last-modified)] of every workbook the library page links, requested or read from from_dir."""
    out = []
    if from_dir:
        for path in sorted(glob.glob(os.path.join(from_dir, "*.xlsx"))):
            meta = {}
            if os.path.exists(path + ".headers.txt"):
                with open(path + ".headers.txt", encoding="utf-8") as f:
                    meta = dict(line.strip().split(": ", 1) for line in f if ": " in line)
            with open(path, "rb") as f:
                when = ip.utc_iso(pd.Timestamp(meta["retrieved"])) if meta.get("retrieved") else ""
                out.append((os.path.basename(path), meta.get("url", CAISO + "/documents/" + os.path.basename(path)), f.read(), meta.get("last-modified", ""), when))
        log(f"  workbooks read from {from_dir}: {len(out)} (requested earlier in this session; no second request)")
        return out
    page = get(PAGE, log, "CAISO production and curtailments page").text
    for path in sorted(set(re.findall(r'"(/documents/[^"]*curtailments?-?data[^"]*\.xlsx)"', page, re.I))):
        r = get(CAISO + path, log, f"CAISO {path}")
        with open(os.path.join(raw_dir, os.path.basename(path)), "wb") as f:
            f.write(r.content)
        out.append((os.path.basename(path), CAISO + path, r.content, r.headers.get("Last-Modified", ""), ip.utc_iso(pd.Timestamp.now(tz="UTC"))))
        time.sleep(1.0)
    return out


def interval_start(day, hour, interval):
    """The UTC start of a 5-minute interval of a Pacific day, and whether its clock time happens twice that day."""
    naive = pd.Timestamp(day) + pd.Timedelta(hours=int(hour) - 1, minutes=(int(interval) - 1) * 5)
    try:
        return naive.tz_localize(TZ).tz_convert("UTC"), False
    except Exception as e:  # noqa: BLE001
        if "mbiguous" in type(e).__name__ or "mbiguous" in str(e):
            return naive.tz_localize(TZ, ambiguous=True).tz_convert("UTC"), True
        raise


def legacy_rows(books, log):
    """The 5-minute rows from FIRST_YEAR and the days the workbooks cover: [(ts, variable, MW, url, vintage, retrieved)],
    {day: (url, vintage, retrieved)}."""
    import openpyxl
    num = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
    rows, days, seen, dup, twice, skipped = [], {}, set(), 0, 0, 0
    for name, url, content, modified, retrieved in books:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        if "Curtailments" not in wb.sheetnames:
            wb.close()
            continue
        vintage = ip.utc_iso(pd.Timestamp(modified).tz_convert("UTC")) if modified else ""
        it = wb["Curtailments"].iter_rows(values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(it)]
        ci = {h: i for i, h in enumerate(head)}
        n, first, last = 0, None, None
        for row in it:
            when = row[ci["Date"]]
            if not isinstance(when, dt.datetime) or when.year < FIRST_YEAR:
                continue
            day = when.date().isoformat()
            first, last = min(first or day, day), max(last or day, day)
            reason = str(row[ci["Reason"]]).strip().lower() if "Reason" in ci and row[ci["Reason"]] else ""
            w, s = row[ci["Wind Curtailment"]], row[ci["Solar Curtailment"]]
            key = (day, row[ci["Hour"]], row[ci["Interval"]], reason, w if num(w) else None, s if num(s) else None)
            if key in seen:
                dup += 1
                continue
            seen.add(key)
            if reason not in ("", "local", "system") or not num(row[ci["Hour"]]) or not num(row[ci["Interval"]]):
                skipped += 1
                continue
            ts, amb = interval_start(day, row[ci["Hour"]], row[ci["Interval"]])
            twice += int(amb)
            for fuel, v in (("wind", w), ("solar", s)):
                if num(v) and round(v, 6) > 0:
                    rows.append((ts, f"curtailed_{fuel}_{reason}_mw" if reason else f"curtailed_{fuel}_mw", float(v), url, vintage, retrieved))
                    n += 1
        wb.close()
        if first:
            for d in pd.date_range(first, min(last, f"{int(last[:4])}-12-31"), freq="D").strftime("%Y-%m-%d"):
                days.setdefault(d, (url, vintage, retrieved))
            # a workbook is a year (or the first months of one): its days run from its first row's day to its last
        log(f"  {name}: {n:,} values from {FIRST_YEAR}, days {first} to {last}")
    log(f"  workbooks: {len(rows):,} values; {dup:,} rows listed in two workbooks read once; {twice} rows in the hour that happens twice on an autumn day, "
        f"dated the first; {skipped} rows with a reason or an hour that could not be read, not written")
    return rows, days


def js_array(html, name):
    m = re.search(r"\b" + re.escape(name) + r"\s*=\s*(?:JSON\.parse\(\[\")?\[([^\]]*)\]", html)
    if not m:
        return None
    out = []
    for x in m.group(1).split(","):
        x = x.strip().strip('"')
        out.append(None if x in ("", "NA", "null") else float(x))
    return out


def report_days(log, since, until):
    index = get(DRR_PAGE, log, "CAISO daily renewable reports index").text
    out = {}
    for m in sorted(set(re.findall(r'"(/library/daily-renewable-reports-[a-z]{3}-\d{4})"', index))):
        mon, year = m.rsplit("-", 2)[1:]
        first = pd.Timestamp(f"{mon} 1 {year}")
        if first + pd.offsets.MonthEnd(0) < pd.Timestamp(since) or first > pd.Timestamp(until):
            continue
        page = get(CAISO + m, log, f"CAISO {m}").text
        for u in set(re.findall(r'"(/documents/daily-renewable-report-([a-z]{3})-(\d{2})-(\d{4})\.html)"', page)):
            d = pd.Timestamp(f"{u[1]} {u[2]} {u[3]}").date().isoformat()
            if since <= d <= until:
                out[d] = CAISO + u[0]
        time.sleep(0.3)
    return out


def report_rows(html, day):
    """One report's hourly rows by fuel and category, or None when the page is not that day's or an array is missing."""
    m = re.search(r"var chart_date = new Date\(Date\.UTC\((\d{4}), (\d+), (\d+)\)\)", html)
    if not m or dt.date(int(m.group(1)), int(m.group(2)) + 1, int(m.group(3))).isoformat() != day:
        return None
    start = pd.Timestamp(day).tz_localize(TZ)
    hours = int(((pd.Timestamp(day) + pd.Timedelta(days=1)).tz_localize(TZ) - start) / pd.Timedelta(hours=1))
    out, totals = [], {"wind": 0.0, "solar": 0.0}
    for fuel in ("solar", "wind"):
        for cat in CATS:
            arr = js_array(html, f"curt_hr_tot_{fuel}_{cat}_mwh")
            if arr is None or len(arr) != hours or any(v is None for v in arr):
                return None
            for i, v in enumerate(arr):
                if round(v, 6) > 0:
                    out.append(((start + pd.Timedelta(hours=i)).tz_convert("UTC"), f"curtailed_{fuel}_{cat}_mwh", float(v)))
            totals[fuel] += sum(arr)
    return out, totals


def main(argv=None):
    ap = argparse.ArgumentParser(description="CAISO wind and solar curtailment by interval and reason, from 2019")
    ap.add_argument("--from-dir", help="read the workbooks saved in this directory instead of requesting them")
    ap.add_argument("--out-dir", help="a trial run: the table and its log under this directory; nothing in warehouse/output")
    ap.add_argument("--no-reports", action="store_true", help="the workbooks only (no request for the daily reports of 2026)")
    ap.add_argument("--reports-dir", help="read the daily reports saved in this directory by an earlier run (drr_<day>.html) instead of requesting them")
    a = ap.parse_args(argv)
    if ip.paused("caiso"):
        print(ip.pause_line("caiso"))
        return 0
    if a.out_dir:
        os.makedirs(a.out_dir, exist_ok=True)
    run_id = pd.Timestamp.now(tz="UTC").strftime("%Y%m%dT%H%M%SZ")
    retrieved = ip.utc_iso(pd.Timestamp.now(tz="UTC"))
    log_dir = os.path.join(a.out_dir, "logs") if a.out_dir else ip.LOG_DIR
    os.makedirs(log_dir, exist_ok=True)
    log = ip.Log(os.path.join(log_dir, f"caiso_curtailment_intervals_{run_id}.log"))
    raw_dir = os.path.join(ip.RAW_DIR if hasattr(ip, "RAW_DIR") else os.path.join(os.path.dirname(ip.OUT_DIR), "raw"), NAME, run_id)
    os.makedirs(raw_dir, exist_ok=True)
    books = workbooks(a.from_dir, raw_dir, log)
    five, days = legacy_rows(books, log)
    rows = []

    def add(ts, variable, value, unit, freq, source, url, vintage, when):
        rows.append(dict(entity="caiso:ISO", variable=variable, ts_utc=ip.utc_iso(ts) if not isinstance(ts, str) else ts, value=value, unit=unit, freq=freq, geo="US-CA",
                         market="", node="", source=source, source_url=url, retrieved_at=when or retrieved, vintage=vintage))
    day_total = {}
    for ts, variable, mw, url, vintage, when in five:
        add(ts, variable, round(mw, 6), "MW", "PT5M", "caiso:production_curtailments", url, vintage, when)
        d = ts.tz_convert(TZ).strftime("%Y-%m-%d")
        k = (d, variable.split("_")[1])
        day_total[k] = day_total.get(k, 0.0) + mw * 5 / 60
    for d, (url, vintage, when) in sorted(days.items()):
        if d >= DRR_FROM:
            continue
        for fuel in ("wind", "solar"):
            add(f"{d}T00:00:00Z", f"curtailed_{fuel}_day_mwh", round(day_total.get((d, fuel), 0.0), 3), "MWh", "P1D", "caiso:production_curtailments", url, vintage, when)
    n_reports = n_bad = 0
    if not a.no_reports:
        until = (pd.Timestamp.now(tz=TZ) - pd.Timedelta(days=1)).strftime("%Y-%m-%d")
        if a.reports_dir:
            saved = {os.path.basename(f)[4:14]: f for f in sorted(glob.glob(os.path.join(a.reports_dir, "drr_*.html")))}
            urls = {d: CAISO + "/documents/daily-renewable-report-" + pd.Timestamp(d).strftime("%b-%d-%Y").lower() + ".html" for d in saved}
            log(f"  daily renewable reports read from {a.reports_dir}: {len(urls)} (requested by an earlier run of this session; no second request)")
        else:
            saved = {}
            urls = report_days(log, DRR_FROM, until)
            log(f"  daily renewable reports {DRR_FROM} to {until}: {len(urls)} listed")
        for d, url in sorted(urls.items()):
            if d in saved:
                with open(saved[d], encoding="utf-8") as f:
                    html = f.read()
                retrieved_d = ip.utc_iso(pd.Timestamp(os.path.getmtime(saved[d]), unit="s", tz="UTC"))
            else:
                html = get(url, log, f"CAISO report {d}").text
                retrieved_d = retrieved
                with open(os.path.join(raw_dir, f"drr_{d}.html"), "w", encoding="utf-8") as f:
                    f.write(html)
            got = report_rows(html, d)
            if got is None:
                n_bad += 1
                log(f"  report {d}: not that day's page, or an hourly array missing or short; day not written")
                continue
            hourly, totals = got
            for ts, variable, v in hourly:
                add(ts, variable, round(v, 6), "MWh", "PT1H", "caiso:daily_renewable_report", url, "", retrieved_d)
            for fuel in ("wind", "solar"):
                add(f"{d}T00:00:00Z", f"curtailed_{fuel}_day_mwh", round(totals[fuel], 3), "MWh", "P1D", "caiso:daily_renewable_report", url, "", retrieved_d)
            n_reports += 1
            if d not in saved:
                time.sleep(0.3)
    out = pd.DataFrame(rows)[ip.SERIES_COLS]
    by = out.groupby("freq").size().to_dict()
    log(f"  rows: {len(out):,} ({by}); daily reports read {n_reports}, not written {n_bad}; ceiling {CEILING:,}")
    if len(out) > CEILING:
        log(f"  STOPPED: {len(out):,} rows would pass the pull's ceiling of {CEILING:,}; nothing written")
        log.close()
        print(f"{NAME}: STOPPED, {len(out):,} rows would pass the ceiling of {CEILING:,}; nothing written")
        return 1
    if a.out_dir:
        ip.set_out_dir(a.out_dir)
    first, last = out[out["freq"] == "P1D"]["ts_utc"].min()[:10], out[out["freq"] == "P1D"]["ts_utc"].max()[:10]
    ip.write_csv(out, NAME, [
        "Energy Research Warehouse (ERW): CAISO wind and solar curtailment at the finest interval CAISO publishes, by reason, from 2019 (session 98)",
        "Shape: series (docs/datastandard.md v0). Entity caiso:ISO. freq PT5M (MW, the interval's average; to 2025-12-31): curtailed_<fuel>_mw (no reason "
        "published), curtailed_<fuel>_local_mw, curtailed_<fuel>_system_mw. freq PT1H (MWh; from 2026-01-01): curtailed_<fuel>_<econ|ss|oi>_<local|system>_mwh "
        "(economic, self-schedule cut, operator instruction). freq P1D (MWh): curtailed_<fuel>_day_mwh, one row for every day covered, zero included; ts_utc is "
        "the Pacific day at 00:00:00Z. ts_utc of an interval is its start, UTC. Only intervals with a curtailment are written. Definitions: "
        "docs/methods/caiso_curtailment_intervals.md.",
        f"Window: {first} to {last}. An approved pull (session 98): from {FIRST_YEAR}, ceiling {CEILING:,} rows; this file holds {len(out):,}.",
        f"Retrieved: {run_id} (UTC) by warehouse/connectors/caiso_curtailment_intervals.py"
        + (f"; the workbooks were requested earlier in the session and read from {a.from_dir}" if a.from_dir else "")
        + (f"; the daily reports were requested by an earlier run and read from {a.reports_dir.replace(os.sep, '/')}" if a.reports_dir else ""),
        f"Run log: warehouse/output/logs/caiso_curtailment_intervals_{run_id}.log", f"Raw files: {os.path.relpath(raw_dir, os.path.dirname(os.path.dirname(ip.OUT_DIR)))} (not in git)",
        f"Source: caiso:production_curtailments California ISO, Production and curtailments data (5-minute), {PAGE}; {len(books)} workbooks",
        f"Source: caiso:daily_renewable_report California ISO, Daily Renewable Report, {DRR_PAGE}; {n_reports} daily reports ({n_bad} not written)",
        f"License: public, with credit to the California ISO. {TERMS}",
    ], log, key=["entity", "variable", "ts_utc"])
    ip.update_sources([dict(source=s, publisher="California ISO (CAISO)", report=rep, report_url=url, document_list=url, license="public", tables=[NAME])
                       for s, (rep, url) in SOURCES.items()])
    log.close()
    print(f"{NAME}: {len(out):,} rows ({by}), {first} to {last}; daily reports read {n_reports}, not written {n_bad}" + (f" (trial, under {a.out_dir})" if a.out_dir else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
