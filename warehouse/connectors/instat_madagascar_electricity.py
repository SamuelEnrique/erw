#!/usr/bin/env python3
"""Madagascar's public grid a month: JIRAMA's electricity production (total, hydro, thermal) and consumption (total,
households, industry and services, street lighting), as the Institut National de la Statistique (INSTAT) gives them in
table 21, "Evolution mensuelle des production et consommation d'electricite" (source: JIRAMA), of the data workbook of
its Tableau de Bord de l'Economie (TBE) no. 63, August 2026 (session 2532, chain 124; the ERW held Madagascar only in
EIA's and OWID's international tables).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/instat_madagascar_electricity_monthly.csv,
from one workbook (no key):

    https://www.instat.mg/documents/upload/main/Donn%C3%A9es%20TBE63_25_08_2026.xlsx

The sheet "SortieElec" holds 24 months, July 2024 to June 2026, one row a month (its date the 15th), then two rows of
half-year totals and their change, which are not taken. Its columns are found by their headers: under "Production
(MWh)" the total, "Hydraulique" and "Thermique"; under "Consommation d'electricite (MWh)" the total, "Menages",
"Industries et services" and "Eclairage public". The average price (Ar/kWh) is not taken (a price in another currency).

    entity      jirama:mg:<item>: total, hydro, thermal (production); total, households, industry_services,
                public_lighting (consumption)
    variable    production, consumption
    unit        MWh (as INSTAT gives them); ts_utc the month's first day; freq P1M; geo MG
    node        INSTAT's column label; x_period the month (YYYY-MM)

The run stops on a sheet without the table, a header not found, a month given twice, a gap between months, a figure that
is not a number, or a total that differs from the sum of its parts by more than 1 MWh. Nothing is filled.

    python warehouse/connectors/instat_madagascar_electricity.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/instat_mg_tbe/ and reused unless --fresh; --offline asks nothing.
www.instat.mg's robots file allows every path with a crawl delay of 3 seconds, which is kept (session 2532). License:
internal. INSTAT's site states "Copyright INSTAT 2026. Tous les droits sont reserves" and links no terms of reuse; the
table is internal until a ruling.
"""

import argparse
import datetime as dt
import io
import os
import sys
import traceback
import unicodedata

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "instat_madagascar_electricity_monthly"
CONNECTOR = "instat_mg_tbe"
SOURCE = "instat_mg:tbe63_table21"
HOST = "www.instat.mg"
URL = "https://www.instat.mg/documents/upload/main/Donn%C3%A9es%20TBE63_25_08_2026.xlsx"
PAGE = "https://www.instat.mg/p/tableau-de-bord-de-leconomie-tbe-n0-63-aout-2026"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
SHEET = "SortieElec"
CITATION = ("Institut National de la Statistique (INSTAT), Madagascar, Tableau de Bord de l'Economie no. 63 (August 2026), data "
            "workbook, table 21: Evolution mensuelle des production et consommation d'electricite (source: JIRAMA)")
# (variable, item, the header under which the column sits, the column's own label)
COLUMNS = [("production", "total", "production (mwh)", "total"), ("production", "hydro", "production (mwh)", "hydraulique"),
           ("production", "thermal", "production (mwh)", "thermique"),
           ("consumption", "total", "consommation d'electricite (mwh)", "total"),
           ("consumption", "households", "consommation d'electricite (mwh)", "menages"),
           ("consumption", "industry_services", "consommation d'electricite (mwh)", "industries et services"),
           ("consumption", "public_lighting", "consommation d'electricite (mwh)", "eclairage public")]
PARTS = {"production": ["hydro", "thermal"], "consumption": ["households", "industry_services", "public_lighting"]}
X_COLS = ["x_period"]


def key(s):
    """A header as the map holds it: accents dropped, spaces collapsed, lower case."""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return " ".join(s.split()).lower()


def parse_rows(rows, base):
    """rows: the sheet's rows as tuples of values."""
    title = next((i for i, r in enumerate(rows) if any("production et consommation d" in key(c) for c in r if c is not None)), None)
    if title is None:
        raise RuntimeError("table 21 not found in the sheet")
    head, sub = rows[title + 1], rows[title + 2]
    groups, cur = {}, None
    for j in range(max(len(head), len(sub))):
        h = head[j] if j < len(head) else None
        if h is not None and key(h) in ("production (mwh)", "consommation d'electricite (mwh)", "prix moyen (ar/kwh)"):
            cur = key(h)
        s = sub[j] if j < len(sub) else None
        if s is not None and cur is not None:
            if (cur, key(s)) in groups:
                raise RuntimeError(f"a header given twice: {cur} / {s}")
            groups[(cur, key(s))] = j
    cols = {}
    for var, item, grp, lab in COLUMNS:
        if (grp, lab) not in groups:
            raise RuntimeError(f"header not found: {grp} / {lab}")
        cols[(var, item)] = groups[(grp, lab)]
    out, seen = [], []
    for r in rows[title + 3:]:
        d = r[1] if len(r) > 1 else None
        if not isinstance(d, dt.datetime):
            if seen and d is not None:
                break
            continue
        month = d.strftime("%Y-%m")
        if month in seen:
            raise RuntimeError(f"month {month} given twice")
        seen.append(month)
        vals = {}
        for (var, item), j in cols.items():
            v = r[j]
            if isinstance(v, bool) or not isinstance(v, (int, float)):
                raise RuntimeError(f"{month} {var} {item}: not a figure: {v!r}")
            vals[(var, item)] = float(v)
        for var, parts in PARTS.items():
            s = sum(vals[(var, p)] for p in parts)
            if abs(s - vals[(var, "total")]) > 1:
                raise RuntimeError(f"{month} {var}: the total {vals[(var, 'total')]} is not the sum of its parts {s}")
        for (var, item), v in vals.items():
            label = next(lab for vv, it, grp, lab in COLUMNS if (vv, it) == (var, item))
            out.append(dict(entity=f"jirama:mg:{item}", variable=var, ts_utc=f"{month}-01T00:00:00Z", value=round(v, 3), unit="MWh",
                            freq="P1M", geo="MG", market="", node=f"{var}: {label}", **base, x_period=month))
    if not seen:
        raise RuntimeError("no month in table 21")
    ms = [pd.Period(m, "M") for m in seen]
    for a, b in zip(ms, ms[1:]):
        if b != a + 1:
            raise RuntimeError(f"a gap or a month out of order: {a} then {b}")
    return pd.DataFrame(out)


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    if SHEET not in wb.sheetnames:
        raise RuntimeError(f"no sheet {SHEET!r} in the workbook")
    rows = [tuple(r) for r in wb[SHEET].iter_rows(values_only=True)]
    wb.close()
    return parse_rows(rows, base)


def main(argv=None):
    ap = argparse.ArgumentParser(description="INSTAT Madagascar: JIRAMA's electricity production and consumption a month")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=3)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t = parse(content, base)
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["variable", "entity", "ts_utc"]).reset_index(drop=True)
        first, last = t.x_period.min(), t.x_period.max()
        y2025 = t[(t.x_period.str[:4] == "2025") & (t.entity == "jirama:mg:total")].groupby("variable").value.sum()
        head = [
            "Energy Research Warehouse (ERW): Institut National de la Statistique (INSTAT), Madagascar, Tableau de Bord de l'Economie "
            "no. 63 (August 2026), table 21: JIRAMA's electricity production and consumption a month (session 2532)",
            "Shape: series (docs/datastandard.md v0). entity jirama:mg:<item> (production: total, hydro, thermal; consumption: total, "
            "households, industry_services, public_lighting); variable production, consumption; unit MWh (as INSTAT gives them); "
            "ts_utc the month's first day; freq P1M; geo MG; node INSTAT's column label; x_period the month.",
            f"{len(t)} rows, {first} to {last}; in 2025 production {y2025.get('production', 0) / 1e6:.3f} TWh and consumption "
            f"{y2025.get('consumption', 0) / 1e6:.3f} TWh. The average price (Ar/kWh) and the half-year totals are not taken; nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/instat_madagascar_electricity.py; the workbook retrieved {base['retrieved_at']}; "
            f"last modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL} (listed on {PAGE})",
            "License: internal. INSTAT's site states \"Copyright INSTAT 2026. Tous les droits sont reserves\" and links no terms of reuse; "
            f"internal until a ruling. Cite: {CITATION}.",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities, {first} to {last}")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="Institut National de la Statistique (INSTAT), Madagascar",
                                    report="Tableau de Bord de l'Economie no. 63 (August 2026), table 21: monthly electricity production "
                                           "and consumption (source: JIRAMA)", report_url=URL,
                                    document_list="https://www.instat.mg/statistiques/energie-et-hydrocarbures", license="internal",
                                    tables=[NAME])])
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
