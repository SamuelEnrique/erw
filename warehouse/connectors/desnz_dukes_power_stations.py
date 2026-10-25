#!/usr/bin/env python3
"""The major power producers' power stations in the United Kingdom operational at the end of May 2026, one row a station, as the
Department for Energy Security and Net Zero (DESNZ) lists them in the Digest of UK Energy Statistics (DUKES), table 5.11 (session
2714, chain 136, the session's own choice; the ERW held the UK's capacity only as totals, by region in
desnz_regional_renewable_capacity_annual and nationally in Energy Trends, and no UK station).

Energy Research Warehouse (ERW) connector. One entities table, warehouse/output/desnz_uk_power_stations.csv, from one workbook
(no key), its sheet "5.11 Full list":

    https://assets.publishing.service.gov.uk/media/6a6a365b0ddb7e4831c629f3/DUKES_5.11.xlsx

    entity_id      dukes:<DESNZ site code> (SOL-000443, WON-001539, ...: unique on the sheet)
    entity_type    plant; name the site name; operator the company name, as DESNZ writes them
    geo            GB-ENG, GB-SCT, GB-WLS or GB-NIR from the sheet's country
    lat, lon       WGS84, converted here from the sheet's X and Y coordinates, which are British National Grid eastings and
                   northings (EPSG:27700) for every station, Northern Ireland's included (checked on Belfast, Ballylumford, Kilroot
                   and Coolkeeragh); empty where a coordinate is not a number
    capacity_mw    the sheet's installed capacity (MW), as stated
    status         operating (the list's stations are those operational at the end of May 2026); status_date that day, read from
                   the sheet's title
    other columns  technology, type, chp, primary_fuel, secondary_fuel, grid_connection_type, country, region, postcode,
                   os_reference, x_coordinate, y_coordinate, year_commissioned, desnz_site_code: the sheet's own, as text

DESNZ's note 1: the list covers the major power producers' stations only; other generators (many renewable sites and
autogenerators) are in the summary table 5.11.A but not listed. Checks: the header as named; every site code unique; every
capacity a positive number; the list's capacity by fuel equal to table 5.11.A's major power producer rows (wind split by the
sheet's type, onshore and offshore; MSW as waste; pumped hydro as pumped storage) and in all to its total, within 0.01 MW. A
check missed, an unknown country or fuel stops the run. Nothing is filled. The earlier editions' sheets (DUKES 2004 to 2025) are
in other layouts and are not read.

    python warehouse/connectors/desnz_dukes_power_stations.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/desnz_dukes/ and reused unless --fresh; --offline asks nothing.
www.gov.uk's robots file closes only print and search paths; assets.publishing.service.gov.uk has no robots file (session 2678).
License: the Open Government Licence v3.0 (DESNZ's verdict on record, session 2167): public.
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "desnz_uk_power_stations"
CONNECTOR = "desnz_dukes"
SOURCE = "desnz:dukes_5_11_power_stations"
HOST = "assets.publishing.service.gov.uk"
URL = "https://assets.publishing.service.gov.uk/media/6a6a365b0ddb7e4831c629f3/DUKES_5.11.xlsx"
PAGE = "https://www.gov.uk/government/statistics/electricity-chapter-5-digest-of-united-kingdom-energy-statistics-dukes"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = "Department for Energy Security and Net Zero, Digest of UK Energy Statistics (DUKES) 2026, table 5.11, Power stations in the United Kingdom"
HEADER = ["Company Name [note 30]", "Site Name", "Technology", "Type", "CHP", "Primary Fuel", "Secondary Fuel", "InstalledCapacity (MW)",
          "Grid Connection Type", "Country", "Region", "Postcode", "OS Reference", "X-Coordinate", "Y-Coordinate", "Year Commissioned",
          "DESNZ site code"]
EXTRA = ["technology", "type", "chp", "primary_fuel", "secondary_fuel", "grid_connection_type", "country", "region", "postcode",
         "os_reference", "x_coordinate", "y_coordinate", "year_commissioned", "desnz_site_code"]
BASE = ["entity_id", "entity_type", "name", "geo", "lat", "lon", "capacity_mw", "status", "status_date", "operator",
        "source", "source_url", "retrieved_at", "vintage"]
GEO = {"England": "GB-ENG", "Scotland": "GB-SCT", "Wales": "GB-WLS", "Northern Ireland": "GB-NIR"}
SUMMARY_FUEL = {"Biomass": "Biomass", "Diesel/Gas Oil": "Diesel/Gas Oil", "Hydro": "Hydro", "MSW": "Waste", "Natural Gas": "Natural Gas",
                "Nuclear": "Nuclear", "Pumped hydro": "Pumped Storage", "Solar": "Solar", "Coal": "Coal", "Sour Gas": "Sour Gas"}
MONTHS = {m: i for i, m in enumerate(["January", "February", "March", "April", "May", "June", "July", "August", "September",
                                      "October", "November", "December"], 1)}


def text(v):
    if v is None:
        return ""
    s = re.sub(r"\s+", " ", str(v)).strip()
    return "" if s.lower() == "nan" else s


def as_at(title):
    """The list's day: the end of the month its title names ("... operational at the end of May 2026 ...")."""
    m = re.search(r"operational at the end of (\w+) (\d{4})", title)
    if not m or m.group(1) not in MONTHS:
        raise RuntimeError(f"the title names no month: {title!r}")
    y, mo = int(m.group(2)), MONTHS[m.group(1)]
    return (dt.date(y + (mo == 12), mo % 12 + 1, 1) - dt.timedelta(days=1)).isoformat()


def to_wgs84(xs, ys):
    """British National Grid eastings and northings to WGS84 (lat, lon), each rounded to 6 places; pyproj imported here only."""
    from pyproj import Transformer
    t = Transformer.from_crs("EPSG:27700", "EPSG:4326", always_xy=True)
    lon, lat = t.transform(xs, ys)
    return [round(a, 6) for a in lat], [round(o, 6) for o in lon]


def summary_mpp(rows):
    """Table 5.11.A's major power producer rows: {fuel: MW}, with 'Total'."""
    out = {}
    for r in rows:
        if r and text(r[0]) == "Major power producer" and len(r) > 2 and isinstance(r[2], (int, float)):
            out[text(r[1])] = float(r[2])
    if "Total" not in out:
        raise RuntimeError("table 5.11.A has no major power producer total")
    return out


def build(rows, summary, base, convert=True):
    title = text(rows[0][0]) if rows and rows[0] else ""
    day = as_at(title)
    h = next((i for i, r in enumerate(rows) if r and text(r[0]).startswith("Company Name")), None)
    if h is None or [text(x) for x in rows[h][:len(HEADER)]] != HEADER:
        raise RuntimeError(f"the header is not {HEADER}")
    out, xs, ys = [], [], []
    for r in rows[h + 1:]:
        if not any(text(x) for x in r):
            continue
        v = dict(zip(HEADER, r))
        code, country, mw = text(v["DESNZ site code"]), text(v["Country"]), v["InstalledCapacity (MW)"]
        if not code:
            raise RuntimeError(f"a station with no site code: {text(v['Site Name'])!r}")
        if country not in GEO:
            raise RuntimeError(f"{code}: a country not in the map: {country!r}")
        if not isinstance(mw, (int, float)) or isinstance(mw, bool) or mw <= 0:
            raise RuntimeError(f"{code}: not a positive capacity: {mw!r}")
        row = {"entity_id": f"dukes:{code}", "entity_type": "plant", "name": text(v["Site Name"]), "geo": GEO[country], "lat": "",
               "lon": "", "capacity_mw": float(mw), "status": "operating", "status_date": day,
               "operator": text(v["Company Name [note 30]"]), **base}
        for c, k in zip(HEADER[2:], ["technology", "type", "chp", "primary_fuel", "secondary_fuel", "_mw", "grid_connection_type",
                                     "country", "region", "postcode", "os_reference", "x_coordinate", "y_coordinate",
                                     "year_commissioned", "desnz_site_code"]):
            if k != "_mw":
                row[k] = text(v[c])
        out.append(row)
        try:
            xs.append(float(row["x_coordinate"]))
            ys.append(float(row["y_coordinate"]))
        except ValueError:
            xs.append(None)
            ys.append(None)
    t = pd.DataFrame(out, columns=BASE + EXTRA)
    twice = t.entity_id[t.entity_id.duplicated()]
    if len(twice):
        raise RuntimeError(f"a site code given twice: {list(twice)[:3]}")
    got = {}
    for fuel, typ, mw in zip(t.primary_fuel, t.type, t.capacity_mw):
        key = ("Onshore Wind" if typ == "Onshore" else "Offshore Wind" if typ == "Offshore" else None) if fuel == "Wind" else SUMMARY_FUEL.get(fuel)
        if key is None:
            raise RuntimeError(f"a fuel not in the map: {fuel!r} ({typ!r})")
        got[key] = got.get(key, 0.0) + mw
    for k, want in summary.items():
        have = sum(got.values()) if k == "Total" else got.get(k, 0.0)
        if abs(have - want) > 0.01:
            raise RuntimeError(f"the list's {k} is {have:,.3f} MW, table 5.11.A's {want:,.3f}")
    extra = set(got) - set(summary)
    if extra:
        raise RuntimeError(f"fuels in the list and not in table 5.11.A: {sorted(extra)}")
    if convert:
        ok = [i for i, (x, y) in enumerate(zip(xs, ys)) if x is not None and y is not None]
        lat, lon = to_wgs84([xs[i] for i in ok], [ys[i] for i in ok])
        t.loc[ok, "lat"], t.loc[ok, "lon"] = lat, lon
    return t.sort_values("entity_id").reset_index(drop=True), day


def main(argv=None):
    ap = argparse.ArgumentParser(description="DESNZ DUKES 5.11: the UK's major power producers' power stations")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        for s in ("5.11 Full list", "5.11 Summary"):
            if s not in wb.sheetnames:
                raise RuntimeError(f"the workbook has no sheet {s!r}: {wb.sheetnames}")
        summary = summary_mpp([tuple(r) for r in wb["5.11 Summary"].iter_rows(values_only=True)])
        t, day = build([tuple(r) for r in wb["5.11 Full list"].iter_rows(values_only=True)], summary, base)
        by = t.groupby("geo").capacity_mw.agg(["size", "sum"])
        fuel = t.groupby("primary_fuel").capacity_mw.sum().sort_values(ascending=False)
        head = [
            "Energy Research Warehouse (ERW): DESNZ, Digest of UK Energy Statistics (DUKES) 2026, table 5.11: the major power "
            f"producers' power stations in the United Kingdom operational at {day}, one row a station (session 2714)",
            "Shape: entities (docs/datastandard.md v0). entity_id dukes:<DESNZ site code>; entity_type plant; name the site; operator "
            "the company; geo GB-ENG, GB-SCT, GB-WLS or GB-NIR; lat and lon WGS84, converted from the sheet's British National Grid "
            "X and Y; capacity_mw installed capacity as stated; status operating, status_date the list's day; the sheet's other "
            "columns kept as text.",
            f"{len(t)} stations, {t.capacity_mw.sum():,.1f} MW, equal to table 5.11.A's major power producers' total and to each of "
            "its fuels; " + "; ".join(f"{g}: {int(r['size'])} stations, {r['sum']:,.0f} MW" for g, r in by.iterrows())
            + "; by primary fuel, MW: " + ", ".join(f"{k} {v:,.0f}" for k, v in fuel.items())
            + ". Other generators (many renewable sites and autogenerators) are not listed (DESNZ's note 1). Nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/desnz_dukes_power_stations.py; the workbook retrieved "
            f"{base['retrieved_at']}; last modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL} (listed on {PAGE})",
            "License: Open Government Licence v3.0. Credit: Department for Energy Security and Net Zero.",
        ]
        ip.write_snapshot(t, NAME, head, log, list(t.columns))
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} stations, {t.capacity_mw.sum():,.1f} MW")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="UK Department for Energy Security and Net Zero (DESNZ)",
                                    report="Digest of UK Energy Statistics (DUKES) 2026, table 5.11, Power stations in the United Kingdom",
                                    report_url=PAGE, document_list=PAGE, license="public", tables=[NAME])])  # OGL v3.0
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
