#!/usr/bin/env python3
"""Electricity generation by fuel and the supply balance a year for each UK nation (Scotland, Wales, Northern Ireland, England)
and the UK, 2004 to 2025, as the Department for Energy Security and Net Zero (DESNZ) publishes them in the data tables of its
Energy Trends special feature "Electricity generation and supply in Scotland, Wales, Northern Ireland and England" (September
2026 edition) (session 2678, chain 134, the session's own choice; the ERW held DESNZ's UK-wide Energy Trends tables 5.1 and 6.1
and not the nations').

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/desnz_uk_nations_electricity_annual.csv, from one
workbook (no key):

    https://assets.publishing.service.gov.uk/media/6aba6c4afe72ed1e2b02f14b/Regional_Electricity_Generation_and_Supply_Timeseries_2026.xlsx

    entity      desnz:<nation>:<item>: nation uk, scotland, wales, northern_ireland, england;
                Table 2a, generation by fuel: mpp_<fuel> and other_<fuel> (the major power producers' and the other
                generators' rows, their totals mpp_total and other_total) (variable generation);
                Table 1, the supply balance: generated_mpp, generated_other, generated_total (generation); own_use_mpp,
                own_use_other (own_use); net_supplied_mpp, net_supplied_other (net_supply); transfer_to_england,
                transfer_to_northern_ireland, transfer_to_europe (net_transfers, net of receipts, as published);
                transfers_other_to_public (transfers); transmission_losses, distribution_losses (losses); consumption_public,
                consumption_autogenerators, consumption_total (consumption); sales (sales); statistical_difference
    unit        MWh (the tables' GWh times 1,000); ts_utc the year's first day; freq P1Y; geo GB (the UK), GB-SCT, GB-WLS,
                GB-NIR or GB-ENG
    node        the row's label as DESNZ gives it (its note marks dropped), and for Table 2a its generator type

Each table is read by its title and its header of "<year> <nation>" columns (a digit after the nation is dropped: Table 1
heads one column "2022 England6"); each row by its label, and a label not in the map stops the run. "[x]" (not available)
and a blank cell have no row; nothing is filled. Not taken: Table 2a's "All generating companies" rows, which DESNZ derives and
which do not add (in the September 2026 file fossil fuels miss coal, oil and gas by up to 1,250 GWh, and the total its parts by
up to 7,875 GWh); Table 2b (shares, per cent). DESNZ's note 9: the nations' totals do not sum exactly to the UK's. Table 1 is
kept as published although its total generated misses its two parts in seven nation-years (Wales +1,235 GWh and England -1,250
in 2020, Northern Ireland +15 in 2020, England and the UK -96 in 2023 and -1,071 in 2024). Check: a Table 2a block's fuels add
to its total within 10 GWh (the largest gap is 7, major power producers in England in 2020).

    python warehouse/connectors/desnz_uk_nations_electricity.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/desnz_regional/ and reused unless --fresh; --offline asks nothing.
www.gov.uk's robots file closes only print and search paths; assets.publishing.service.gov.uk has no robots file (session 2678).
License: the Open Government Licence v3.0 (the publication page: "All content is available under the Open Government Licence
v3.0, except where otherwise stated"): public.
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "desnz_uk_nations_electricity_annual"
CONNECTOR = "desnz_regional"
SOURCE = "desnz:electricity_generation_supply_nations_2026"
HOST = "assets.publishing.service.gov.uk"
URL = ("https://assets.publishing.service.gov.uk/media/6aba6c4afe72ed1e2b02f14b/"
       "Regional_Electricity_Generation_and_Supply_Timeseries_2026.xlsx")
PAGE = ("https://www.gov.uk/government/statistics/energy-trends-september-2026-special-feature-article-electricity-generation-"
        "and-supply-in-scotland-wales-northern-ireland-and-england-2021-to-2025")
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = ("Department for Energy Security and Net Zero, Electricity generation and supply in Scotland, Wales, Northern Ireland and "
            "England, 2004 to 2025: data tables (Energy Trends, September 2026)")
NATIONS = {"UK Total": ("uk", "GB"), "Scotland": ("scotland", "GB-SCT"), "Wales": ("wales", "GB-WLS"),
           "Northern Ireland": ("northern_ireland", "GB-NIR"), "England": ("england", "GB-ENG")}
FUELS = {"Coal": "coal", "Oil": "oil", "Gas": "gas", "Nuclear": "nuclear", "Hydro natural flow": "hydro", "Wind": "wind",
         "Solar": "solar", "Shoreline wave and tidal": "wave_tidal", "Bioenergy": "bioenergy", "Other fuels": "other_fuels",
         "Hydro pumped storage": "pumped_storage", "Battery storage": "battery_storage",
         "Total major power producers": "total", "Total other generators": "total"}
BLOCKS = {"Major power producers": "mpp", "Other generators": "other"}
SUPPLY = {"Electricity generated by Major Power Producers (MPPs)": ("generated_mpp", "generation"),
          "Electricity generated by other generators": ("generated_other", "generation"),
          "Total electricity generated": ("generated_total", "generation"),
          "Own use by other generators": ("own_use_other", "own_use"),
          "Net electricity supplied by other generators": ("net_supplied_other", "net_supply"),
          "Used in energy storage and other own use by MPPs": ("own_use_mpp", "own_use"),
          "Net electricity supplied by MPPs": ("net_supplied_mpp", "net_supply"),
          "Electricity transferred to England (net of receipts)": ("transfer_to_england", "net_transfers"),
          "Electricity transferred to Northern Ireland (net of receipts)": ("transfer_to_northern_ireland", "net_transfers"),
          "Electricity transferred to Europe (net of receipts)": ("transfer_to_europe", "net_transfers"),
          "Transfers from other generators to public supply": ("transfers_other_to_public", "transfers"),
          "Transmission losses": ("transmission_losses", "losses"),
          "Distribution losses and theft": ("distribution_losses", "losses"),
          "Consumption from public supply": ("consumption_public", "consumption"),
          "Consumption by autogenerators": ("consumption_autogenerators", "consumption"),
          "Total electricity consumed": ("consumption_total", "consumption"),
          "Electricity sales (public supply)": ("sales", "sales"),
          "Statistical difference between calculated consumption": ("statistical_difference", "statistical_difference")}
X_COLS = []


def label(v):
    s = re.sub(r"\s*\[[^\]]*\]", "", " ".join(str(v).split())) if v is not None else ""
    return s.strip()


def header_cols(row):
    cols = {}
    for j, x in enumerate(row):
        m = re.fullmatch(r"(\d{4}) (.+?)\d*", label(x))  # Table 1 heads one column "2022 England6"
        if m:
            if m.group(2) not in NATIONS:
                raise RuntimeError(f"a nation not in the map: {m.group(2)!r}")
            cols[j] = (int(m.group(1)), m.group(2))
    years = sorted({y for y, _ in cols.values()})
    if not years or years[0] != 2004 or years != list(range(2004, years[-1] + 1)) or len(cols) != 5 * len(years):
        raise RuntimeError(f"the header's years and nations are not whole: {years[:2]} ... {years[-2:]}, {len(cols)} columns")
    return cols


def value(v):
    if v is None or (isinstance(v, str) and v.strip() in ("", "[x]")):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    raise RuntimeError(f"not a number: {v!r}")


def parse_fuel(rows):
    t = next((i for i, r in enumerate(rows) if label(r[0]).startswith("Table 2a: Electricity generated by fuel")), None)
    if t is None:
        raise RuntimeError("Table 2a not found")
    cols = header_cols(rows[t + 1])
    out, seen = {}, set()
    for r in rows[t + 2:]:
        g = label(r[0])
        if g.startswith("Table 2b") or g == "All generating companies":
            break
        if g not in BLOCKS:
            raise RuntimeError(f"a generator type not in the map: {g!r}")
        f = label(r[1])
        if f not in FUELS:
            raise RuntimeError(f"a fuel not in the map: {f!r}")
        item = f"{BLOCKS[g]}_{FUELS[f]}"
        if item in seen:
            raise RuntimeError(f"a row given twice: {g} {f}")
        seen.add(item)
        for j, (y, n) in cols.items():
            v = value(r[j])
            if v is not None:
                out[(NATIONS[n][0], item, y)] = (v, f"{g}: {f}", n)
    for b in BLOCKS.values():
        if f"{b}_total" not in seen:
            raise RuntimeError(f"no total row for {b}")
    for (n, item, y), (v, _, _) in list(out.items()):
        if item.endswith("_total"):
            b = item[:-6]
            s = sum(x for (nn, it, yy), (x, _, _) in out.items() if nn == n and yy == y and it.startswith(b + "_") and it != item)
            if abs(s - v) > 10:
                raise RuntimeError(f"{n} {y} {b}: the fuels make {s:,.1f} GWh, the total {v:,.1f}")
    return out


def parse_supply(rows):
    t = next((i for i, r in enumerate(rows) if label(r[0]) == "Components of electricity generation and supply"), None)
    if t is None:
        raise RuntimeError("Table 1's header not found")
    cols = header_cols(rows[t])
    out, seen = {}, set()
    for r in rows[t + 1:]:
        lab = label(r[0])
        if not lab:
            continue
        key = next((k for k in SUPPLY if lab == k or lab.startswith(k)), None)
        if key is None:
            raise RuntimeError(f"a row not in the map: {lab!r}")
        item = SUPPLY[key][0]
        if item in seen:
            raise RuntimeError(f"a row given twice: {lab!r}")
        seen.add(item)
        for j, (y, n) in cols.items():
            v = value(r[j])
            if v is not None:
                out[(NATIONS[n][0], item, y)] = (v, lab, n)
    missing = sorted(set(v[0] for v in SUPPLY.values()) - seen)
    if missing:
        raise RuntimeError(f"rows missing: {missing}")
    return out  # kept as published: its total generated misses its parts in seven nation-years (the docstring)


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    for s in ("Generation and supply", "Electricity generation by fuel"):
        if s not in wb.sheetnames:
            raise RuntimeError(f"no sheet {s!r}: {wb.sheetnames}")
    fuel = parse_fuel([tuple(r) for r in wb["Electricity generation by fuel"].iter_rows(values_only=True)])
    supply = parse_supply([tuple(r) for r in wb["Generation and supply"].iter_rows(values_only=True)])
    var = {f"{b}_{f}": "generation" for b in BLOCKS.values() for f in FUELS.values()}
    var.update({v[0]: v[1] for v in SUPPLY.values()})
    geo = {v[0]: v[1] for v in NATIONS.values()}
    rows = [dict(entity=f"desnz:{n}:{item}", variable=var[item], ts_utc=f"{y}-01-01T00:00:00Z", value=round(v * 1000, 3),
                 unit="MWh", freq="P1Y", geo=geo[n], market="", node=lab, **base)
            for (n, item, y), (v, lab, _) in sorted({**fuel, **supply}.items())]
    return pd.DataFrame(rows)


def main(argv=None):
    ap = argparse.ArgumentParser(description="DESNZ: electricity generation by fuel and supply for each UK nation, a year")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t = parse(content, base)
        if t.duplicated(["entity", "variable", "ts_utc"]).any():
            raise RuntimeError("a key given twice")
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["entity", "ts_utc"]).reset_index(drop=True)
        g = t[t.entity.str.endswith(":generated_total") & (t.ts_utc.str[:4] == t.ts_utc.max()[:4])].set_index("entity").value
        head = [
            "Energy Research Warehouse (ERW): DESNZ, Electricity generation and supply in Scotland, Wales, Northern Ireland and "
            f"England: generation by fuel and the supply balance for each UK nation, {t.ts_utc.min()[:4]} to {t.ts_utc.max()[:4]} "
            "(session 2678)",
            "Shape: series (docs/datastandard.md v0). entity desnz:<nation>:<item>; variable generation, own_use, net_supply, "
            "net_transfers, transfers, losses, consumption, sales or statistical_difference; unit MWh (the tables' GWh times 1,000); "
            "ts_utc the year's first day; freq P1Y; geo GB or GB-SCT, GB-WLS, GB-NIR, GB-ENG; node the row's label.",
            f"{len(t)} rows; total generated in {t.ts_utc.max()[:4]}, TWh: "
            + ", ".join(f"{e.split(':')[1]} {v / 1e6:.1f}" for e, v in g.items())
            + ". The major power producers' and the other generators' blocks only (the derived all-companies rows do not add); "
            "the nations do not sum exactly to the UK (DESNZ's note 9); nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/desnz_uk_nations_electricity.py; the workbook retrieved "
            f"{base['retrieved_at']}; last modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL} (listed on {PAGE})",
            "License: Open Government Licence v3.0 (the publication page). Credit: Department for Energy Security and Net Zero.",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="UK Department for Energy Security and Net Zero (DESNZ)",
                                    report="Electricity generation and supply in Scotland, Wales, Northern Ireland and England, 2004 to "
                                           "2025: data tables (Energy Trends, September 2026)", report_url=PAGE, document_list=PAGE,
                                    license="public", tables=[NAME])])  # Open Government Licence v3.0
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
