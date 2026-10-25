#!/usr/bin/env python3
"""Hawaii's electricity by island: generation by source and island (2024), the utilities' customers and sales by island
(2025), and Hawaiian Electric's service on Oahu a year (customers, net input, sales, generating capability, 1994 to 2025),
as the Hawaii State Department of Business, Economic Development and Tourism (DBEDT) gives them in tables 17.07, 17.10 and
17.13 of its State of Hawaii Data Book 2025 (session 2580, chain 127, a source of the session's own choosing; the ERW held
Hawaii only as a state in EIA's tables).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/dbedt_hawaii_electricity_annual.csv, from one
workbook (no key):

    https://files.hawaii.gov/dbedt/economic/databook/db2025/DB2025_section17.xlsx

    entity      dbedt:hi:<island>:<source> (17.07: generation; island state_total, hawaii, maui, lanai, molokai, oahu, kauai,
                not_specified; source total, petroleum, coal, wind, biomass, geothermal, hydro, other_gases, solar, others);
                dbedt:hi:<island> (17.10: residential_customers, residential_sales, other_customers, other_sales);
                dbedt:hi:heco_oahu (17.13: customers, residential_customers, net_input, sales, generating_capability)
    unit        MWh (17.07's million kWh times 1,000; 17.10's million kWh times 1,000; 17.13's 1,000 kWh as MWh), MW (17.13's
                1,000 kW), count; ts_utc the year's first day; freq P1Y; geo US-HI
    node        DBEDT's label; x_table the table

Table 17.07's own note says its components may not sum to the total because of rounding; each island's sources are checked
against its total within 0.5 million kWh (nine rounded components), and the islands against the state total within 0.5
million kWh. Its "Others"
column holds negative figures (energy storage's losses), kept as given. Table 17.10's rates and revenues (USD) and average
use per customer, and 17.13's average use and rate, are not taken. Nothing is filled.

    python warehouse/connectors/dbedt_hawaii_electricity.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/dbedt_hi/ and reused unless --fresh; --offline asks nothing.
files.hawaii.gov has no robots file; dbedt.hawaii.gov's asks a crawl delay of 10 seconds (session 2580). License: internal.
The State's pages read "Copyright (c) 2026, State of Hawaii. All rights reserved."; its Terms of Use page is drawn by
script; internal until a ruling. Table 17.07's source is EIA's Form EIA-923, table 17.10's the utilities' monthly financial
reports, 17.13's Hawaiian Electric's records, each with DBEDT's calculations.
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "dbedt_hawaii_electricity_annual"
CONNECTOR = "dbedt_hi"
SOURCE = "dbedt_hi:data_book_2025_section_17"
HOST = "files.hawaii.gov"
URL = "https://files.hawaii.gov/dbedt/economic/databook/db2025/DB2025_section17.xlsx"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = ("Hawaii State Department of Business, Economic Development and Tourism (DBEDT), The State of Hawaii Data Book 2025, "
            "section 17, tables 17.07, 17.10 and 17.13")
SOURCES = [("Total", "total"), ("Petroleum", "petroleum"), ("Coal", "coal"), ("Wind", "wind"), ("Biomass [a]", "biomass"),
           ("Geothermal", "geothermal"), ("Hydro", "hydro"), ("Other gases [b]", "other_gases"), ("Solar", "solar"), ("Others [c]", "others")]
ISLANDS = [("State total", "state_total"), ("Hawaii", "hawaii"), ("Maui", "maui"), ("Lanai", "lanai"), ("Molokai", "molokai"),
           ("Oahu", "oahu"), ("Kauai", "kauai"), ("Not specified", "not_specified")]
UTIL_ISLANDS = [("State Total", "state_total"), ("Oahu", "oahu"), ("Hawaii", "hawaii"), ("Kauai", "kauai"), ("Lanai", "lanai"),
                ("Maui", "maui"), ("Molokai", "molokai")]
HECO = [(1, "Total number of customers (as of Dec. 31)", "customers", "count", 1),
        (2, "Residential customers", "residential_customers", "count", 1),
        (3, "Net input (1,000 kWh) [a]", "net_input", "MWh", 1),
        (4, "Electricity sales (1,000 kWh)", "sales", "MWh", 1),
        (7, "Generating capability (as of Dec. 31) (1,000 kW) [c]", "generating_capability", "MW", 1)]
X_COLS = ["x_table"]


def lab(v):
    return " ".join(str(v).split()) if v is not None else ""


def num(v, where, neg=False):
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        raise RuntimeError(f"{where}: not a number: {v!r}")
    if v < 0 and not neg:
        raise RuntimeError(f"{where}: negative: {v!r}")
    return float(v)


def row(entity, var, year, v, unit, node, table, base):
    return dict(entity=entity, variable=var, ts_utc=f"{year}-01-01T00:00:00Z", value=round(v, 3), unit=unit, freq="P1Y", geo="US-HI",
                market="", node=node, **base, x_table=table)


def year_of(title, key):
    m = re.fullmatch(re.escape(key) + r".*: (\d{4})", title)
    if not m:
        raise RuntimeError(f"{key}: no single year in its title: {title!r}")
    return int(m.group(1))


def parse_1707(rows, base):
    y = year_of(lab(rows[0][0]), "Table 17.07-- ELECTRICITY PRODUCTION, BY SOURCE AND BY ISLAND")
    if lab(rows[1][0]) != "[Million kWh. Components may not sum to total due to rounding]":
        raise RuntimeError(f"17.07: the unit line is not million kWh: {rows[1][0]!r}")
    head = [lab(v) for v in rows[2]]
    if head[0] != "Island" or head[1:1 + len(SOURCES)] != [s[0] for s in SOURCES]:
        raise RuntimeError(f"17.07: the header is not the sources expected: {head}")
    out, vals = [], {}
    body = {lab(r[0]): r for r in rows[3:3 + len(ISLANDS)]}
    if list(body) != [i[0] for i in ISLANDS]:
        raise RuntimeError(f"17.07: the islands are not those expected: {list(body)}")
    for name, isl in ISLANDS:
        r = body[name]
        v = [num(r[j + 1], f"17.07 {name} {s}", neg=(key == "others")) for j, (s, key) in enumerate(SOURCES)]
        if abs(sum(v[1:]) - v[0]) > 0.5:
            raise RuntimeError(f"17.07 {name}: the sources make {sum(v[1:])}, the total {v[0]}")
        vals[isl] = v
        out += [row(f"dbedt:hi:{isl}:{key}", "generation", y, x * 1000, "MWh", f"{name}: {s}", "17.07", base)
                for (s, key), x in zip(SOURCES, v)]
    for j in range(len(SOURCES)):
        s = sum(vals[i][j] for _, i in ISLANDS[1:])
        if abs(s - vals["state_total"][j]) > 0.5:
            raise RuntimeError(f"17.07 {SOURCES[j][0]}: the islands make {s}, the state {vals['state_total'][j]}")
    return out


def parse_1710(rows, base):
    y = year_of(lab(rows[0][0]), "Table 17.10-- ELECTRIC UTILITIES, BY ISLAND")
    out = []
    blocks = [i for i, r in enumerate(rows) if lab(r[0]) == "Island"]
    if len(blocks) != 2:
        raise RuntimeError(f"17.10: two blocks expected, {len(blocks)} found")
    for b, kind in zip(blocks, ("residential", "other")):
        h = [lab(v) for v in rows[b]]
        lead = "Residential" if kind == "residential" else "Other-use"
        if not (h[1].startswith(f"Number of {lead.lower()} consumer") or h[1].startswith(f"Number of {lead} consumer")) \
                or h[2] != f"{lead} power sold (million kWh)":
            raise RuntimeError(f"17.10: the {kind} block's header is not as expected: {h[:3]}")
        got = {lab(r[0]): r for r in rows[b + 1:b + 1 + len(UTIL_ISLANDS)]}
        if list(got) != [i[0] for i in UTIL_ISLANDS]:
            raise RuntimeError(f"17.10: the {kind} block's islands are not those expected: {list(got)}")
        cust = {i: num(got[n][1], f"17.10 {n} {kind} customers") for n, i in UTIL_ISLANDS}
        sold = {i: num(got[n][2], f"17.10 {n} {kind} sold") for n, i in UTIL_ISLANDS}
        for d, what in ((cust, "customers"), (sold, "sales")):
            s = sum(v for k, v in d.items() if k != "state_total")
            tol = 1 if what == "customers" else 3
            if abs(s - d["state_total"]) > tol:
                raise RuntimeError(f"17.10 {kind} {what}: the islands make {s}, the state {d['state_total']}")
        for n, i in UTIL_ISLANDS:
            out.append(row(f"dbedt:hi:{i}", f"{kind}_customers", y, cust[i], "count", f"{n}: {h[1]}", "17.10", base))
            out.append(row(f"dbedt:hi:{i}", f"{kind}_sales", y, sold[i] * 1000, "MWh", f"{n}: {h[2]}", "17.10", base))
    return out


def parse_1713(rows, base):
    if not lab(rows[0][0]).startswith("Table 17.13-- SERVICE PROVIDED BY HAWAIIAN ELECTRIC COMPANY, INC. ON OAHU"):
        raise RuntimeError(f"17.13: the title is not as expected: {rows[0][0]!r}")
    h = [lab(v) for v in rows[1]]
    for j, name, *_ in HECO:
        if h[j] != name:
            raise RuntimeError(f"17.13: column {j} is {h[j]!r}, not {name!r}")
    out, years = [], []
    for r in rows[2:]:
        y = r[0]
        if not isinstance(y, int):
            break
        years.append(y)
        for j, name, var, unit, _ in HECO:
            out.append(row("dbedt:hi:heco_oahu", var, y, num(r[j], f"17.13 {y} {var}"), unit, name, "17.13", base))
    if not years or years != list(range(years[0], years[-1] + 1)):
        raise RuntimeError(f"17.13: the years are not a run: {years}")
    return out


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    sheets = {}
    for s in ("17.07", "17.10", "17.13"):
        if s not in wb.sheetnames:
            raise RuntimeError(f"no sheet {s}")
        sheets[s] = [list(r) + [None] * 12 for r in wb[s].iter_rows(values_only=True) if any(v is not None for v in r)]
    return pd.DataFrame(parse_1707(sheets["17.07"], base) + parse_1710(sheets["17.10"], base) + parse_1713(sheets["17.13"], base))


def main(argv=None):
    ap = argparse.ArgumentParser(description="DBEDT: Hawaii's electricity by island and Hawaiian Electric's service on Oahu")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t = parse(content, base)
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["variable", "entity", "ts_utc"]).reset_index(drop=True)
        v = {(x, e, ts[:4]): val for x, e, ts, val in zip(t.variable, t.entity, t.ts_utc, t.value)}
        g = t[t.variable == "generation"].ts_utc.str[:4].max()
        head = [
            "Energy Research Warehouse (ERW): Hawaii State Department of Business, Economic Development and Tourism (DBEDT), State of "
            "Hawaii Data Book 2025, tables 17.07 (generation by source and island), 17.10 (utilities by island) and 17.13 (Hawaiian "
            "Electric on Oahu) (session 2580)",
            "Shape: series (docs/datastandard.md v0). entity dbedt:hi:<island>:<source>, dbedt:hi:<island> or dbedt:hi:heco_oahu; "
            "variable generation, residential_customers, residential_sales, other_customers, other_sales, customers, net_input, "
            "sales, generating_capability; unit MWh, MW or count; ts_utc the year's first day; freq P1Y; geo US-HI; x_table the table.",
            f"{len(t)} rows; in {g} generation {v[('generation', 'dbedt:hi:state_total:total', g)] / 1e3:,.1f} GWh, Oahu "
            f"{v[('generation', 'dbedt:hi:oahu:total', g)] / 1e3:,.1f} GWh. 17.07's components may not sum to its totals (its own note); "
            "its Others column's negative figures kept as given; rates, revenues and averages not taken; nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/dbedt_hawaii_electricity.py; the workbook retrieved {base['retrieved_at']}; "
            f"last modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL}",
            "License: internal. The State's pages read \"Copyright (c) 2026, State of Hawaii. All rights reserved.\"; internal until a "
            f"ruling. Cite: {CITATION}.",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="Hawaii State Department of Business, Economic Development and Tourism (DBEDT)",
                                    report="The State of Hawaii Data Book 2025, section 17, tables 17.07, 17.10 and 17.13", report_url=URL,
                                    document_list="https://dbedt.hawaii.gov/economic/databook/db2025/", license="internal", tables=[NAME])])
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
