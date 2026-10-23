#!/usr/bin/env python3
"""EPA's Greenhouse Gas Reporting Program: every reporting facility's total reported emissions by year, in every sector
the program's summary covers (session 2072, chain 067: the emitters beside the power plants of epa_ghgrp_power_co2).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/epa_ghgrp_facility_emissions.csv, from the
workbook ghgp_data_by_year_<year>.xlsx inside EPA's newest data summary spreadsheets zip (the one the program's data sets
page links: https://www.epa.gov/ghgreporting/data-sets). The workbook has one sheet a segment of the program; six are
emissions and are taken:

    Direct Point Emitters          emissions_direct                 (power plants, refineries, chemicals, waste ...)
    Onshore Oil & Gas Prod.        emissions_onshore_production     (subpart W onshore production, by basin)
    Gathering & Boosting           emissions_gathering_boosting
    Transmission Pipelines         emissions_transmission_pipelines
    LDC - Direct Emissions         emissions_ldc                    (local distribution companies)
    SF6 from Elec. Equip.          emissions_sf6_equipment          (SF6 from electrical transmission equipment)

Not taken: "Suppliers" (the greenhouse gas quantity in fuels and gases supplied, not emitted), "CO2 Injection" and
"Geologic Sequestration of CO2" (quantities injected and stored), the industry type key and the FAQs.

    entity      ghgrp:<GHGRP facility id>; for transmission pipelines and local distribution companies, which EPA lists
                a row a state where the emissions occur, ghgrp:<facility id>:<state> (x_emission_state; a row that
                names no such state keeps ghgrp:<facility id>, its geo the reported state); node the facility's name
    variable    emissions_<segment>, metric tons of CO2 equivalent (tCO2e, Decision 61), as EPA totals them with the
                IPCC's Fourth Assessment Report global warming potentials
    ts_utc      the reporting year's 1 January; freq P1Y; geo US-<state> (the state where the emissions occur where the
                sheet gives it, else the facility's state; several states written as a list)
    x_*         the segment, the FRS id, the primary NAICS code, the latest reported industry type by subpart and by
                sector (where the sheet gives it), the basin (onshore production and gathering), the county

From 2015 (the workbook holds 2011 on: --first-year). A blank cell has no row; nothing is filled. A sheet without its
"Facility Id" header, a column of years not in EPA's form, a value that is not a number or a facility given twice in a
year and segment stops the run.

    python warehouse/connectors/ghgrp_facilities.py [--offline] [--first-year 2015] [--out-dir DIR]

Ceiling: two requests (the data sets page, fresh each run unless --offline, and the zip it links, reused from
warehouse/raw/ghgrp/, which warehouse/connectors/ghgrp_power.py shares). www.epa.gov's robots file disallows only Drupal's
own paths and search pages. License: public (EPA, record 55 of license_reads.csv; cite "U.S. EPA, Greenhouse Gas
Reporting Program, data summary <year>").
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback
import urllib.parse
import zipfile

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "epa_ghgrp_facility_emissions"
CONNECTOR = "ghgrp_facilities"
RAW = "ghgrp"   # the raw store shared with ghgrp_power.py
SOURCE = "epa:ghgrp_facilities"
PAGE = "https://www.epa.gov/ghgreporting/data-sets"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
COLS = ["entity", "variable", "ts_utc", "value", "unit", "freq", "geo", "market", "node", "source", "source_url",
        "retrieved_at", "vintage", "x_segment", "x_frs_id", "x_naics", "x_subparts", "x_sectors", "x_basin", "x_county", "x_emission_state"]
SEGMENTS = {"Direct Point Emitters": "direct", "Onshore Oil & Gas Prod.": "onshore_production",
            "Gathering & Boosting": "gathering_boosting", "Transmission Pipelines": "transmission_pipelines",
            "LDC - Direct Emissions": "ldc", "SF6 from Elec. Equip.": "sf6_equipment"}
BY_STATE = ("Transmission Pipelines", "LDC - Direct Emissions")
YEAR_RE = re.compile(r"^(\d{4}) Total reported (?:direct )?emissions")
STATE_RE = re.compile(r"^[A-Z]{2}$")
NAMES = {"Alabama": "AL", "Alaska": "AK", "Arizona": "AZ", "Arkansas": "AR", "California": "CA", "Colorado": "CO",
         "Connecticut": "CT", "Delaware": "DE", "District of Columbia": "DC", "Florida": "FL", "Georgia": "GA",
         "Hawaii": "HI", "Idaho": "ID", "Illinois": "IL", "Indiana": "IN", "Iowa": "IA", "Kansas": "KS", "Kentucky": "KY",
         "Louisiana": "LA", "Maine": "ME", "Maryland": "MD", "Massachusetts": "MA", "Michigan": "MI", "Minnesota": "MN",
         "Mississippi": "MS", "Missouri": "MO", "Montana": "MT", "Nebraska": "NE", "Nevada": "NV", "New Hampshire": "NH",
         "New Jersey": "NJ", "New Mexico": "NM", "New York": "NY", "North Carolina": "NC", "North Dakota": "ND",
         "Ohio": "OH", "Oklahoma": "OK", "Oregon": "OR", "Pennsylvania": "PA", "Rhode Island": "RI",
         "South Carolina": "SC", "South Dakota": "SD", "Tennessee": "TN", "Texas": "TX", "Utah": "UT", "Vermont": "VT",
         "Virginia": "VA", "Washington": "WA", "West Virginia": "WV", "Wisconsin": "WI", "Wyoming": "WY",
         "Puerto Rico": "PR", "Guam": "GU", "American Samoa": "AS", "Virgin Islands": "VI",
         "Northern Mariana Islands": "MP"}


def text(v):
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def pick(cols, *names):
    for n in names:
        if n in cols:
            return cols.index(n)
    return None


def code_of(name):
    """A state's two-letter code from EPA's "State where Emissions Occur", a full name or already a code; a name not in
    the map stops."""
    if STATE_RE.match(name) and name in NAMES.values():
        return name
    if name not in NAMES:
        raise RuntimeError(f"a state name not in the map: {name!r}")
    return NAMES[name]


def geo_of(state):
    parts = [p.strip() for p in re.split(r"[,;/ ]+", state) if p.strip()]
    if parts and all(STATE_RE.match(p) for p in parts):
        return ",".join(f"US-{p}" for p in parts)
    return ""


def sheet_rows(ws, segment, first_year, url, retrieved, vintage):
    rows = ws.iter_rows(values_only=True)
    head = None
    for r in rows:
        if r and text(r[0]) == "Facility Id":
            head = [text(c) for c in r]
            break
    if head is None:
        raise RuntimeError(f"{segment}: no 'Facility Id' header")
    years = {}
    for j, c in enumerate(head):
        if re.match(r"^\d{4} ", c):
            m = YEAR_RE.match(c)
            if not m:
                raise RuntimeError(f"{segment}: a column of years not in EPA's form: {c!r}")
            years[j] = int(m.group(1))
    i_id, i_frs, i_name = pick(head, "Facility Id"), pick(head, "FRS Id"), pick(head, "Facility Name")
    by_state = "State where Emissions Occur" in head
    i_state = pick(head, "State where Emissions Occur", "State", "Reported State")
    i_naics, i_sub = pick(head, "Primary NAICS Code"), pick(head, "Latest Reported Industry Type (subparts)")
    i_sec, i_basin = pick(head, "Latest Reported Industry Type (sectors)"), pick(head, "Basin")
    i_county = pick(head, "County", "Reported County")
    var = f"emissions_{SEGMENTS[segment]}"
    out = []
    for r in rows:
        if not r or not text(r[i_id]):
            continue
        cell = lambda i: text(r[i]) if i is not None and i < len(r) else ""
        fid, state, where = cell(i_id), cell(i_state), ""
        if by_state:
            where = code_of(state) if state not in ("", "None") else ""
            state = where or cell(pick(head, "Reported State"))
        eid = f"ghgrp:{fid}:{where}" if where else f"ghgrp:{fid}"
        for j, year in years.items():
            if year < first_year:
                continue
            v = cell(j)
            if v == "":
                continue
            try:
                value = float(v)
            except ValueError:
                raise RuntimeError(f"{segment}: {fid} {year}: not a number: {v!r}")
            out.append({"entity": eid, "variable": var, "ts_utc": f"{year}-01-01T00:00:00Z", "value": value,
                        "unit": "tCO2e", "freq": "P1Y", "geo": geo_of(state), "market": "", "node": cell(i_name),
                        "source": SOURCE, "source_url": url, "retrieved_at": retrieved, "vintage": vintage,
                        "x_segment": segment, "x_frs_id": cell(i_frs), "x_naics": cell(i_naics), "x_subparts": cell(i_sub),
                        "x_sectors": cell(i_sec), "x_basin": cell(i_basin), "x_county": cell(i_county),
                        "x_emission_state": where})
    return out


def build(workbook_bytes, first_year, url, retrieved, vintage):
    wb = openpyxl.load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    missing = [s for s in SEGMENTS if s not in wb.sheetnames]
    if missing:
        raise RuntimeError(f"sheets missing: {missing}")
    rows = []
    for segment in SEGMENTS:
        rows.extend(sheet_rows(wb[segment], segment, first_year, url, retrieved, vintage))
    t = pd.DataFrame(rows, columns=COLS)
    twice = t[t.duplicated(["entity", "variable", "ts_utc"], keep=False)]
    if len(twice):
        raise RuntimeError(f"a facility given twice in a year and segment: {twice[['entity', 'variable', 'ts_utc']].head(3).values.tolist()}")
    return t.sort_values(["variable", "entity", "ts_utc"]).reset_index(drop=True)


def main(argv=None):
    ap = argparse.ArgumentParser(description="EPA GHGRP: every facility's total reported emissions by year")
    ap.add_argument("--offline", action="store_true", help="the saved files only, no request")
    ap.add_argument("--first-year", type=int, default=2015)
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(RAW, run_id)
    results = []
    try:
        page, _ = ip.fetch_raw(RAW, PAGE, log, offline=args.offline, fresh=not args.offline, headers=UA, timeout=120)
        links = re.findall(r'href="([^"]*?(\d{4})_data_summary_spreadsheets\.zip)"', page.decode("utf-8", errors="replace"))
        if not links:
            raise RuntimeError(f"{PAGE} links no data summary zip")
        rel, year = max(links, key=lambda x: int(x[1]))
        url = urllib.parse.urljoin(PAGE, rel)
        content, rec = ip.fetch_raw(RAW, url, log, offline=args.offline, headers=UA, timeout=600)
        z = zipfile.ZipFile(io.BytesIO(content))
        name = f"ghgp_data_by_year_{year}.xlsx"
        if name not in z.namelist():
            raise RuntimeError(f"{url} holds no {name}")
        t = build(z.read(name), args.first_year, url, rec.get("retrieved_at") or "", ip.vintage_of(rec))
        by = t.groupby("x_segment").agg(facilities=("entity", "nunique"), rows=("entity", "size"))
        header = [
            f"Energy Research Warehouse (ERW): U.S. EPA, Greenhouse Gas Reporting Program, data summary {year}: every reporting facility's total reported emissions by year, "
            f"{args.first_year} to {year}, in the six emitting segments of the summary (session 2072)",
            "Shape: series (docs/datastandard.md v0). entity ghgrp:<facility id>, or ghgrp:<facility id>:<state> for transmission pipelines and local distribution companies (a row a state where the emissions occur); variable emissions_<segment> (direct, onshore_production, gathering_boosting, "
            "transmission_pipelines, ldc, sf6_equipment); unit tCO2e (metric tons of CO2 equivalent, AR4 global warming potentials, Decision 61); ts_utc the year's 1 January; freq P1Y.",
            f"{len(t)} rows: " + "; ".join(f"{s} {r.facilities} {'facility-states' if s in BY_STATE else 'facilities'}, {r.rows} rows" for s, r in by.iterrows())
            + ". Suppliers, CO2 injection and sequestration are quantities, not emissions, and are not taken. A blank cell has no row; nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/ghgrp_facilities.py; the zip retrieved {rec.get('retrieved_at') or ''}, last modified {ip.vintage_of(rec)}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{RAW}/ (not in git; shared with ghgrp_power.py)",
            f"Source: {SOURCE} U.S. Environmental Protection Agency, Greenhouse Gas Reporting Program, data summary spreadsheets, {name}, {PAGE}",
            f"  access: {url}",
            f"License: public (a work of the U.S. government). Cite: U.S. EPA, Greenhouse Gas Reporting Program, data summary {year}.",
        ]
        ip.write_csv(t, NAME, header, log, cols=COLS)
        if not args.out_dir:
            ip.update_sources([dict(source=SOURCE, publisher="U.S. Environmental Protection Agency (EPA)",
                                    report=f"Greenhouse Gas Reporting Program, data summary spreadsheets ({name})", report_url=PAGE,
                                    document_list=url, license="public", tables=[NAME])])
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} facilities, {args.first_year} to {year}")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"FAILED, no output file written:\n{tb}")
        print(f"{CONNECTOR} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
