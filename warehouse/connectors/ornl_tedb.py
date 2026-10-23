#!/usr/bin/env python3
"""Oak Ridge National Laboratory's Transportation Energy Data Book: transportation's energy use by mode and fuel, the
sectors' consumption, and light-duty vehicle sales by powertrain, a year a row (session 2048, chain 059: transport's
energy as electric vehicles grow).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/ornl_tedb_transportation_annual.csv, from
five tables of the latest edition's spreadsheets (one zip of all 267 tables and figures, linked from the book's data
page https://tedb.ornl.gov/data/):

    Table 2.2  U.S. consumption of total energy by end-use sector (quadrillion Btu, written here as TBtu: x 1,000)
    Table 2.4  distribution of transportation energy consumption by source (trillion Btu, and each source's share)
    Table 2.7  domestic consumption of transportation energy by mode and fuel type, the latest year (trillion Btu)
    Table 2.8  transportation energy use by mode, the latest two years (trillion Btu)
    Table 6.2  light-duty hybrid and plug-in electric vehicle sales (thousands, written here as a count: x 1,000)

    entity      ornl_tedb:<table>:<row>, e.g. ornl_tedb:t2_7:highway.light_duty.cars; node the row's label as printed
    variable    energy_consumption (2.2), energy_use and share_<source> (2.4; a share in percent: TEDB's fraction x 100),
                energy_use_<fuel> (2.7), energy_use (2.8), sales_hybrid, sales_plugin_hybrid, sales_battery_electric and
                sales_light_duty (6.2)
    ts_utc      the year's first day, freq P1Y; geo US

A row label is read against a closed map (its footnote letter dropped: "Light-duty trucksb"), and a mode's place in the
hierarchy from its capitals and indentation. A dash or an empty cell is no row. Table 2.4's and 6.2's per-cent change
rows are not taken, nor the shares of 2.2 and 2.8, nor 6.2's shares (each a ratio of two rows held). Values as TEDB
publishes them, but for the stated conversions; nothing is filled. A label not in the map, a header not where it is
expected or a year twice stops the run.

    python warehouse/connectors/ornl_tedb.py [--offline] [--out-dir DIR]

Ceiling: two requests (the data page, fresh each run, and the zip it links, reused from warehouse/raw/ornl_tedb/ when
saved). tedb.ornl.gov's robots file allows every path but /admin/. License: the book's pages state no terms of reuse
(its citation page, session 2048); ORNL is managed by UT-Battelle for the Department of Energy. Table 6.2 is compiled by
Argonne's Transportation Research Center. Flagged internal until a ruling (session 2048).
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback
import urllib.parse
import zipfile

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "ornl_tedb_transportation_annual"
CONNECTOR = "ornl_tedb"
SOURCE = "ornl:tedb"
PAGE = "https://tedb.ornl.gov/data/"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
COLS = ["entity", "variable", "ts_utc", "value", "unit", "freq", "geo", "market", "node", "source", "source_url",
        "retrieved_at", "vintage", "x_table", "x_edition"]
# the modes of Tables 2.7 and 2.8, by their label lowercased without its footnote letter: (slug, level); level 0 is a
# capitalised total, 1 a mode, 2 a part of a mode, 3 a part of a part
MODES = {"highway": ("highway", 0), "nonhighway": ("nonhighway", 0), "hwy & nonhwy total": ("total", 0),
         "total hwy & nonhwy": ("total", 0), "off-highway": ("off_highway", 0),
         "light-duty vehicles": ("light_duty", 1), "cars": ("cars", 2), "light-duty trucks": ("trucks", 2),
         "motorcycles": ("motorcycles", 2), "buses": ("buses", 1), "transit": ("transit", 2), "intercity": ("intercity", 2),
         "school": ("school", 2), "medium/heavy-duty trucks": ("medium_heavy_trucks", 1), "class 3-6 trucks": ("class_3_6", 2),
         "class 7-8 trucks": ("class_7_8", 2), "air": ("air", 1), "general aviation": ("general_aviation", 2),
         "domestic air carriers": ("domestic_air_carriers", 2), "international air": ("international_air", 2),
         "international air carriers": ("international_air", 2), "water": ("water", 1), "freight": ("freight", 2),
         "recreational": ("recreational", 2), "pipeline": ("pipeline", 1), "rail": ("rail", 1),
         "freight (class i)": ("freight_class_i", 2), "passenger": ("passenger", 2), "commuter": ("commuter", 3)}
FUELS = {"gasoline": "gasoline", "diesel fuel": "diesel", "liquefied petroleum gas": "lpg", "jet fuel": "jet_fuel",
         "residual fuel oil": "residual_fuel_oil", "natural gas": "natural_gas", "electricity": "electricity", "total": "total"}
SOURCES_24 = {"petroleum": "petroleum", "natural gas": "natural_gas", "coal": "coal", "renewables": "renewables",
              "electricity": "electricity", "total": "total"}
SECTORS_22 = {"transportation": "transportation", "industrial": "industrial", "commercial": "commercial",
              "residential": "residential", "total": "total"}
SALES_62 = ["sales_hybrid", "sales_plugin_hybrid", "sales_battery_electric", "sales_light_duty"]


def text(v):
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def label(v, known):
    """A row or column label against a closed map, its footnote letter dropped: (clean key) or a stop."""
    s = text(v).lower()
    if s in known:
        return s
    if s[:-1] in known and s[-1] in "abcdefg":
        return s[:-1]
    raise RuntimeError(f"a label not in the map: {text(v)!r}")


def num(v):
    s = text(v)
    if s in ("", "-", "b", "a"):
        return None
    try:
        return float(s)
    except ValueError:
        raise RuntimeError(f"not a number: {s!r}")


def rows_of(content):
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    while rows and all(not r or text(r[0]) == "" for r in rows):  # TEDB's sheets begin in column B
        rows = [r[1:] for r in rows]
    return rows


def first_row(rows, test, what):
    for i, r in enumerate(rows):
        if r and test(r):
            return i
    raise RuntimeError(f"{what} not found")


def year_rows(rows, start):
    """The rows from start whose first cell is a year, stopping at the first that is not."""
    out = []
    for r in rows[start:]:
        y = text(r[0] if r else "")
        if not re.fullmatch(r"\d{4}", y.replace(".0", "")):
            break
        out.append((int(float(y)), r))
    return out


def t2_2(rows):
    h = first_row(rows, lambda r: text(r[0]) == "Year" and "Transportation" in text(r[1]), "Table 2.2's header")
    keys = {}
    for j, c in enumerate(rows[h]):
        s = text(c).lower()
        if s.startswith("percentage") or j == 0 or not s:
            continue
        keys[j] = SECTORS_22[label(c, SECTORS_22)]
    out = []
    for year, r in year_rows(rows, h + 1):
        for j, key in keys.items():
            v = num(r[j])
            if v is not None:
                out.append((f"t2_2:{key}", "energy_consumption", year, round(v * 1000, 6), "TBtu", text(rows[h][j])))
    return out


def t2_4(rows):
    h = first_row(rows, lambda r: text(r[0]) == "Year" and "Petroleum" in text(r[1]), "Table 2.4's header")
    keys = {j: SOURCES_24[label(c, SOURCES_24)] for j, c in enumerate(rows[h]) if j > 0 and text(c)}
    out = []
    for year, r in year_rows(rows, h + 2):
        for j, key in keys.items():
            v = num(r[j])
            if v is None:
                continue
            if key == "total":
                out.append(("t2_4:transportation", "energy_use", year, v, "TBtu", "Total"))
            else:
                out.append(("t2_4:transportation", f"share_{key}", year, round(v * 100, 6), "pct", text(rows[h][j])))
    return out


def mode_rows(rows, start, stop_words=("Source:",)):
    """(path, label, row) of the mode rows from start, the path from each label's level."""
    out, stack = [], {}
    for r in rows[start:]:
        raw = r[0] if r else None
        if raw is None or not text(raw):
            if out:
                break
            continue
        if text(raw) in stop_words:
            break
        key = label(raw, MODES)
        slug, level = MODES[key]
        indent = len(str(raw)) - len(str(raw).lstrip(" "))
        if level >= 2 and indent >= 6:
            level = 3
        stack[level] = slug
        for k in list(stack):
            if k > level:
                del stack[k]
        path = ".".join(stack[k] for k in sorted(stack) if k <= level)
        out.append((path, text(raw), r))
    return out


def t2_7(rows):
    h = first_row(rows, lambda r: text(r[1]) == "Gasoline", "Table 2.7's header")
    year_line = " ".join(text(c) for r in rows[:h] for c in r if c)
    years = re.findall(r"\b(20\d\d)\b", year_line)
    keys = {j: FUELS[label(c, FUELS)] for j, c in enumerate(rows[h]) if j > 0 and text(c)}
    out = []
    for path, lab, r in mode_rows(rows, h + 1):
        for j, fuel in keys.items():
            v = num(r[j])
            if v is not None:
                out.append((f"t2_7:{path}", f"energy_use_{fuel}", None, v, "TBtu", lab))
    return out, years


def t2_8(rows):
    h = first_row(rows, lambda r: len(r) > 2 and re.fullmatch(r"20\d\d", text(r[1])), "Table 2.8's years")
    ycols = [(j, int(text(c))) for j, c in enumerate(rows[h][:3]) if re.fullmatch(r"20\d\d", text(c))]
    out = []
    for path, lab, r in mode_rows(rows, h + 1):
        for j, year in ycols:
            v = num(r[j])
            if v is not None:
                out.append((f"t2_8:{path}", "energy_use", year, v, "TBtu", lab))
    return out


def t6_2(rows):
    h = first_row(rows, lambda r: text(r[0]) == "Calendar" and "Hybrid vehicle sales" in text(r[1]), "Table 6.2's header")
    heads = [text(c).lower() for c in rows[h][1:5]]
    want = ["hybrid vehicle sales", "plug-in hybrid vehicle sales", "all-electric vehicle sales", "all light-duty vehicle sales"]
    if not all(a.startswith(b) for a, b in zip(heads, want)):
        raise RuntimeError(f"Table 6.2's columns are not as expected: {heads}")
    out = []
    for year, r in year_rows(rows, h + 2):
        for j, var in enumerate(SALES_62, start=1):
            v = num(r[j])
            if v is not None:
                out.append(("t6_2:light_duty", var, year, round(v * 1000), "count", text(rows[h][j])))
    return out


def build(files, url, retrieved, vintage, edition):
    """files: {table number: workbook bytes}. Returns the series table."""
    got = []
    got += [r + ("Table 2.2",) for r in t2_2(rows_of(files["2_02"]))]
    got += [r + ("Table 2.4",) for r in t2_4(rows_of(files["2_04"]))]
    t27, years = t2_7(rows_of(files["2_07"]))
    t28 = t2_8(rows_of(files["2_08"]))
    latest = max(y for _, _, y, *_ in t28)
    if years and str(latest) not in years:
        raise RuntimeError(f"Table 2.7's year {years} is not Table 2.8's latest, {latest}")
    got += [(e, v, latest, x, u, n, "Table 2.7") for e, v, _, x, u, n in t27]
    got += [r + ("Table 2.8",) for r in t28]
    got += [r + ("Table 6.2",) for r in t6_2(rows_of(files["6_02"]))]
    t = pd.DataFrame([{"entity": f"ornl_tedb:{e}", "variable": v, "ts_utc": f"{y}-01-01T00:00:00Z", "value": float(x),
                       "unit": u, "freq": "P1Y", "geo": "US", "market": "", "node": n, "source": SOURCE, "source_url": url,
                       "retrieved_at": retrieved, "vintage": vintage, "x_table": tab, "x_edition": edition}
                      for e, v, y, x, u, n, tab in got], columns=COLS)
    twice = t[t.duplicated(["entity", "variable", "ts_utc"], keep=False)]
    if len(twice):
        raise RuntimeError(f"a key twice: {twice[['entity', 'variable', 'ts_utc']].head(3).values.tolist()}")
    return t.sort_values(["entity", "variable", "ts_utc"]).reset_index(drop=True)


def main(argv=None):
    ap = argparse.ArgumentParser(description="ORNL's Transportation Energy Data Book, five tables")
    ap.add_argument("--offline", action="store_true", help="the saved files only, no request")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        page, _ = ip.fetch_raw(CONNECTOR, PAGE, log, offline=args.offline, fresh=not args.offline, headers=UA, timeout=120)
        links = re.findall(r'href="([^"]*Spreadsheets_Ed_(\d+)\.zip)"', page.decode("utf-8", errors="replace"))
        if not links:
            raise RuntimeError(f"{PAGE} links no spreadsheets zip")
        rel, edition = max(links, key=lambda x: int(x[1]))
        url = urllib.parse.urljoin(PAGE, rel)
        content, rec = ip.fetch_raw(CONNECTOR, url, log, offline=args.offline, headers=UA, timeout=300)
        z = zipfile.ZipFile(io.BytesIO(content))
        files = {}
        for key in ("2_02", "2_04", "2_07", "2_08", "6_02"):
            names = [n for n in z.namelist() if re.fullmatch(rf"Table{key[0]}_{key[2:]}_\d{{4}}-\d{{2}}\.xlsx", n)]
            if len(names) != 1:
                raise RuntimeError(f"Table {key}: {len(names)} files in the zip")
            files[key] = z.read(names[0])
        t = build(files, url, rec.get("retrieved_at") or "", ip.vintage_of(rec), f"Edition {edition}")
        by = t.groupby("x_table").size()
        header = [
            f"Energy Research Warehouse (ERW): Oak Ridge National Laboratory, Transportation Energy Data Book, Edition {edition}: Tables 2.2, 2.4, 2.7, 2.8 and 6.2 (session 2048)",
            "Shape: series (docs/datastandard.md v0). entity ornl_tedb:<table>:<row> (a mode's path, e.g. t2_7:highway.light_duty.cars); variables energy_consumption "
            "(2.2, TEDB's quadrillion Btu x 1,000 as TBtu), energy_use and share_<source> (2.4; a share in percent, TEDB's fraction x 100), energy_use_<fuel> (2.7), "
            "energy_use (2.8), sales_* (6.2, TEDB's thousands x 1,000 as a count); ts_utc the year's first day, freq P1Y; node the row's label as printed.",
            f"{len(t)} rows: " + "; ".join(f"{k} {v}" for k, v in by.items()) + ". The per-cent change rows and the shares that are ratios of rows held are not taken. "
            "Values as TEDB publishes them, but for the stated conversions; nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/ornl_tedb.py; the zip retrieved {rec.get('retrieved_at') or ''}, last modified {ip.vintage_of(rec)}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} Davis, S. C., Dahal, P. M., Moore, A. M., Patwary, L., Boundy, R. G., Transportation Energy Data Book: Edition {edition}, Oak Ridge National "
            "Laboratory, https://tedb.ornl.gov ; Table 6.2 compiled by the Transportation Research Center at Argonne National Laboratory",
            f"  access: {url}",
            "License: no terms of reuse stated on the book's pages (session 2048); held internal until a ruling.",
        ]
        ip.write_csv(t, NAME, header, log, cols=COLS)
        if not args.out_dir:
            ip.update_sources([dict(source=SOURCE, publisher="Oak Ridge National Laboratory (ORNL), Transportation Energy Data Book",
                                    report=f"Transportation Energy Data Book, Edition {edition}, Tables 2.2, 2.4, 2.7, 2.8 and 6.2",
                                    report_url="https://tedb.ornl.gov/", document_list=url, license="internal", tables=[NAME])])
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, Edition {edition} ({by.to_dict()})")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"FAILED, no output file written:\n{tb}")
        print(f"{CONNECTOR} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
