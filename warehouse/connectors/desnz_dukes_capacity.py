#!/usr/bin/env python3
"""The United Kingdom's de-rated generating capacity at each year's end, 1996 to 2025, by kind of company and fuel, and its
thermal stations by technology, as the Department for Energy Security and Net Zero (DESNZ) gives them in the Digest of UK Energy
Statistics (DUKES), table 5.7 (session 2734, chain 137, the session's third own choice; the companion of chain 136's DUKES 5.11
station list, which holds installed capacity, and of the ERW's DESNZ regional renewable capacity).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/desnz_uk_derated_capacity_annual.csv, from one
workbook (no key), its sheet "5.7", two tables one above another:

    https://assets.publishing.service.gov.uk/media/6a6a3622cceb23e8678976bf/DUKES_5.7.xlsx

    entity      desnz_cap:<group>:<item>: group mpp (major power producers), other (other generators: autogenerators, domestic
                generation and smaller renewable generators), all (all generating companies); item the row's fuel or technology
                as a slug (gas_fired, nuclear_stations, onshore_wind, ccgt_stations, ...), or total, total_renewables,
                total_low_carbon, total_fossil_fuels, total_thermal
    variable    derated_capacity (table 5.7.A, capacity by fuel) or thermal_capacity (table 5.7.B, thermal stations by type)
    unit        MW; ts_utc the year's first day; freq P1Y; geo GB; node the sheet's two labels as written

DESNZ's notes: grid export capacity where available, else installed capacity (note 1); small-scale hydro, wind and solar are on
a declared net capability basis, de-rated for intermittency by 0.365, 0.43 and 0.17 (note 4), so these figures are not the
installed capacity of DUKES 5.11; other generators' data before 2006 are declared net capacity (note 6), and their falls in
2007, 2010 and 2012 are capacity re-classified to the major power producers (note 7); conventional steam stations leave out
bioenergy and waste (note 8); renewables are wind, wave and tidal, solar, bioenergy and natural flow hydro, not pumped hydro
(note 9); low carbon is renewables and nuclear (note 10); fossil fuels are coal, oil, gas, mixed or dual fired and other fossil
fuels (note 11).

Each table is found by its title and its year row (each year from 1996 in turn); each row's two labels (note marks read off) are
mapped to an entity, and a group or a repeated item stops the run. A blank cell has no row; a zero stays a zero. Checks, within
0.5 MW in every year (but DESNZ's own 99 MW in all companies' total and fossil fuels total of 2015 and 2016, named in KNOWN
and kept as published): each group's parts add to its total; the major power producers and the other generators add to all
companies, item by item; all companies' renewables, low carbon and fossil totals equal their parts, and fossil, low carbon and
pumped hydro make the total. A check missed stops the run. Nothing is filled.

    python warehouse/connectors/desnz_dukes_capacity.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/desnz_dukes/ and reused unless --fresh; --offline asks nothing.
www.gov.uk's robots file closes only print and search paths; assets.publishing.service.gov.uk has no robots file (session 2678).
License: the Open Government Licence v3.0 (DESNZ's verdict on record, session 2167): public.
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "desnz_uk_derated_capacity_annual"
CONNECTOR = "desnz_dukes"
SOURCE = "desnz:dukes_5_7_plant_capacity"
HOST = "assets.publishing.service.gov.uk"
URL = "https://assets.publishing.service.gov.uk/media/6a6a3622cceb23e8678976bf/DUKES_5.7.xlsx"
PAGE = "https://www.gov.uk/government/statistics/electricity-chapter-5-digest-of-united-kingdom-energy-statistics-dukes"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = "Department for Energy Security and Net Zero, Digest of UK Energy Statistics (DUKES) 2026, table 5.7, Plant capacity: United Kingdom"
TABLES = [("Table 5.7.A Capacity by fuel, MW", "derated_capacity"), ("Table 5.7.B Capacity of Thermal Stations by Type, MW", "thermal_capacity")]
GROUPS = {"Major power producers": "mpp", "Other generators": "other", "All generating companies": "all"}
TOTALS = {"Total transmission entry capacity": "total", "Total capacity of own generating plant": "total", "Total capacity": "total",
          "Total renewables capacity": "total_renewables", "Total low carbon capacity": "total_low_carbon",
          "Total fossil fuels capacity": "total_fossil_fuels", "Total capacity of MPP thermal stations": "total",
          "Total capacity of other thermal stations": "total", "Total capacity of all thermal stations": "total"}
ALIASES = {"Mixed or dual fuelled": "Mixed or dual fired"}
FIRST_YEAR = 1996
TOL = 0.5
# DESNZ's sheet itself: in 2015 and 2016 all companies' total and fossil fuels total are each 99 MW above their parts (and the
# total above the major power producers' and the others' totals), every fuel row adding exactly; kept as published, named here,
# and any other difference still stops the run
KNOWN = {(2015, "all:total"): 99.0, (2015, "all:total_fossil_fuels"): 99.0, (2016, "all:total"): 99.0,
         (2016, "all:total_fossil_fuels"): 99.0}
X_COLS = []


def label(v):
    s = " ".join(str(v).replace(chr(0xA0), " ").split()) if v is not None else ""
    return re.sub(r"\s*\[note \d+\]", "", s).strip()


def slug(s):
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", s.lower())).strip("_")


def item_of(name):
    name = ALIASES.get(name, name)
    return TOTALS.get(name, slug(name))


def parse_sheet(rows):
    titles = {label(r[0]): i for i, r in enumerate(rows) if r and r[0] is not None}
    out, nodes = {}, {}
    for title, var in TABLES:
        if title not in titles:
            raise RuntimeError(f"no table titled {title!r}")
        i = titles[title]
        head = [label(x) for x in rows[i + 1]]
        if head[:2] not in (["Generator Type", "Fuel"], ["Generator Type", "Technology"]):
            raise RuntimeError(f"{title}: the header is {head[:2]}")
        years = head[2:]
        while years and not years[-1]:
            years.pop()
        if not years or years != [str(FIRST_YEAR + k) for k in range(len(years))]:
            raise RuntimeError(f"{title}: the years are not each year from {FIRST_YEAR}")
        for r in rows[i + 2:]:
            g = label(r[0]) if r else ""
            if not g or g.startswith("Table 5.7"):
                break
            if g not in GROUPS:
                raise RuntimeError(f"{title}: a group not in the map: {g!r}")
            name = label(r[1])
            ent = f"{GROUPS[g]}:{item_of(name)}"
            if (ent, var) in nodes:
                raise RuntimeError(f"{title}: {ent} given twice")
            nodes[(ent, var)] = f"{g}: {name}"
            for k, y in enumerate(years):
                v = r[2 + k] if 2 + k < len(r) else None
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                if not isinstance(v, (int, float)) or isinstance(v, bool) or v < 0:
                    raise RuntimeError(f"{title} {g} {name} {y}: not a quantity: {v!r}")
                out[(int(y), ent, var)] = float(v)
    check(out)
    return out, nodes


def miss(y, ent, have, want):
    """True for a difference the checks must stop on: beyond TOL and not DESNZ's known one for this year and row."""
    d = want - have
    return abs(d) > TOL and not (abs(d - KNOWN.get((y, ent), float("nan"))) <= 0.05)


def check(out):
    years = sorted({k[0] for k in out})
    for var in {k[2] for k in out}:
        ents = {k[1] for k in out if k[2] == var}
        for y in years:
            g = lambda e: out.get((y, e, var))
            for grp in ("mpp", "other", "all"):
                parts = [e for e in ents if e.startswith(grp + ":") and not e.split(":")[1].startswith("total")]
                tot = g(f"{grp}:total")
                vals = [g(e) for e in parts]
                if tot is not None and vals and None not in vals and miss(y, f"{grp}:total", sum(vals), tot):
                    raise RuntimeError(f"{y} {var} {grp}: the parts make {sum(vals):,.2f} MW, the total {tot:,.2f}")
            for e in ents:
                if e.startswith("all:"):
                    a, m, o = g(e), g("mpp:" + e[4:]), g("other:" + e[4:])
                    if a is not None and (m is not None or o is not None) and not e[4:].startswith("total_"):
                        s = (m or 0.0) + (o or 0.0)
                        if miss(y, e, s, a):
                            raise RuntimeError(f"{y} {var} {e}: the major power producers and the others make {s:,.2f} MW, all {a:,.2f}")
            if var == "derated_capacity":
                ren = ["onshore_wind", "offshore_wind", "wave_and_tidal", "solar", "bioenergy_and_waste", "hydro_natural_flow"]
                fos = ["coal_fired", "oil_fired", "gas_fired", "mixed_or_dual_fired", "other_fossil_fuels"]
                sums = {"total_renewables": [g("all:" + x) for x in ren],
                        "total_low_carbon": [g("all:total_renewables"), g("all:nuclear_stations")],
                        "total_fossil_fuels": [g("all:" + x) for x in fos],
                        "total": [g("all:total_fossil_fuels"), g("all:total_low_carbon"), g("all:pumped_hydro")]}
                for k, vals in sums.items():
                    tot = g("all:" + k)
                    if tot is not None and None not in vals and miss(y, "all:" + k, sum(vals), tot):
                        raise RuntimeError(f"{y} all:{k}: its parts make {sum(vals):,.2f} MW, the sheet {tot:,.2f}")


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    if "5.7" not in wb.sheetnames:
        raise RuntimeError(f"the workbook has no sheet 5.7: {wb.sheetnames}")
    vals, nodes = parse_sheet([tuple(r) for r in wb["5.7"].iter_rows(values_only=True)])
    rows = [dict(entity=f"desnz_cap:{e}", variable=v, ts_utc=f"{y}-01-01T00:00:00Z", value=round(x, 4), unit="MW", freq="P1Y",
                 geo="GB", market="", node=nodes[(e, v)], **base) for (y, e, v), x in sorted(vals.items())]
    return pd.DataFrame(rows)


def main(argv=None):
    ap = argparse.ArgumentParser(description="DESNZ DUKES 5.7: the UK's de-rated plant capacity a year")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t = parse(content, base)
        if t.duplicated(["entity", "variable", "ts_utc"]).any():
            raise RuntimeError("a key given twice")
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["entity", "variable", "ts_utc"]).reset_index(drop=True)
        tot = t[(t.entity == "desnz_cap:all:total") & (t.variable == "derated_capacity")].set_index("ts_utc").value
        head = [
            "Energy Research Warehouse (ERW): DESNZ, Digest of UK Energy Statistics (DUKES) 2026, table 5.7: the UK's de-rated "
            f"generating capacity by kind of company and fuel, and its thermal stations by type, {t.ts_utc.min()[:4]} to "
            f"{t.ts_utc.max()[:4]} (session 2734)",
            "Shape: series (docs/datastandard.md v0). entity desnz_cap:<mpp|other|all>:<item>; variable derated_capacity (5.7.A) or "
            "thermal_capacity (5.7.B); unit MW (wind, solar and small hydro de-rated by 0.43, 0.17 and 0.365, DESNZ's note 4); "
            "ts_utc the year's first day; freq P1Y; geo GB; node the sheet's labels.",
            f"{len(t)} rows; all companies' de-rated capacity, MW: " + ", ".join(f"{k[:4]} {v:,.0f}" for k, v in tot.items()
                                                                                if k[:4] in ("1996", "2010", "2020", "2025"))
            + ". Every total checked against its parts and the groups against all companies; nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/desnz_dukes_capacity.py; the workbook retrieved "
            f"{base['retrieved_at']}; last modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL} (listed on {PAGE})",
            "License: Open Government Licence v3.0. Credit: Department for Energy Security and Net Zero.",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="UK Department for Energy Security and Net Zero (DESNZ)",
                                    report="Digest of UK Energy Statistics (DUKES) 2026, table 5.7, Plant capacity: United Kingdom",
                                    report_url=PAGE, document_list=PAGE, license="public", tables=[NAME])])  # OGL v3.0
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
