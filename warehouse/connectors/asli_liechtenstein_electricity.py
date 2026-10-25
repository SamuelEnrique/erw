#!/usr/bin/env python3
"""Liechtenstein's electricity a year since 1985: domestic production, imports, exports and domestic consumption, as the Amt fuer
Statistik Liechtenstein gives them in Table 5 ("Stromproduktion und -verbrauch seit 1985") of its Energie 2025 tables, published
1 July 2026 under CC BY 4.0 (session 2622, chain 130, the session's own choice; the ERW held Liechtenstein in no table).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/asli_liechtenstein_electricity_annual.csv, from
one workbook (no key):

    https://www.statistikportal.li/statistikportal/publications/551-energie/2025/01/1/551.2025.01.1_02_energie-2025-tabellen.xlsx

    entity      asli:li:production (variable generation), asli:li:imports (imports), asli:li:exports (exports),
                asli:li:consumption (consumption)
    unit        MWh; ts_utc the year's first day; freq P1Y; geo LI
    node        the table's column heading
    x_pv_method old (to 2021) or new (from 2022): the Office changed how it counts photovoltaic production in 2023 and gives
                2022 by both methods, on two rows
    x_old_method_value  on 2022's rows only, the figure by the old method

The sheet "Tabelle 5" is read by its title cell ("Stromproduktion und -verbrauch seit 1985") and its two header rows (the
columns, then "Jahr" and four "MWh"); every later row with a year in the first column is a year, to the first row without one.
The years run from 1985 without a gap; only 2022 may be given twice. Of its two rows the first is taken as the old method and
the second as the new, the order of the table's note ("nach der alten und neuen Methode"); the new is written, as 2023 on are
counted by it, and the old kept in x_old_method_value. Check: production plus imports less exports equals consumption within
1 MWh each year (both 2022 rows). A cell "-" is the workbook's zero and "." not available (not written). Nothing is filled.

    python warehouse/connectors/asli_liechtenstein_electricity.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/asli_li/ and reused unless --fresh; --offline asks nothing.
www.statistikportal.li's robots file names * with no rule (session 2622). License: public (the workbook's Nutzungsbedingungen:
CC BY 4.0). Cite: Amt fuer Statistik Liechtenstein, Energie 2025, Tabelle 5.
"""

import argparse
import datetime as dt
import io
import os
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "asli_liechtenstein_electricity_annual"
CONNECTOR = "asli_li"
SOURCE = "asli_li:energie_2025_tabelle5"
HOST = "www.statistikportal.li"
URL = ("https://www.statistikportal.li/statistikportal/publications/551-energie/2025/01/1/"
       "551.2025.01.1_02_energie-2025-tabellen.xlsx")
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = "Amt fuer Statistik Liechtenstein, Energie 2025, Tabelle 5: Stromproduktion und -verbrauch seit 1985 (CC BY 4.0)"
TITLE = "Stromproduktion und -verbrauch seit 1985"
COLS = [("Produktion im Inland", "production", "generation"), ("Import", "imports", "imports"), ("Export", "exports", "exports"),
        ("Verbrauch im Inland", "consumption", "consumption")]
X_COLS = ["x_pv_method", "x_old_method_value"]


def cell(v, where):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).strip() if v is not None else ""
    if s == "-":
        return 0.0
    if s == ".":
        return None
    raise RuntimeError(f"{where}: a cell not in the key: {v!r}")


def parse_rows(rows):
    at = [k for k, r in enumerate(rows) if r and r[0] is not None and str(r[0]).strip() == TITLE]
    if len(at) != 1:
        raise RuntimeError(f"the title found {len(at)} times")
    hd = [k for k, r in enumerate(rows) if r and r[0] is None and r[1] is not None and str(r[1]).strip() == COLS[0][0]]
    if len(hd) != 1:
        raise RuntimeError("the header row not found once")
    h = [str(v).strip() if v is not None else "" for v in rows[hd[0]][1:5]]
    if h != [c[0] for c in COLS]:
        raise RuntimeError(f"the columns are {h}")
    u = [str(v).strip() if v is not None else "" for v in rows[hd[0] + 1][:5]]
    if u != ["Jahr", "MWh", "MWh", "MWh", "MWh"]:
        raise RuntimeError(f"the unit row is {u}")
    years = []
    for r in rows[hd[0] + 2:]:
        if not isinstance(r[0], (int, float)) or isinstance(r[0], bool):
            break
        years.append((int(r[0]), [cell(r[j + 1], f"{int(r[0])} {COLS[j][0]}") for j in range(4)]))
    ys = [y for y, _ in years]
    if not ys or ys[0] != 1985:
        raise RuntimeError("the years do not begin in 1985")
    if sorted(set(ys)) != list(range(1985, ys[-1] + 1)):
        raise RuntimeError("the years have a gap")
    twice = sorted({y for y in ys if ys.count(y) > 1})
    if twice not in ([], [2022]) or ys.count(2022) > 2:
        raise RuntimeError(f"years given twice: {twice}")
    for y, v in years:
        if None not in v and abs(v[0] + v[1] - v[2] - v[3]) > 1:
            raise RuntimeError(f"{y}: production, imports and exports make {v[0] + v[1] - v[2]:.1f} MWh, consumption {v[3]}")
    return years


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    if "Tabelle 5" not in wb.sheetnames:
        raise RuntimeError(f"no sheet Tabelle 5: {wb.sheetnames}")
    years = parse_rows([tuple(r) for r in wb["Tabelle 5"].iter_rows(values_only=True)])
    by = {}
    for y, v in years:
        by.setdefault(y, []).append(v)
    out = []
    for y, vs in sorted(by.items()):
        v, old = (vs[1], vs[0]) if len(vs) == 2 else (vs[0], None)
        for j, (label, item, variable) in enumerate(COLS):
            if v[j] is None:
                continue
            out.append(dict(entity=f"asli:li:{item}", variable=variable, ts_utc=f"{y}-01-01T00:00:00Z", value=round(v[j], 3), unit="MWh",
                            freq="P1Y", geo="LI", market="", node=label, x_pv_method="new" if y >= 2022 else "old",
                            x_old_method_value="" if old is None or old[j] is None else round(old[j], 3), **base))
    return pd.DataFrame(out)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Amt fuer Statistik: Liechtenstein's electricity production, trade and consumption a year")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t = parse(content, base)
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["entity", "ts_utc"]).reset_index(drop=True)
        last = t.ts_utc.max()[:4]
        cons = t[(t.entity == "asli:li:consumption") & (t.ts_utc.str[:4] == last)].value.iloc[0]
        imp = t[(t.entity == "asli:li:imports") & (t.ts_utc.str[:4] == last)].value.iloc[0]
        head = [
            "Energy Research Warehouse (ERW): Amt fuer Statistik Liechtenstein, Energie 2025, Tabelle 5: Liechtenstein's electricity "
            "production, imports, exports and consumption a year since 1985 (session 2622)",
            "Shape: series (docs/datastandard.md v0). entity asli:li:<item> (production, imports, exports, consumption); variable "
            "generation, imports, exports or consumption; unit MWh; ts_utc the year's first day; freq P1Y; geo LI; node the table's "
            "column; x_pv_method old (to 2021) or new (from 2022); x_old_method_value 2022's figure by the old method.",
            f"{len(t)} rows; in {last} consumption {cons:,.0f} MWh, {100 * imp / cons:.1f}% of it imported. Production plus imports less "
            "exports checked against consumption each year; nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/asli_liechtenstein_electricity.py; the workbook retrieved "
            f"{base['retrieved_at']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL}",
            f"License: public, CC BY 4.0 (the workbook's Nutzungsbedingungen). Cite: {CITATION}.",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="Amt fuer Statistik Liechtenstein",
                                    report="Energie 2025, Tabelle 5: Stromproduktion und -verbrauch seit 1985", report_url=URL,
                                    document_list="https://www.statistikportal.li/de/themen/raum-umwelt-und-energie/energie",
                                    license="public", tables=[NAME])])
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
