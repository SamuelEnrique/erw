#!/usr/bin/env python3
"""The UK's renewable electricity capacity and generation by technology, a quarter, from the Department for Energy
Security and Net Zero's Energy Trends table 6.1 (session 2174, chain 094: the UK beside Eurostat's renewable capacity,
which leaves the UK out, and beside the UK's generation by fuel of desnz_uk_electricity_quarterly; a third source of the
session's own choosing).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/desnz_uk_renewables_quarterly.csv, from
the Energy Trends section 6 page's table ET 6.1 (an xlsx on assets.publishing.service.gov.uk), its sheet "Quarter":

    https://www.gov.uk/government/statistics/energy-trends-section-6-renewables     (the page, for the file's address)

The sheet holds four blocks one under another; two are taken, each row a technology:

    CUMULATIVE INSTALLED CAPACITY (MW)   -> variable capacity, MW, at the end of the quarter (DESNZ's note 1)
    ELECTRICITY GENERATED (GWh)          -> variable generated, MWh (DESNZ's GWh times 1,000); actual figures where
                                            DESNZ has them, otherwise estimated from a typical load factor (its note 8,
                                            the feed-in tariff schemes among them)

The load factors and the shares are DESNZ's ratios of these and are not taken.

    entity      desnz_et61:<technology>, in snake case as the sheet names it, its note markers left out (onshore_wind,
                offshore_wind_seabed, solar_photovoltaics, plant_biomass, total ...); node DESNZ's words. TOTAL and the
                rows DESNZ prints beside it (co-firing, non-biodegradable wastes) are DESNZ's own: a sum over the
                technologies must leave all three out (session 2214). The two blocks name some technologies
                differently (offshore wind seabed and floating, large and small scale hydro, co-firing in capacity;
                offshore wind, hydro, co-firing with fossil fuels in generation): a join by entity misses them
    ts_utc      the quarter's first day at 00:00:00Z; freq P3M; geo GB (the United Kingdom); x_provisional yes for a
                quarter DESNZ marks provisional

"[x]" (no figure) has no row; nothing is filled. A sheet without the two blocks, a column header not "<year> <n>th
quarter", a value that is neither a number nor "[x]", or a key given twice stops the run.

    python warehouse/connectors/desnz_renewables.py [--offline] [--out-dir DIR]

Ceiling: two requests (the page and the file it links), reused from warehouse/raw/desnz_renewables/ with --offline.
License: the Open Government Licence v3.0 (GOV.UK's licence for its content, record of session 2167); attribution:
"Contains public sector information licensed under the Open Government Licence v3.0".
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "desnz_uk_renewables_quarterly"
CONNECTOR = "desnz_renewables"
SOURCE = "desnz:energy_trends_6_1"
PAGE = "https://www.gov.uk/government/statistics/energy-trends-section-6-renewables"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
BLOCKS = {"CUMULATIVE INSTALLED CAPACITY (MW)": ("capacity", "MW", 1.0),
          "ELECTRICITY GENERATED (GWh)": ("generated", "MWh", 1000.0)}
QUARTER = re.compile(r"^(\d{4}) (1st|2nd|3rd|4th) quarter( \[provisional\])?$")
ORD = {"1st": 1, "2nd": 2, "3rd": 3, "4th": 4}
COLS = ["entity", "variable", "ts_utc", "value", "unit", "freq", "geo", "market", "node", "source", "source_url",
        "retrieved_at", "vintage", "x_technology", "x_provisional"]


def text(v):
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def label(v):
    return re.sub(r"\s*\[note[^\]]*\]", "", text(v), flags=re.I).strip()


def slug(s):
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


def block_of(title):
    t = label(title)
    for k in BLOCKS:
        if t.upper().startswith(k.upper()):
            return k
    return None


def build(workbook, url, retrieved, vintage):
    wb = openpyxl.load_workbook(io.BytesIO(workbook), read_only=True, data_only=True)
    if "Quarter" not in wb.sheetnames:
        raise RuntimeError(f"no sheet 'Quarter': {wb.sheetnames}")
    rows = [list(r) for r in wb["Quarter"].iter_rows(values_only=True)]
    out, found = [], set()
    i = 0
    while i < len(rows):
        k = block_of(rows[i][0]) if rows[i] and rows[i][0] else None
        if k:
            var, unit, factor = BLOCKS[k]
            found.add(k)
            cols = []
            for j, h in enumerate(rows[i][1:], start=1):
                h = text(h)
                if not h:
                    continue
                m = QUARTER.fullmatch(h)
                if not m:
                    raise RuntimeError(f"{k}: a column header not '<year> <n>th quarter': {h!r}")
                y, q = int(m.group(1)), ORD[m.group(2)]
                cols.append((j, f"{y}-{3 * q - 2:02d}-01T00:00:00Z", "yes" if m.group(3) else "no"))
            i += 1
            while i < len(rows) and rows[i] and text(rows[i][0]):
                tech = label(rows[i][0])
                for j, ts, prov in cols:
                    v = rows[i][j] if j < len(rows[i]) else None
                    if v is None or text(v) == "[x]":
                        continue
                    if not isinstance(v, (int, float)):
                        raise RuntimeError(f"{k}: {tech}, {ts[:7]}: not a number: {v!r}")
                    out.append({"entity": f"desnz_et61:{slug(tech)}", "variable": var, "ts_utc": ts,
                                "value": round(float(v) * factor, 3), "unit": unit, "freq": "P3M", "geo": "GB", "market": "",
                                "node": tech, "source": SOURCE, "source_url": url, "retrieved_at": retrieved,
                                "vintage": vintage, "x_technology": tech, "x_provisional": prov})
                i += 1
            continue
        i += 1
    missing = set(BLOCKS) - found
    if missing:
        raise RuntimeError(f"the sheet lacks {sorted(missing)}")
    t = pd.DataFrame(out, columns=COLS)
    twice = t[t.duplicated(["entity", "variable", "ts_utc"], keep=False)]
    if len(twice):
        raise RuntimeError(f"a key given twice: {twice.entity.iloc[0]} {twice.variable.iloc[0]} {twice.ts_utc.iloc[0]}")
    return t.sort_values(["variable", "entity", "ts_utc"]).reset_index(drop=True)


def main(argv=None):
    ap = argparse.ArgumentParser(description="DESNZ Energy Trends 6.1: UK renewable capacity and generation, a quarter")
    ap.add_argument("--offline", action="store_true", help="the saved files only, no request")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        if not args.offline and (ip.paused("desnz") or ip.paused_host("www.gov.uk") or ip.paused_host("assets.publishing.service.gov.uk")):
            raise RuntimeError("DESNZ is paused (warehouse/metadata/paused_sources.csv): not requested")
        page, _ = ip.fetch_raw(CONNECTOR, PAGE, log, offline=args.offline, fresh=not args.offline, headers=UA, timeout=120)
        links = re.findall(r'href="(https://assets\.publishing\.service\.gov\.uk/[^"]*/ET_6\.1_[^"]*\.xlsx)"', page.decode("utf-8", errors="replace"))
        if not links:
            raise RuntimeError(f"{PAGE} links no ET 6.1 workbook")
        url = links[0]
        content, rec = ip.fetch_raw(CONNECTOR, url, log, offline=args.offline, headers=UA, timeout=300)
        t = build(content, url, rec.get("retrieved_at") or "", ip.vintage_of(rec))
        cap = t[(t.entity == "desnz_et61:total") & (t.variable == "capacity")]
        header = [
            f"Energy Research Warehouse (ERW): UK Department for Energy Security and Net Zero (DESNZ), Energy Trends table 6.1, renewable electricity capacity and generation by technology, a quarter, {t.ts_utc.min()[:7]} to {t.ts_utc.max()[:7]} (session 2174)",
            "Shape: series (docs/datastandard.md v0). entity desnz_et61:<technology>; variable capacity (MW, cumulative at the quarter's end) or generated (MWh, DESNZ's GWh times 1,000, estimated where no actual figure exists); "
            "ts_utc the quarter's first day; freq P3M; geo GB. TOTAL and the rows beside it are DESNZ's own: a sum over the technologies must leave out TOTAL, co_firing (capacity) and non_biodegradable_wastes (generation). "
            "Capacity and generation name offshore wind, hydro and co-firing differently: a join by entity misses them.",
            f"{len(t)} rows, {t.entity.nunique()} technology rows over both blocks (capacity {t[t.variable == 'capacity'].entity.nunique()}, generated {t[t.variable == 'generated'].entity.nunique()}); renewable capacity at the end of {cap.ts_utc.max()[:7]}: {cap.value.iloc[-1]:,.0f} MW. "
            f"{(t.x_provisional == 'yes').sum()} rows of quarters DESNZ marks provisional. '[x]' (no figure) has no row; nothing is filled. Load factors and shares are not taken.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/desnz_renewables.py; the workbook retrieved {rec.get('retrieved_at') or ''}"
            + (f", last modified {ip.vintage_of(rec)}" if ip.vintage_of(rec) else ""),
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} UK Department for Energy Security and Net Zero, Energy Trends section 6: renewables, table ET 6.1, {PAGE}",
            f"  access: {url}",
            "License: public. Contains public sector information licensed under the Open Government Licence v3.0 (https://www.nationalarchives.gov.uk/doc/open-government-licence/version/3/).",
        ]
        ip.write_csv(t, NAME, header, log, cols=COLS)
        if not args.out_dir:
            ip.update_sources([dict(source=SOURCE, publisher="UK Department for Energy Security and Net Zero (DESNZ)",
                                    report="Energy Trends section 6: renewables, table ET 6.1 (renewable electricity capacity and generation)",
                                    report_url=PAGE, document_list=url, license="public", tables=[NAME])])
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.ts_utc.min()[:7]} to {t.ts_utc.max()[:7]}")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"FAILED, no output file written:\n{tb}")
        print(f"{CONNECTOR} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
