#!/usr/bin/env python3
"""Natural resource layers for "Where the resources are" (session 146).

What this connector does
  --pull    downloads each publisher's own file into the raw store, under one ceiling for all of them
            (4 GB). The size is asked first; a download that would pass the ceiling is refused before a
            byte of it is read. Every file, and every terms page, is one row of downloads.csv.
  --build   reduces each raw file to a web-sized layer in site/data/resources/layers/ and describes it in
            site/data/resources/manifest.json (the contract with the page). The raw file keeps the full
            resolution; nothing is filled, smoothed or interpolated.
  --only    one layer id (or several, comma separated).
  --out-dir a trial folder for the web layers and the manifest, instead of the site's.

Self-contained: no import from another connector. The geospatial libraries it needs for --build are in
resource_layers.requirements.txt (not in requirements.txt: the daily runner does not build these layers).

Run it on downloaded files with python -I (the archives are untrusted data).
The only contact string this connector ever sends is its User-Agent, below. No e-mail address.
"""
import argparse
import base64
import csv
import datetime as dt
import hashlib
import json
import math
import os
import re
import sys
import zipfile

UA = "ERW research project, github.com/SamuelEnrique/erw"
CEILING_BYTES = 4 * 1024 ** 3
HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
RAW_DIR = os.environ.get("ERW_RESOURCES_RAW") or os.path.join(REPO, "warehouse", "raw", "resources")
WEB_DIR = os.path.join(REPO, "site", "data", "resources", "layers")  # not under public/: a locked page's files are not open to visitors (the page reads them through /resources/layer/<name>)
MANIFEST = os.path.join(REPO, "site", "data", "resources", "manifest.json")
PAUSED_FILE = os.path.join(REPO, "warehouse", "metadata", "paused_sources.csv")
LEDGER_COLS = ["source", "layer", "url", "file", "bytes", "sha256", "retrieved_at_utc", "http_status",
               "terms_url", "terms_file", "terms_sha256"]
# never requested, whatever a source list says (CLAUDE.md, COMMON.md rule 6)
FORBIDDEN = re.compile(r"(?i)misoenergy\.org|pjm\.com")
NODATA = 65535
LEVELS = [0.2, 0.1, 0.05, 0.025]


def now_utc():
    return dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def log(msg):
    print(f"[resource_layers {now_utc()}] {msg}", flush=True)


# ---------------------------------------------------------------------------------------------------------
# The raw store: the ledger, the ceiling, the downloads
# ---------------------------------------------------------------------------------------------------------

class CeilingRefused(Exception):
    """A download that would pass the ceiling. Raised before the download starts."""


class Paused(Exception):
    """A publisher that is never requested."""


def refuse_paused(url):
    if FORBIDDEN.search(url):
        raise Paused(f"never requested: {url}")
    if os.path.exists(PAUSED_FILE):
        with open(PAUSED_FILE, newline="", encoding="utf-8") as f:
            for r in csv.DictReader(f):
                if r.get("match") and re.search(r["match"], url):
                    raise Paused(f"paused publisher ({r.get('scope')}): {url}")


def ledger_path(raw_dir):
    return os.path.join(raw_dir, "downloads.csv")


def read_ledger(raw_dir):
    p = ledger_path(raw_dir)
    if not os.path.exists(p):
        return []
    with open(p, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def ledger_total(rows):
    """Every byte read from a publisher counts, a terms page and an aborted read included."""
    return sum(int(r["bytes"] or 0) for r in rows)


def append_ledger(raw_dir, row):
    os.makedirs(raw_dir, exist_ok=True)
    p = ledger_path(raw_dir)
    new = not os.path.exists(p)
    with open(p, "a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=LEDGER_COLS)
        if new:
            w.writeheader()
        w.writerow({c: row.get(c, "") for c in LEDGER_COLS})


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def http_session():
    import requests
    s = requests.Session()
    s.headers.clear()
    s.headers.update({"User-Agent": UA, "Accept": "*/*", "Accept-Encoding": "identity"})
    return s


def ask_size(sess, url):
    """The size the server states for url, asked without reading the file: (bytes or None, status)."""
    r = sess.head(url, allow_redirects=True, timeout=60)
    n = r.headers.get("Content-Length", "")
    if r.status_code == 200 and n.isdigit() and int(n) > 0:
        return int(n), r.status_code
    r = sess.get(url, headers={"Range": "bytes=0-0"}, stream=True, allow_redirects=True, timeout=60)
    try:
        m = re.match(r"bytes \d+-\d+/(\d+)", r.headers.get("Content-Range", ""))
        if m:
            return int(m.group(1)), 200
        n = r.headers.get("Content-Length", "")
        if r.status_code == 200 and n.isdigit() and int(n) > 0:
            return int(n), r.status_code
        return None, r.status_code
    finally:
        r.close()


def fetch(sess, raw_dir, source, layer, url, file, terms=None, ceiling=CEILING_BYTES):
    """Download url to <raw_dir>/<source>/<file> under the ceiling. Returns the path.

    The size is asked first and the download is refused (CeilingRefused) when the bytes already in the
    ledger plus this file would pass the ceiling. A server that states no size is read with the remaining
    budget as a hard stop: the read ends before the byte that would pass the ceiling.
    """
    refuse_paused(url)
    rows = read_ledger(raw_dir)
    dest = os.path.join(raw_dir, source, file)
    held = [r for r in rows if r["source"] == source and r["file"] == file and r["http_status"] == "200"]
    if held and os.path.exists(dest) and sha256_file(dest) == held[-1]["sha256"]:
        log(f"held: {source}/{file} ({held[-1]['bytes']} bytes)")
        return dest
    spent = ledger_total(rows)
    size, status = ask_size(sess, url)
    if size is not None and spent + size > ceiling:
        raise CeilingRefused(f"{source}/{file}: {size} bytes asked, {spent} already read, ceiling {ceiling}")
    if size is None and status not in (200, 206):
        raise RuntimeError(f"{source}/{file}: HTTP {status} from {url}")
    budget = ceiling - spent
    if budget <= 0:
        raise CeilingRefused(f"{source}/{file}: nothing left under the ceiling ({spent} of {ceiling})")
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    if os.path.exists(dest):  # never overwrite a raw file: a changed file is kept beside the old one
        stem, ext = os.path.splitext(file)
        file = f"{stem}.{dt.datetime.now(dt.timezone.utc).strftime('%Y%m%dT%H%M%SZ')}{ext}"
        dest = os.path.join(raw_dir, source, file)
    part = dest + ".part"
    got = 0
    h = hashlib.sha256()
    r = sess.get(url, stream=True, allow_redirects=True, timeout=120)
    try:
        if r.status_code != 200:
            raise RuntimeError(f"{source}/{file}: HTTP {r.status_code} from {url}")
        with open(part, "wb") as f:
            for chunk in r.iter_content(1 << 20):
                if got + len(chunk) > budget:
                    f.close()
                    os.remove(part)
                    append_ledger(raw_dir, {"source": source, "layer": layer, "url": url, "file": file,
                                            "bytes": got, "retrieved_at_utc": now_utc(),
                                            "http_status": "stopped at the ceiling"})
                    raise CeilingRefused(f"{source}/{file}: stopped at {got} bytes, the ceiling")
                f.write(chunk)
                h.update(chunk)
                got += len(chunk)
    finally:
        r.close()
    os.replace(part, dest)
    row = {"source": source, "layer": layer, "url": url, "file": file, "bytes": got, "sha256": h.hexdigest(),
           "retrieved_at_utc": now_utc(), "http_status": "200"}
    if terms:
        row.update({"terms_url": terms["url"], "terms_file": terms["file"], "terms_sha256": terms["sha256"]})
    append_ledger(raw_dir, row)
    log(f"pulled: {source}/{file} {got} bytes (stated {size}); total {spent + got} of {ceiling}")
    return dest


def fetch_terms(sess, raw_dir, source, key, ceiling=CEILING_BYTES):
    """The publisher's terms page, saved as it came, with its hash. Returns {"url", "file", "sha256"}."""
    t = TERMS[key]
    file = os.path.join("terms", t["file"]).replace("\\", "/")
    path = fetch(sess, raw_dir, source, "terms", t["url"], file, ceiling=ceiling)
    rel = os.path.relpath(path, raw_dir).replace("\\", "/")
    return {"url": t["url"], "file": rel, "sha256": sha256_file(path)}


def unpack(zip_path):
    """Unpack an archive into its own new folder beside it (untrusted: no path may leave the folder)."""
    out = os.path.join(os.path.dirname(zip_path), "unpacked", os.path.splitext(os.path.basename(zip_path))[0])
    if os.path.isdir(out) and os.listdir(out):
        return out
    os.makedirs(out, exist_ok=True)
    root = os.path.realpath(out)
    with zipfile.ZipFile(zip_path) as z:
        for m in z.infolist():
            target = os.path.realpath(os.path.join(out, m.filename))
            if not (target == root or target.startswith(root + os.sep)):
                raise RuntimeError(f"archive member leaves its folder: {m.filename}")
        z.extractall(out)
    return out


def find_file(folder, pattern):
    rx = re.compile(pattern, re.I)
    hits = []
    for d, _, files in os.walk(folder):
        if "__MACOSX" in d:
            continue
        for f in files:
            if rx.search(f) and not f.startswith("._"):
                hits.append(os.path.join(d, f))
    if not hits:
        raise FileNotFoundError(f"no file matching {pattern} under {folder}")
    return sorted(hits)[0]


def html_text(path):
    """The visible text of a saved page, whitespace collapsed: what a terms quote is checked against."""
    import html
    with open(path, "rb") as f:
        s = f.read().decode("utf-8", "replace")
    s = re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", s)
    s = re.sub(r"(?s)<[^>]+>", " ", s)
    s = html.unescape(s).replace(" ", " ")
    s = s.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    s = re.sub(r"\s+", " ", s).strip()
    return re.sub(r" ([.,;:])", r"\1", s)


# ---------------------------------------------------------------------------------------------------------
# The web layers: the manifest, the grid encoding, the shape and point writers
# ---------------------------------------------------------------------------------------------------------

def write_json_atomic(path, obj):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="\n") as f:
        json.dump(obj, f, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    os.replace(tmp, path)
    return os.path.getsize(path)


def read_manifest(path):
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    return {"built_at_utc": now_utc(), "layers": [], "missing": []}


def write_manifest(path, man):
    man["built_at_utc"] = now_utc()
    order = {k: i for i, k in enumerate(LAYER_ORDER)}
    man["layers"].sort(key=lambda l: order.get(l["id"], 999))
    man["missing"] = [m for m in man.get("missing", []) if m["id"] not in {l["id"] for l in man["layers"]}]
    man["missing"].sort(key=lambda l: order.get(l["id"], 999))
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="\n") as f:
        json.dump(man, f, ensure_ascii=False, indent=1, allow_nan=False)
        f.write("\n")
    os.replace(tmp, path)


def manifest_put(path, layer):
    man = read_manifest(path)
    man["layers"] = [l for l in man["layers"] if l["id"] != layer["id"]] + [layer]
    write_manifest(path, man)
    log(f"manifest: {layer['id']} ({layer['kind']}) written")


def manifest_missing(path, id, group, title, reason):
    man = read_manifest(path)
    if any(l["id"] == id for l in man["layers"]):
        return
    man["missing"] = [m for m in man.get("missing", []) if m["id"] != id] + [
        {"id": id, "group": group, "title": title, "reason": reason}]
    write_manifest(path, man)


def encode_grid(values, scale, offset=0.0):
    """A float array (NaN where empty) as uint16 little-endian base64, row-major, north to south."""
    import numpy as np
    v = np.asarray(values, dtype="float64")
    q = np.rint((v - offset) / scale)
    ok = np.isfinite(v)
    if ok.any() and (q[ok].min() < 0 or q[ok].max() >= NODATA):
        raise ValueError(f"values do not fit uint16 at scale {scale}: {q[ok].min()} to {q[ok].max()}")
    out = np.where(ok, q, NODATA).astype("<u2")
    return base64.b64encode(out.tobytes()).decode("ascii")


def decode_grid(obj):
    """The grid file back to a float array (NaN where empty)."""
    import numpy as np
    raw = np.frombuffer(base64.b64decode(obj["values"]), dtype="<u2").reshape(obj["nrows"], obj["ncols"])
    v = raw.astype("float64") * obj["scale"] + obj["offset"]
    v[raw == obj["nodata"]] = np.nan
    return v


def allowed_levels(source_cell_deg):
    """The pyramid's cell sizes for a source of this cell size: never finer than the source."""
    return [c for c in LEVELS if c + 1e-12 >= source_cell_deg]


def source_cell_deg(ds):
    """The source raster's cell size in degrees. For a projected raster: the wider of its north to south
    size and its east to west size at the northern edge of what it covers (where a degree of longitude is
    shortest), so that no level is finer than the source anywhere."""
    import pyproj
    if ds.crs.is_geographic:
        return max(abs(ds.transform.a), abs(ds.transform.e))
    tr = pyproj.Transformer.from_crs(ds.crs, 4326, always_xy=True)
    b = ds.bounds
    lats = [tr.transform(x, y)[1] for x in (b.left, (b.left + b.right) / 2, b.right) for y in (b.top, b.bottom)]
    north = min(max(lats), 89.0)
    unit = ds.crs.linear_units_factor[1] if ds.crs.linear_units_factor else 1.0
    dx = abs(ds.transform.a) * unit
    dy = abs(ds.transform.e) * unit
    return max(dy / 111320.0, dx / (111320.0 * math.cos(math.radians(north))))


def snap(v, step, up):
    k = v / step
    k = math.ceil(k - 1e-9) if up else math.floor(k + 1e-9)
    return round(k * step, 6)


def accumulate_raster(path, band, lon0, lat0, cell, ncols, nrows, valid=None, chunk_rows=256):
    """Assign every source cell centre to a cell of the target grid; return (sum, n_valid, n_cells, stats).

    n_cells counts every source cell whose centre falls in the target cell, valid or not, so that the
    half rule can be applied. Read by windows: memory stays small whatever the raster's size.
    """
    import numpy as np
    import pyproj
    import rasterio
    from rasterio.windows import Window
    n = ncols * nrows
    s = np.zeros(n)
    c = np.zeros(n)
    t = np.zeros(n)
    st = {"n_valid": 0, "sum": 0.0, "min": math.inf, "max": -math.inf, "n_cells": 0,
          "n_valid_in_grid": 0, "sum_in_grid": 0.0, "min_in_grid": math.inf, "max_in_grid": -math.inf}
    with rasterio.open(path) as ds:
        tr = None if ds.crs.is_geographic else pyproj.Transformer.from_crs(ds.crs, 4326, always_xy=True)
        nod = ds.nodatavals[band - 1]
        T = ds.transform
        cols = np.arange(ds.width) + 0.5
        for r0 in range(0, ds.height, chunk_rows):
            h = min(chunk_rows, ds.height - r0)
            a = ds.read(band, window=Window(0, r0, ds.width, h)).astype("float64")
            ok = np.isfinite(a)
            if nod is not None and not (isinstance(nod, float) and math.isnan(nod)):
                ok &= a != nod
            if valid is not None:
                ok &= valid(a)
            rows = np.arange(r0, r0 + h) + 0.5
            cc, rr = np.meshgrid(cols, rows)
            x = T.a * cc + T.b * rr + T.c
            y = T.d * cc + T.e * rr + T.f
            if tr is not None:
                x, y = tr.transform(x, y)
            ci = np.floor((x - lon0) / cell).astype("int64")
            ri = np.floor((lat0 - y) / cell).astype("int64")
            inside = (ci >= 0) & (ci < ncols) & (ri >= 0) & (ri < nrows)
            idx = ri * ncols + ci
            t += np.bincount(idx[inside], minlength=n)
            m = inside & ok
            c += np.bincount(idx[m], minlength=n)
            s += np.bincount(idx[m], weights=a[m], minlength=n)
            st["n_cells"] += a.size
            if ok.any():
                st["n_valid"] += int(ok.sum())
                st["sum"] += float(a[ok].sum())
                st["min"] = min(st["min"], float(a[ok].min()))
                st["max"] = max(st["max"], float(a[ok].max()))
            if m.any():
                st["n_valid_in_grid"] += int(m.sum())
                st["sum_in_grid"] += float(a[m].sum())
                st["min_in_grid"] = min(st["min_in_grid"], float(a[m].min()))
                st["max_in_grid"] = max(st["max_in_grid"], float(a[m].max()))
    return s.reshape(nrows, ncols), c.reshape(nrows, ncols), t.reshape(nrows, ncols), st


def coarsen(a, k):
    """Sum blocks of k by k cells."""
    r, c = a.shape
    return a.reshape(r // k, k, c // k, k).sum(axis=(1, 3))


def nice_stops(lo, hi, n=6):
    if not (math.isfinite(lo) and math.isfinite(hi)) or hi <= lo:
        return [lo, hi]
    raw = (hi - lo) / (n - 1)
    mag = 10 ** math.floor(math.log10(raw))
    step = min((1, 2, 2.5, 5, 10), key=lambda m: abs(m * mag - raw)) * mag
    first = math.ceil(lo / step - 1e-9) * step
    out = []
    v = first
    while v <= hi + 1e-9:
        out.append(round(v, 10))
        v += step
    return out


def build_grid(id, raster, band, web_dir, bounds, scale, valid=None, level_cells=None, crop=False,
               notice=None):
    """One raster to a pyramid of grid files. Returns (levels, legend, stats, source_res_deg).

    bounds: (west, south, east, north) in degrees, snapped outward here to 0.2 degrees.
    The mean in each cell is the mean of the source's own valid cells whose centres fall in it; a cell with
    fewer than half of its source cells valid is empty. No interpolation, no filling, no smoothing.
    """
    import numpy as np
    import rasterio
    with rasterio.open(raster) as ds:
        res = source_cell_deg(ds)
    cells = level_cells or allowed_levels(res)
    if not cells:
        raise RuntimeError(f"{id}: the source's cells ({res:.4f} degrees) are coarser than every level")
    base = min(cells)
    k_top = int(round(0.2 / base))
    w, s_, e, n_ = bounds
    lon0, lat0 = snap(w, 0.2, False), snap(n_, 0.2, True)
    ncols = int(round((snap(e, 0.2, True) - lon0) / 0.2)) * k_top
    nrows = int(round((lat0 - snap(s_, 0.2, False)) / 0.2)) * k_top
    S, C, T, st = accumulate_raster(raster, band, lon0, lat0, base, ncols, nrows, valid=valid)
    if crop:  # keep the 0.2 degree blocks from the first to the last that hold a source value
        top = coarsen(C, k_top)
        rr = np.nonzero(top.sum(axis=1))[0]
        cc = np.nonzero(top.sum(axis=0))[0]
        r0, r1, c0, c1 = rr[0] * k_top, (rr[-1] + 1) * k_top, cc[0] * k_top, (cc[-1] + 1) * k_top
        S, C, T = S[r0:r1, c0:c1], C[r0:r1, c0:c1], T[r0:r1, c0:c1]
        lon0, lat0 = round(lon0 + c0 * base, 6), round(lat0 - r0 * base, 6)
    if not st["n_valid_in_grid"]:
        raise RuntimeError(f"{id}: the source holds no value inside {bounds}")
    src_mean = st["sum_in_grid"] / st["n_valid_in_grid"]
    log(f"{id}: source {st['n_valid']} valid cells of {st['n_cells']} (whole file: mean "
        f"{st['sum'] / st['n_valid']:.6g}, min {st['min']:.6g}, max {st['max']:.6g}); cell {res:.5f} degrees; "
        f"inside the grid {st['n_valid_in_grid']} valid cells, mean {src_mean:.6g}, "
        f"min {st['min_in_grid']:.6g}, max {st['max_in_grid']:.6g}")
    levels, checks = [], []
    finest = None
    for cell in sorted(cells, reverse=True):
        k = int(round(cell / base))
        s, c, t = coarsen(S, k), coarsen(C, k), coarsen(T, k)
        keep = (c > 0) & (c * 2 >= t)
        v = np.full(s.shape, np.nan)
        v[keep] = s[keep] / c[keep]
        tag = ("%g" % cell).replace(".", "p")
        file = f"{id}_{tag}.json"
        obj = {"lon0": lon0, "lat0": lat0, "dlon": cell, "dlat": cell, "ncols": int(v.shape[1]),
               "nrows": int(v.shape[0]), "scale": scale, "offset": 0.0, "nodata": NODATA,
               "encoding": "uint16-le-base64", "values": encode_grid(v, scale)}
        if notice:
            obj["notice"] = notice
        size = write_json_atomic(os.path.join(web_dir, file), obj)
        back = decode_grid(obj)
        okb = np.isfinite(back)
        wmean = float((back[okb] * c[okb]).sum() / c[okb].sum())
        kept_share = float(c[keep].sum() / st["n_valid_in_grid"]) if st["n_valid_in_grid"] else float("nan")
        kept_src_mean = float(s[keep].sum() / c[keep].sum())
        checks.append({"cell_deg": cell, "cells_with_value": int(okb.sum()), "mean_of_cells": float(back[okb].mean()),
                       "mean_weighted_by_source_cells": wmean, "source_mean_of_kept_cells": kept_src_mean,
                       "source_mean_all": src_mean, "share_of_source_cells_kept": kept_share,
                       "min": float(back[okb].min()), "max": float(back[okb].max())})
        log(f"{id} {cell}: {v.shape[1]}x{v.shape[0]}, {int(okb.sum())} cells with a value, {size} bytes; "
            f"mean weighted by source cells {wmean:.6g} against the source's {kept_src_mean:.6g} over the same "
            f"cells and {src_mean:.6g} over all its cells in the grid ({kept_share:.4%} of them kept); "
            f"plain mean of cells {back[okb].mean():.6g}")
        if abs(wmean - kept_src_mean) > scale:
            raise RuntimeError(f"{id} {cell}: the level's mean {wmean} is not the source's {kept_src_mean}")
        levels.append({"file": file, "cell_deg": cell, "ncols": int(v.shape[1]), "nrows": int(v.shape[0]),
                       "bytes": size})
        finest = back
    okf = np.isfinite(finest)
    lo, hi = float(finest[okf].min()), float(finest[okf].max())
    p2, p98 = (float(x) for x in np.percentile(finest[okf], [2, 98]))
    legend = {"min": round(lo, 6), "max": round(hi, 6), "stops": nice_stops(p2, p98)}
    st["mean"] = src_mean
    st["checks"] = checks
    last = levels[0]
    st["grid"] = {"lon0": lon0, "lat0": lat0, "west": lon0, "north": lat0,
                  "east": round(lon0 + last["ncols"] * last["cell_deg"], 6),
                  "south": round(lat0 - last["nrows"] * last["cell_deg"], 6)}
    return levels, legend, st, res


def round_geom(geom, tolerance, digits=3):
    """Simplify one geometry with its topology kept, then put its coordinates on a 10^-digits grid."""
    import shapely
    g = shapely.simplify(geom, tolerance, preserve_topology=True) if tolerance else geom
    g = shapely.set_precision(g, 10 ** -digits)
    if g.is_empty:
        g = shapely.set_precision(geom, 10 ** -digits)
    if not g.is_valid:  # a source polygon that was not valid itself (shells that touch or nest)
        g = shapely.make_valid(g)
        if g.geom_type == "GeometryCollection":
            g = shapely.union_all([x for x in g.geoms if x.geom_type in ("Polygon", "MultiPolygon")])
        g = shapely.set_precision(g, 10 ** -digits)
    return g


def geom_json(g, digits=3):
    import shapely
    obj = json.loads(shapely.to_geojson(g))

    def rnd(x):
        if isinstance(x, (list, tuple)):
            return [rnd(v) for v in x]
        return round(x, digits)
    if "coordinates" in obj:
        obj["coordinates"] = rnd(obj["coordinates"])
    else:
        for gg in obj.get("geometries", []):
            gg["coordinates"] = rnd(gg["coordinates"])
    return obj


def clean(v):
    """A source field as JSON: numbers stay numbers, missing is null, text is stripped."""
    import numpy as np
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, np.integer)):
        return int(v)
    if isinstance(v, (float, np.floating)):
        return None if not math.isfinite(float(v)) else float(v)
    if isinstance(v, (dt.date, dt.datetime)) or hasattr(v, "isoformat"):
        try:
            return v.isoformat()
        except Exception:
            return str(v)
    s = str(v).strip()
    return s if s else None


def write_shapes(web_dir, file, gdf, props_fn, tolerance=0.01, coverage=False, digits=3, notice=None):
    """A GeoDataFrame to one GeoJSON FeatureCollection in WGS84. Returns (bytes, features, bounds).

    coverage: the polygons tile an area and share edges (classes, counties). They are simplified together
    (shapely.coverage_simplify), so that neighbours still meet along one line. Refused when the source's
    polygons are not a valid coverage: the caller then simplifies one feature at a time.
    """
    import shapely
    g = gdf.to_crs(4326) if gdf.crs is not None and gdf.crs.to_epsg() != 4326 else gdf
    g = g[~(g.geometry.isna() | g.geometry.is_empty)]
    if coverage:
        geoms = shapely.make_valid(g.geometry.values) if not g.geometry.is_valid.all() else g.geometry.values
        if not shapely.coverage_is_valid(geoms):
            raise ValueError(f"{file}: the source's polygons are not a valid coverage")
        g = g.set_geometry(shapely.coverage_simplify(geoms, tolerance), crs=g.crs)
        tolerance = 0
    feats = []
    for _, row in g.iterrows():
        if row.geometry is None or row.geometry.is_empty:
            continue
        geom = round_geom(row.geometry, tolerance, digits)
        if geom.is_empty:
            continue
        feats.append({"type": "Feature", "properties": props_fn(row), "geometry": geom_json(geom, digits)})
    fc = {"type": "FeatureCollection", "features": feats}
    if notice:
        fc["notice"] = notice
    size = write_json_atomic(os.path.join(web_dir, file), fc)
    b = [round(float(x), 3) for x in g.total_bounds]
    return size, len(feats), b


def write_points(web_dir, file, columns, rows):
    """A points file: {"columns": [...], "rows": [[lon, lat, name, value, ...]]}. Returns its size."""
    if columns[:4] != ["lon", "lat", "name", "value"]:
        raise ValueError("a points file starts with lon, lat, name, value")
    return write_json_atomic(os.path.join(web_dir, file), {"columns": columns, "rows": rows})


def source_fields(row, skip=("geometry",)):
    return {k: clean(v) for k, v in row.items() if k not in skip}


EXPECTED = {
    "wind_speed_100m": ("wind", "Wind speed at 100 m"),
    "wind_capacity_factor": ("wind", "Wind capacity factor (supply curve sites)"),
    "wind_gross_capacity_factor": ("wind", "Gross capacity factor"),   # not held, and not fetched (NOT_HELD)
    "solar_ghi": ("solar", "Global horizontal irradiance"),
    "solar_dni": ("solar", "Direct normal irradiance"),
    "geothermal_hydrothermal_sites": ("geothermal", "Identified hydrothermal systems"),
    "geothermal_hydrothermal_favorability": ("geothermal", "Hydrothermal favorability (western states)"),
    "geothermal_egs_favorability": ("geothermal", "Deep enhanced geothermal favorability"),
    "oil_gas_basins": ("oil and gas", "Sedimentary basins"),
    "oil_gas_plays": ("oil and gas", "Tight oil and shale gas plays"),
    # session 159 (the owner's ruling of 8 October 2026): hydropower in the form its publisher publishes it
    "hydropower_npd": ("hydropower", "Hydropower potential at non-powered dams"),
    "hydropower_nsd": ("hydropower", "Hydropower potential of new stream-reach development"),
    "biomass": ("biomass", "Solid biomass resources by county"),
    "offshore_wind_leases": ("offshore wind", "Offshore wind lease areas"),
    "offshore_wind_planning_areas": ("offshore wind", "Offshore wind planning areas"),
}
LAYER_ORDER = list(EXPECTED)
# An id an earlier session wrote into the manifest's missing list and that no longer stands for a layer: session
# 146's one greyed "hydropower_potential" is two layers of Oak Ridge's since session 159.
RETIRED_IDS = {"hydropower_potential"}
# A layer that was looked for and left, with the reason the page shows on its greyed toggle. (Session 146 left
# hydropower here: Oak Ridge's download page puts a form asking for a name and an e-mail address in front of its
# files. The owner ruled on 8 October 2026 that the published files are fetched by their own addresses with a
# plain request, no form filled and no name or address sent; SOURCES below holds them since session 159.)
NOT_HELD = {
    # the owner's ruling of 8 October 2026: "a gross capacity factor layer that needs a personal key is not
    # fetched". One sentence, shown on the greyed toggle's hover and in the Method note. No key was asked for.
    "wind_gross_capacity_factor": (
        "Not fetched: no gross capacity factor was found among the files the laboratory (NLR, until 2025 NREL) "
        "offers without a key, its developer service needs a key issued to a named person, and no key is "
        "asked for."),
}
# Recorded and left (COMMON.md rule 5): the addresses are kept here so that a ruling needs no new search.
LEFT_SOURCES = {
    "ornl_nsd_ea": {
        "publisher": "Oak Ridge National Laboratory (HydroSource), for the U.S. Department of Energy",
        "title": "Hydropower Potential from New Stream-Reach Development for the Conterminous United States "
                 "(2014): the environmental attributes (the files named EA)",
        "landing": "https://hydrosource.ornl.gov/data/datasets/"
                   "hydropower-potential-new-stream-reach-development-conterminous-united-states/",
        "files": {"ORNL_NHAAP_NSD_EA_All_v1.zip": 564645432, "ORNL_NHAAP_NSD_EA_All_v1.xlsx": 575614},
        "terms": "https://hydrosource.ornl.gov/data-use-policy/",
        "why_left": "they describe the environment of each watershed, not the hydropower resource, so the map "
                    "has no layer for them; the page's file links also open a form that asks for a name, an "
                    "e-mail address, a company and an occupation, which is never filled"},
    "doe_billion_ton_2023": {
        "publisher": "Oak Ridge National Laboratory (Bioenergy Knowledge Discovery Framework), for the U.S. "
                     "Department of Energy",
        "title": "2023 Billion-Ton Report, county data downloads",
        "landing": "https://bioenergykdf.ornl.gov/bt23-data-portal",
        "why_left": "the county downloads need an account (\"Log in to download data\")"},
}

# The publishers' terms pages. The quotes used in the manifest are checked word for word against the saved
# page by --check-terms (and by the tests when the raw store is on the machine).
DOI_QUOTE = ("Generally, materials produced by federal agencies are in the public domain and may be reproduced "
             "without permission. However, not all materials appearing on this web site are in the public domain.")
TERMS = {
    "eia": {"url": "https://www.eia.gov/about/copyrights_reuse.php", "file": "eia_copyrights_reuse.html",
            "quote": "U.S. government publications are in the public domain and are not subject to copyright "
                     "protection. You may use and/or distribute any of our data, files, databases, reports, "
                     "graphs, charts, and other information products that are on our website or that you "
                     "receive through our email distribution service. However, if you use or reproduce any of "
                     "our information products, you should use an acknowledgment, which includes the "
                     "publication date, such as: \"Source: U.S. Energy Information Administration (Oct 2008).\""},
    "nlr": {"url": "https://www.nlr.gov/disclaimer", "file": "nlr_disclaimer.html",
            "quote": "The user is granted the right, without any fee or cost, to use or copy the Data, provided "
                     "that this entire notice appears in all copies of the Data. Further, the user agrees to "
                     "credit the U.S. Department of Energy (DOE)/NLR/ALLIANCE in any publication that results "
                     "from the use of the Data."},
    "usgs": {"url": "https://www.usgs.gov/information-policies-and-instructions/copyrights-and-credits",
             "file": "usgs_copyrights_and_credits.html",
             "quote": ["USGS-authored or produced data and information are considered to be in the U.S. Public "
                       "Domain.",
                       "When using information from USGS information products, publications, or websites, we ask "
                       "that proper credit be given."]},
    "openei_8314": {"url": "https://data.openei.org/submissions/8314", "file": "openei_submission_8314.html",
                    "quote": "Content is available under Creative Commons Attribution 4.0 unless otherwise noted."},
    "doi": {"url": "https://www.doi.gov/copyright", "file": "doi_copyright.html", "quote": DOI_QUOTE},
    "boem": {"url": "https://www.boem.gov/renewable-energy/mapping-and-data/renewable-energy-gis-data",
             "file": "boem_renewable_energy_gis_data.html", "also": ["doi"],
             "quote": "Note to users: Data downloaded from this site is to be used for informational and "
                      "planning purposes only."},
    # session 159: HydroSource's data use policy, its two paragraphs in full (use, then citation)
    "ornl": {"url": "https://hydrosource.ornl.gov/data-use-policy/", "file": "hydrosource_data_use_policy.html",
             "quote": ["Data hosted on HydroSource is openly shared, without restriction, in accordance with "
                       "Department of Energy's Public Access Plan.",
                       "Bibliographic citations should be included in the References section of publications "
                       "and other media to acknowledge those who have created the data, services, and tools "
                       "provided by HydroSource. Proper citations include the authors, title, publisher, and "
                       "Digital Object Identifier (DOI) and will allow the products to be discovered and "
                       "re-used by others. The proper citation for each HydroSource product is provided on its "
                       "landing page."]},
}
SOURCES = {
    "eia_basins": {"source": "eia", "layers": ["oil_gas_basins"], "terms": "eia",
                   "url": "https://www.eia.gov/maps/map_data/SedimentaryBasins_US_EIA.zip",
                   "file": "SedimentaryBasins_US_EIA.zip"},
    "eia_plays": {"source": "eia", "layers": ["oil_gas_plays"], "terms": "eia",
                  "url": "https://www.eia.gov/maps/map_data/TightOil_ShaleGas_Plays_Lower48_EIA.zip",
                  "file": "TightOil_ShaleGas_Plays_Lower48_EIA.zip"},
    "boem_shapefiles": {"source": "boem", "layers": ["offshore_wind_leases", "offshore_wind_planning_areas"],
                        "terms": "boem",
                        "url": "https://www.boem.gov/renewable-energy/boem-renewable-energy-shapefiles",
                        "file": "boem-renewable-energy-shapefiles.zip"},
    "boem_geodatabase": {"source": "boem", "layers": ["offshore_wind_planning_areas"], "terms": "boem",
                         "url": "https://www.boem.gov/renewable-energy/boem-renewable-energy-geodatabase",
                         "file": "boem-renewable-energy-geodatabase.zip"},
    "nlr_wind": {"source": "nlr", "layers": ["wind_speed_100m"], "terms": "nlr",
                 "url": "https://www.nlr.gov/docs/libraries/gis/us-wind-data.zip", "file": "us-wind-data.zip"},
    "nlr_ghi": {"source": "nlr", "layers": ["solar_ghi"], "terms": "nlr",
                "url": "https://www.nlr.gov/docs/libraries/gis/nsrdbv3_ghi.zip", "file": "nsrdbv3_ghi.zip"},
    "nlr_dni": {"source": "nlr", "layers": ["solar_dni"], "terms": "nlr",
                "url": "https://www.nlr.gov/docs/libraries/gis/nsrdbv3_dni.zip", "file": "nsrdbv3_dni.zip"},
    "boem_planning_meta": {
        "source": "boem", "layers": ["offshore_wind_planning_areas"], "terms": "boem",
        "url": "https://services7.arcgis.com/G5Ma95RzqJRPKsWL/ArcGIS/rest/services/"
               "Wind_Planning_Area_Boundaries__BOEM_/FeatureServer/0?f=json",
        "file": "Wind_Planning_Area_Boundaries_layer0.json"},
    "boem_planning": {
        "source": "boem", "layers": ["offshore_wind_planning_areas"], "terms": "boem",
        "url": "https://services7.arcgis.com/G5Ma95RzqJRPKsWL/ArcGIS/rest/services/"
               "Wind_Planning_Area_Boundaries__BOEM_/FeatureServer/0/query"
               "?where=1%3D1&outFields=*&outSR=4326&f=geojson",
        "file": "Wind_Planning_Area_Boundaries_layer0.geojson"},
}
SB = "https://www.sciencebase.gov/catalog/file/get/"
for _stem, _item, _layer, _parts in (
    ("IdentifiedGeothermalSystems", "6606f534d34e4df16bd58277", "geothermal_hydrothermal_sites", {
        "shp": "ae%2Fbd%2F32%2Faebd32c1cf3b64ce4f19dc6e2535e5a5bc3382da",
        "dbf": "a4%2Ffb%2F25%2Fa4fb257343cfb36fec23c0da4bc1ba243d5c132c",
        "shx": "ec%2Fdb%2Fbe%2Fecdbbe23f01f970066c35c233071b5ff5ecaef8f",
        "prj": "84%2Fcc%2F68%2F84cc6899edc8b186b89318986329c3976fdc3394",
        "xml": "f3%2F34%2F82%2Ff33482fb63f32b2027fd4fb9bb50889a697ea27e"}),
    ("FavorabilitySurface", "6606ed51d34e4df16bd58251", "geothermal_hydrothermal_favorability", {
        "shp": "79%2Fb1%2Fce%2F79b1ce18796d6fc4ae5b9292c58bd8aaee4846e5",
        "dbf": "fc%2Fc5%2Fc7%2Ffcc5c7ec4d94360183d66dd9438b26a40fd6f74c",
        "shx": "59%2F53%2Fe2%2F5953e29dadb4ec6c1dc42e27a8888289d7052a0a",
        "prj": "55%2Fce%2F4a%2F55ce4a0577798f1bc84dbe91a4ad54cef8c3b663",
        "xml": "b0%2F3d%2Fd8%2Fb03dd8396bbe0fc5ed9270e5e3665a9928de7660"}),
):
    for _ext, _disk in _parts.items():
        SOURCES[f"usgs_{_stem}_{_ext}"] = {
            "source": "usgs", "layers": [_layer], "terms": "usgs",
            "url": f"{SB}{_item}?f=__disk__{_disk}", "file": f"{_stem}/{_stem}.{_ext}"}
SOURCES.update({
    "nlr_egs": {"source": "nlr", "layers": ["geothermal_egs_favorability"], "terms": "nlr",
                "url": "https://www.nlr.gov/docs/libraries/gis/egs.zip", "file": "egs.zip"},
    "nlr_biomass": {"source": "nlr", "layers": ["biomass"], "terms": "nlr",
                    "url": "https://www.nlr.gov/docs/libraries/gis/solid-biomass.zip", "file": "solid-biomass.zip"},
    "openei_wind_sc": {
        "source": "openei", "layers": ["wind_capacity_factor"], "terms": "openei_8314",
        "url": "https://data.openei.org/files/8314/lbw_open_access_2035_moderate_115hh_170rd_supply_curve.csv",
        "file": "lbw_open_access_2035_moderate_115hh_170rd_supply_curve.csv"},
    "openei_wind_sc_lookup": {
        "source": "openei", "layers": ["wind_capacity_factor"], "terms": "openei_8314",
        "url": "https://data.openei.org/files/8314/lbw_column_lookup.csv", "file": "lbw_column_lookup.csv"},
})
# Session 159, the owner's ruling of 8 October 2026: Oak Ridge National Laboratory's hydropower assessments on
# HydroSource, each file by its own address (the addresses the dataset pages themselves print), with a plain
# request. The pages' download buttons open a form (name, e-mail, company, occupation): it is never filled and
# nothing but the User-Agent above is sent. A file whose address refuses a plain request is recorded and left.
ORNL = "https://hydrosource.s3.us-east-2.amazonaws.com/files/data/datasets/"
ORNL_NPD = ORNL + "hydropower-capacity-us-npd/"
ORNL_NSD = ORNL + "hydropower-potential-new-stream-reach-development-conterminous-united-states/"
SOURCES.update({
    "ornl_npd": {"source": "ornl", "layers": ["hydropower_npd"], "terms": "ornl",
                 "url": ORNL_NPD + "TechPotentialNPDs.csv", "file": "npd/TechPotentialNPDs.csv"},
    "ornl_npd_fields": {"source": "ornl", "layers": ["hydropower_npd"], "terms": "ornl",
                        "url": ORNL_NPD + "TechPotentialNPDs_field_descriptions.csv",
                        "file": "npd/TechPotentialNPDs_field_descriptions.csv"},
    "ornl_npd_readme": {"source": "ornl", "layers": ["hydropower_npd"], "terms": "ornl",
                        "url": ORNL_NPD + "TechPotentialNPDs_Readme.txt", "file": "npd/TechPotentialNPDs_Readme.txt"},
    "ornl_nsd_xlsx": {"source": "ornl", "layers": ["hydropower_nsd"], "terms": "ornl",
                      "url": ORNL_NSD + "ORNL_NHAAP_NSD_SR_All_v1.xlsx", "file": "nsd/ORNL_NHAAP_NSD_SR_All_v1.xlsx"},
    "ornl_nsd_zip": {"source": "ornl", "layers": ["hydropower_nsd"], "terms": "ornl",
                     "url": ORNL_NSD + "ORNL_NHAAP_NSD_SR_All_v1.zip", "file": "nsd/ORNL_NHAAP_NSD_SR_All_v1.zip"},
})
BUILDERS = {}


def builder(id):
    def deco(fn):
        BUILDERS[id] = fn
        return fn
    return deco


def held(raw_dir, key):
    """The ledger row of a pulled source file, with its path. Fails loudly when the file is not held."""
    src = SOURCES[key]
    rows = [r for r in read_ledger(raw_dir)
            if r["source"] == src["source"] and r["file"] == src["file"] and r["http_status"] == "200"]
    if not rows:
        raise FileNotFoundError(f"{key}: not in the raw store ({raw_dir}); run --pull first")
    r = dict(rows[-1])
    r["path"] = os.path.join(raw_dir, r["source"], r["file"])
    r["raw_dir"] = raw_dir
    if not os.path.exists(r["path"]):
        raise FileNotFoundError(r["path"])
    return r


def base_layer(id, group, title, kind, row, key, **more):
    """The manifest fields every layer carries, from its source's ledger row."""
    t = TERMS[SOURCES[key]["terms"]]
    quotes = t["quote"] if isinstance(t["quote"], list) else [t["quote"]]
    layer = {"id": id, "group": group, "title": title, "kind": kind,
             "source_url": row["url"], "terms_url": row["terms_url"], "terms_quote": " [...] ".join(quotes),
             "terms_quotes": quotes,
             "terms_file": row["terms_file"], "terms_sha256": row["terms_sha256"],
             "retrieved_at_utc": row["retrieved_at_utc"],
             "source_file": f"{row['source']}/{row['file']}", "source_bytes": int(row["bytes"]),
             "source_sha256": row["sha256"]}
    for k in t.get("also", []):  # a second terms page of the same publisher, saved and hashed like the first
        a = TERMS[k]
        rel = f"{row['source']}/terms/{a['file']}"
        path = os.path.join(row["raw_dir"], rel)
        layer.setdefault("terms_also", []).append({
            "url": a["url"], "file": rel, "quote": a["quote"],
            "sha256": sha256_file(path) if os.path.exists(path) else ""})
    layer.update(more)
    return layer


def shape_stats(values):
    v = [x for x in values if isinstance(x, (int, float)) and math.isfinite(x)]
    if not v:
        return None
    return {"n": len(v), "min": min(v), "mean": sum(v) / len(v), "max": max(v)}


def raster_bounds(path):
    """The raster's extent in longitude and latitude: (west, south, east, north)."""
    import numpy as np
    import pyproj
    import rasterio
    with rasterio.open(path) as ds:
        b = ds.bounds
        if ds.crs.is_geographic:
            return b.left, b.bottom, b.right, b.top
        tr = pyproj.Transformer.from_crs(ds.crs, 4326, always_xy=True)
        xs = np.linspace(b.left, b.right, 200)
        ys = np.linspace(b.bottom, b.top, 200)
        edge_x = np.concatenate([xs, xs, np.full(200, b.left), np.full(200, b.right)])
        edge_y = np.concatenate([np.full(200, b.bottom), np.full(200, b.top), ys, ys])
        lon, lat = tr.transform(edge_x, edge_y)
        return float(lon.min()), float(lat.min()), float(lon.max()), float(lat.max())


CONUS = (-125.0, 24.4, -66.8, 49.6)  # the contiguous states with their coastal waters, on 0.2 degree lines


def grid_layer(raw_dir, web_dir, manifest, *, id, key, raster, band=1, bounds=None, scale, valid=None,
               level_cells=None, extras=None, **fields):
    """Build one grid layer's pyramid and write its manifest entry. bounds None: all the raster covers,
    cut to the 0.2 degree blocks that hold a value."""
    row = held(raw_dir, key)
    notice = fields.get("terms_notice")
    levels, legend, st, res = build_grid(id, raster, band, web_dir, bounds or raster_bounds(raster), scale,
                                         valid=valid, level_cells=level_cells, crop=bounds is None,
                                         notice=notice)
    layer = base_layer(id, fields.pop("group"), fields.pop("title"), "grid", row, key,
                       levels=levels, legend=legend, source_cell_deg=round(res, 6),
                       source_stats={"n_valid": st["n_valid_in_grid"], "min": st["min_in_grid"],
                                     "mean": st["mean"], "max": st["max_in_grid"],
                                     "n_valid_whole_file": st["n_valid"], "n_cells_whole_file": st["n_cells"]},
                       checks=st["checks"], grid=st["grid"], **fields)
    # what the source covers beyond the main grid, in files of their own (the page draws the main grid first)
    for tag, (b, extent) in (extras or {}).items():
        lv, lg, sx, _ = build_grid(f"{id}_{tag}", raster, band, web_dir, b, scale, valid=valid,
                                   level_cells=level_cells, notice=notice)
        layer.setdefault("other_extents", []).append({
            "extent": extent, "levels": lv, "legend": lg, "grid": sx["grid"], "checks": sx["checks"],
            "source_stats": {"n_valid": sx["n_valid_in_grid"], "min": sx["min_in_grid"], "mean": sx["mean"],
                             "max": sx["max_in_grid"]}})
    manifest_put(manifest, layer)
    return layer


GRID_REDUCTION = ("each source cell is assigned by its centre to a cell of a longitude and latitude grid (WGS84); "
                  "a cell's value is the mean of the source's own valid cells in it; a cell with fewer than half "
                  "of its source cells valid is empty; no interpolation, no filling, no smoothing; no level finer "
                  "than the source's own cell")

NLR_PUBLISHER = "National Laboratory of the Rockies (NLR), until 2025 the National Renewable Energy Laboratory (NREL)"
NLR_CREDIT = "U.S. Department of Energy (DOE)/NLR/ALLIANCE"


def nlr_notice(raw_dir):
    """NLR's whole Data and Software notice, from the saved page: its terms ask that it travel with the data."""
    t = html_text(os.path.join(raw_dir, "nlr", "terms", TERMS["nlr"]["file"]))
    a = t.index("Access to or use of any data or software made available on this server")
    b = t.index("ACCESS, USE OR PERFORMANCE OF THE DATA.") + len("ACCESS, USE OR PERFORMANCE OF THE DATA.")
    return t[a:b]


@builder("wind_speed_100m")
def build_wind_speed(raw_dir, web_dir, manifest):
    row = held(raw_dir, "nlr_wind")
    tif = find_file(unpack(row["path"]), r"wtk_conus_100m_mean_masked\.tif$")
    grid_layer(
        raw_dir, web_dir, manifest, id="wind_speed_100m", key="nlr_wind", raster=tif, scale=0.001,
        group="wind", title="Wind speed at 100 m", unit="m/s", value_label="annual mean wind speed",
        publisher=NLR_PUBLISHER,
        source_title="Wind Integration National Dataset (WIND) Toolkit, Multi-year Annual Average "
                     "(wtk_conus_100m_mean_masked.tif)",
        vintage="2007-2013 (the file's metadata: \"Multi-year (2007-2013) Annual Average Wind Speed, meters "
                "per second, at 100 meters above surface level\"); publication date 2015",
        extent="contiguous United States, onshore and offshore", source_resolution="2 km (\"2km x 2km\")",
        reduction=GRID_REDUCTION, credit=NLR_CREDIT, terms_notice=nlr_notice(raw_dir),
        notes_for_method="The publisher's words: \"This data provides modeled annual average wind speed for "
                         "the contiguous United States both onshore and offshore for the period 2007-2013.\" "
                         "\"Raster value: Multi-year (2007-2013) Annual Average Wind Speed, meters per second, "
                         "at 100 meters above surface level.\" It is a model's long-run average, not a "
                         "measurement at a mast and not a forecast; it says nothing of a turbine's output, of "
                         "land that may be used or of the grid. The source's cell is 2 km in a Lambert "
                         "projection, which is wider than 0.025 degrees of longitude in the north, so the "
                         "finest level is 0.05 degrees. The landing page that listed the file (NREL's wind "
                         "resource maps page) no longer exists since the laboratory's site moved to nlr.gov; "
                         "the file itself is still on the laboratory's server.")


SOLAR_EXTRAS = {"hawaii": ((-160.6, 18.6, -154.6, 22.4), "Hawaii"),
                "alaska": ((-180.0, 51.0, -129.8, 60.0),
                           "Alaska south of 60 degrees north, with the part of Canada inside the same box")}


def solar_layer(raw_dir, web_dir, manifest, id, key, short, long_name):
    row = held(raw_dir, key)
    tif = find_file(unpack(row["path"]), rf"nsrdb3_{short}\.tif$")
    grid_layer(
        raw_dir, web_dir, manifest, id=id, key=key, raster=tif, scale=0.001, bounds=CONUS, extras=SOLAR_EXTRAS,
        group="solar", title=long_name, unit="kWh/m2/day",
        value_label="annual average daily total solar resource", publisher=NLR_PUBLISHER,
        source_title=f"Physical Solar Model version 3 {long_name} Multi-year Annual Average "
                     f"(nsrdb3_{short}.tif), National Solar Radiation Database",
        vintage="1998-2016 (the file's metadata: \"The data are averaged from hourly model output over 19 "
                "years (1998-2016)\"); publication date 2018",
        extent="the box of the contiguous United States, longitude -125.0 to -66.8, latitude 24.4 to 49.6 "
               "(the source also holds values for the parts of Canada and Mexico inside the box, and they are "
               "kept); Hawaii and Alaska south of 60 degrees north are in files of their own",
        source_resolution="0.04 degrees in the file (the metadata: \"surface cells of 0.038 degrees in both "
                          "latitude and longitude, or nominally 4 km in size\")",
        reduction=GRID_REDUCTION, credit=NLR_CREDIT, terms_notice=nlr_notice(raw_dir),
        notes_for_method="The publisher's words: \"This data provides annual average daily total solar "
                         "resource averaged over surface cells of 0.038 degrees in both latitude and longitude, "
                         "or nominally 4 km in size. The solar radiation values represent the resource "
                         "available to solar energy systems.\" \"The data are averaged from hourly model output "
                         "over 19 years (1998-2016).\" The unit: the metadata's line reads \"Raster value: "
                         "solar irradiance in kWh/m2/year\", but its description says annual average daily "
                         "total, and the values are daily totals (a few kWh a square metre); the layer is "
                         "labelled kWh/m2/day and the values are the file's own, unconverted. It is a "
                         "satellite-based model's long-run average, not a measurement on the ground, and not "
                         "the output of a solar plant. The source's file covers most of the Americas up to 60 "
                         "degrees north; this layer keeps the box of the contiguous states. The landing page "
                         "that listed the file (NREL's solar resource maps page) no longer exists since the "
                         "laboratory's site moved to nlr.gov; the file itself is still on the laboratory's "
                         "server.")


@builder("solar_ghi")
def build_ghi(raw_dir, web_dir, manifest):
    solar_layer(raw_dir, web_dir, manifest, "solar_ghi", "nlr_ghi", "ghi", "Global Horizontal Irradiance")


@builder("solar_dni")
def build_dni(raw_dir, web_dir, manifest):
    solar_layer(raw_dir, web_dir, manifest, "solar_dni", "nlr_dni", "dni", "Direct Normal Irradiance")


def xml_field(path, tag):
    """The text of the first <tag> of a metadata file, whitespace collapsed."""
    import html
    with open(path, encoding="utf-8", errors="replace") as f:
        m = re.search(rf"<{tag}[^>]*>(.*?)</{tag}>", f.read(), re.S)
    if not m:
        return ""
    t = re.sub(r"<[^>]+>", " ", html.unescape(m.group(1)))
    return re.sub(r"\s+", " ", t.replace("\ufffd", " ")).strip()


def shapes_any(web_dir, file, gdf, props, tolerance, coverage, notice=None):
    """Simplify as a coverage when the source is one; say which was done."""
    if coverage:
        try:
            size, n, b = write_shapes(web_dir, file, gdf, props, tolerance=tolerance, coverage=True,
                                      notice=notice)
            return size, n, b, (f"the polygons simplified together as one coverage, so that neighbours still "
                                f"meet along one line (tolerance {tolerance} degrees); coordinates to 3 decimals; "
                                f"no feature dropped")
        except ValueError as e:
            log(f"{file}: {e}; simplified one feature at a time")
    size, n, b = write_shapes(web_dir, file, gdf, props, tolerance=tolerance, notice=notice)
    return size, n, b, (f"each polygon simplified to {tolerance} degrees with its topology kept (Douglas-Peucker, "
                        f"one feature at a time: a thin gap or overlap can show between neighbours when zoomed "
                        f"far in); coordinates to 3 decimals; no feature dropped")


USGS_PUBLISHER = "U.S. Geological Survey (USGS)"
USGS_CREDIT = "Credit: U.S. Geological Survey"


@builder("geothermal_hydrothermal_sites")
def build_geo_sites(raw_dir, web_dir, manifest):
    import geopandas as gpd
    row = held(raw_dir, "usgs_IdentifiedGeothermalSystems_shp")
    for ext in ("dbf", "shx", "prj", "xml"):
        held(raw_dir, f"usgs_IdentifiedGeothermalSystems_{ext}")
    g = gpd.read_file(row["path"]).to_crs(4326)
    extra = ["State", "Temp_C_Li", "Temp_C_Mn", "Temp_C_Mx", "Vol_km3_Li", "Vol_km3_Mn", "Vol_km3_Mx",
             "MWe_P95", "MWe_P5", "Lat_83", "Lon_83"]
    rows = [[round(r.geometry.x, 4), round(r.geometry.y, 4), clean(r["Name"]), clean(r["MWe_Mean"])]
            + [clean(r[c]) for c in extra] for _, r in g.iterrows()]
    size = write_points(web_dir, "geothermal_hydrothermal_sites.json", ["lon", "lat", "name", "value"] + extra, rows)
    vals = [r[3] for r in rows]
    st = shape_stats(vals)
    meta = row["path"][:-4] + ".xml"
    manifest_put(manifest, base_layer(
        "geothermal_hydrothermal_sites", "geothermal", "Identified hydrothermal systems", "points", row,
        "usgs_IdentifiedGeothermalSystems_shp", unit="MWe",
        value_label="Mean Megawatts electric (MWe), the source's MWe_Mean", publisher=USGS_PUBLISHER,
        source_title="Identified Moderate and High Temperature Geothermal Systems "
                     "(IdentifiedGeothermalSystems.shp), DeAngelo and Williams, 2010, from the 2008 Assessment "
                     "of moderate- and high-temperature geothermal resources of the United States",
        vintage="the 2008 assessment (USGS Fact Sheet 2008-3082); the data file is dated 2010 and was released "
                "again on ScienceBase on 29 March 2024 (metadata date 20240329)",
        extent="United States: the western states, Alaska and Hawaii",
        source_resolution="one point for each identified system",
        reduction="none: every row of the source, at the source's own longitude and latitude (NAD83 read as "
                  "WGS84, coordinates to 4 decimals)",
        file="geothermal_hydrothermal_sites.json", bytes=size, rows=len(rows),
        legend={"min": st["min"], "max": st["max"], "stops": [1, 10, 100, 1000]}, stats=st,
        total_mwe_mean=round(sum(v for v in vals if v is not None), 1),
        columns_meaning={"value": "Mean Megawatts electric (MWe)", "MWe_P95": "95% likelyhood Megawatts electric "
                         "(MWe)", "MWe_P5": "5% likelyhood Megawatts electric (MWe)",
                         "Temp_C_Li": "Most likely temperature (C) of reservoir",
                         "Temp_C_Mn": "Minimum temperature (C) of reservoir",
                         "Temp_C_Mx": "Maximum temperature (C) of reservoir",
                         "Vol_km3_Li": "Most likely volume of reservoir (km3)",
                         "Vol_km3_Mn": "Minimum volume of reservoir (km3)",
                         "Vol_km3_Mx": "Maximum volume of reservoir (km3)"},
        credit=USGS_CREDIT, terms_in_file="Use constraints: " + xml_field(meta, "useconst") + ". "
                                          + xml_field(meta, "accconst"),
        notes_for_method="Hydrothermal: the natural hot-water systems already identified. USGS's words: \""
                         + xml_field(meta, "abstract") + "\" \"" + xml_field(meta, "purpose") + "\" Each point "
                         "is one identified system with USGS's estimate of the electric power it could "
                         "generate (the mean, and the 95 and 5 percent values). It is an estimate of a "
                         "reservoir's potential, not a power plant and not what is installed. Undiscovered "
                         "systems and enhanced geothermal are not in this layer."))


@builder("geothermal_hydrothermal_favorability")
def build_geo_fav(raw_dir, web_dir, manifest):
    import geopandas as gpd
    row = held(raw_dir, "usgs_FavorabilitySurface_shp")
    for ext in ("dbf", "shx", "prj", "xml"):
        held(raw_dir, f"usgs_FavorabilitySurface_{ext}")
    g = gpd.read_file(row["path"])

    def props(r):
        p = {"name": clean(r["Descript"]), "kind": "relative favorability class",
             "value": clean(r["GRIDCODE"]), "value_unit": "class (the source's GRIDCODE, 1 to 10)"}
        p.update(source_fields(r))
        return p
    size, n, b, how = shapes_any(web_dir, "geothermal_hydrothermal_favorability.json", g, props, 0.03, True)
    meta = row["path"][:-4] + ".xml"
    classes = [{"value": int(r["GRIDCODE"]), "label": clean(r["Descript"])}
               for _, r in g.sort_values("GRIDCODE").iterrows()]
    manifest_put(manifest, base_layer(
        "geothermal_hydrothermal_favorability", "geothermal", "Hydrothermal favorability (western states)",
        "shapes", row, "usgs_FavorabilitySurface_shp", unit="class",
        value_label="relative favorability class (the source's GRIDCODE, with its Descript range)",
        publisher=USGS_PUBLISHER,
        source_title="Geothermal Favorability Map Derived From Logistic Regression Models "
                     "(FavorabilitySurface.shp), DeAngelo and Williams, 2010",
        vintage="the 2008 assessment; the data file is dated 2010 and was released again on ScienceBase on "
                "29 March 2024 (metadata date 20240329)",
        extent="western United States", source_resolution="polygon contours of a favorability surface",
        reduction=how, file="geothermal_hydrothermal_favorability.json", bytes=size, features=n, bounds=b,
        classes=classes, legend={"kinds": [c["label"] for c in classes]}, stats=None,
        credit=USGS_CREDIT, terms_in_file="Use constraints: " + xml_field(meta, "useconst") + ". "
                                          + xml_field(meta, "accconst"),
        notes_for_method="Hydrothermal, undiscovered: where such systems are more likely to be found. USGS's "
                         "words: \"" + xml_field(meta, "abstract") + "\" \"" + xml_field(meta, "purpose") + "\" "
                         "The file gives ten classes (GRIDCODE 1 to 10), each with a range in its Descript "
                         "field (from \"< 0.1\" to \"> 15\"); the file's metadata does not say what unit those "
                         "ranges are in, so the layer shows the class and the range as written and nothing "
                         "more. It covers the western states only. It is not enhanced geothermal."))


EGS_CLASSES = {1: "Class 1 (most favorable)", 2: "Class 2", 3: "Class 3", 4: "Class 4",
               5: "Class 5 (least favorable)",
               999: "Class 999 (temperatures less than 150 C at 10 km depth: not assessed for deep EGS potential)"}


@builder("geothermal_egs_favorability")
def build_egs(raw_dir, web_dir, manifest):
    import geopandas as gpd
    row = held(raw_dir, "nlr_egs")
    shp = find_file(unpack(row["path"]), r"GeothermalLCOE_NoExclusionsforAtlas\.shp$")
    g = gpd.read_file(shp)

    def props(r):
        c = int(r["CLASS"])
        return {"name": EGS_CLASSES.get(c, f"Class {c}"), "kind": "deep enhanced geothermal favorability class",
                "value": c, "value_unit": "class (1 most favorable, 5 least favorable, 999 not assessed)",
                "CLASS": c}
    size, n, b, how = shapes_any(web_dir, "geothermal_egs_favorability.json", g, props, 0.02, True,
                                notice=nlr_notice(raw_dir))
    meta = shp + ".xml"
    counts = {int(k): int(v) for k, v in g["CLASS"].value_counts().items()}
    classes = [{"value": k, "label": EGS_CLASSES.get(k, f"Class {k}"), "features": counts.get(k, 0)}
               for k in sorted(counts)]
    manifest_put(manifest, base_layer(
        "geothermal_egs_favorability", "geothermal", "Deep enhanced geothermal favorability", "shapes", row,
        "nlr_egs", unit="class",
        value_label="relative favorability of the deep enhanced geothermal resource (the source's CLASS)",
        publisher=NLR_PUBLISHER,
        source_title="Geothermal Resource of the United States: favorability of deep enhanced geothermal "
                     "systems (GeothermalLCOE_NoExclusionsforAtlas.shp)",
        vintage="2009 (the file's metadata: analyses \"performed by NREL (2009)\" on temperature at depth from "
                "Southern Methodist University Geothermal Laboratory, Blackwell and Richards, 2009); metadata "
                "dated 23 February 2012",
        extent="contiguous United States (the source: \"Temperature at depth data for deep EGS in Alaska and "
               "Hawaii not available\")",
        source_resolution="polygons of classes", reduction=how,
        file="geothermal_egs_favorability.json", bytes=size, features=n, bounds=b,
        classes=classes, legend={"kinds": [c["label"] for c in classes]}, stats=None,
        credit="NREL (the file's notice); " + NLR_CREDIT, terms_notice=nlr_notice(raw_dir),
        terms_in_file=xml_field(meta, "useconst"),
        notes_for_method="Enhanced geothermal (EGS): heat in deep rock that would need an engineered reservoir, "
                         "kept apart from the hydrothermal layers. The publisher's words: \""
                         + xml_field(meta, "abstract") + "\" The classes rank relative favorability; they are "
                         "not megawatts, and the ranking rests on a 2009 estimate of the levelized cost of "
                         "electricity. The landing page that listed the file (NREL's geothermal maps page) no "
                         "longer exists since the laboratory's site moved to nlr.gov; the file itself is still "
                         "on the laboratory's server."))


@builder("biomass")
def build_biomass(raw_dir, web_dir, manifest):
    import geopandas as gpd
    row = held(raw_dir, "nlr_biomass")
    shp = find_file(unpack(row["path"]), r"SolidBiomass\.shp$")
    g = gpd.read_file(shp)
    parts = ["CropRes", "ForestRes", "PrimMill", "SecMill", "UrbanWood"]

    def props(r):
        p = {"name": f"{clean(r['CNTY_NAME'])}, {clean(r['STATE_NAME'])}", "kind": "county",
             "value": clean(r["Total"]), "value_unit": "dry metric tons/year"}
        p.update(source_fields(r))
        return p
    size, n, b, how = shapes_any(web_dir, "biomass.json", g, props, 0.015, True, notice=nlr_notice(raw_dir))
    meta = shp[:-4] + ".xml"
    vals = [clean(x) for x in g["Total"]]
    st = shape_stats(vals)
    import numpy as np
    q = [float(x) for x in np.percentile([v for v in vals if v is not None], [50, 75, 90, 95, 99])]
    manifest_put(manifest, base_layer(
        "biomass", "biomass", "Solid biomass resources by county", "shapes", row, "nlr_biomass",
        unit="dry metric tons/year",
        value_label="all solid biomass resources of the county (the source's Total)", publisher=NLR_PUBLISHER,
        source_title="Solid biomass resources by county (SolidBiomass.shp), Anelia Milbrandt, NREL",
        vintage="data for 2012 (the file's metadata: \"Data for 2012, in dry metric tons/year\"); published "
                "30 October 2014",
        extent="United States by county: the contiguous states, Alaska and Hawaii",
        source_resolution="county", reduction=how,
        file="biomass.json", bytes=size, features=n, bounds=b,
        legend={"min": st["min"], "max": st["max"], "stops": [round(x, -3) for x in q],
                "stops_are": "the 50th, 75th, 90th, 95th and 99th percentile of the counties"},
        stats=st, national_total=float(sum(v for v in vals if v is not None)),
        parts={c: float(g[c].sum()) for c in parts},
        credit="NREL (the file's notice); " + NLR_CREDIT, terms_notice=nlr_notice(raw_dir),
        terms_in_file=xml_field(meta, "useconst"),
        notes_for_method="This is the Energy Department laboratory's county file, not the Billion-Ton study "
                         "(whose county downloads need an account, so they were left). The publisher's words: \""
                         + xml_field(meta, "abstract") + "\" The Total field: \"This field combines all solid "
                         "biomass resources by county: crop residues, forest residues, primary mill residues, "
                         "secondary mill residues, and urban wood waste. Data for 2012, in dry metric "
                         "tons/year.\" The value is a county's total, so a large county shows more than a small "
                         "one with the same resource a square mile. The outlines are simplified for the web: a "
                         "very small county or independent city (a few square miles) can lose much of its "
                         "shape, and its value is unchanged. Of primary mill residues the file says: "
                         "\"Note that most of this resource is currently utilized.\" It is a 2012 estimate of "
                         "what is generated, not of what is available to a new plant at a price. The landing "
                         "page that listed the file no longer exists since the laboratory's site moved to "
                         "nlr.gov; the file itself is still on the laboratory's server."))


@builder("wind_capacity_factor")
def build_wind_cf(raw_dir, web_dir, manifest):
    row = held(raw_dir, "openei_wind_sc")
    look = held(raw_dir, "openei_wind_sc_lookup")
    with open(look["path"], encoding="utf-8-sig", newline="") as f:
        lookup = {r["reV Column"]: r for r in csv.DictReader(f)}
    cf_words = lookup["capacity_factor_ac"]["Description"].strip()
    extra = ["resource", "capacity_ac_mw", "area_developable_sq_km"]
    rows, vals = [], []
    with open(row["path"], encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            cf = float(r["capacity_factor_ac"])
            rows.append([round(float(r["longitude"]), 3), round(float(r["latitude"]), 3), r["sc_point_gid"],
                         round(cf, 4), round(float(r["resource"]), 2), round(float(r["capacity_ac_mw"]), 1),
                         round(float(r["area_developable_sq_km"]), 1)])
            vals.append(cf)
    size = write_points(web_dir, "wind_capacity_factor.json", ["lon", "lat", "name", "value"] + extra, rows)
    st = shape_stats(vals)
    manifest_put(manifest, base_layer(
        "wind_capacity_factor", "wind", "Wind capacity factor (supply curve sites)", "points", row,
        "openei_wind_sc", unit=lookup["capacity_factor_ac"]["Units"].strip(),
        value_label=cf_words + " (the source's capacity_factor_ac)", publisher=NLR_PUBLISHER,
        source_title="United States Land-based Wind Supply Curves 2024, open access siting, 2035 moderate "
                     "technology, 115 m hub height, 170 m rotor diameter "
                     "(lbw_open_access_2035_moderate_115hh_170rd_supply_curve.csv)",
        vintage="the 2024 edition, published 2025-01-01 on the Open Energy Data Initiative; the file models a "
                "2035 turbine (\"2035 moderate\", 115 m hub height, 170 m rotor diameter); the weather years "
                "behind the mean are not stated in the files",
        extent="contiguous United States, land only, where the open access scenario leaves developable area",
        source_resolution="11.5 km (\"Centroid latitude of the 11.5km grid-cell\")",
        reduction="none: every row of the source at the source's own centroid (coordinates to 3 decimals, the "
                  "capacity factor to 4); the state, county, cost and transmission columns are left out of the web "
                  "file",
        file="wind_capacity_factor.json", bytes=size, rows=len(rows),
        legend={"min": round(st["min"], 4), "max": round(st["max"], 4), "stops": [0.1, 0.2, 0.3, 0.4, 0.5]},
        stats=st, point_spacing_km=11.5,
        columns_meaning={c: lookup[k]["Description"].strip() + " (" + lookup[k]["Units"].strip() + ")"
                         for c, k in (("value", "capacity_factor_ac"), ("resource", "resource"),
                                      ("capacity_ac_mw", "capacity_ac_mw"),
                                      ("area_developable_sq_km", "area_developable_sq_km"))},
        credit=NLR_CREDIT + "; Creative Commons Attribution 4.0",
        notes_for_method="No gross capacity factor raster was found among the files the laboratory offers "
                         "without a key, so this is not one, and it is not named one. It is the capacity factor "
                         "column of the "
                         "laboratory's 2024 land-based wind supply curve, in the source's words: \"" + cf_words
                         + "\", unit \"" + lookup["capacity_factor_ac"]["Units"].strip() + "\". The column's name "
                         "is capacity_factor_ac; the file's column list does not say whether losses are taken "
                         "off, and the technical report (NREL/TP-6A20-91900) was not read, so the layer does "
                         "not say gross or net. It is the modelled output of a 2035 turbine (115 m hub, 170 m "
                         "rotor), one value for each 11.5 km cell that has land left to build on under the "
                         "\"open access\" siting scenario: a cell with no developable land is absent, which is "
                         "an absence of land in the model, not of wind. The resource column is the source's "
                         "\"" + lookup["resource"]["Description"].strip() + "\"."))


EIA_ACK = "Source: U.S. Energy Information Administration"


@builder("oil_gas_basins")
def build_basins(raw_dir, web_dir, manifest):
    import geopandas as gpd
    row = held(raw_dir, "eia_basins")
    shp = find_file(unpack(row["path"]), r"\.shp$")
    g = gpd.read_file(shp)

    def props(r):
        p = {"name": clean(r["NAME"]).title(), "kind": "sedimentary basin",
             "value": clean(r["Area_sq_mi"]), "value_unit": "square miles"}
        p.update(source_fields(r))
        return p
    size, n, b = write_shapes(web_dir, "oil_gas_basins.json", g, props, tolerance=0.01)
    manifest_put(manifest, base_layer(
        "oil_gas_basins", "oil and gas", "Sedimentary basins", "shapes", row, "eia_basins",
        unit="square miles", value_label="area of the basin (the source's Area_sq_mi)",
        publisher="U.S. Energy Information Administration (EIA)",
        source_title="U.S. Sedimentary Basins (SedimentaryBasins_US_May2011_v2.shp)",
        vintage="as of 5-6-2011 (the file's metadata: \"Sedimentary basins associated with the EIA shale "
                "plays as of 5-6-2011\"); file modified 10 March 2016",
        extent="contiguous United States", source_resolution="polygons, no scale stated by the source",
        reduction="each polygon simplified to 0.01 degrees with its topology kept (Douglas-Peucker, one "
                  "feature at a time); coordinates to 3 decimals; no feature dropped",
        file="oil_gas_basins.json", bytes=size, features=n, bounds=b,
        legend={"kinds": ["sedimentary basin"]},
        stats=shape_stats([clean(x) for x in g["Area_sq_mi"]]),
        acknowledgment=EIA_ACK + " (May 2011)",
        notes_for_method="EIA's words: \"Sedimentary basins associated with the EIA shale plays as of "
                         "5-6-2011. Sedimentary basin which do not have shale plays as of publication date are "
                         "not included in this file. Sources for the basins are mostly from the US Geological "
                         "Survey and state agencies such as the WY Geological Survey.\" It is an outline of where "
                         "the basins lie, not a measure of oil or gas in place or of what can be recovered. The "
                         "file's own limit: \"These data and related graphics, if available, are not legal "
                         "documents and are not intended to be used as such. The information contained in these "
                         "data is dynamic and may change over time.\""))


@builder("oil_gas_plays")
def build_plays(raw_dir, web_dir, manifest):
    import geopandas as gpd
    row = held(raw_dir, "eia_plays")
    shp = find_file(unpack(row["path"]), r"\.shp$")
    g = gpd.read_file(shp)

    def props(r):
        p = {"name": clean(r["Shale_play"]), "kind": "tight oil and shale gas play",
             "value": clean(r["Area_sq_mi"]), "value_unit": "square miles"}
        p.update(source_fields(r, skip=("geometry", "Shape_Leng", "Shape_Area", "Age_color")))
        return p
    size, n, b = write_shapes(web_dir, "oil_gas_plays.json", g, props, tolerance=0.01)
    manifest_put(manifest, base_layer(
        "oil_gas_plays", "oil and gas", "Tight oil and shale gas plays", "shapes", row, "eia_plays",
        unit="square miles", value_label="area of the play (the source's Area_sq_mi)",
        publisher="U.S. Energy Information Administration (EIA)",
        source_title="Tight Oil and Shale Gas Plays in the U.S. (ShalePlays_US_EIA_Dec2021.shp)",
        vintage="updated December 2021 (the file's metadata: \"Updated December 2021 with addition of "
                "Spraberry and Wolfcamp Plays in the Permian Basin\")",
        extent="contiguous United States (the source's Lower 48)",
        source_resolution="polygons, \"General areas for shale plays in the U.S.\"",
        reduction="each polygon simplified to 0.01 degrees with its topology kept (Douglas-Peucker, one "
                  "feature at a time); coordinates to 3 decimals; no feature dropped",
        file="oil_gas_plays.json", bytes=size, features=n, bounds=b,
        legend={"kinds": ["tight oil and shale gas play"]},
        stats=shape_stats([clean(x) for x in g["Area_sq_mi"]]),
        acknowledgment=EIA_ACK + " (Dec 2021)",
        notes_for_method="EIA's words: \"General areas for shale plays in the U.S. Source: EIA based on Enverus "
                         "DrillingInfo Inc., the U.S. Geological Survey, publicly available peer-reviewed "
                         "research papers, academic theses, and publications of State Geological Agencies. "
                         "Updated December 2021 with addition of Spraberry and Wolfcamp Plays in the Permian "
                         "Basin.\" A play's outline is a general area, not a reserve estimate and not a lease "
                         "map. The file's own limit: \"None (public use). Users are advised to thoroughly review "
                         "the metadata to understand the appropriate use and limitations of the data.\""))


@builder("offshore_wind_leases")
def build_leases(raw_dir, web_dir, manifest):
    import geopandas as gpd
    row = held(raw_dir, "boem_shapefiles")
    shp = find_file(unpack(row["path"]), r"Offshore_Wind_Leases_outlines\.shp$")
    g = gpd.read_file(shp)

    def props(r):
        lt = clean(r["LEASE_TYPE"])
        p = {"name": clean(r["LEASE_NU_1"]) or clean(r["LEASE_NUMB"]),
             "kind": f"lease area ({lt})" if lt else "lease area",
             "value": clean(r["ACRES"]) or None, "value_unit": "acres"}
        p.update(source_fields(r, skip=("geometry", "Shape_Leng", "Shape_Area")))
        return p
    size, n, b = write_shapes(web_dir, "offshore_wind_leases.json", g, props, tolerance=0, digits=4)
    kinds = sorted({f"lease area ({clean(x)})" for x in g["LEASE_TYPE"] if clean(x)})
    manifest_put(manifest, base_layer(
        "offshore_wind_leases", "offshore wind", "Offshore wind lease areas", "shapes", row,
        "boem_shapefiles", unit="acres", value_label="area of the lease (the source's ACRES)",
        publisher="Bureau of Ocean Energy Management (BOEM)",
        source_title="Renewable Energy Leases and Planning Areas, All Shapefile "
                     "(Offshore_Wind_Leases_outlines.shp)",
        vintage="last file update on 02/05/2025 (BOEM's page); the archive on BOEM's server is dated "
                "6 August 2025",
        extent="US Outer Continental Shelf: Atlantic, Gulf and Pacific",
        source_resolution="lease outlines built from Outer Continental Shelf blocks",
        reduction="coordinates to 4 decimals (about 10 m), so that a narrow easement keeps its shape; no "
                  "other simplification; no feature dropped",
        file="offshore_wind_leases.json", bytes=size, features=n, bounds=b,
        legend={"kinds": kinds},
        stats=shape_stats([clean(x) for x in g["ACRES"] if x]),
        zero_acres=int((g["ACRES"] == 0).sum()),
        notes_for_method="BOEM's words: \"Boundaries of renewable energy lease areas, wind planning areas, and "
                         "marine hydrokinetic planning areas (last file update on 02/05/2025).\" These are the "
                         "outlines of leases, easements and research leases BOEM has issued; a lease is a right "
                         "to propose a project, not a wind farm and not a measure of the wind. A lease's status "
                         "may have changed since the file's date. Where the source's ACRES field is 0 (some "
                         "easements and cable routes) no area is shown: the field is kept as the source wrote "
                         "it and the value is left empty."))


@builder("offshore_wind_planning_areas")
def build_planning(raw_dir, web_dir, manifest):
    import geopandas as gpd
    row = held(raw_dir, "boem_planning")
    meta_row = held(raw_dir, "boem_planning_meta")
    with open(meta_row["path"], encoding="utf-8") as f:
        meta = json.load(f)
    g = gpd.read_file(row["path"])
    if len(g) == 0:
        raise RuntimeError("the planning area service returned no feature")
    edit = meta.get("editingInfo", {}).get("dataLastEditDate")
    edited = dt.datetime.fromtimestamp(edit / 1000, dt.timezone.utc).strftime("%Y-%m-%d") if edit else None

    def props(r):
        cat = clean(r.get("CATEGORY1"))
        p = {"name": clean(r.get("ADDITIONAL_INFORMATION")) or clean(r.get("PROTRACTION_NUMBER")),
             "kind": f"planning area ({cat})" if cat else "planning area",
             "value": None, "value_unit": None}
        p.update(source_fields(r, skip=("geometry", "Shape__Area", "Shape__Length")))
        return p
    size, n, b = write_shapes(web_dir, "offshore_wind_planning_areas.json", g, props, tolerance=0, digits=4)
    kinds = sorted({f"planning area ({clean(x)})" for x in g["CATEGORY1"] if clean(x)}) or ["planning area"]
    manifest_put(manifest, base_layer(
        "offshore_wind_planning_areas", "offshore wind",
        "Offshore wind planning areas (rescinded July 30, 2025)", "shapes", row, "boem_planning",
        unit="", value_label="no quantity in the source",
        publisher="Bureau of Ocean Energy Management (BOEM)",
        source_title=meta.get("name", "Offshore Wind Planning Area Outlines"),
        vintage=f"data last edited {edited} (the service's own date); BOEM names the layer "
                f"\"{meta.get('name')}\"",
        extent="US Outer Continental Shelf",
        source_resolution="planning area outlines built from Outer Continental Shelf blocks",
        reduction="coordinates to 4 decimals (about 10 m), so that a narrow easement keeps its shape; no "
                  "other simplification; no feature dropped",
        file="offshore_wind_planning_areas.json", bytes=size, features=n, bounds=b,
        legend={"kinds": kinds}, stats=None, source_rows=int(len(g)),
        area_status={str(k): int(v) for k, v in g["AREA_STATUS"].value_counts().items()},
        notes_for_method="BOEM's words: \"These data are an outline version of BOEM's offshore wind planning "
                         "areas. These areas represent the current investigations by BOEM for new areas of "
                         "interest in wind energy development. Individual blocks within each wind planning area "
                         "dissolved into single polygons to generate the area outlines.\" BOEM's own name for "
                         f"the layer is \"{meta.get('name')}\": the planning areas were rescinded, and are shown "
                         "as the record of where BOEM had been looking, not as areas open to leasing. The "
                         "archive BOEM offers for download holds the leases only; the planning areas come from "
                         "BOEM's public map service, which BOEM's page links. The service holds 19 rows; one "
                         "has no shape and no name and is not drawn. Each area's AREA_STATUS (Active or "
                         "Inactive) is the source's own field and was last edited before the rescission."))


# ---------------------------------------------------------------------------------------------------------
# Session 159: hydropower, in the form its publisher publishes it. Oak Ridge National Laboratory's two
# assessments for the Energy Department, on HydroSource. Nothing is re-estimated here and nothing is spread
# over an area it was not published for: a dam is a point with the file's own capacity, generation and
# capacity factor; new stream-reach development is the file's own total a HUC10 watershed.
# ---------------------------------------------------------------------------------------------------------

ORNL_PUBLISHER = "Oak Ridge National Laboratory (HydroSource), for the U.S. Department of Energy"
ORNL_FETCHED = ("The file was fetched by its own address, the one HydroSource's dataset page prints, with a plain "
                "request; the download form on that page (a name, an e-mail address, a company, an occupation) "
                "was not filled and nothing was sent to it (the owner's ruling of 8 October 2026).")


def readme_section(path, heading):
    """The lines under a heading of a publisher's readme, up to the next blank line, as one string."""
    with open(path, encoding="utf-8", errors="replace") as f:
        lines = [x.strip() for x in f.read().splitlines()]
    out, on = [], False
    for x in lines:
        if on and not x:
            if out:
                break
            continue
        if on:
            out.append(x)
        if x.rstrip(":").strip().lower() == heading.lower():
            on = True
    return " ".join(out).strip()


def as_number(text):
    """A cell of a publisher's file as the number it writes; an empty cell is None (never a zero)."""
    s = (text or "").strip()
    if not s:
        return None
    v = float(s)
    return int(v) if v.is_integer() and "." not in s and "e" not in s.lower() else v


@builder("hydropower_npd")
def build_hydro_npd(raw_dir, web_dir, manifest):
    row = held(raw_dir, "ornl_npd")
    fields_row = held(raw_dir, "ornl_npd_fields")
    readme_row = held(raw_dir, "ornl_npd_readme")
    with open(fields_row["path"], encoding="utf-8-sig", newline="") as f:
        fields = {r["Field Name"].strip(): r for r in csv.DictReader(f)}
    with open(row["path"], encoding="utf-8-sig", newline="") as f:
        src = list(csv.DictReader(f))
    if not src:
        raise RuntimeError("the non-powered dams file holds no row")
    # the file's own column for longitude is "long"; its field descriptions call it "lon"
    lon_col = "long" if "long" in src[0] else "lon"
    # the owner's name (the file's dam_owner) is left out of the web file: some owners are private persons
    extra = ["gen_mwh_yr", "cf_yr", "state_abbr", "waterway", "nididfull", "county", "head_ft_yr", "q30_cfs",
             "inputdata"]
    numeric = {"gen_mwh_yr", "cf_yr", "head_ft_yr", "q30_cfs"}
    rows, caps, gens = [], [], []
    for r in src:
        cap = as_number(r["cap_mw"])
        rows.append([round(float(r[lon_col]), 4), round(float(r["lat"]), 4), clean(r["dam_name"]), cap]
                    + [(as_number(r[c]) if c in numeric else clean(r[c])) for c in extra])
        if cap is not None:
            caps.append(cap)
        g = as_number(r["gen_mwh_yr"])
        if g is not None:
            gens.append(g)
    size = write_points(web_dir, "hydropower_npd.json", ["lon", "lat", "name", "value"] + extra, rows)
    st = shape_stats(caps)

    def meaning(c):
        d = fields[c]
        u = d["Units"].strip()
        return d["Description"].strip() + (f" ({u})" if u and u != "unitless" else "")
    abstract = readme_section(readme_row["path"], "Abstract")
    method = readme_section(readme_row["path"], "Methodology")
    citation = readme_section(readme_row["path"], "Citation")
    manifest_put(manifest, base_layer(
        "hydropower_npd", "hydropower", "Hydropower potential at non-powered dams", "points", row, "ornl_npd",
        unit=fields["cap_mw"]["Units"].strip(),
        value_label=fields["cap_mw"]["Description"].strip() + " (the source's cap_mw)",
        publisher=ORNL_PUBLISHER,
        source_title="Technical Potential for Hydropower Capacity at Non-powered Dams (TechPotentialNPDs.csv), "
                     "Oak Ridge National Laboratory and Idaho National Laboratory, "
                     "doi:10.21951/HydroCapacity_NPD/2570407",
        vintage="published 3 October 2024, revised 25 June 2025 (HydroSource's page; the file on its server is "
                "dated 25 June 2025); the file's readme: \"These estimates represent the conditions over the "
                "historical period of 1980-2015\"",
        extent="conterminous United States (the readme: \"2,616 NPDs in the conterminous US\")",
        source_resolution="one row for each dam, at the dam's own latitude and longitude",
        reduction="none: every row of the source at the source's own longitude and latitude (North American "
                  "Datum 83 read as WGS84, coordinates to 4 decimals); the capacity, the generation and the "
                  "capacity factor are the file's own figures, unconverted; the monthly columns are left out of "
                  "the web file",
        file="hydropower_npd.json", bytes=size, rows=len(rows),
        legend={"min": st["min"], "max": st["max"], "stops": [0.1, 1, 10, 100]}, stats=st,
        total_cap_mw=round(sum(caps), 3), total_gen_mwh_yr=round(sum(gens), 3),
        # what the hover lists under the capacity: the file's own fields, named in the file's own words
        hover_fields=[["gen_mwh_yr", fields["gen_mwh_yr"]["Description"].strip() + ", "
                       + fields["gen_mwh_yr"]["Units"].strip()],
                      ["cf_yr", fields["cf_yr"]["Description"].split(" (")[0].strip()],
                      ["waterway", "River"], ["state_abbr", "State"],
                      ["nididfull", "National Inventory of Dams id"]],
        columns_meaning={"value": meaning("cap_mw"), "gen_mwh_yr": meaning("gen_mwh_yr"), "cf_yr": meaning("cf_yr"),
                         "state_abbr": meaning("state_abbr"), "waterway": meaning("waterway"),
                         "nididfull": meaning("nididfull"), "county": meaning("county"),
                         "head_ft_yr": meaning("head_ft_yr"), "q30_cfs": meaning("q30_cfs"),
                         "inputdata": meaning("inputdata")},
        credit="Citation, as the publisher asks: " + citation,
        notes_for_method="Oak Ridge's words (the file's readme): \"" + abstract + "\" \"" + method + "\" A "
                         "non-powered dam is a dam that exists and has no power plant. Each point is one such "
                         "dam with the laboratory's own estimate of the capacity a retrofit could have, the "
                         "generation of an average year and the capacity factor, from a model of 1980 to 2015 "
                         "flows; nothing is estimated again here. It is a technical potential: not a plant, not "
                         "a proposal, not a finding that a retrofit would pay, and not every dam in the country "
                         "(the file holds the dams an earlier assessment gave at least 100 kW). Alaska and "
                         "Hawaii are not in the file. The publisher is Oak Ridge National Laboratory, with "
                         "Idaho National Laboratory; it is not a USGS layer. " + ORNL_FETCHED))


def workbook_sheets(path, names):
    """The named sheets of a publisher's workbook as lists of rows, each cell in its own column (an empty cell
    is None; empty cells at the end of a row and wholly empty rows are dropped)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for n in names:
        rows = []
        for r in wb[n].iter_rows(values_only=True):
            r = list(r)
            while r and r[-1] is None:
                r.pop()
            if r:
                rows.append(r)
        out[n] = rows
    wb.close()
    return out


NSD_NO_VALUE = (-9999, -999)   # the file's own codes where it gives no figure for a watershed


@builder("hydropower_nsd")
def build_hydro_nsd(raw_dir, web_dir, manifest):
    """New stream-reach development, at the unit Oak Ridge publishes it at: the HUC10 watershed, whose outlines
    are in the publisher's own shapefile. A watershed's figure is the file's own total for it (P_MW_Sum); a
    watershed for which the file gives no capacity (its code -9999) is not drawn. Nothing is computed here."""
    import geopandas as gpd
    import numpy as np
    row = held(raw_dir, "ornl_nsd_zip")
    book_row = held(raw_dir, "ornl_nsd_xlsx")
    shp = find_file(unpack(row["path"]), r"ORNL_NHAAP_NSD_SR_All_v1\.shp$")
    sheets = workbook_sheets(book_row["path"], ["NSD", "NSD_Columns", "NSD_Metadata"])
    cols = {r[0]: {"long": r[1], "units": r[2], "how": r[3], "words": re.sub(r"\s+", " ", str(r[4])).strip()}
            for r in sheets["NSD_Columns"][1:] if len(r) >= 5 and all(x is not None for x in r[:5])}
    # the metadata sheet's contact and creator rows hold persons' e-mail addresses: they are never read into a layer
    meta = {str(r[0]).strip(): re.sub(r"\s+", " ", str(r[1])).strip() for r in sheets["NSD_Metadata"]
            if len(r) >= 2 and r[0] is not None and r[1] is not None and str(r[0]).strip() not in ("contact", "creator")}
    # every watershed's attributes (no outline read): the counts of what the file gives and does not give
    import pyogrio
    att = pyogrio.read_dataframe(shp, read_geometry=False)
    with_value = att[att["NUMREACH"] > 0]
    if not (with_value["P_MW_Sum"] > 0).all() or (att[att["NUMREACH"] <= 0]["P_MW_Sum"] > 0).any():
        raise RuntimeError("the watershed file's capacity and its count of stream-reaches do not agree")
    # the workbook beside the archive lists the same watersheds with the same totals: read both, compare
    book = [r for r in sheets["NSD"] if len(r) >= 4 and isinstance(r[0], (int, float)) and isinstance(r[3], (int, float))]
    book_total = float(sum(r[3] for r in book))
    if len(book) != len(with_value) or abs(book_total - float(with_value["P_MW_Sum"].sum())) > 1e-3:
        raise RuntimeError(f"the workbook ({len(book)} watersheds, {book_total} MW) and the shapefile "
                           f"({len(with_value)}, {float(with_value['P_MW_Sum'].sum())} MW) do not agree")
    g = gpd.read_file(shp, where="NUMREACH > 0")
    if len(g) != len(with_value):
        raise RuntimeError(f"{len(g)} outlines read, {len(with_value)} watersheds hold a capacity")
    keep = ["HUC10", "HUC10_NAME", "NUMREACH", "P_MW_Sum", "E_MWh_Sm", "Cf_yr", "H_ft_Avg", "Q30cfsAg"]

    def props(r):
        p = {"name": f"{clean(r['HUC10_NAME'])} ({clean(r['HUC10'])})", "kind": "HUC10 watershed",
             "value": clean(r["P_MW_Sum"]), "value_unit": cols["P_MW_Sum"]["units"]}
        for c in keep:
            v = clean(r[c])
            p[c] = None if isinstance(v, (int, float)) and v in NSD_NO_VALUE else v
        return p
    size, n, b, how = shapes_any(web_dir, "hydropower_nsd.json", g, props, 0.01, True)
    vals = [float(x) for x in g["P_MW_Sum"]]
    st = shape_stats(vals)
    energy = [float(x) for x in g["E_MWh_Sm"] if x is not None and np.isfinite(x) and x not in NSD_NO_VALUE]
    # a watershed whose annual energy the shapefile leaves empty, and what the workbook writes in the same cell
    no_energy = sorted(str(h) for h, x in zip(g["HUC10"], g["E_MWh_Sm"]) if x is None or not np.isfinite(x))
    book_e = {str(int(r[0])).zfill(10): r[4] for r in book if len(r) >= 5}
    book_says = [("%g" % book_e[h]) if isinstance(book_e.get(h), (int, float)) else "nothing" for h in no_energy]
    label = {c: f"{cols[c]['long']}" + (f", {cols[c]['units']}" if cols[c]["units"] not in ("n/a", "ratio") and
                                        cols[c]["units"].lower() != cols[c]["long"].lower() else "")
             for c in ("NUMREACH", "E_MWh_Sm", "Cf_yr", "H_ft_Avg", "Q30cfsAg")}
    manifest_put(manifest, base_layer(
        "hydropower_nsd", "hydropower", "Hydropower potential of new stream-reach development", "shapes", row,
        "ornl_nsd_zip", unit=cols["P_MW_Sum"]["units"],
        value_label=cols["P_MW_Sum"]["words"] + " (the source's P_MW_Sum, a HUC10 total)",
        publisher=ORNL_PUBLISHER,
        source_title=meta["title"].rstrip(".") + " (ORNL_NHAAP_NSD_SR_All_v1.shp, version " + meta["version"]
                     + "), Oak Ridge National Laboratory",
        vintage="published 31 December 2014 (HydroSource's page); the file's own metadata is dated 10 January "
                "2014 and its workbook was created on " + meta["creation date"][:10] + ", version "
                + meta["version"] + "; the years of flow behind the estimate are not stated in the files",
        extent="conterminous United States, by HUC10 watershed (the file's words: \"" + meta["title"] + "\")",
        source_resolution="the HUC10 watershed of the Watershed Boundary Dataset: " + f"{len(att):,}"
                          + " watersheds in the file, each a total over its stream-reaches of more than 1 MW",
        reduction=how + f"; only the {n:,} watersheds for which the file gives a capacity are in the web file",
        file="hydropower_nsd.json", bytes=size, features=n, bounds=b,
        legend={"min": st["min"], "max": st["max"], "stops": [10, 25, 50, 100, 250]}, stats=st,
        hover_fields=[[c, label[c]] for c in ("NUMREACH", "E_MWh_Sm", "Cf_yr")],
        watersheds_in_file=int(len(att)), watersheds_with_capacity=int(len(with_value)),
        watersheds_no_reach=int((att["NUMREACH"] == 0).sum()),
        watersheds_coded_minus_999=int((att["NUMREACH"] == -999).sum()),
        total_p_mw=round(float(sum(vals)), 3), total_reaches=int(with_value["NUMREACH"].sum()),
        total_e_mwh=round(float(sum(energy)), 3), watersheds_with_energy=len(energy),
        watersheds_without_energy=no_energy,
        columns_meaning={c: f"{cols[c]['words']} ({cols[c]['units']}, {cols[c]['how']})" for c in keep[2:]},
        credit="The file asks: \"The ORNL should be acknowledged as the data source in products derived from "
               "these data.\" Source: Oak Ridge National Laboratory, New Stream-reach Development Resource "
               "Assessment, 2014. " + meta["dataset credit"],
        terms_in_file=meta["use constraints"],
        notes_for_method="Oak Ridge's words (the file's metadata): \"" + meta["description"] + "\" New "
                         "stream-reach development is the hydropower that could be built on stretches of "
                         "river that have no dam and no plant today. Oak Ridge publishes it summed to the "
                         "HUC10 watershed (a drainage area of the federal Watershed Boundary Dataset), and that "
                         "is what is drawn: each watershed is shaded by the file's own total for it, in the "
                         "file's words \"" + cols["P_MW_Sum"]["words"] + "\". The figure belongs to the "
                         "watershed as a whole: the file does not say where in it the reaches lie, so nothing "
                         "is drawn at a site and nothing is spread or estimated again here. A larger watershed "
                         "can show more than a smaller one with the same rivers. Of the "
                         + f"{len(att):,} watersheds in the file, {len(with_value):,} hold a capacity and are "
                         f"drawn; {int((att['NUMREACH'] == 0).sum()):,} count no stream-reach above 1 MW and "
                         f"{int((att['NUMREACH'] == -999).sum()):,} carry the file's code -999 in that count "
                         "(the file does not say what the code stands for); both carry -9999 as their "
                         "capacity and are not drawn. "
                         + (f"For {len(no_energy):,} of the watersheds drawn ({', '.join(no_energy)}) the "
                            "shapefile holds no annual energy and no capacity factor, where the publisher's "
                            f"workbook writes {', '.join(sorted(set(book_says)))}: none is shown. "
                            if no_energy else "")
                         + "It is a potential before any study of a site. The file's own limit: \""
                         + meta["use constraints"] + "\" Alaska and Hawaii are not in this file. The publisher "
                         "is Oak Ridge National Laboratory; it is not a USGS layer. The outlines are "
                         "simplified for the web. " + ORNL_FETCHED))


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--pull", action="store_true")
    ap.add_argument("--build", action="store_true")
    ap.add_argument("--check-terms", action="store_true")
    ap.add_argument("--method-doc", action="store_true",
                    help="rewrite the layer sections of docs/methods/resources.md from the manifest")
    ap.add_argument("--only", default="")
    ap.add_argument("--tables", action="store_true",
                    help="the vector layers as entities tables, a trial into --out-dir")
    ap.add_argument("--registry", action="store_true",
                    help="add this connector's sources to warehouse/metadata/sources.csv of this copy")
    ap.add_argument("--out-dir", default="")
    ap.add_argument("--raw-dir", default=RAW_DIR)
    args = ap.parse_args()
    only = [x for x in args.only.split(",") if x]
    web_dir = os.path.join(args.out_dir, "resources-data") if args.out_dir else WEB_DIR
    manifest = os.path.join(args.out_dir, "manifest.json") if args.out_dir else MANIFEST
    failed = []
    if args.pull:
        sess = http_session()
        for key, src in SOURCES.items():
            if only and not (set(only) & set(src["layers"])) and key not in only:
                continue
            try:
                for more in TERMS[src["terms"]].get("also", []):
                    fetch_terms(sess, args.raw_dir, src["source"], more)
                terms = fetch_terms(sess, args.raw_dir, src["source"], src["terms"])
                fetch(sess, args.raw_dir, src["source"], ";".join(src["layers"]), src["url"], src["file"], terms)
            except Exception as e:  # a pull that fails is recorded with its exact error and left
                log(f"FAILED pull {key}: {type(e).__name__}: {e}")
                failed.append(key)
        log(f"ledger total: {ledger_total(read_ledger(args.raw_dir))} bytes of {CEILING_BYTES}")
    if args.build:
        man = read_manifest(manifest)
        if any(m["id"] in RETIRED_IDS for m in man.get("missing", [])):  # an id that stands for no layer now
            man["missing"] = [m for m in man["missing"] if m["id"] not in RETIRED_IDS]
            write_manifest(manifest, man)
        held_ids = {l["id"] for l in man["layers"]}
        for id, (group, title) in EXPECTED.items():
            if id not in held_ids:
                manifest_missing(manifest, id, group, title, NOT_HELD.get(id, "not built yet"))
        for id, fn in BUILDERS.items():
            if only and id not in only:
                continue
            try:
                fn(args.raw_dir, web_dir, manifest)
            except Exception as e:
                import traceback
                traceback.print_exc()
                log(f"FAILED build {id}: {type(e).__name__}: {e}")
                failed.append(id)
    if args.check_terms:
        failed += check_terms(args.raw_dir, manifest)
    if args.method_doc:
        write_method_doc(manifest)
    if args.registry:
        write_registry(args.raw_dir)
    if args.tables:
        if not args.out_dir:
            log("--tables is a trial: give --out-dir (the tables are not written into warehouse/output here)")
            return 2
        build_tables(args.raw_dir, args.out_dir)
    if failed:
        log(f"failed: {failed}")
        return 1
    return 0


METHOD_DOC = os.path.join(REPO, "docs", "methods", "resources.md")
METHOD_MARK = "<!-- The sections below are written from the manifest by resource_layers.py --method-doc -->"


def fmt_num(x):
    if x is None:
        return "none"
    if isinstance(x, int) or float(x).is_integer():
        return f"{int(x):,}"
    return f"{x:,.4g}" if abs(x) < 1000 else f"{x:,.0f}"


def write_method_doc(manifest, path=METHOD_DOC):
    """One section a layer, from the manifest, under the hand-written head of docs/methods/resources.md."""
    man = read_manifest(manifest)
    with open(path, encoding="utf-8") as f:
        head = f.read().split(METHOD_MARK)[0].rstrip() + "\n\n"
    out = [head + METHOD_MARK + "\n"]
    for l in man["layers"]:
        out.append(f"## {l['title']} (`{l['id']}`)\n")
        out.append(f"- **What it is.** {l['notes_for_method']}")
        out.append(f"- **Publisher.** {l['publisher']}. {l['source_title']}.")
        out.append(f"- **File.** `{l['source_file']}` in the raw store, {l['source_bytes']:,} bytes, sha256 "
                   f"`{l['source_sha256'][:16]}`, retrieved {l['retrieved_at_utc']} from <{l['source_url']}>.")
        out.append(f"- **Vintage.** {l['vintage']}.")
        out.append(f"- **Extent.** {l['extent']}. Source resolution: {l['source_resolution']}.")
        out.append(f"- **Reduction.** {l['reduction']}.")
        if l["kind"] == "grid":
            st = l["source_stats"]
            out.append(f"- **Unit.** {l['unit']} ({l['value_label']}). The source's own cells: minimum "
                       f"{fmt_num(st['min'])}, mean {fmt_num(st['mean'])}, maximum {fmt_num(st['max'])} over "
                       f"{st['n_valid']:,} cells with a value.")
            out.append("- **Levels.**\n")
            out.append("  | cell (degrees) | file | columns x rows | cells with a value | bytes | mean weighted by "
                       "source cells | the source's mean over the same cells | share of the source's cells kept |")
            out.append("  |---|---|---|---|---|---|---|---|")
            for v in l["levels"]:
                c = [x for x in l["checks"] if x["cell_deg"] == v["cell_deg"]][0]
                out.append(f"  | {v['cell_deg']} | `{v['file']}` | {v['ncols']} x {v['nrows']} | "
                           f"{c['cells_with_value']:,} | {v['bytes']:,} | {c['mean_weighted_by_source_cells']:.5g} | "
                           f"{c['source_mean_of_kept_cells']:.5g} | {c['share_of_source_cells_kept']:.2%} |")
            for x in l.get("other_extents", []):
                files = ", ".join(f"`{v['file']}` ({v['cell_deg']} degrees, {v['bytes']:,} bytes)"
                                  for v in x["levels"])
                sx = x["source_stats"]
                out.append(f"- **Also held: {x['extent']}.** {files}. The source's own cells there: minimum "
                           f"{fmt_num(sx['min'])}, mean {fmt_num(sx['mean'])}, maximum {fmt_num(sx['max'])} over "
                           f"{sx['n_valid']:,} cells.")
        else:
            n = l.get("features", l.get("rows"))
            what = "features" if l["kind"] == "shapes" else "rows"
            line = f"- **Web file.** `{l['file']}`, {l['bytes']:,} bytes, {n:,} {what}."
            st = l.get("stats")
            if st:
                line += (f" {l['value_label']}, in {l['unit']}: minimum {fmt_num(st['min'])}, mean "
                         f"{fmt_num(st['mean'])}, maximum {fmt_num(st['max'])} over {st['n']:,} {what}.")
            out.append(line)
        if l.get("classes"):
            out.append("- **Classes, in the source's words.** " + "; ".join(
                f"{c['value']}: {c['label']}" for c in l["classes"]) + ".")
        out.append(f"- **Terms.** <{l['terms_url']}> (saved as `{l['terms_file']}`, sha256 "
                   f"`{l['terms_sha256'][:16]}`): \"{l['terms_quote']}\"")
        for t in l.get("terms_also", []):
            out.append(f"  Also <{t['url']}> (saved as `{t['file']}`, sha256 `{t['sha256'][:16]}`): "
                       f"\"{t['quote']}\"")
        if l.get("terms_in_file"):
            out.append(f"- **The file's own notice.** \"{l['terms_in_file']}\"")
        if l.get("terms_notice"):
            out.append("- **The laboratory's notice** travels with the data: it is printed in full at the end of "
                       "this document.")
        if l.get("credit"):
            out.append(f"- **Credit.** {l['credit']}.")
        if l.get("acknowledgment"):
            out.append(f"- **Acknowledgment.** {l['acknowledgment']}.")
        out.append("")
    if man.get("missing"):
        out.append("## Not held\n")
        for m in man["missing"]:
            out.append(f"- **{m['title']}** (`{m['id']}`): {m['reason']}")
        out.append("")
    out.append("## Looked for and left\n")
    out.append("A source that asks for a login, an account or a person's details before it hands over a file is "
               "recorded and left. Nothing below was downloaded.\n")
    for k, x in LEFT_SOURCES.items():
        files = "; ".join(f"`{f}` ({b:,} bytes by the server's own count)" for f, b in x.get("files", {}).items())
        out.append(f"- **{x['title']}.** {x['publisher']}. <{x['landing']}>. Left because {x['why_left']}."
                   + (f" Files: {files}." if files else "")
                   + (f" Its data use policy: <{x['terms']}>." if x.get("terms") else ""))
    out.append("")
    notice = next((l["terms_notice"] for l in man["layers"] if l.get("terms_notice")), None)
    if notice:
        out.append("## The laboratory's notice, in full\n")
        out.append("The National Laboratory of the Rockies (until 2025 the National Renewable Energy Laboratory) "
                   "grants the use of its data \"provided that this entire notice appears in all copies of the "
                   f"Data\". The notice, from <{TERMS['nlr']['url']}> as saved on the day of the pull:\n")
        out.append("> " + notice)
        out.append("")
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(out))
    os.replace(tmp, path)
    log(f"method document: {len(man['layers'])} layers, {len(man.get('missing', []))} not held")


# The source registry's rows (warehouse/metadata/sources.csv), one a source file. No table yet: the layers
# are web files; the tables column is filled when the vector layers become entities tables.
REGISTRY = {
    "eia:maps:sedimentary_basins": ("eia_basins", "U.S. Energy Information Administration (EIA)",
                                    "U.S. Sedimentary Basins, shapefile (EIA maps, layer information for "
                                    "interactive state maps)"),
    "eia:maps:tight_oil_shale_gas_plays": ("eia_plays", "U.S. Energy Information Administration (EIA)",
                                           "Tight Oil and Shale Gas Plays in the U.S., Lower 48, shapefile, "
                                           "updated December 2021"),
    "boem:renewable_energy_leases": ("boem_shapefiles", "Bureau of Ocean Energy Management (BOEM)",
                                     "Renewable Energy Leases and Planning Areas, all shapefiles (last file "
                                     "update 02/05/2025)"),
    "boem:wind_planning_area_outlines": ("boem_planning", "Bureau of Ocean Energy Management (BOEM)",
                                         "Offshore Wind Planning Area Outlines (Rescinded July 30, 2025), "
                                         "public feature service"),
    "nlr:gis:wind_toolkit_mean_wind_speed": ("nlr_wind", "National Laboratory of the Rockies (NLR, formerly NREL)",
                                             "WIND Toolkit, Multi-year (2007-2013) Annual Average Wind Speed at "
                                             "all heights, GeoTIFF, 2 km"),
    "nlr:gis:nsrdb_psm3_ghi": ("nlr_ghi", "National Laboratory of the Rockies (NLR, formerly NREL)",
                               "NSRDB Physical Solar Model version 3 Global Horizontal Irradiance, multi-year "
                               "(1998-2016) annual and monthly averages, GeoTIFF"),
    "nlr:gis:nsrdb_psm3_dni": ("nlr_dni", "National Laboratory of the Rockies (NLR, formerly NREL)",
                               "NSRDB Physical Solar Model version 3 Direct Normal Irradiance, multi-year "
                               "(1998-2016) annual and monthly averages, GeoTIFF"),
    "nlr:gis:deep_egs_favorability": ("nlr_egs", "National Laboratory of the Rockies (NLR, formerly NREL)",
                                      "Favorability of deep enhanced geothermal systems, lower 48, shapefile "
                                      "(2009)"),
    "nlr:gis:solid_biomass": ("nlr_biomass", "National Laboratory of the Rockies (NLR, formerly NREL)",
                              "Solid biomass resources by county, shapefile (data for 2012, published 2014)"),
    "nlr:oedi:land_based_wind_supply_curves_2024": (
        "openei_wind_sc", "National Laboratory of the Rockies (NLR, formerly NREL)",
        "United States Land-based Wind Supply Curves 2024, open access, 2035 moderate, 115 m hub height, 170 m "
        "rotor diameter (Open Energy Data Initiative submission 8314)"),
    "usgs:identified_geothermal_systems": (
        "usgs_IdentifiedGeothermalSystems_shp", "U.S. Geological Survey (USGS)",
        "Identified Moderate and High Temperature Geothermal Systems (2008 assessment; ScienceBase, "
        "doi:10.5066/P1YJWMFC)"),
    "usgs:geothermal_favorability": (
        "usgs_FavorabilitySurface_shp", "U.S. Geological Survey (USGS)",
        "Geothermal Favorability Map Derived From Logistic Regression Models (2008 assessment; ScienceBase, "
        "doi:10.5066/P137NMXE)"),
    # session 159: Oak Ridge's two hydropower assessments (HydroSource)
    "ornl:hydrosource:npd_technical_potential": (
        "ornl_npd", "Oak Ridge National Laboratory (HydroSource), for the U.S. Department of Energy",
        "Technical Potential for Hydropower Capacity at Non-powered Dams (2024, revised 25 June 2025; "
        "doi:10.21951/HydroCapacity_NPD/2570407)"),
    "ornl:hydrosource:nsd_huc10": (
        "ornl_nsd_zip", "Oak Ridge National Laboratory (HydroSource), for the U.S. Department of Energy",
        "Hydropower Potential from New Stream-Reach Development for the Conterminous United States, aggregated "
        "summary of HUC10 watersheds (2014, version 1)"),
}
REGISTRY_PAGES = {
    "ornl:hydrosource:npd_technical_potential": "https://hydrosource.ornl.gov/data/datasets/hydropower-capacity-us-npd/",
    "ornl:hydrosource:nsd_huc10": "https://hydrosource.ornl.gov/data/datasets/"
                                  "hydropower-potential-new-stream-reach-development-conterminous-united-states/",
    "usgs:identified_geothermal_systems": "https://www.sciencebase.gov/catalog/item/6606f534d34e4df16bd58277",
    "usgs:geothermal_favorability": "https://www.sciencebase.gov/catalog/item/6606ed51d34e4df16bd58251",
    "nlr:oedi:land_based_wind_supply_curves_2024": "https://data.openei.org/submissions/8314",
    "boem:renewable_energy_leases": "https://www.boem.gov/renewable-energy/mapping-and-data/renewable-energy-gis-data",
    "boem:wind_planning_area_outlines":
        "https://www.boem.gov/renewable-energy/mapping-and-data/renewable-energy-gis-data",
}


def write_registry(raw_dir, path=os.path.join(REPO, "warehouse", "metadata", "sources.csv")):
    """Add this connector's sources to the registry of the copy it runs in. A row already there keeps its
    first_seen and its tables; nothing is removed."""
    with open(path, encoding="utf-8", newline="") as f:
        rd = csv.DictReader(f)
        cols = rd.fieldnames
        rows = {r["source"]: r for r in rd}
    n = 0
    for sid, (key, publisher, report) in REGISTRY.items():
        try:
            h = held(raw_dir, key)
        except FileNotFoundError:
            continue
        day = h["retrieved_at_utc"][:10]
        old = rows.get(sid, {})
        rows[sid] = {"source": sid, "publisher": publisher, "report": report,
                     "report_url": REGISTRY_PAGES.get(sid, h["url"]), "document_list": h["url"],
                     "license": "public", "tables": old.get("tables", ""),
                     "first_seen": old.get("first_seen", day), "last_seen": day}
        n += 1
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for sid in sorted(rows):
            w.writerow(rows[sid])
    os.replace(tmp, path)
    log(f"source registry: {n} rows of this connector in {path}")


# ---------------------------------------------------------------------------------------------------------
# The vector layers as entities tables (docs/datastandard.md shape b). A trial: --tables needs --out-dir.
# The entity types below (basin, play, lease_area, planning_area, geothermal_system, county) are not in the
# standard's vocabulary yet, so the validator blocks these tables until the owner adds them (the session's
# report, "To finish").
# ---------------------------------------------------------------------------------------------------------

def iso_date(text):
    """MM/DD/YYYY as the source writes it to YYYY-MM-DD; anything else is left empty."""
    m = re.fullmatch(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*", text or "")
    if not m:
        return ""
    try:
        return dt.date(int(m.group(3)), int(m.group(1)), int(m.group(2))).isoformat()
    except ValueError:
        return ""


def cell(v):
    v = clean(v)
    if v is None:
        return ""
    if isinstance(v, float):
        return repr(round(v, 6)).rstrip("0").rstrip(".") if v != int(v) else str(int(v))
    return str(v)


def unique_ids(ids):
    """Make repeated ids unique by a counter in the source's row order (the source has no key of its own)."""
    seen, out = {}, []
    many = {i for i in ids if ids.count(i) > 1}
    for i in ids:
        if i in many:
            seen[i] = seen.get(i, 0) + 1
            out.append(f"{i}#{seen[i]}")
        else:
            out.append(i)
    return out


def write_table(out_dir, name, row, title, cols, rows, note=""):
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, name + ".csv")
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="") as f:
        f.write(f"# Energy Research Warehouse (ERW): {title}\n")
        f.write(f"# Source file: {row['source']}/{row['file']} ({row['bytes']} bytes, sha256 {row['sha256']})\n")
        f.write(f"# Source URL: {row['url']}\n")
        f.write(f"# Retrieved: {row['retrieved_at_utc']}\n")
        f.write(f"# Terms: {row['terms_url']}\n")
        f.write("# Shape: entities (docs/datastandard.md). One row a row of the source; x_ columns are the "
                "source's own fields.\n")
        if note:
            f.write(f"# {note}\n")
        f.write("# Written by warehouse/connectors/resource_layers.py --tables (session 146)\n")
        w = csv.writer(f, lineterminator="\n")
        w.writerow(cols)
        for r in rows:
            w.writerow([cell(r.get(c)) for c in cols])
    os.replace(tmp, path)
    log(f"table: {name}.csv {len(rows)} rows")
    return path


STD = ["entity_id", "entity_type", "name", "geo", "lat", "lon", "operator", "source"]


def build_tables(raw_dir, out_dir):
    import geopandas as gpd
    done = []
    row = held(raw_dir, "eia_basins")
    g = gpd.read_file(find_file(unpack(row["path"]), r"\.shp$"))
    rows = [{"entity_id": i, "entity_type": "basin", "name": r["NAME"], "source": "eia:maps:sedimentary_basins",
             "x_area_sq_mi": r["Area_sq_mi"], "x_area_sq_km": r["Area_sq_km"]}
            for i, (_, r) in zip(unique_ids([f"eia:basin:{clean(x)}" for x in g["NAME"]]), g.iterrows())]
    done.append(write_table(out_dir, "eia_maps_sedimentary_basins", row, "EIA, U.S. Sedimentary Basins (May 2011)",
                            ["entity_id", "entity_type", "name", "source", "x_area_sq_mi", "x_area_sq_km"], rows))
    row = held(raw_dir, "eia_plays")
    g = gpd.read_file(find_file(unpack(row["path"]), r"\.shp$"))
    rows = [{"entity_id": i, "entity_type": "play", "name": r["Shale_play"],
             "source": "eia:maps:tight_oil_shale_gas_plays", "x_basin": r["Basin"], "x_lithology": r["Lithology"],
             "x_age_shale": r["Age_shale"], "x_area_sq_mi": r["Area_sq_mi"], "x_area_sq_km": r["Area_sq_km"],
             "x_references": r["References"]}
            for i, (_, r) in zip(unique_ids([f"eia:play:{clean(x)}" for x in g["ID"]]), g.iterrows())]
    done.append(write_table(out_dir, "eia_maps_shale_plays", row,
                            "EIA, Tight Oil and Shale Gas Plays in the U.S. (December 2021)",
                            ["entity_id", "entity_type", "name", "source", "x_basin", "x_lithology", "x_age_shale",
                             "x_area_sq_mi", "x_area_sq_km", "x_references"], rows))
    row = held(raw_dir, "boem_shapefiles")
    g = gpd.read_file(find_file(unpack(row["path"]), r"Offshore_Wind_Leases_outlines\.shp$"))
    ids = unique_ids([f"boem:{clean(n)}:{clean(t) or 'lease'}" for n, t in zip(g["LEASE_NUMB"], g["LEASE_TYPE"])])
    rows = [{"entity_id": i, "entity_type": "lease_area", "name": r["LEASE_NU_1"], "operator": r["COMPANY"],
             "source": "boem:renewable_energy_leases", "x_lease_number": r["LEASE_NUMB"],
             "x_lease_type": r["LEASE_TYPE"], "x_lease_date": iso_date(clean(r["LEASE_DATE"])),
             "x_lease_date_as_written": r["LEASE_DATE"], "x_lease_term": r["LEASE_TERM"], "x_acres": r["ACRES"],
             "x_state_as_written": r["STATE"], "x_project_name": r["PROJECT_NA"]}
            for i, (_, r) in zip(ids, g.iterrows())]
    done.append(write_table(out_dir, "boem_all_wind_lease_outlines", row,
                            "BOEM, offshore wind lease outlines (file update 02/05/2025)",
                            ["entity_id", "entity_type", "name", "operator", "source", "x_lease_number",
                             "x_lease_type", "x_lease_date", "x_lease_date_as_written", "x_lease_term", "x_acres",
                             "x_state_as_written", "x_project_name"], rows,
                            note="A lease with several outlines of one type has one row an outline, numbered #1, "
                                 "#2 in the source's order. x_lease_date is LEASE_DATE (MM/DD/YYYY) as a date."))
    row = held(raw_dir, "boem_planning")
    g = gpd.read_file(row["path"])
    g = g[~(g.geometry.isna() | g.geometry.is_empty)]
    rows = [{"entity_id": f"boem:planning_area:{clean(r['OBJECTID'])}", "entity_type": "planning_area",
             "name": r["ADDITIONAL_INFORMATION"], "source": "boem:wind_planning_area_outlines",
             "x_category": r["CATEGORY1"], "x_area_status": r["AREA_STATUS"],
             "x_protraction_number": r["PROTRACTION_NUMBER"], "x_url": r["URL1"]} for _, r in g.iterrows()]
    done.append(write_table(out_dir, "boem_all_wind_planning_areas", row,
                            "BOEM, Offshore Wind Planning Area Outlines (Rescinded July 30, 2025)",
                            ["entity_id", "entity_type", "name", "source", "x_category", "x_area_status",
                             "x_protraction_number", "x_url"], rows,
                            note="BOEM names the layer rescinded; x_area_status is the source's own field."))
    row = held(raw_dir, "usgs_IdentifiedGeothermalSystems_shp")
    g = gpd.read_file(row["path"])
    ids = unique_ids([f"usgs:geothermal_system:{clean(s_)}:{clean(n)}" for s_, n in zip(g["State"], g["Name"])])
    extra = ["Temp_C_Li", "Temp_C_Mn", "Temp_C_Mx", "Vol_km3_Li", "Vol_km3_Mn", "Vol_km3_Mx", "MWe_Mean",
             "MWe_P95", "MWe_P5"]
    rows = []
    for i, (_, r) in zip(ids, g.iterrows()):
        d = {"entity_id": i, "entity_type": "geothermal_system", "name": r["Name"], "geo": f"US-{clean(r['State'])}",
             "lat": r["Lat_83"], "lon": r["Lon_83"], "source": "usgs:identified_geothermal_systems"}
        d.update({"x_" + c.lower(): r[c] for c in extra})
        rows.append(d)
    done.append(write_table(out_dir, "usgs_all_geothermal_systems", row,
                            "USGS, Identified Moderate and High Temperature Geothermal Systems (2008 assessment)",
                            ["entity_id", "entity_type", "name", "geo", "lat", "lon", "source"]
                            + ["x_" + c.lower() for c in extra], rows,
                            note="lat and lon are the source's Lat_83 and Lon_83 (NAD83). x_mwe_mean is USGS's "
                                 "mean estimate of electric power generation potential, not an installed "
                                 "capacity, so capacity_mw is left empty."))
    return done


def check_terms(raw_dir, manifest):
    """Every terms quote in the manifest is in the saved terms page, word for word."""
    bad = []
    man = read_manifest(manifest)
    for l in man["layers"]:
        t = l.get("terms_file")
        if not t or not os.path.exists(os.path.join(raw_dir, t)):
            log(f"terms check {l['id']}: the saved page is not on this machine")
            bad.append(l["id"])
            continue
        text = html_text(os.path.join(raw_dir, t))
        for q in l.get("terms_quotes", [l["terms_quote"]]):
            if re.sub(r"\s+", " ", q).strip() not in text:
                log(f"terms check {l['id']}: NOT FOUND word for word: {q[:80]}")
                bad.append(l["id"])
        for a in l.get("terms_also", []):
            p2 = os.path.join(raw_dir, a["file"])
            if not os.path.exists(p2) or re.sub(r"\s+", " ", a["quote"]).strip() not in html_text(p2):
                log(f"terms check {l['id']}: NOT FOUND in {a['file']}: {a['quote'][:80]}")
                bad.append(l["id"])
    if not bad:
        log(f"terms check: every quote of {len(man['layers'])} layers is in its saved page")
    return bad


if __name__ == "__main__":
    sys.exit(main())
