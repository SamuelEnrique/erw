#!/usr/bin/env python3
"""The United Kingdom's electricity storage a year, 1996 to 2025: electricity out of and into storage and the storage's power
capacity, by type (pumped storage, grid-scale batteries, other storage), as the Department for Energy Security and Net Zero
(DESNZ) gives them in the Digest of UK Energy Statistics (DUKES), table 5.16 (session 2720, chain 136, the session's third own
choice; the ERW's storage tables were American, and it held the UK's storage only inside its stations list, session 2714).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/desnz_uk_storage_annual.csv, from one workbook (no
key), its sheet "5.16", three tables one above another:

    https://assets.publishing.service.gov.uk/media/6a6a36a0bf4a7aca8f8976c6/DUKES_5.16.xlsx

    entity      desnz_st:<type>: pumped_storage, battery_grid_scale (grid-scale batteries only: DESNZ's note 2 leaves out
                batteries behind the meter), other_storage, total
    variable    output (5.16.A, electricity out of storage) and input (5.16.B, electricity into it), MWh (the sheet's GWh times
                1,000); power_capacity (5.16.C, the rate at which the sites may supply their stored energy, DESNZ's note 1), MW
    ts_utc      the year's first day; freq P1Y; geo GB; node the table and the type as DESNZ writes them

Each table is found by its title ("Table 5.16.A Electricity output from energy storage facilities, GWh", and B and C as named
below), its year row (each year from 1996 in turn) and its four type rows in order. A blank cell has no row; a zero stays a
zero (no storage of that type that year, as DESNZ writes it). Checks, within 0.01 GWh or MW: the three types add to the total
in every table and year. A title, year or type not in this form, or a value not a number, stops the run. Nothing is filled.

    python warehouse/connectors/desnz_dukes_storage.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/desnz_dukes/ and reused unless --fresh; --offline asks nothing.
www.gov.uk's robots file closes only print and search paths; assets.publishing.service.gov.uk has no robots file (session 2678).
License: the Open Government Licence v3.0 (DESNZ's verdict on record, session 2167): public.
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "desnz_uk_storage_annual"
CONNECTOR = "desnz_dukes"
SOURCE = "desnz:dukes_5_16_energy_storage"
HOST = "assets.publishing.service.gov.uk"
URL = "https://assets.publishing.service.gov.uk/media/6a6a36a0bf4a7aca8f8976c6/DUKES_5.16.xlsx"
PAGE = "https://www.gov.uk/government/statistics/electricity-chapter-5-digest-of-united-kingdom-energy-statistics-dukes"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = ("Department for Energy Security and Net Zero, Digest of UK Energy Statistics (DUKES) 2026, table 5.16, Energy storage "
            "input electricity, output electricity, and capacity overview")
TABLES = [("Table 5.16.A Electricity output from energy storage facilities, GWh", "output", "MWh", 1000),
          ("Table 5.16.B Electricity input to energy storage facilities, GWh", "input", "MWh", 1000),
          ("Table 5.16.C Power capacity of energy storage facilities, MW", "power_capacity", "MW", 1)]
TYPES = [("Pumped storage", "pumped_storage"), ("Battery storage - grid scale", "battery_grid_scale"),
         ("Other energy storage", "other_storage"), ("Total", "total")]
FIRST_YEAR = 1996
X_COLS = []


def label(v):
    s = " ".join(str(v).replace(chr(0xA0), " ").split()) if v is not None else ""
    return re.sub(r"\s*\[note \d+\]", "", s)


def parse_sheet(rows):
    out = {}
    titles = {label(r[0]): i for i, r in enumerate(rows) if r and r[0] is not None}
    for title, var, unit, scale in TABLES:
        if title not in titles:
            raise RuntimeError(f"no table titled {title!r}")
        i = titles[title]
        head = [label(x) for x in rows[i + 1]]
        if head[0] != "Storage type":
            raise RuntimeError(f"{title}: the year row is {head[:3]}")
        years = [h for h in head[1:] if h]
        if not years or years != [str(FIRST_YEAR + k) for k in range(len(years))]:
            raise RuntimeError(f"{title}: the years are not each year from {FIRST_YEAR}: {years[:4]} .. {years[-2:]}")
        for k, (name, key) in enumerate(TYPES):
            r = rows[i + 2 + k]
            if label(r[0]) != name:
                raise RuntimeError(f"{title}: row {k + 1} is {label(r[0])!r}, not {name!r}")
            for j, y in enumerate(years, 1):
                v = r[j] if j < len(r) else None
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                if not isinstance(v, (int, float)) or isinstance(v, bool):
                    raise RuntimeError(f"{title} {name} {y}: not a number: {v!r}")
                out[(int(y), key, var)] = (float(v), unit, scale, f"{title[:11]}: {name}")
        for y in years:
            parts = [out.get((int(y), key, var), (None,))[0] for _, key in TYPES[:3]]
            tot = out.get((int(y), "total", var), (None,))[0]
            if tot is not None and None not in parts and abs(sum(parts) - tot) > 0.01:
                raise RuntimeError(f"{title} {y}: the types add to {sum(parts):,.3f}, the total {tot:,.3f}")
    return out


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    if "5.16" not in wb.sheetnames:
        raise RuntimeError(f"the workbook has no sheet 5.16: {wb.sheetnames}")
    vals = parse_sheet([tuple(r) for r in wb["5.16"].iter_rows(values_only=True)])
    rows = [dict(entity=f"desnz_st:{key}", variable=var, ts_utc=f"{y}-01-01T00:00:00Z", value=round(v * scale, 3), unit=unit,
                 freq="P1Y", geo="GB", market="", node=node, **base) for (y, key, var), (v, unit, scale, node) in sorted(vals.items())]
    return pd.DataFrame(rows)


def main(argv=None):
    ap = argparse.ArgumentParser(description="DESNZ DUKES 5.16: the UK's electricity storage a year")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t = parse(content, base)
        if t.duplicated(["entity", "variable", "ts_utc"]).any():
            raise RuntimeError("a key given twice")
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["entity", "variable", "ts_utc"]).reset_index(drop=True)
        bat = t[(t.entity == "desnz_st:battery_grid_scale") & (t.variable == "power_capacity")].set_index("ts_utc").value
        head = [
            "Energy Research Warehouse (ERW): DESNZ, Digest of UK Energy Statistics (DUKES) 2026, table 5.16: the UK's electricity "
            f"out of and into storage and its power capacity, by type, {t.ts_utc.min()[:4]} to {t.ts_utc.max()[:4]} (session 2720)",
            "Shape: series (docs/datastandard.md v0). entity desnz_st:<type> (pumped_storage, battery_grid_scale, other_storage, "
            "total); variable output and input (MWh, the sheet's GWh times 1,000) or power_capacity (MW); ts_utc the year's first "
            "day; freq P1Y; geo GB; node the table and the type.",
            f"{len(t)} rows; grid-scale battery power capacity, MW: " + ", ".join(f"{k[:4]} {v:,.0f}" for k, v in bat.items()
                                                                                if k[:4] in ("2015", "2020", "2023", "2024", "2025"))
            + ". The types checked against each total; batteries behind the meter are not counted (DESNZ's note 2); nothing is "
            "filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/desnz_dukes_storage.py; the workbook retrieved "
            f"{base['retrieved_at']}; last modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL} (listed on {PAGE})",
            "License: Open Government Licence v3.0. Credit: Department for Energy Security and Net Zero.",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="UK Department for Energy Security and Net Zero (DESNZ)",
                                    report="Digest of UK Energy Statistics (DUKES) 2026, table 5.16, Energy storage overview",
                                    report_url=PAGE, document_list=PAGE, license="public", tables=[NAME])])  # OGL v3.0
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
