#!/usr/bin/env python3
"""ERCOT day-ahead ancillary service prices: the Market Clearing Price for Capacity (MCPC), hourly, from 2018.

Energy Research Warehouse (ERW) connector, session 65 (docs/methods/capacity_and_ancillary.md). Writes one series
table, warehouse/output/ercot_as_prices.csv:

    entity    ercot:<service as ERCOT names it>   ercot:REGUP, ercot:REGDN, ercot:RRS, ercot:NSPIN, ercot:ECRS
    variable  mcpc_dam
    unit      USD/MW-hour   (US dollars per MW of capacity held for one hour, as ERCOT settles it)
    freq      PT1H; ts_utc is the start of the delivery hour, UTC
    market    ercot_dam
    node      the service

    python warehouse/connectors/ercot_as_prices.py                      # 2018-01-01 to the newest day published
    python warehouse/connectors/ercot_as_prices.py --out-dir C:/scratch # a trial run: nothing in warehouse/output
    python warehouse/connectors/ercot_as_prices.py --offline            # from the raw files only, no request

Sources, both on ERCOT's public reports site, the route iso_prices.py already uses for ERCOT's price history:

  NP4-181-ER  Historical DAM Clearing Prices for Capacity: one zip per operating year, a CSV with a row per
              delivery hour and a column per service. The current year's file is republished weekly. From
              4 October 2026 the current year's zip holds one Excel workbook in place of the CSV (session 113):
              the same columns in the same order, under a logo and a title, each price a number to the cent.
              parse_year reads either form, and a zip that holds neither, or both, is an error. The zips of 2018
              to 2025 each held a CSV on 5 October 2026.
  NP4-188-CD  DAM Clearing Prices for Capacity: one zip per day-ahead market run, for the days after the newest
              yearly file ends.

Services: Regulation Up (REGUP), Regulation Down (REGDN), Responsive Reserve (RRS), Non-Spin (NSPIN) and the ERCOT
Contingency Reserve Service (ECRS), which ERCOT began buying in June 2023: its column is empty before, and an empty
cell is not a price, so no row is written for it.

Never fill. A (service, local delivery day) is written only when every hour of that day is there: 24, or 23 on the
spring clock change and 25 on the autumn one (ERCOT's repeated-hour flag marks the second 01:00 hour). A day short of
that is left out, logged, and reported as a gap in the run's status. The raw files stay in
warehouse/raw/ercot_as_prices/ and a file already there is not requested again (iso_prices.fetch_raw).

Ceiling: 500,000 rows (session 65 ruling). A pull that would pass it writes nothing.

License: public. ERCOT's terms (https://www.ercot.com/help/terms, item 5): "raw data provided in public portions of
this website may be used, reproduced, and redistributed in compilations, charts, and analyses".
"""

import argparse
import datetime as dt
import io
import json
import os
import sys
import traceback
import zipfile

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "ercot_as_prices"
CONNECTOR = "ercot_as_prices"
FIRST_YEAR = 2018
CEILING = 500_000
TZ = "America/Chicago"
SERVICES = ["REGUP", "REGDN", "RRS", "NSPIN", "ECRS"]
REPORTS = {  # source id: (report name, report type id)
    "ercot:NP4-181-ER": ("Historical DAM Clearing Prices for Capacity", 13091),
    "ercot:NP4-188-CD": ("DAM Clearing Prices for Capacity", 12329),
}
DOC_LIST = "https://www.ercot.com/misapp/servlets/IceDocListJsonWS?reportTypeId={}"
DOWNLOAD = "https://www.ercot.com/misdownload/servlets/mirDownload?doclookupId={}"
PAGE = "https://www.ercot.com/mp/data-products/data-product-details?id={}"
PAUSE = 3  # seconds between requests


def to_utc(dates, hour_ending, repeated):
    """ERCOT's delivery date and hour ending (01:00 to 24:00, Central) as the hour's start in UTC. `repeated` is the
    repeated-hour flag: Y marks the second 01:00 to 02:00 hour of the autumn clock change."""
    he = hour_ending.str.strip().str.extract(r"^(\d{1,2}):00$")[0]
    if he.isna().any():
        raise RuntimeError(f"hour ending not HH:00: {hour_ending[he.isna()].unique()[:5]}")
    he = he.astype(int)
    if not he.between(1, 24).all():
        raise RuntimeError("hour ending outside 1 to 24")
    local = pd.to_datetime(dates.str.strip(), format="%m/%d/%Y") + pd.to_timedelta(he - 1, unit="h")
    flag = repeated.str.strip().str.upper()
    if not flag.isin(["Y", "N"]).all():
        raise RuntimeError(f"repeated-hour flag not Y or N: {flag[~flag.isin(['Y', 'N'])].unique()[:5]}")
    # ambiguous=True means daylight time, the first of the two hours: the one ERCOT does not flag
    return local.dt.tz_localize(TZ, ambiguous=(flag == "N").values, nonexistent="raise").dt.tz_convert("UTC")


YEAR_HEAD = ["Delivery Date", "Hour Ending", "Repeated Hour Flag"]


def year_workbook(data):
    """The yearly file as ERCOT posts it since 4 October 2026: one Excel workbook, one sheet, the header row under a
    logo and a title. Returned as the CSV reads, every cell text, so both forms pass the same checks. A price is the
    number the workbook stores, written by repr (the shortest text that reads back as the same number): nothing is
    rounded. An empty cell is "", as in the CSV."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if len(wb.worksheets) != 1:
        raise RuntimeError(f"NP4-181-ER workbook holds sheets {wb.sheetnames}, expected one")
    rows = [r for r in wb.worksheets[0].iter_rows(values_only=True)]
    heads = [i for i, r in enumerate(rows) if [str(c).strip() for c in r[:3] if c is not None] == YEAR_HEAD]
    if len(heads) != 1:
        raise RuntimeError(f"NP4-181-ER workbook: {len(heads)} header rows, expected one: layout changed")
    head = rows[heads[0]]
    if any(c is None or not str(c).strip() for c in head):
        raise RuntimeError(f"NP4-181-ER workbook header {list(head)}: a column without a name")
    body = [r for r in rows[heads[0] + 1:] if any(c is not None for c in r)]  # a wholly empty row is no hour
    cells = []
    for r in body:
        if not all(isinstance(c, str) for c in r[:3]):
            raise RuntimeError(f"NP4-181-ER workbook row {list(r)}: delivery date, hour ending or flag is not text")
        prices = []
        for c in r[3:]:
            if c is None:
                prices.append("")
            elif isinstance(c, (int, float)) and not isinstance(c, bool):
                prices.append(repr(c))
            else:
                raise RuntimeError(f"NP4-181-ER workbook row {list(r)}: a price that is not a number")
        cells.append(list(r[:3]) + prices)
    return pd.DataFrame(cells, columns=[str(c) for c in head], dtype=str)


def parse_year(raw):
    """A yearly NP4-181-ER zip, in either form ERCOT has posted (one CSV, or since 4 October 2026 one Excel
    workbook): long rows (service, ts, value, day). Empty cells are omitted, never filled."""
    z = zipfile.ZipFile(io.BytesIO(raw))
    csvs = [n for n in z.namelist() if n.lower().endswith(".csv")]
    books = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
    if len(csvs) == 1 and not books:
        x = pd.read_csv(z.open(csvs[0]), dtype=str, keep_default_na=False)
    elif len(books) == 1 and not csvs:
        x = year_workbook(z.read(books[0]))
    else:
        raise RuntimeError(f"NP4-181-ER zip holds {z.namelist()}, expected one CSV or one Excel workbook")
    x.columns = [c.strip() for c in x.columns]
    need = YEAR_HEAD
    if (list(x.columns[:3]) != need or not set(x.columns[3:]) <= set(SERVICES) or "REGUP" not in x.columns
            or x.columns.duplicated().any()):
        raise RuntimeError(f"NP4-181-ER columns {list(x.columns)}: layout changed")
    if x.empty:
        raise RuntimeError("NP4-181-ER file holds no hour")
    ts = to_utc(x["Delivery Date"], x["Hour Ending"], x["Repeated Hour Flag"])
    out = []
    for s in x.columns[3:]:
        v = x[s].str.strip()
        keep = v != ""
        out.append(pd.DataFrame({"service": s, "ts": ts[keep], "value": pd.to_numeric(v[keep], errors="raise"),
                                 "day": x.loc[keep, "Delivery Date"].str.strip()}))
    return pd.concat(out, ignore_index=True)


def parse_day(raw):
    """A daily NP4-188-CD zip: long rows (service, ts, value, day)."""
    z = zipfile.ZipFile(io.BytesIO(raw))
    names = [n for n in z.namelist() if n.lower().endswith(".csv")]
    if len(names) != 1:
        raise RuntimeError(f"NP4-188-CD zip holds {z.namelist()}, expected one CSV")
    x = pd.read_csv(z.open(names[0]), dtype=str, keep_default_na=False)
    x.columns = [c.strip() for c in x.columns]
    if list(x.columns) != ["DeliveryDate", "HourEnding", "AncillaryType", "MCPC", "DSTFlag"]:
        raise RuntimeError(f"NP4-188-CD columns {list(x.columns)}: layout changed")
    svc = x["AncillaryType"].str.strip()
    if not svc.isin(SERVICES).all():
        raise RuntimeError(f"NP4-188-CD: unknown service {sorted(set(svc) - set(SERVICES))}")
    v = x["MCPC"].str.strip()
    keep = v != ""
    ts = to_utc(x["DeliveryDate"], x["HourEnding"], x["DSTFlag"])
    return pd.DataFrame({"service": svc[keep], "ts": ts[keep], "value": pd.to_numeric(v[keep], errors="raise"),
                         "day": x.loc[keep, "DeliveryDate"].str.strip()}).reset_index(drop=True)


def hours_in(day):
    """The number of real hours in a Central local day given as MM/DD/YYYY: 24, or 23 or 25 at a clock change."""
    d = pd.Timestamp(dt.datetime.strptime(day, "%m/%d/%Y")).tz_localize(TZ)
    return int(((d + pd.DateOffset(days=1)) - d) / pd.Timedelta(hours=1))


def complete_days(rows, log):
    """Keep only (service, day) groups holding every hour of the local day exactly once; the rest are gaps."""
    if rows.duplicated(["service", "ts"]).any():
        raise RuntimeError(f"{int(rows.duplicated(['service', 'ts']).sum())} rows repeat a (service, hour)")
    n = rows.groupby(["service", "day"])["ts"].transform("size")
    want = rows["day"].map({d: hours_in(d) for d in rows["day"].unique()})
    bad = rows[n != want]
    gaps = []
    for (s, d), g in bad.groupby(["service", "day"]):
        gaps.append(f"{s} {dt.datetime.strptime(d, '%m/%d/%Y').date()}: {len(g)} of {hours_in(d)} hours")
    for g in gaps:
        log(f"  gap, left out: {g}")
    return rows[n == want].reset_index(drop=True), gaps


def doc_list(rtid, log, offline):
    raw, _ = ip.fetch_raw(CONNECTOR, DOC_LIST.format(rtid), log, offline=offline, fresh=True, pause=PAUSE)
    docs = json.loads(raw)["ListDocsByRptTypeRes"]["DocumentList"]
    return [d["Document"] for d in docs]


def pull(log, offline=False):
    frames = []
    this_year = pd.Timestamp.now(tz=TZ).year
    yearly = {}
    for d in doc_list(13091, log, offline):
        m = pd.Series([d["FriendlyName"]]).str.extract(r"^DAMASMCPC_(\d{4})$")[0][0]
        if pd.notna(m):
            yearly[int(m)] = d
    missing = [y for y in range(FIRST_YEAR, this_year + 1) if y not in yearly]
    if missing:
        raise RuntimeError(f"NP4-181-ER lists no file for {missing}")
    for year in range(FIRST_YEAR, this_year + 1):
        d = yearly[year]
        url = DOWNLOAD.format(d["DocID"])
        raw, rec = ip.fetch_raw(CONNECTOR, url, log, offline=offline, pause=PAUSE)
        f = parse_year(raw)
        f["source"], f["source_url"], f["retrieved_at"] = "ercot:NP4-181-ER", url, rec["retrieved_at"]
        f["vintage"] = ip.utc_iso(pd.Timestamp(d["PublishDate"]))
        log(f"  {year}: {d['ConstructedName']}, published {d['PublishDate']}: {len(f)} prices, "
            f"{f['ts'].min()} to {f['ts'].max()}")
        frames.append(f)
    hist = pd.concat(frames, ignore_index=True)
    last = hist["ts"].max()
    # the days after the yearly files end: one document per day-ahead market run
    daily = [d for d in doc_list(12329, log, offline) if d["ConstructedName"].lower().endswith("_csv.zip")]
    for d in sorted(daily, key=lambda d: d["PublishDate"]):
        # a run published on day D clears delivery day D + 1; skip runs whose delivery day the yearly files hold
        if pd.Timestamp(d["PublishDate"]).tz_convert("UTC") + pd.Timedelta(days=2) <= last:
            continue
        url = DOWNLOAD.format(d["DocID"])
        raw, rec = ip.fetch_raw(CONNECTOR, url, log, offline=offline, pause=PAUSE)
        f = parse_day(raw)
        f = f[f["ts"] > last]
        if f.empty:
            continue
        f["source"], f["source_url"], f["retrieved_at"] = "ercot:NP4-188-CD", url, rec["retrieved_at"]
        f["vintage"] = ip.utc_iso(pd.Timestamp(d["PublishDate"]))
        log(f"  daily {d['ConstructedName']}: {len(f)} prices for {sorted(f['day'].unique())}")
        frames.append(f)
    rows = pd.concat(frames, ignore_index=True)
    return rows[rows["ts"] >= pd.Timestamp(f"{FIRST_YEAR}-01-01", tz=TZ).tz_convert("UTC")].reset_index(drop=True)


def to_series(rows):
    s = pd.DataFrame({
        "entity": "ercot:" + rows["service"], "variable": "mcpc_dam", "ts_utc": rows["ts"].map(ip.utc_iso),
        "value": rows["value"], "unit": "USD/MW-hour", "freq": "PT1H", "geo": "US-TX", "market": "ercot_dam",
        "node": rows["service"], "source": rows["source"], "source_url": rows["source_url"],
        "retrieved_at": rows["retrieved_at"], "vintage": rows["vintage"],
    })
    return s.sort_values(["entity", "ts_utc"]).reset_index(drop=True)[ip.SERIES_COLS]


def main(argv=None):
    ap = argparse.ArgumentParser(description="ERW: ERCOT day-ahead ancillary service prices (MCPC)")
    ap.add_argument("--out-dir", help="write the table, log, registry and status under this directory instead of "
                    "the repository (raw files stay in warehouse/raw)")
    ap.add_argument("--offline", action="store_true", help="read the raw files only; make no request")
    a = ap.parse_args(argv)
    raw_dir = ip.RAW_DIR
    if a.out_dir:
        ip.set_out_dir(a.out_dir)
        ip.RAW_DIR = raw_dir  # source documents are kept once, in warehouse/raw, whatever the output directory
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        rows, gaps = complete_days(pull(log, a.offline), log)
        if len(rows) > CEILING:
            raise RuntimeError(f"{len(rows)} rows would pass the ceiling of {CEILING}; nothing written")
        s = to_series(rows)
        per = rows.groupby("service")["ts"].agg(["size", "min", "max"])
        for svc, r in per.iterrows():
            log(f"  {svc}: {r['size']} hours, {ip.utc_iso(r['min'])} to {ip.utc_iso(r['max'])}")
        header = [
            "Energy Research Warehouse (ERW): ERCOT day-ahead market clearing prices for capacity (MCPC), the "
            "ancillary services, hourly",
            "Shape: series (docs/datastandard.md v0). Unit USD/MW-hour: US dollars per MW of capacity held for one "
            "hour. ts_utc is the start of the delivery hour, UTC (ERCOT publishes the hour ending, Central time).",
            "Services: REGUP Regulation Up, REGDN Regulation Down, RRS Responsive Reserve, NSPIN Non-Spin, ECRS "
            "ERCOT Contingency Reserve Service (bought from June 2023; no row before it).",
            "Never filled: a (service, local day) is written only when every hour of the day is there; "
            f"{len(gaps)} such days are left out of this run" + (": " + "; ".join(gaps[:20]) if gaps else "") + ".",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/ercot_as_prices.py",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git; each run's manifest.csv lists each file and URL)",
        ] + [f"Source: {sid} ERCOT, {name}, {PAGE.format(sid.split(':')[1])}" for sid, (name, _) in REPORTS.items()] + [
            "License: public. ERCOT terms of use, item 5: raw data in public portions of the website may be used, "
            "reproduced, and redistributed in compilations, charts, and analyses.",
            "Method: docs/methods/capacity_and_ancillary.md",
        ]
        ip.write_csv(s, NAME, header, log)
        ip.update_sources([dict(source=sid, publisher=ip.ISO_PUBLISHERS["ercot"], report=name,
                                report_url=PAGE.format(sid.split(":")[1]), document_list=DOC_LIST.format(rtid),
                                tables=[NAME]) for sid, (name, rtid) in REPORTS.items()])
        results.append(dict(table=NAME, market="DAM", status="ok", detail=f"{len(s)} rows of {CEILING}"))
        for g in gaps:
            results.append(dict(table=NAME, market="DAM", status="gap", detail=g))
    except Exception:
        tb = ip.redact(traceback.format_exc())
        last = tb.strip().splitlines()[-1]
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{CONNECTOR} {NAME} FAILED, no output file written: {last}", file=sys.stderr)
        results.append(dict(table=NAME, market="DAM", status="failed", detail=last[:300]))
    ip.write_status(CONNECTOR, run_id, results)
    failures = sum(r["status"] == "failed" for r in results)
    log(f"done, failures={failures}")
    log.close()
    print(f"{CONNECTOR} run log: {os.path.relpath(log.path, ip.ROOT)}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
