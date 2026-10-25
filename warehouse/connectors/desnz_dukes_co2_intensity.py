#!/usr/bin/env python3
"""The estimated carbon dioxide intensity of the United Kingdom's electricity supplied a year, 1996 to 2025, by fuel (coal, gas,
all non-renewable fuels, all fuels), as the Department for Energy Security and Net Zero (DESNZ) gives it in the Digest of UK
Energy Statistics (DUKES), table 5.14 (session 2729, chain 137, the session's second own choice; the ERW's carbon intensity tables
were American).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/desnz_uk_co2_intensity_annual.csv, from one workbook
(no key), its sheet "5.14":

    https://assets.publishing.service.gov.uk/media/6a6a3684bf4a7aca8f8976c4/DUKES_5.14.xlsx

    entity      desnz_ci:<fuel>: coal, gas, non_renewable (coal, oil, gas and other solid fuels, non-renewable waste included,
                DESNZ's note 5), all_fuels (every fuel of DUKES 5.6, nuclear and renewables included, net pumped storage and net
                imports left out, note 6)
    variable    co2_intensity
    unit        kgCO2/MWh (the sheet's tonnes of CO2 per GWh supplied, the same number); ts_utc the year's first day; freq P1Y;
                geo GB; node the sheet's row as written

DESNZ's notes: the numerator is the power stations' emissions with an estimate for autogeneration, the denominator the
electricity supplied by all generators (note 1); the figures differ from the government's conversion factors for company
reporting, which lag two years (note 2); 2025's emissions are provisional (note 3); coal generation ceased in 2024, so coal has no
figure from 2025 (note 4); only the generators' operational emissions are counted, not the fuel supply chain's (note 7).

The table is found by its title, its year row (each year from 1996 in turn, note marks read off) and its four fuel rows in order.
A blank cell has no row. Checks: every value a positive number, and all fuels' intensity no greater than the non-renewable fuels'
in any year. A title, year or row not in this form stops the run. Nothing is filled.

    python warehouse/connectors/desnz_dukes_co2_intensity.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/desnz_dukes/ and reused unless --fresh; --offline asks nothing.
www.gov.uk's robots file closes only print and search paths; assets.publishing.service.gov.uk has no robots file (session 2678).
License: the Open Government Licence v3.0 (DESNZ's verdict on record, session 2167): public.
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "desnz_uk_co2_intensity_annual"
CONNECTOR = "desnz_dukes"
SOURCE = "desnz:dukes_5_14_co2_intensity"
HOST = "assets.publishing.service.gov.uk"
URL = "https://assets.publishing.service.gov.uk/media/6a6a3684bf4a7aca8f8976c4/DUKES_5.14.xlsx"
PAGE = "https://www.gov.uk/government/statistics/electricity-chapter-5-digest-of-united-kingdom-energy-statistics-dukes"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = ("Department for Energy Security and Net Zero, Digest of UK Energy Statistics (DUKES) 2026, table 5.14, Estimated carbon "
            "dioxide intensity of electricity supplied")
TITLE = "Table 5.14 estimated carbon dioxide emission intensity of electricity supplied, tonnes of CO2 per GWh supplied"
ROWS = [("Coal", "coal"), ("Gas", "gas"), ("All non renewable fuels", "non_renewable"),
        ("All fuels (including nuclear and renewables)", "all_fuels")]
FIRST_YEAR = 1996
X_COLS = []


def label(v):
    s = " ".join(str(v).replace(chr(0xA0), " ").split()) if v is not None else ""
    return re.sub(r"\s*\[note \d+\]", "", s).strip()


def parse_sheet(rows):
    if not rows or label(rows[0][0]) != TITLE:
        raise RuntimeError(f"the title is {label(rows[0][0]) if rows else ''!r}")
    h = next((i for i, r in enumerate(rows) if r and label(r[0]) == "Fuel"), None)
    if h is None:
        raise RuntimeError("no header row Fuel")
    years = [label(x) for x in rows[h][1:]]
    while years and not years[-1]:
        years.pop()
    if not years or years != [str(FIRST_YEAR + k) for k in range(len(years))]:
        raise RuntimeError(f"the years are not each year from {FIRST_YEAR}: {years[:3]} .. {years[-2:]}")
    out = {}
    for k, (name, key) in enumerate(ROWS):
        r = rows[h + 1 + k]
        if label(r[0]) != name:
            raise RuntimeError(f"row {k + 1} is {label(r[0])!r}, not {name!r}")
        for j, y in enumerate(years, 1):
            v = r[j] if j < len(r) else None
            if v is None or (isinstance(v, str) and v.strip() in ("", "[x]")):
                continue
            if not isinstance(v, (int, float)) or isinstance(v, bool) or v <= 0:
                raise RuntimeError(f"{name} {y}: not a positive number: {v!r}")
            out[(int(y), key)] = (float(v), label(r[0]))
    for y in {k[0] for k in out}:
        a, n = out.get((y, "all_fuels"), (None,))[0], out.get((y, "non_renewable"), (None,))[0]
        if None not in (a, n) and a > n:
            raise RuntimeError(f"{y}: all fuels' intensity {a} above the non-renewable fuels' {n}")
    return out


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    if "5.14" not in wb.sheetnames:
        raise RuntimeError(f"the workbook has no sheet 5.14: {wb.sheetnames}")
    vals = parse_sheet([tuple(r) for r in wb["5.14"].iter_rows(values_only=True)])
    rows = [dict(entity=f"desnz_ci:{key}", variable="co2_intensity", ts_utc=f"{y}-01-01T00:00:00Z", value=round(v, 4),
                 unit="kgCO2/MWh", freq="P1Y", geo="GB", market="", node=node, **base) for (y, key), (v, node) in sorted(vals.items())]
    return pd.DataFrame(rows)


def main(argv=None):
    ap = argparse.ArgumentParser(description="DESNZ DUKES 5.14: the CO2 intensity of the UK's electricity supplied a year")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t = parse(content, base)
        if t.duplicated(["entity", "variable", "ts_utc"]).any():
            raise RuntimeError("a key given twice")
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["entity", "ts_utc"]).reset_index(drop=True)
        allf = t[t.entity == "desnz_ci:all_fuels"].set_index("ts_utc").value
        head = [
            "Energy Research Warehouse (ERW): DESNZ, Digest of UK Energy Statistics (DUKES) 2026, table 5.14: the estimated CO2 "
            f"intensity of the UK's electricity supplied, by fuel, {t.ts_utc.min()[:4]} to {t.ts_utc.max()[:4]} (session 2729)",
            "Shape: series (docs/datastandard.md v0). entity desnz_ci:<fuel> (coal, gas, non_renewable, all_fuels); variable "
            "co2_intensity; unit kgCO2/MWh (the sheet's tonnes per GWh supplied); ts_utc the year's first day; freq P1Y; geo GB; "
            "node the sheet's row.",
            f"{len(t)} rows; all fuels, kgCO2/MWh: " + ", ".join(f"{k[:4]} {v:,.1f}" for k, v in allf.items()
                                                               if k[:4] in ("1996", "2010", "2020", "2024", "2025"))
            + ". 2025's emissions are provisional (DESNZ's note 3); coal has no figure from 2025 (note 4); nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/desnz_dukes_co2_intensity.py; the workbook retrieved "
            f"{base['retrieved_at']}; last modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL} (listed on {PAGE})",
            "License: Open Government Licence v3.0. Credit: Department for Energy Security and Net Zero.",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="UK Department for Energy Security and Net Zero (DESNZ)",
                                    report="Digest of UK Energy Statistics (DUKES) 2026, table 5.14, Carbon dioxide intensity of electricity supplied",
                                    report_url=PAGE, document_list=PAGE, license="public", tables=[NAME])])  # OGL v3.0
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
