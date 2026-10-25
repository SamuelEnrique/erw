#!/usr/bin/env python3
"""The United Kingdom's installed generating capacity at each year's end, 2011 to 2025, by the network it connects to
(transmission or distribution) and by fuel, for Great Britain, Northern Ireland and the UK, as the Department for Energy Security
and Net Zero (DESNZ) gives it in the Digest of UK Energy Statistics (DUKES), table 5.12 (session 2742, chain 138, the session's
own choice; the companion of chain 137's de-rated capacity, DUKES 5.7, and chain 136's station list, DUKES 5.11).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/desnz_uk_capacity_by_connection_annual.csv, from one
workbook (no key), its sheet "5.12", three tables one above another:

    https://assets.publishing.service.gov.uk/media/6a6a3668862aaf18d9c629ed/DUKES_5.12.xlsx

    entity      desnz_conn:<area>:<network>:<fuel>: area gb (5.12.A), ni (5.12.B), uk (5.12.C); network transmission, distribution,
                or all (the UK table's sum of the two); fuel the row as a slug (coal, oil, ocgt_and_conventional_thermal_gas,
                ccgt_gas, nuclear_magnox, nuclear_pwr, nuclear_agr, hydro_natural_flow, wind_onshore, wind_offshore, solar,
                wave_and_tidal, bioenergy, other_fuels, pumped_storage, total)
    variable    installed_capacity
    unit        MW; ts_utc the year's first day (the capacity is at 31 December, DESNZ's note 2); freq P1Y; geo GB, GB-NIR or
                GB for the UK; node the sheet's two labels as written

DESNZ's notes: installed capacity where available, else grid export capacity (note 1), so these figures are not the de-rated
capacity of DUKES 5.7; coal includes mixed fuel and co-firing coal stations (note 4); OCGT oil capacity is under oil (note 5);
other fuels are municipal solid waste for the major power producers, and coke oven gas, blast furnace gas and chemical waste for
other generators (note 6); pumped storage was revised up from earlier editions (note 7).

Each table is found by its title and its year row (each year from 2011 in turn); a blank network label carries the one above
(the UK table's wave and tidal row has none). A network or fuel row given twice, or a network not in the map, stops the run. A
blank cell has no row; a zero stays a zero. Checks, within 0.5 MW in every year: each area's and network's fuels add to its total;
Great Britain and Northern Ireland add to the UK, network by network and fuel by fuel; the UK's transmission and distribution add
to its all. A check missed stops the run. Nothing is filled.

    python warehouse/connectors/desnz_dukes_connection.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/desnz_dukes/ and reused unless --fresh; --offline asks nothing.
www.gov.uk's robots file closes only print and search paths; assets.publishing.service.gov.uk has no robots file (session 2678).
License: the Open Government Licence v3.0 (DESNZ's verdict on record, session 2167): public.
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "desnz_uk_capacity_by_connection_annual"
CONNECTOR = "desnz_dukes"
SOURCE = "desnz:dukes_5_12_capacity_by_connection"
HOST = "assets.publishing.service.gov.uk"
URL = "https://assets.publishing.service.gov.uk/media/6a6a3668862aaf18d9c629ed/DUKES_5.12.xlsx"
PAGE = "https://www.gov.uk/government/statistics/electricity-chapter-5-digest-of-united-kingdom-energy-statistics-dukes"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = ("Department for Energy Security and Net Zero, Digest of UK Energy Statistics (DUKES) 2026, table 5.12, Plant installed "
            "capacity, by connection")
TABLES = [("Table 5.12.A Electricity network of Great Britain, MW", "gb", "GB"),
          ("Table 5.12.B Electricity network of Northern Ireland, MW", "ni", "GB-NIR"),
          ("Table 5.12.C Electricity network UK total, MW", "uk", "GB")]
NETWORKS = {"Transmission": "transmission", "Distribution": "distribution", "All": "all"}
FIRST_YEAR = 2011
TOL = 0.5
X_COLS = []


def label(v):
    s = " ".join(str(v).replace(chr(0xA0), " ").split()) if v is not None else ""
    return re.sub(r"\s*\[note \d+\]", "", s).strip()


def slug(s):
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", s.lower())).strip("_")


def parse_sheet(rows):
    titles = {label(r[0]): i for i, r in enumerate(rows) if r and r[0] is not None}
    out, nodes = {}, {}
    for title, area, geo in TABLES:
        if title not in titles:
            raise RuntimeError(f"no table titled {title!r}")
        i = titles[title]
        head = [label(x) for x in rows[i + 1]]
        if head[:2] != ["Network type", "Generator type"]:
            raise RuntimeError(f"{title}: the header is {head[:2]}")
        years = head[2:]
        while years and not years[-1]:
            years.pop()
        if not years or years != [str(FIRST_YEAR + k) for k in range(len(years))]:
            raise RuntimeError(f"{title}: the years are not each year from {FIRST_YEAR}")
        net = None
        for r in rows[i + 2:]:
            if not r or not any(label(x) for x in r[:2]):
                continue
            if label(r[0]).startswith("Table 5.12"):
                break
            if label(r[0]):
                if label(r[0]) not in NETWORKS:
                    raise RuntimeError(f"{title}: a network not in the map: {label(r[0])!r}")
                net = NETWORKS[label(r[0])]
            if net is None:
                raise RuntimeError(f"{title}: a row before any network")
            fuel = label(r[1])
            ent = f"{area}:{net}:{slug(fuel)}"
            if ent in nodes:
                raise RuntimeError(f"{title}: {ent} given twice")
            nodes[ent] = (f"{label(r[0]) or net.title()}: {fuel}", geo)
            for k, y in enumerate(years):
                v = r[2 + k] if 2 + k < len(r) else None
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                if not isinstance(v, (int, float)) or isinstance(v, bool) or v < 0:
                    raise RuntimeError(f"{title} {net} {fuel} {y}: not a quantity: {v!r}")
                out[(int(y), ent)] = float(v)
    check(out)
    return out, nodes


def check(out):
    for y in sorted({k[0] for k in out}):
        g = {k[1]: v for k, v in out.items() if k[0] == y}
        groups = {e.rsplit(":", 1)[0] for e in g}
        for grp in groups:
            tot = g.get(grp + ":total")
            parts = [v for e, v in g.items() if e.startswith(grp + ":") and e != grp + ":total"]
            if tot is not None and parts and abs(sum(parts) - tot) > TOL:
                raise RuntimeError(f"{y} {grp}: the fuels make {sum(parts):,.2f} MW, the total {tot:,.2f}")
        for e, v in g.items():
            if e.startswith("uk:") and not e.startswith("uk:all:"):
                s = g.get("gb:" + e[3:], 0.0) + g.get("ni:" + e[3:], 0.0)
                if ("gb:" + e[3:] in g or "ni:" + e[3:] in g) and abs(s - v) > TOL:
                    raise RuntimeError(f"{y} {e[3:]}: Great Britain and Northern Ireland make {s:,.2f} MW, the UK {v:,.2f}")
            if e.startswith("uk:all:"):
                f = e[len("uk:all:"):]
                s = g.get("uk:transmission:" + f, 0.0) + g.get("uk:distribution:" + f, 0.0)
                if abs(s - v) > TOL:
                    raise RuntimeError(f"{y} {f}: the UK's transmission and distribution make {s:,.2f} MW, its all {v:,.2f}")


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    if "5.12" not in wb.sheetnames:
        raise RuntimeError(f"the workbook has no sheet 5.12: {wb.sheetnames}")
    vals, nodes = parse_sheet([tuple(r) for r in wb["5.12"].iter_rows(values_only=True)])
    rows = [dict(entity=f"desnz_conn:{e}", variable="installed_capacity", ts_utc=f"{y}-01-01T00:00:00Z", value=round(x, 4), unit="MW",
                 freq="P1Y", geo=nodes[e][1], market="", node=nodes[e][0], **base) for (y, e), x in sorted(vals.items())]
    return pd.DataFrame(rows)


def main(argv=None):
    ap = argparse.ArgumentParser(description="DESNZ DUKES 5.12: the UK's installed capacity by connection a year")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t = parse(content, base)
        if t.duplicated(["entity", "variable", "ts_utc"]).any():
            raise RuntimeError("a key given twice")
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["entity", "ts_utc"]).reset_index(drop=True)
        d = t[t.entity == "desnz_conn:uk:distribution:total"].set_index("ts_utc").value
        tr = t[t.entity == "desnz_conn:uk:transmission:total"].set_index("ts_utc").value
        head = [
            "Energy Research Warehouse (ERW): DESNZ, Digest of UK Energy Statistics (DUKES) 2026, table 5.12: the UK's installed "
            f"capacity at year end by connection (transmission or distribution) and fuel, {t.ts_utc.min()[:4]} to "
            f"{t.ts_utc.max()[:4]}, for Great Britain, Northern Ireland and the UK (session 2742)",
            "Shape: series (docs/datastandard.md v0). entity desnz_conn:<gb|ni|uk>:<transmission|distribution|all>:<fuel>; variable "
            "installed_capacity; unit MW (at 31 December; installed capacity, not de-rated); ts_utc the year's first day; freq P1Y; "
            "geo GB or GB-NIR; node the sheet's labels.",
            f"{len(t)} rows; the UK's transmission and distribution totals, MW: " + ", ".join(
                f"{k[:4]} {tr[k]:,.0f} and {d[k]:,.0f}" for k in tr.index if k[:4] in ("2011", "2018", "2025") and k in d.index)
            + ". Every total and the areas checked against their parts; nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/desnz_dukes_connection.py; the workbook retrieved "
            f"{base['retrieved_at']}; last modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL} (listed on {PAGE})",
            "License: Open Government Licence v3.0. Credit: Department for Energy Security and Net Zero.",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="UK Department for Energy Security and Net Zero (DESNZ)",
                                    report="Digest of UK Energy Statistics (DUKES) 2026, table 5.12, Plant installed capacity, by connection",
                                    report_url=PAGE, document_list=PAGE, license="public", tables=[NAME])])  # OGL v3.0
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
