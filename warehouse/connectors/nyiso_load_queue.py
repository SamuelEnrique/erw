#!/usr/bin/env python3
"""NYISO load interconnection requests, by request and by zone (session 140, approved pull).

Energy Research Warehouse (ERW) connector. One entities table, warehouse/output/nyiso_load_queue.csv, and the page's
file site/data/nyiso_load_queue.json, from the two load sheets of NYISO's interconnection queue workbook:

    https://www.nyiso.com/documents/20142/1407078/NYISO-Interconnection-Queue.xlsx      (linked from
    https://www.nyiso.com/interconnections; the same workbook iso_queues.py reads for the generator sheets)
    sheet "Load Projects"            one row a load interconnection request, 21 columns, then an End-Use Key and NOTES
    sheet "Load Project Tracking"    NYISO's own summary of that sheet: megawatts by zone, by status, end use and year

    python warehouse/connectors/nyiso_load_queue.py --pull            # one request: the workbook, saved with a manifest
    python warehouse/connectors/nyiso_load_queue.py --write           # the table and the page's file (needs the data lock)
    python warehouse/connectors/nyiso_load_queue.py --pull --write    # both, as the weekly queue refresh runs it
    python warehouse/connectors/nyiso_load_queue.py --site            # the page's file only, from the saved workbook
    python warehouse/connectors/nyiso_load_queue.py --terms           # one request: NYISO's legal notice, saved, the quoted sentences checked
    python warehouse/connectors/nyiso_load_queue.py --write --out-dir DIR     # a trial: table, log, registry, status and the page's file under DIR
    [--site-file PATH] [--ceiling ROWS]

WHAT IS WRITTEN. One row a load request, in the shape and with the column names of nyiso_interconnection_queue
(entities, docs/datastandard.md), every column the sheet prints kept:

    entity_id nyiso_load_queue:<Queue Number as printed>; entity_type project; name "Project: Project Name";
    operator "Developer Name"; capacity_mw "Peak MW load" as printed; queue_id; queue_date "IR Submission Date";
    end_use and end_use_words (NYISO's code, and its label on the tracking sheet); record_type; type_fuel; county;
    state; zone (NYISO's letter, A to K) and zone_entity (the id the ERW's NYISO price and load tables use,
    nyiso:WEST and so on); poi "Points of Interconnection"; transmission_owner "CTO/Utility";
    affected_transmission_owner; iso_status "Project Status #" as printed and iso_status_words, the words of the
    sheet's own key ("12=Under Construction"); sis_bundle; last_updated; availability_of_studies; ia_tender;
    fs_completion; proposed_initial_backfeed "Proposed Initial Backfeed Date"; sheet and sheet_row say where it is.

NOTHING IS INFERRED. A blank cell stays blank. A date column comes twice: <x>_printed is the cell as the sheet shows
it (a date cell in the cell's own format, "11-02-05"; a text cell as typed, "I/S", "11-2026"), <x>_date is YYYY-MM-DD
only where the cell holds a full date. The proposed backfeed is printed as month and year ("11-2026"): it is kept as
printed, and proposed_initial_backfeed_month holds YYYY-MM where it reads so; no day is added. White space inside a
cell is collapsed to one space; nothing else is changed (county "St; Lawrence" and state "New York" stay as typed).
geo is US-<state> where the state is a two-letter code or a state's name. status is the harmonized value the queue
tables use (active, withdrawn, completed): 0 is withdrawn, 14 (In Service Commercial) is completed, every other code
of the sheet's key is active; a code that is not in the sheet's key fails the run. status_date is empty: the sheet
prints no date of withdrawal or of entering service.

THE TWO SHEETS. "Load Project Tracking" holds no requests: its cells are formulas over "Load Projects" (megawatts by
zone for each status, end use and submission year). So the table's rows are the rows of "Load Projects" (column sheet
says so), and the tracking sheet's printed cells are kept whole in the page's file (nyiso_summary), with every
difference between its cells and the sums of the rows listed there (checks) and in the log.

IN LINE. The page's file counts a request as in line when its status is in the sheet's key and its words say neither
"Withdrawn" nor "In Service" (so not 0, 13, 14 or 15). Withdrawn and in-service requests, and a request with no
status printed, are listed and counted apart, never added to "in line". Counts and megawatts are sums within NYISO's
one list, by zone.

REQUESTS. --pull makes one request, the workbook. When the weekly queue run (iso_queues.py) saved the same workbook
on this machine within the last seven days, that copy is used and no request is made (the manifest says which).
--terms makes one request, the legal notice. --write and --site make none.

CEILING. 50,000 rows (--ceiling): a workbook that would write more fails with no file written.

TERMS. License: public, with a caution (the standing reading of NYISO's notice since session 65, kept by sessions
85, 136 and 138). The workbook is on www.nyiso.com, the site the notice covers
(https://www.nyiso.com/legal-notice, fetched and read 7 October 2026): "Access to this Web site does not confer any
license or ownership interest in either the form or content of the Web site, including any confidential or
proprietary information or intellectual property of any kind or nature, and the NYISO hereby expressly reserves such
rights and property in its entirety." and "Downloading, republishing, retransmitting, reproducing, or other use of
any image or video on this website as a stand-alone file is strictly prohibited". No license is granted and NYISO
reserves its rights; the one prohibition names images and video, not data. The notice neither permits nor forbids
showing rows of the queue on a public page. A person can rule otherwise.

SESSION 149, THE RULING APPLIED (the owner, 7 October 2026: NYISO's terms are quoted, and the rows are shown only if
the words allow it). Read by its words alone, the notice does not allow it: it confers no license in the content of
the site, "expressly reserves such rights and property in its entirety", and closes "All Rights Reserved"; no sentence
grants a reader leave to copy, redistribute or display anything. So from session 149 the table is INTERNAL (its
header, the registry and so the Redivis dataset it goes to), and the page's file (site/data/nyiso_load_queue.json)
holds the source, the notice's sentences and the reading, and NO request, zone or megawatt: the page reads "NYISO's
terms do not allow it" where the megawatts in line stood. SHOWN below is the one switch; a person turns it.

MISO and PJM are never requested here. NYISO is asked for only if iso_prices.paused("nyiso") says it is not paused.
"""

import argparse
import csv
import datetime as dt
import decimal
import glob
import hashlib
import html
import json
import os
import re
import sys
import traceback
from email.utils import parsedate_to_datetime

import pandas as pd
import requests

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "nyiso_load_queue"
CONNECTOR = "nyiso_load_queue"
ISO = "nyiso"
SOURCE = "nyiso:load_queue"
PUBLISHER = "New York Independent System Operator (NYISO)"
URL = "https://www.nyiso.com/documents/20142/1407078/NYISO-Interconnection-Queue.xlsx"
PAGE = "https://www.nyiso.com/interconnections"
TERMS_URL = "https://www.nyiso.com/legal-notice"
TERMS_READ = "2026-10-07"  # the day the notice was fetched and these sentences were checked in it word for word (--terms)
TERMS_SHORT = "Access to this Web site does not confer any license or ownership interest in either the form or content of the Web site"
TERMS_QUOTES = [
    "The NYISO maintains this Web site for the benefit of its Market Participants and other authorized users.",
    TERMS_SHORT + ", including any confidential or proprietary information or intellectual property of any kind or nature, "
    "and the NYISO hereby expressly reserves such rights and property in its entirety.",
    "Downloading, republishing, retransmitting, reproducing, or other use of any image or video on this website as a stand-alone file is strictly prohibited",
    # session 149: the two sentences on NYISO's marks and the notice's closing line, so that every sentence that bears
    # on copying, redistribution and display is quoted (the apostrophes and the sign are the page's own)
    "The NYISO\u2019s trademarks (including its logo) are owned by the NYISO and may only be used with the NYISO\u2019s prior written permission.",
    "Even if prior written permission is obtained, the NYISO may revoke permission to use the NYISO\u2019s trademarks at any time.",
    "Copyright \u00a9 2026 New York Independent System Operator. All Rights Reserved.",
]
TERMS_READING_140 = ("The notice grants no license and reserves NYISO's rights; what it forbids by name is republishing an image or a "
                     "video as a stand-alone file, not data. It neither permits nor forbids showing rows of the queue on a public page. "
                     "The ERW's standing reading since session 65: public, with a caution. A person can rule otherwise.")
# Session 149, the owner's ruling of 7 October 2026 ("quoted, and the rows shown if they allow it"), read by the words alone
TERMS_READING = ("By its words the notice does not allow showing the rows. It confers no license in the form or content of the site, "
                 "reserves NYISO's rights and property in their entirety, and closes with all rights reserved; no sentence grants leave "
                 "to copy, redistribute or display. What it forbids by name is republishing an image or a video as a stand-alone file, "
                 "and that is not a grant for anything else. So the requests are held internal and are not shown. A person can rule otherwise.")
SHOWN = False        # the one switch: True puts the requests, the zones and their megawatts back in the page's file
LICENSE = "public" if SHOWN else "internal"
HELD_WORDS = "NYISO's terms do not allow it"
HELD_WHY = ("NYISO's legal notice confers no license in the content of its site and reserves its rights in their entirety. "
            "The load requests are held and not shown.")
SHEET = "Load Projects"
TRACKING = "Load Project Tracking"
RAW = os.path.join(ip.ROOT, "warehouse", "raw", CONNECTOR)  # the saved workbook; a trial reads the same file
MANIFEST = os.path.join(RAW, "manifest.csv")
MANIFEST_COLS = ["file", "url", "answered_at", "status", "bytes", "sha256", "last_modified", "retrieved_at", "origin"]
SITE_FILE = os.path.join(ip.ROOT, "site", "data", "nyiso_load_queue.json")
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
REUSE_DAYS = 7
CEILING = 50000

# NYISO's load zone letters and the zone names its price and load files use (the ERW's entity is nyiso:<name>, as in
# nyiso_dam_zone_prices and nyiso_zone_load_hourly). NYISO's own pairing: A West, B Genesee, C Central, D North,
# E Mohawk Valley, F Capital, G Hudson Valley, H Millwood, I Dunwoodie, J New York City, K Long Island.
ZONES = {"A": "WEST", "B": "GENESE", "C": "CENTRL", "D": "NORTH", "E": "MHK VL", "F": "CAPITL", "G": "HUD VL",
         "H": "MILLWD", "I": "DUNWOD", "J": "N.Y.C.", "K": "LONGIL"}

# the sheet's header -> the table's column. A header that is not here, or one that is missing, fails the run.
TEXT_COLS = {"Queue Number": "queue_id", "Developer Name": "operator", "Project: Project Name": "name",
             "Peak MW load": "capacity_mw", "End-Use": "end_use", "Record Type Name": "record_type",
             "Type/Fuel": "type_fuel", "County": "county", "State": "state", "NYISO Zone": "zone",
             "Points of Interconnection": "poi", "CTO/Utility": "transmission_owner",
             "Affected Transmission Owner (ATO)": "affected_transmission_owner", "Project Status #": "iso_status",
             "SIS Bundle": "sis_bundle", "Availability of Studies": "availability_of_studies"}
DATE_COLS = {"IR Submission Date": "queue_date", "Last Updated Date": "last_updated", "IA Tender Date": "ia_tender",
             "FS Completion Date": "fs_completion", "Proposed Initial Backfeed Date": "proposed_initial_backfeed"}
EM_DASH = chr(0x2014)
BULLETS = (chr(0x25CF), chr(0x2022))  # the marks NYISO's notes begin with
HEADERS = ["Queue Number", "Developer Name", "Project: Project Name", "IR Submission Date", "Peak MW load", "End-Use",
           "Record Type Name", "Type/Fuel", "County", "State", "NYISO Zone", "Points of Interconnection",
           "CTO/Utility", "Affected Transmission Owner (ATO)", "Project Status #", "SIS Bundle", "Last Updated Date",
           "Availability of Studies", "IA Tender Date", "FS Completion Date", "Proposed Initial Backfeed Date"]
STATES = {"New York": "NY", "New Jersey": "NJ", "Pennsylvania": "PA", "Connecticut": "CT", "Massachusetts": "MA",
          "Vermont": "VT"}

ENTITY_COLS = ["entity_id", "entity_type", "name", "geo", "lat", "lon", "capacity_mw", "status", "status_date",
               "operator", "source"]
EXTRA_COLS = ["source_url", "retrieved_at", "vintage", "workbook_last_modified", "sheet", "sheet_row", "queue_id",
              "county", "state", "poi", "zone", "zone_entity", "transmission_owner", "affected_transmission_owner",
              "end_use", "end_use_words", "record_type", "type_fuel", "iso_status", "iso_status_words", "sis_bundle",
              "availability_of_studies", "queue_date_printed", "queue_date", "last_updated_printed",
              "last_updated_date", "ia_tender_printed", "ia_tender_date", "fs_completion_printed",
              "fs_completion_date", "proposed_initial_backfeed_printed", "proposed_initial_backfeed_date",
              "proposed_initial_backfeed_month"]
COLS = ENTITY_COLS + EXTRA_COLS

WITHDRAWN, IN_SERVICE, IN_LINE, NO_STATUS = "withdrawn", "in service", "in line", "no status printed"


# ---------------------------------------------------------------------------
# The saved file and its manifest
# ---------------------------------------------------------------------------

def now_utc():
    return dt.datetime.now(dt.timezone.utc)


def rel(path):
    """A path as the repository names it, with forward slashes; the path itself where it is on another drive."""
    try:
        return os.path.relpath(path, ip.ROOT).replace(os.sep, "/")
    except ValueError:
        return path


def read_manifest():
    if not os.path.exists(MANIFEST):
        return []
    with open(MANIFEST, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def append_manifest(row):
    os.makedirs(RAW, exist_ok=True)
    new = not os.path.exists(MANIFEST)
    with open(MANIFEST, "a", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=MANIFEST_COLS, lineterminator="\n")
        if new:
            w.writeheader()
        w.writerow(row)


def daily_copy(now, raw_root=None):
    """The newest copy of the same workbook the weekly queue run (iso_queues.py) saved on this machine in the last
    REUSE_DAYS days, as (bytes, its manifest row, its path), or None."""
    raw_root = raw_root or os.path.join(ip.ROOT, "warehouse", "raw")
    best = None
    for manifest in glob.glob(os.path.join(raw_root, "iso_queues", "*", "manifest.csv")):
        with open(manifest, encoding="utf-8", newline="") as f:
            for r in csv.DictReader(f):
                url = r.get("url", "")
                if "nyiso.com" not in url or "Interconnection-Queue" not in url or r.get("status") != "200":
                    continue
                path = os.path.join(os.path.dirname(manifest), r["file"])
                if not os.path.exists(path):
                    continue
                got = dt.datetime.strptime(r["retrieved_at"], "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=dt.timezone.utc)
                if dt.timedelta(0) <= now - got <= dt.timedelta(days=REUSE_DAYS) and (best is None or got > best[0]):
                    best = (got, r, path)
    if best is None:
        return None
    with open(best[2], "rb") as f:
        content = f.read()
    if hashlib.sha256(content).hexdigest() != best[1]["sha256"] or content[:2] != b"PK":
        return None
    return content, best[1], best[2]


def pull(log, raw_root=None):
    """One request for the workbook, or none when the weekly queue run's copy on this machine is seven days old or
    less. The file and its manifest line go under warehouse/raw/nyiso_load_queue/."""
    now = now_utc()
    stamp = now.strftime("%Y%m%dT%H%M%SZ")
    hit = daily_copy(now, raw_root)
    if hit:
        content, r, path = hit
        origin = "copied from " + rel(path)
        row = dict(url=URL, answered_at=r["url"], status="200", last_modified=r.get("last_modified", ""),
                   retrieved_at=r["retrieved_at"], origin=origin)
        log(f"pull: no request made. The weekly queue run's copy of the workbook is used ({origin}, retrieved {r['retrieved_at']})")
        made = 0
    else:
        log(f"pull: GET {URL}")
        resp = requests.get(URL, headers=UA, timeout=180)
        content = resp.content
        row = dict(url=URL, answered_at=resp.url, status=str(resp.status_code),
                   last_modified=resp.headers.get("Last-Modified", ""),
                   retrieved_at=now_utc().strftime("%Y-%m-%dT%H:%M:%SZ"), origin="requested by nyiso_load_queue.py --pull")
        made = 1
        log(f"pull: HTTP {resp.status_code}, {len(content)} bytes, content type {resp.headers.get('Content-Type', '')!r}")
    ok = row["status"] == "200" and content[:2] == b"PK"
    fname = f"{stamp}_NYISO-Interconnection-Queue.xlsx" if ok else f"{stamp}_refused_response.bin"
    os.makedirs(RAW, exist_ok=True)
    with open(os.path.join(RAW, fname), "wb") as f:
        f.write(content)
    row.update(file=fname, bytes=str(len(content)), sha256=hashlib.sha256(content).hexdigest())
    append_manifest(row)
    log(f"pull: saved warehouse/raw/{CONNECTOR}/{fname} ({row['bytes']} bytes, sha256 {row['sha256']}); requests made: {made}")
    print(f"{CONNECTOR} pull: {fname}, {row['bytes']} bytes, requests made: {made}")
    if not ok:
        raise RuntimeError(f"the answer is not the workbook (HTTP {row['status']}, first bytes {content[:8]!r}); kept as {fname}, nothing is written from it")
    return row


def page_text(content):
    """A web page's visible text, white space collapsed (scripts, styles and tags dropped)."""
    t = content.decode("utf-8", "replace")
    t = re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", t)
    t = html.unescape(re.sub(r"(?s)<[^>]+>", " ", t))
    return re.sub(r"\s+", " ", t).strip()


def terms(log):
    """One request: NYISO's legal notice, saved beside the workbook. Each sentence this connector quotes must be in
    the page word for word (white space collapsed), or the stage fails."""
    stamp = now_utc().strftime("%Y%m%dT%H%M%SZ")
    log(f"terms: GET {TERMS_URL}")
    resp = requests.get(TERMS_URL, headers=UA, timeout=120)
    content = resp.content
    fname = f"{stamp}_legal-notice.html"
    os.makedirs(RAW, exist_ok=True)
    with open(os.path.join(RAW, fname), "wb") as f:
        f.write(content)
    append_manifest(dict(file=fname, url=TERMS_URL, answered_at=resp.url, status=str(resp.status_code), bytes=str(len(content)),
                         sha256=hashlib.sha256(content).hexdigest(), last_modified=resp.headers.get("Last-Modified", ""),
                         retrieved_at=now_utc().strftime("%Y-%m-%dT%H:%M:%SZ"), origin="requested by nyiso_load_queue.py --terms"))
    log(f"terms: HTTP {resp.status_code}, {len(content)} bytes, saved warehouse/raw/{CONNECTOR}/{fname}; requests made: 1")
    if resp.status_code != 200:
        raise RuntimeError(f"the legal notice answered HTTP {resp.status_code}; kept as {fname}")
    text = page_text(content)
    missing = [q for q in TERMS_QUOTES if q not in text]
    for q in TERMS_QUOTES:
        log(f"terms: {'found' if q not in missing else 'NOT FOUND'} word for word: \"{q}\"")
        print(f"{CONNECTOR} terms: {'found' if q not in missing else 'NOT FOUND'}: \"{q}\"")
    if missing:
        raise RuntimeError(f"{len(missing)} quoted sentence(s) are not in the legal notice as saved ({fname}); read it and correct the quote")
    return fname


def saved_workbook():
    """The newest saved workbook whose bytes match its manifest line, as (path, manifest row)."""
    for r in reversed(read_manifest()):
        path = os.path.join(RAW, r["file"])
        if r["status"] != "200" or not r["file"].endswith(".xlsx") or not os.path.exists(path):
            continue
        with open(path, "rb") as f:
            if hashlib.sha256(f.read()).hexdigest() == r["sha256"]:
                return path, r
    raise RuntimeError(f"no saved workbook under warehouse/raw/{CONNECTOR}/ (run --pull first)")


def terms_saved(log):
    """No request. Where --terms saved the legal notice on this machine, each quoted sentence must still be in that
    page word for word, or the run fails; where it did not (another machine), the log says the quotes were not
    checked here."""
    for r in reversed(read_manifest()):
        path = os.path.join(RAW, r["file"])
        if r["url"] == TERMS_URL and r["status"] == "200" and os.path.exists(path):
            with open(path, "rb") as f:
                text = page_text(f.read())
            missing = [q for q in TERMS_QUOTES if q not in text]
            if missing:
                raise RuntimeError(f"{len(missing)} quoted sentence(s) of NYISO's legal notice are not in the page saved {r['retrieved_at']} ({r['file']}): {missing[0][:80]!r}")
            log(f"  terms: the {len(TERMS_QUOTES)} quoted sentences are in the legal notice saved {r['retrieved_at']} ({r['file']}), word for word")
            return r["retrieved_at"]
    log("  terms: no saved legal notice on this machine; the quoted sentences were not checked in this run (--terms fetches and checks)")
    return ""


# ---------------------------------------------------------------------------
# Reading the two sheets
# ---------------------------------------------------------------------------

def squeeze(s):
    return re.sub(r"\s+", " ", str(s)).strip()


def number_text(v):
    """A number as the sheet prints it in a General cell: 40, 50.2, 14.38."""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else repr(v)
    return squeeze(v)


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return number_text(v)


def date_printed(v, number_format):
    """A date cell as the sheet shows it, in the cell's own number format; (text, True) when the format is one this
    reader knows, else the ISO date and False."""
    f = (number_format or "").lower()
    if f == "mm-dd-yy":
        return v.strftime("%m-%d-%y"), True
    if f == "m/d/yy;@" or f == "m/d/yy":
        return f"{v.month}/{v.day}/{v.strftime('%y')}", True
    if f == "mmm-yy":
        return v.strftime("%b-%y"), True
    if f == "mm/dd/yyyy":
        return v.strftime("%m/%d/%Y"), True
    if f == "yyyy-mm-dd":
        return v.strftime("%Y-%m-%d"), True
    return v.strftime("%Y-%m-%d"), False


def date_cell(v, number_format):
    """(printed, YYYY-MM-DD or "", YYYY-MM or "", format known). A full date only where the cell holds one; a month
    where the text reads MM-YYYY or MM/YYYY; nothing else is read into a text cell."""
    if v is None:
        return "", "", "", True
    if isinstance(v, (dt.datetime, dt.date)):
        printed, known = date_printed(v, number_format)
        return printed, v.strftime("%Y-%m-%d"), v.strftime("%Y-%m"), known
    t = squeeze(v) if isinstance(v, str) else number_text(v)
    m = re.fullmatch(r"(\d{1,2})[-/](\d{4})", t)
    if m and 1 <= int(m.group(1)) <= 12:
        return t, "", f"{m.group(2)}-{int(m.group(1)):02d}", True
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", t)
    if m:
        try:
            d = dt.date(int(m.group(3)), int(m.group(1)), int(m.group(2)))
            return t, d.strftime("%Y-%m-%d"), d.strftime("%Y-%m"), True
        except ValueError:
            pass
    return t, "", "", True


def status_key(notes):
    """The sheet's own key of "Project Status #", code -> words, read from its NOTES ("0=Withdrawn, 1=Scoping Meeting
    Pending, ..."). The key runs from the line that names it to the next bullet."""
    start = next((i for i, n in enumerate(notes) if re.search(r"project status\s*#\s*key", n, re.I)), None)
    if start is None:
        raise RuntimeError(f'"{SHEET}": the NOTES hold no "Project status # Key"; the status codes cannot be put in NYISO\'s words')
    text = notes[start].split(":", 1)[1] if ":" in notes[start] else notes[start]
    for n in notes[start + 1:]:
        if n.lstrip()[:1] in BULLETS:  # the next note begins
            break
        text += " " + n
    marks = list(re.finditer(r"(?:^|(?<=[\s,]))([0-9]{1,2}[A-Z]?|P)=", text))
    key = {}
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
        words = squeeze(text[m.end():end]).strip(" ,")
        if m.group(1) in key and key[m.group(1)] != words:
            raise RuntimeError(f'"{SHEET}": the status key gives code {m.group(1)} twice with different words')
        key[m.group(1)] = words
    if not key:
        raise RuntimeError(f'"{SHEET}": the status key could not be read')
    return key


def group_of(code, key):
    """Where a request stands for the page's count, from NYISO's own words for its status code."""
    if code == "":
        return NO_STATUS
    words = key[code].lower()
    if "withdrawn" in words:
        return WITHDRAWN
    if "in service" in words or "in-service" in words:
        return IN_SERVICE
    return IN_LINE


def harmonized(code, key):
    """The queue tables' status vocabulary (iso_queues.py): withdrawn, completed (in service, commercial), else active."""
    if code == "":
        return ""
    words = key[code].lower()
    if "withdrawn" in words:
        return "withdrawn"
    if words == "in service commercial":
        return "completed"
    return "active"


def read_projects(ws, log):
    """The "Load Projects" sheet as (rows, status key, the End-Use Key's lines, the NOTES lines). A row is a dict of
    the table's columns without the provenance columns."""
    grid = list(ws.iter_rows())
    if not grid:
        raise RuntimeError(f'"{SHEET}" is empty')
    head = {}
    for c in grid[0]:
        if c.value is not None and squeeze(c.value):
            head[squeeze(c.value)] = c.column
    unknown = [h for h in head if h not in HEADERS]
    missing = [h for h in HEADERS if h not in head]
    if unknown or missing:
        raise RuntimeError(f'"{SHEET}": the header row changed. New column(s) {unknown}, missing column(s) {missing}; '
                           "every column the sheet prints is kept, so the connector must be told its name")
    rows, lines, unknown_formats = [], [], {}
    for cells in grid[1:]:
        r = cells[0].row
        at = {c.column: c for c in cells}
        vals = {h: at[i] for h, i in head.items()}
        beyond = [c for c in cells if c.column not in head.values() and c.value is not None and squeeze(c.value)]
        if beyond:
            raise RuntimeError(f'"{SHEET}" row {r}: a value under no header ({beyond[0].coordinate}); not written')
        filled = [h for h, c in vals.items() if c.value is not None and squeeze(c.value) != ""]
        if not filled:
            continue
        if len(filled) == 1 and filled[0] in ("Queue Number", "Developer Name"):
            lines.append(squeeze(vals[filled[0]].value))  # a line of the End-Use Key or of the NOTES, not a request
            continue
        qid = cell_text(vals["Queue Number"].value)
        if qid == "":
            raise RuntimeError(f'"{SHEET}" row {r}: a row with data and no queue number; not written')
        row = {"sheet": SHEET, "sheet_row": str(r)}
        for h, col in TEXT_COLS.items():
            row[col] = cell_text(vals[h].value)
        for h, stem in DATE_COLS.items():
            printed, iso, month, known = date_cell(vals[h].value, vals[h].number_format)
            if not known:
                unknown_formats[vals[h].number_format] = unknown_formats.get(vals[h].number_format, 0) + 1
            row[stem + "_printed"], row[stem if stem.endswith("_date") else stem + "_date"] = printed, iso
            if stem == "proposed_initial_backfeed":
                row[stem + "_month"] = month
        rows.append(row)
    for f, n in unknown_formats.items():
        log(f'  "{SHEET}": {n} date cell(s) in a number format this reader does not know ({f!r}); printed as YYYY-MM-DD')
    if not rows:
        raise RuntimeError(f'"{SHEET}": no request rows')
    cut = next((i for i, n in enumerate(lines) if n.upper().startswith("NOTES")), len(lines))
    end_use_key = [n for n in lines[:cut] if n]
    notes = [n for n in lines[cut:] if n]
    key = status_key(notes)
    bad = sorted({r["iso_status"] for r in rows if r["iso_status"] and r["iso_status"] not in key})
    if bad:
        raise RuntimeError(f'"{SHEET}": status code(s) {bad} are not in the sheet\'s own key {sorted(key)}; not guessed, nothing written')
    mw_bad = [r["queue_id"] for r in rows if r["capacity_mw"] and not re.fullmatch(r"-?\d+(\.\d+)?", r["capacity_mw"])]
    if mw_bad:
        raise RuntimeError(f'"{SHEET}": "Peak MW load" is not a number for queue number(s) {mw_bad[:5]}')
    zone_bad = sorted({r["zone"] for r in rows if r["zone"] and r["zone"] not in ZONES})
    if zone_bad:
        raise RuntimeError(f'"{SHEET}": zone(s) {zone_bad} are not NYISO\'s letters A to K')
    return rows, key, end_use_key, notes


def read_tracking(ws, log):
    """The "Load Project Tracking" sheet as printed: its title, its blocks (a header row "MW By ...", one row a
    status, end use or year with megawatts under each zone letter and NYCA, and a Total row) and its notes. The
    megawatts are the values the workbook holds for its formula cells (what the sheet shows)."""
    out = {"sheet": TRACKING, "title": "", "blocks": [], "notes": []}
    block = None
    for row in ws.iter_rows():
        cells = {c.column: c.value for c in row if c.value is not None and squeeze(c.value) != ""}
        if not cells:
            continue
        r = row[0].row
        texts = [squeeze(v) for v in cells.values() if isinstance(v, str)]
        lead = next((t for t in texts if t.startswith("MW By")), None)
        if lead:
            cols = {i: squeeze(v) for i, v in cells.items() if isinstance(v, str) and (squeeze(v) in ZONES or squeeze(v) in ("NYCA", "Running Total"))}
            block = {"block": lead, "sheet_row": r, "columns": list(cols.values()), "rows": [], "_cols": cols}
            out["blocks"].append(block)
            continue
        if block is not None and any(i in block["_cols"] for i in cells):
            labels = [cell_text(v) for i, v in sorted(cells.items()) if i < min(block["_cols"])]
            mw = {}
            for i, name in block["_cols"].items():
                v = cells.get(i)
                if v is None:
                    mw[name] = None
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    mw[name] = as_number(decimal.Decimal(repr(v)).quantize(decimal.Decimal("0.0001")).normalize())
                else:
                    raise RuntimeError(f'"{TRACKING}" row {r}: {name} is not a number ({v!r}); the workbook holds no computed value')
            is_total = bool(labels) and labels[-1].lower() == "total"
            block["rows"].append({"sheet_row": r, "label": labels[-1] if labels else "",
                                  "code": labels[0] if len(labels) > 1 and block["block"] != "MW By IR Submission Date" else "",
                                  "total_row": is_total, "mw": mw})
            if is_total:
                block = None
            continue
        if len(texts) == 1 and len(cells) == 1:
            if not out["title"] and not out["blocks"]:
                out["title"] = texts[0]
            else:
                out["notes"].append({"sheet_row": r, "text": texts[0]})
    for b in out["blocks"]:
        del b["_cols"]
    if not any(b["block"] == "MW By Status" for b in out["blocks"]):
        raise RuntimeError(f'"{TRACKING}": no "MW By Status" block; the sheet changed')
    log(f'  "{TRACKING}": {len(out["blocks"])} blocks ({", ".join(b["block"] + " " + str(len(b["rows"])) + " rows" for b in out["blocks"])}), {len(out["notes"])} notes')
    return out


def end_use_labels(tracking):
    """NYISO's label for each end-use code, from the tracking sheet's "MW By End-Use" block (DAT-AI: "AI Datacenter")."""
    for b in tracking["blocks"]:
        if b["block"] == "MW By End-Use":
            return {r["code"]: r["label"] for r in b["rows"] if r["code"] and not r["total_row"]}
    return {}


def parse_workbook(path, log):
    """Both sheets of a saved workbook: (rows, status key, End-Use Key lines, NOTES lines, the tracking sheet)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)  # data_only: a formula cell gives the value the sheet shows
    for s in (SHEET, TRACKING):
        if s not in wb.sheetnames:
            raise RuntimeError(f"the workbook has no sheet {s!r}; its sheets are {wb.sheetnames}")
    rows, key, end_use_key, notes = read_projects(wb[SHEET], log)
    tracking = read_tracking(wb[TRACKING], log)
    labels = end_use_labels(tracking)
    for r in rows:
        r["end_use_words"] = labels.get(r["end_use"], "")
        r["iso_status_words"] = key.get(r["iso_status"], "")
    unlabelled = sorted({r["end_use"] for r in rows if r["end_use"] and not r["end_use_words"]})
    if unlabelled:
        log(f"  end-use code(s) {unlabelled} have no label on the tracking sheet; end_use_words left empty")
    return rows, key, end_use_key, notes, tracking


# ---------------------------------------------------------------------------
# The table
# ---------------------------------------------------------------------------

def build_table(rows, key, rec, log):
    got = rec["retrieved_at"]
    lm = ""
    if rec.get("last_modified"):
        try:
            lm = ip.utc_iso(pd.Timestamp(parsedate_to_datetime(rec["last_modified"])))
        except (TypeError, ValueError):
            lm = ""
    out, seen = [], {}
    for r in rows:
        seen[r["queue_id"]] = seen.get(r["queue_id"], 0) + 1
    count = {}
    for r in rows:
        eid = f"{NAME}:{r['queue_id']}"
        if seen[r["queue_id"]] > 1:  # as iso_queues.py: one queue number on several rows gets :<n> in the sheet's order
            count[r["queue_id"]] = count.get(r["queue_id"], 0) + 1
            eid += f":{count[r['queue_id']]}"
        st = STATES.get(r["state"], r["state"])
        row = dict(r)
        row.update(entity_id=eid, entity_type="project", geo=f"US-{st}" if re.fullmatch(r"[A-Z]{2}", st) else "", lat="", lon="",
                   status=harmonized(r["iso_status"], key), status_date="", source=SOURCE, source_url=URL, retrieved_at=got,
                   vintage=got[:10], workbook_last_modified=lm, zone_entity=f"nyiso:{ZONES[r['zone']]}" if r["zone"] else "")
        out.append(row)
    dup = [q for q, n in seen.items() if n > 1]
    if dup:
        log(f"  {len(dup)} queue number(s) are on more than one row ({dup[:5]}); ids get :<n> in the sheet's order")
    df = pd.DataFrame(out, columns=COLS).fillna("")
    log(f"  {len(df)} requests on \"{SHEET}\": " + ", ".join(f"{k or '(no status)'} {v}" for k, v in df["status"].value_counts().items()))
    return df


# ---------------------------------------------------------------------------
# The page's file
# ---------------------------------------------------------------------------

def dec(text):
    return decimal.Decimal(text) if text else decimal.Decimal(0)


def as_number(d):
    """A sum of megawatts as a JSON number, exact to the decimals the sheet prints."""
    d = decimal.Decimal(d)
    return int(d) if d == d.to_integral_value() else float(str(d.normalize()))


def tally(rows):
    return {"requests": len(rows), "mw": as_number(sum((dec(r["capacity_mw"]) for r in rows), decimal.Decimal(0))),
            "requests_without_mw": sum(1 for r in rows if not r["capacity_mw"])}


def status_order(code):
    m = re.match(r"(\d+)(.*)", code)
    return (int(m.group(1)), m.group(2)) if m else (999, code)


def by_status(rows, key):
    out = []
    for code in sorted({r["iso_status"] for r in rows}, key=status_order):
        some = [r for r in rows if r["iso_status"] == code]
        t = tally(some)
        out.append({"status_code": code, "status": key.get(code, ""), "requests": t["requests"], "mw": t["mw"]})
    return out


def compare_tracking(rows, key, tracking):
    """Every cell of the tracking sheet's "MW By Status" block that differs from the sum of the rows of "Load
    Projects" with that status and zone, and every status the rows hold that the block has no line for."""
    block = next(b for b in tracking["blocks"] if b["block"] == "MW By Status")
    codes, checks = set(), []
    for line in block["rows"]:
        if line["total_row"]:
            continue
        code = {"-": "0", "I/S": "14"}.get(line["code"], line["code"])  # the sheet's own marks for its withdrawn and in-service lines
        if code not in key:
            checks.append({"what": f'the tracking sheet\'s line "{line["label"]}" (row {line["sheet_row"]}) has a mark, {line["code"]!r}, that is not a status code of the key', "sheet_row": line["sheet_row"]})
            continue
        codes.add(code)
        for z in ZONES:
            mine = as_number(sum((dec(r["capacity_mw"]) for r in rows if r["iso_status"] == code and r["zone"] == z), decimal.Decimal(0)))
            theirs = line["mw"].get(z)
            if theirs is not None and decimal.Decimal(str(theirs)) != decimal.Decimal(str(mine)):
                checks.append({"what": f'status {code} ("{line["label"]}"), zone {z}: the tracking sheet prints {theirs} MW; the rows of "{SHEET}" sum to {mine} MW',
                               "sheet_row": line["sheet_row"], "zone": z, "status_code": code, "tracking_sheet_mw": theirs, "rows_mw": mine})
    for code in sorted({r["iso_status"] for r in rows if r["iso_status"]} - codes, key=status_order):
        some = [r for r in rows if r["iso_status"] == code]
        checks.append({"what": f'status {code} ("{key[code]}"): {len(some)} request(s), {tally(some)["mw"]} MW, on "{SHEET}"; the tracking sheet\'s "MW By Status" block has no line for it',
                       "status_code": code, "rows_mw": tally(some)["mw"], "requests": len(some)})
    # each block's Total line against the megawatts in line by this file's rule, zone by zone and for the state
    line = [r for r in rows if group_of(r["iso_status"], key) == IN_LINE]
    for b in tracking["blocks"]:
        for t in (x for x in b["rows"] if x["total_row"]):
            for z in list(ZONES) + ["NYCA"]:
                mine = as_number(sum((dec(r["capacity_mw"]) for r in line if z == "NYCA" or r["zone"] == z), decimal.Decimal(0)))
                theirs = t["mw"].get(z)
                if theirs is not None and decimal.Decimal(str(theirs)) != decimal.Decimal(str(mine)):
                    checks.append({"what": f'"{b["block"]}", Total, {"New York Control Area (NYCA)" if z == "NYCA" else "zone " + z}: the tracking sheet prints '
                                           f'{theirs} MW; the rows in line by this file\'s rule sum to {mine} MW',
                                   "block": b["block"], "sheet_row": t["sheet_row"], "zone": z, "tracking_sheet_mw": theirs, "rows_in_line_mw": mine})
    return checks


def build_site(rows, key, end_use_key, notes, tracking, rec, built):
    """The page's file. Every count and sum is within NYISO's one list, the sheet "Load Projects"."""
    groups = {code: group_of(code, key) for code in key}
    for r in rows:
        r["_group"] = group_of(r["iso_status"], key)
    line = [r for r in rows if r["_group"] == IN_LINE]
    zones = []
    for z, zname in ZONES.items():
        mine = [r for r in line if r["zone"] == z]
        t = tally(mine)
        zones.append({"zone": z, "zone_name": zname, "zone_entity": f"nyiso:{zname}", "requests": t["requests"], "mw": t["mw"],
                      "requests_without_mw": t["requests_without_mw"], "by_status": by_status(mine, key),
                      "not_in_line": {g: {k: v for k, v in tally([r for r in rows if r["zone"] == z and r["_group"] == g]).items() if k != "requests_without_mw"}
                                      for g in (WITHDRAWN, IN_SERVICE, NO_STATUS)}})
    nozone = [r for r in line if not r["zone"]]
    total = tally(line)
    out_rows = []
    for r in rows:
        out_rows.append({
            "queue_position": r["queue_id"], "name": r["name"], "zone": r["zone"],
            "zone_name": ZONES.get(r["zone"], ""), "zone_entity": f"nyiso:{ZONES[r['zone']]}" if r["zone"] else "",
            "county": r["county"], "state": r["state"],
            "mw": as_number(dec(r["capacity_mw"])) if r["capacity_mw"] else None,
            "request_date": r["queue_date"], "request_date_printed": r["queue_date_printed"],
            "proposed_in_service_printed": r["proposed_initial_backfeed_printed"],
            "proposed_in_service_month": r["proposed_initial_backfeed_month"],
            "status_code": r["iso_status"], "status": r["iso_status_words"],
            "end_use": r["end_use"], "end_use_words": r["end_use_words"], "utility": r["transmission_owner"],
            "in_line": r["_group"] == IN_LINE, "group": r["_group"], "sheet": r["sheet"], "sheet_row": int(r["sheet_row"])})
    used = sorted({r["iso_status"] for r in rows if r["iso_status"]}, key=status_order)
    doc = {
        "built": built,
        "source": {
            "name": "NYISO Interconnection Queue (workbook), load interconnection requests",
            "publisher": PUBLISHER, "url": URL, "report_page": PAGE,
            "sheet_names": [SHEET, TRACKING],
            "retrieved": rec["retrieved_at"], "workbook_last_modified": rec.get("last_modified", ""),
            "sha256": rec["sha256"], "bytes": int(rec["bytes"]),
            "terms": {"url": TERMS_URL, "read": TERMS_READ, "quotes": TERMS_QUOTES, "reading": TERMS_READING,
                      "license": "public, with a caution" if SHOWN else "internal"},
        },
        "rule": {
            "in_line": (f'A request is counted as in line when it is on the sheet "{SHEET}", its "Project Status #" is in the sheet\'s own key, '
                        'and NYISO\'s words for that status say neither "Withdrawn" nor "In Service". Withdrawn requests, requests in service '
                        "and a request with no status printed are counted apart and are never added to in line."),
            "sums_cover_sheet": SHEET,
            "megawatts": 'the sheet\'s column "Peak MW load", as printed, summed within NYISO\'s one list',
            "proposed_in_service": 'the sheet\'s column "Proposed Initial Backfeed Date", printed as month and year; no day is added',
            "statuses_in_line": [{"status_code": c, "status": key[c], "on_the_sheet_now": c in used}
                                 for c in sorted(key, key=status_order) if groups[c] == IN_LINE],
            "statuses_not_in_line": [{"status_code": c, "status": key[c], "group": groups[c], "on_the_sheet_now": c in used}
                                     for c in sorted(key, key=status_order) if groups[c] != IN_LINE],
        },
        "zones": zones,
        "total": {"requests": total["requests"], "mw": total["mw"], "requests_without_mw": total["requests_without_mw"],
                  "requests_without_zone": len(nozone)},
        "not_in_line": {g: {k: v for k, v in tally([r for r in rows if r["_group"] == g]).items() if k != "requests_without_mw"}
                        for g in (WITHDRAWN, IN_SERVICE, NO_STATUS)},
        "all_requests_on_sheet": len(rows),
        "rows": out_rows,
        "end_use_key": end_use_key,
        "notes": notes,
        "nyiso_summary": dict(tracking, relation=(
            f'"{TRACKING}" holds no requests. Its cells are formulas over "{SHEET}": megawatts by zone for each status, end use and '
            "submission year. They are kept here as the sheet prints them; where a cell differs from the sum of the rows, the difference is in checks.")),
        "checks": compare_tracking(rows, key, tracking),
    }
    for r in rows:
        del r["_group"]
    if sum(z["requests"] for z in zones) + len(nozone) != total["requests"]:
        raise RuntimeError("the zones' requests do not add to the total; the page's file is not written")
    if sum((decimal.Decimal(str(z["mw"])) for z in zones), decimal.Decimal(0)) + sum((dec(r["capacity_mw"]) for r in nozone), decimal.Decimal(0)) != decimal.Decimal(str(total["mw"])):
        raise RuntimeError("the zones' megawatts do not add to the total; the page's file is not written")
    return doc


def held_site(doc):
    """Session 149: the page's file when the rows are not shown: where the list is, the notice's sentences and the
    reading, and the placeholder's words. No request, no zone, no megawatt, no count of requests in line."""
    return {"built": doc["built"], "shown": False, "license": "internal", "words": HELD_WORDS, "why": HELD_WHY,
            "source": {k: doc["source"][k] for k in ("name", "publisher", "url", "report_page", "sheet_names", "terms")}}


def write_site(doc, path, log):
    text = json.dumps(doc, indent=1, ensure_ascii=False) + "\n"
    if EM_DASH in text:  # house rule: no em dash in a file of this repository; the table keeps the cell as typed
        log(f"  {text.count(EM_DASH)} em dash(es) in NYISO's text are written as a comma in the page's file")
        text = text.replace(EM_DASH, ", ")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)
    os.replace(tmp, path)
    if doc.get("shown") is False:
        line = f"{os.path.basename(path)}: the requests are held internal and not shown ({doc['words']}); the file holds no request and no megawatt"
    else:
        line = (f"{os.path.basename(path)}: in line {doc['total']['requests']} requests, {doc['total']['mw']} MW; "
                f"{doc['all_requests_on_sheet']} requests on the sheet; {len(doc['checks'])} difference(s) from the tracking sheet")
    log("  " + line + f" -> {path}")
    print(line)


# ---------------------------------------------------------------------------

def header_lines(run_id, rec, df, doc, wb_path):
    t, n = doc["total"], doc["not_in_line"]
    return [
        "Energy Research Warehouse (ERW): NYISO load interconnection requests",
        "Shape: entities (docs/datastandard.md v0). One row per load interconnection request; entity_type project; "
        "status harmonized (active, withdrawn, completed), iso_status NYISO's own code and iso_status_words the words of the sheet's key.",
        f"Retrieved: {rec['retrieved_at']} (UTC); built {run_id} by warehouse/connectors/nyiso_load_queue.py",
        f"Run log: {rel(os.path.join(ip.LOG_DIR, CONNECTOR + '_' + run_id + '.log'))}",
        f"Raw file: warehouse/raw/{CONNECTOR}/{os.path.basename(wb_path)} (not in git), {rec['bytes']} bytes, sha256 {rec['sha256']}, "
        f"Last-Modified {rec.get('last_modified') or 'not given'} ({rec['origin']})",
        f"Source: {SOURCE} NYISO Interconnection Queue workbook, sheets \"{SHEET}\" and \"{TRACKING}\", {URL}",
        f"  report page: {PAGE}",
        f"Rows: the {len(df)} requests of the sheet \"{SHEET}\". \"{TRACKING}\" holds no requests: its cells are NYISO's formulas over "
        f"\"{SHEET}\" (megawatts by zone for each status, end use and year); they are kept as printed in site/data/nyiso_load_queue.json.",
        "capacity_mw: \"Peak MW load\" as printed. Dates: <x>_printed is the cell as the sheet shows it, <x>_date is YYYY-MM-DD only where the "
        "cell holds a full date; the proposed backfeed is printed as month and year and proposed_initial_backfeed_month is YYYY-MM, no day added. "
        "A blank cell stays blank; nothing is inferred. status_date is empty: the sheet prints no date of withdrawal or of entering service.",
        f"In line by the page's rule (status in the key, its words neither Withdrawn nor In Service): {t['requests']} requests, {t['mw']} MW. "
        f"Apart: withdrawn {n[WITHDRAWN]['requests']} requests, {n[WITHDRAWN]['mw']} MW; in service {n[IN_SERVICE]['requests']} requests, {n[IN_SERVICE]['mw']} MW.",
        "Snapshot: the table holds the sheet as retrieved; the next run replaces it.",
        ("License: public, with a caution (sessions 65, 85, 136, 138). NYISO's legal notice grants no license and forbids republishing "
         "\"any image or video on this website as a stand-alone file\", which names images and video, not data (https://www.nyiso.com/legal-notice). "
         "A person can rule otherwise.") if SHOWN else
        ("License: internal. Session 149, on the owner's ruling of 7 October 2026 (the rows are shown only if NYISO's terms allow it): NYISO's legal "
         "notice confers no license (\"" + TERMS_SHORT + "\") and \"expressly reserves such rights and property in its entirety\" "
         "(https://www.nyiso.com/legal-notice, read 7 October 2026). Not published, in no public dataset, no download and no page. "
         "A person can rule otherwise."),
    ]


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--pull", action="store_true", help="fetch the workbook once and save it with a manifest line")
    ap.add_argument("--write", action="store_true", help="build the table and the page's file from the saved workbook")
    ap.add_argument("--site", action="store_true", help="build the page's file only")
    ap.add_argument("--terms", action="store_true", help="fetch NYISO's legal notice once and check the quoted sentences")
    ap.add_argument("--out-dir", help="a trial: everything but the saved workbook goes under this directory")
    ap.add_argument("--site-file", help="where the page's file goes (default site/data/nyiso_load_queue.json, or <out-dir>/site/ in a trial)")
    ap.add_argument("--ceiling", type=int, default=CEILING, help="most rows this run may write")
    args = ap.parse_args(argv)
    if not (args.pull or args.write or args.site or args.terms):
        ap.error("name a stage: --pull, --write, --site or --terms")
    if args.out_dir:
        ip.set_out_dir(args.out_dir)
    site_file = os.path.abspath(args.site_file) if args.site_file else (
        os.path.join(ip.OUT_DIR, "site", "nyiso_load_queue.json") if args.out_dir else SITE_FILE)
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = now_utc().strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    results = []
    if ip.paused(ISO):  # session 89: no request, and the table stays as it is
        log(f"{NAME}: {ip.pause_line(ISO)}")
        print(f"{CONNECTOR} {ip.pause_line(ISO)}")
        ip.write_status(CONNECTOR, run_id, [dict(table=NAME, market="queue", status="skipped", detail=ip.pause_line(ISO)[:300])])
        log.close()
        return 0
    try:
        if args.terms:
            terms(log)
        if args.pull:
            pull(log)
        if args.write or args.site:
            wb_path, rec = saved_workbook()
            log(f"{NAME}: reading {rel(wb_path)} (retrieved {rec['retrieved_at']}, {rec['origin']}); no request made")
            rows, key, end_use_key, notes, tracking = parse_workbook(wb_path, log)
            if len(rows) > args.ceiling:
                raise RuntimeError(f"{len(rows):,} rows would pass this run's ceiling of {args.ceiling:,}; nothing is written")
            log(f"  {len(rows)} rows against a ceiling of {args.ceiling:,}")
            df = build_table(rows, key, rec, log)
            terms_saved(log)
            doc = build_site(rows, key, end_use_key, notes, tracking, rec, ip.utc_iso(pd.Timestamp.now(tz="UTC")))
            for c in doc["checks"]:
                log("  differs from the tracking sheet: " + c["what"])
            for z in doc["zones"]:
                log(f"  in line, zone {z['zone']} ({z['zone_entity']}): {z['requests']} requests, {z['mw']} MW")
            log(f"  in line, all zones: {doc['total']['requests']} requests, {doc['total']['mw']} MW")
            if args.write:
                ip.write_snapshot(df, NAME, header_lines(run_id, rec, df, doc, wb_path), log, COLS)
                ip.update_sources([dict(
                    source=SOURCE, publisher=PUBLISHER,
                    report=f'NYISO Interconnection Queue workbook, sheets "{SHEET}" and "{TRACKING}" (load interconnection requests) '
                           f'[terms: "{TERMS_SHORT}"; the notice forbids republishing "any image or video on this website as a stand-alone file", '
                           f"not data ({TERMS_URL}, read 7 October 2026)"
                           + ("]" if SHOWN else "; held internal from session 149: the notice confers no license and reserves all rights]"),
                    report_url=PAGE, document_list=URL, license=LICENSE, tables=[NAME])])
                results.append(dict(table=NAME, market="queue", status="ok", detail=""))
            write_site(doc if SHOWN else held_site(doc), site_file, log)
    except Exception:
        tb = ip.redact(traceback.format_exc())
        last = tb.strip().splitlines()[-1]
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{CONNECTOR} {NAME} FAILED, no output file written: {last}", file=sys.stderr)
        results = [dict(table=NAME, market="queue", status="failed", detail=last[:300])]
    if results:
        ip.write_status(CONNECTOR, run_id, results)
    failures = sum(r["status"] != "ok" for r in results)
    log(f"done, failures={failures}")
    log.close()
    print(f"{CONNECTOR} run log: {rel(log.path)}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
