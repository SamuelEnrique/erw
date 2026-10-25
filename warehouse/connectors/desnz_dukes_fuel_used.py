#!/usr/bin/env python3
"""The fuel the United Kingdom's power stations used in generation a year, 1996 to 2025, by kind of company and fuel, and the
major power producers' by plant type, as the Department for Energy Security and Net Zero (DESNZ) gives them in the Digest of UK
Energy Statistics (DUKES), table 5.3 (session 2748, chain 138, the session's third own choice; the fuel side of chain 137's
DUKES 5.7 capacity and 5.14 carbon intensity).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/desnz_uk_fuel_used_generation_annual.csv, from one
workbook (no key), its sheet "5.3", tables 5.3.B and 5.3.C (both in million tonnes of oil equivalent):

    https://assets.publishing.service.gov.uk/media/6a6a35ef0ddb7e4831c629f0/DUKES_5.3.xlsx

    entity      desnz_fuel:<group>:<fuel>: group mpp (major power producers), other_transport and other_industrial (other
                generators, undertakings in transport and in the industrial and commercial sectors), other (their total), all (all
                generating companies); fuel the row as a slug (coal, oil, gas, nuclear, hydro_natural_flow, wind, solar,
                wind_wave_and_solar_photovoltaics, other_renewables, other_fuels), or total (a group's "All" row); for 5.3.C,
                desnz_fuel:mpp:conventional_thermal_and_other_stations and desnz_fuel:mpp:combined_cycle_gas_turbine_stations
    variable    fuel_used (5.3.B, by fuel) or fuel_used_by_plant (5.3.C, by plant type)
    unit        MWh: the sheet's million tonnes of oil equivalent times 11,630,000 (DESNZ's 1 ktoe = 11.63 GWh, stated on the
                DUKES 5.11 workbook's cover sheet); ts_utc the year's first day; freq P1Y; geo GB; node the sheet's two labels
    stated in   the header, with the sheet's MTOE figure for the totals

DESNZ's notes: oil includes orimulsion, oil in gas turbines and diesel plant and oil for lighting up coal boilers (note 3),
refinery gas (note 4); coal includes coke oven coke (note 5); gas includes colliery methane (note 6); the major power producers'
other fuels are non-biodegradable waste, in other renewables before 2013, and other generators' mainly that, coke oven gas, blast
furnace gas and chemical waste, in other renewables before 2007 (note 8); net imports' fuel equivalent was removed when the UK
became a net exporter in 2022 (note 9); 5.3.C's conventional thermal stations include gas turbines, oil engines, coal and
renewable plants other than hydro, and exclude nuclear (note 10); nuclear and the renewables follow the IEA's primary energy
conventions (note 11). Table 5.3.A (fuels in their original units, million tonnes and GWh) is not taken.

Each table is found by its title and its year row (each year from 1996 in turn); each row's two labels (note marks read off) are
mapped to an entity; a group not in the map or an entity given twice stops the run. A blank cell has no row; a zero stays a
zero. Checks in every year, within 0.02 MTOE (the sheet's rows and totals differ by up to 0.0153, a little beyond rounding): each group's fuels add to its total; the transport and industrial generators add
to the other generators; the major power producers' and the other generators' totals add to all companies'. A check missed stops
the run. Nothing is filled.

    python warehouse/connectors/desnz_dukes_fuel_used.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/desnz_dukes/ and reused unless --fresh; --offline asks nothing.
www.gov.uk's robots file closes only print and search paths; assets.publishing.service.gov.uk has no robots file (session 2678).
License: the Open Government Licence v3.0 (DESNZ's verdict on record, session 2167): public.
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "desnz_uk_fuel_used_generation_annual"
CONNECTOR = "desnz_dukes"
SOURCE = "desnz:dukes_5_3_fuel_used"
HOST = "assets.publishing.service.gov.uk"
URL = "https://assets.publishing.service.gov.uk/media/6a6a35ef0ddb7e4831c629f0/DUKES_5.3.xlsx"
PAGE = "https://www.gov.uk/government/statistics/electricity-chapter-5-digest-of-united-kingdom-energy-statistics-dukes"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = "Department for Energy Security and Net Zero, Digest of UK Energy Statistics (DUKES) 2026, table 5.3, Fuel used in generation"
MWH_PER_MTOE = 11_630_000  # 1 ktoe = 11.63 GWh (DESNZ)
TABLES = [("Table 5.3.B Fuel used in generation, million tonnes of oil equivalent (MTOE)", "fuel_used"),
          ("Table 5.3.C Fuel used in generation in Combined Cycle Gas Turbines (CCGT) and conventional steam plants (MTOE)",
           "fuel_used_by_plant")]
GROUPS = {"Major power producers": ("mpp", False), "Total major power producers": ("mpp", True),
          "Other generators undertakings in transport": ("other_transport", False),
          "Other generators undertakings in industrial and commercial sectors": ("other_industrial", False),
          "Total other generators": ("other", True), "All generating companies": ("all", False),
          "Total all generating companies": ("all", True)}
FIRST_YEAR = 1996
TOL = 0.02  # DESNZ's rows and totals differ by up to 0.0153 MTOE (all companies, measured in session 2748), beyond rounding
X_COLS = []


def label(v):
    s = " ".join(str(v).replace(chr(0xA0), " ").split()) if v is not None else ""
    return re.sub(r"\s*\[note \d+\]", "", s).strip()


def slug(s):
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", s.lower())).strip("_")


def parse_sheet(rows):
    titles = {label(r[0]): i for i, r in enumerate(rows) if r and r[0] is not None}
    out, nodes = {}, {}
    for title, var in TABLES:
        if title not in titles:
            raise RuntimeError(f"no table titled {title!r}")
        i = titles[title]
        head = [label(x) for x in rows[i + 1]]
        if head[:2] != ["Generator type", "Fuel"]:
            raise RuntimeError(f"{title}: the header is {head[:2]}")
        years = head[2:]
        while years and not years[-1]:
            years.pop()
        if not years or years != [str(FIRST_YEAR + k) for k in range(len(years))]:
            raise RuntimeError(f"{title}: the years are not each year from {FIRST_YEAR}")
        for r in rows[i + 2:]:
            if not r or not any(label(x) for x in r[:2]) or label(r[0]).startswith("Table 5.3"):
                break
            g = label(r[0])
            if g not in GROUPS:
                raise RuntimeError(f"{title}: a group not in the map: {g!r}")
            grp, is_total = GROUPS[g]
            item = "total" if is_total else slug(label(r[1]))
            ent = f"{grp}:{item}"
            if (ent, var) in nodes:
                raise RuntimeError(f"{title}: {ent} given twice")
            nodes[(ent, var)] = f"{g}: {label(r[1])}"
            for k, y in enumerate(years):
                v = r[2 + k] if 2 + k < len(r) else None
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                if not isinstance(v, (int, float)) or isinstance(v, bool) or v < 0:
                    raise RuntimeError(f"{title} {g} {label(r[1])} {y}: not a quantity: {v!r}")
                out[(int(y), ent, var)] = float(v)
    check(out)
    return out, nodes


def check(out):
    for y in sorted({k[0] for k in out}):
        g = {k[1]: v for k, v in out.items() if k[0] == y and k[2] == "fuel_used"}
        for grp in ("mpp", "all"):
            parts = [v for e, v in g.items() if e.startswith(grp + ":") and e != grp + ":total"]
            tot = g.get(grp + ":total")
            if tot is not None and parts and abs(sum(parts) - tot) > TOL:
                raise RuntimeError(f"{y} {grp}: the fuels make {sum(parts):.4f} MTOE, the total {tot:.4f}")
        parts = [v for e, v in g.items() if e.startswith(("other_transport:", "other_industrial:"))]
        tot = g.get("other:total")
        if tot is not None and parts and abs(sum(parts) - tot) > TOL:
            raise RuntimeError(f"{y} other: the transport and industrial generators make {sum(parts):.4f} MTOE, the total {tot:.4f}")
        m, o, a = g.get("mpp:total"), g.get("other:total"), g.get("all:total")
        if None not in (m, o, a) and abs(m + o - a) > TOL:
            raise RuntimeError(f"{y}: the major power producers and the others make {m + o:.4f} MTOE, all companies {a:.4f}")


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    if "5.3" not in wb.sheetnames:
        raise RuntimeError(f"the workbook has no sheet 5.3: {wb.sheetnames}")
    vals, nodes = parse_sheet([tuple(r) for r in wb["5.3"].iter_rows(values_only=True)])
    rows = [dict(entity=f"desnz_fuel:{e}", variable=v, ts_utc=f"{y}-01-01T00:00:00Z", value=round(x * MWH_PER_MTOE, 1), unit="MWh",
                 freq="P1Y", geo="GB", market="", node=nodes[(e, v)], **base) for (y, e, v), x in sorted(vals.items())]
    return pd.DataFrame(rows), vals


def main(argv=None):
    ap = argparse.ArgumentParser(description="DESNZ DUKES 5.3: the fuel used in UK electricity generation a year")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t, vals = parse(content, base)
        if t.duplicated(["entity", "variable", "ts_utc"]).any():
            raise RuntimeError("a key given twice")
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["entity", "variable", "ts_utc"]).reset_index(drop=True)
        mtoe = {y: x for (y, e, v), x in vals.items() if e == "all:total" and v == "fuel_used"}
        head = [
            "Energy Research Warehouse (ERW): DESNZ, Digest of UK Energy Statistics (DUKES) 2026, table 5.3: the fuel used in UK "
            f"electricity generation by kind of company and fuel, and the major power producers' by plant type, {t.ts_utc.min()[:4]} "
            f"to {t.ts_utc.max()[:4]} (session 2748)",
            "Shape: series (docs/datastandard.md v0). entity desnz_fuel:<group>:<fuel>; variable fuel_used (5.3.B) or "
            "fuel_used_by_plant (5.3.C); unit MWh, the sheet's million tonnes of oil equivalent times 11,630,000 (DESNZ's 1 ktoe = "
            "11.63 GWh); ts_utc the year's first day; freq P1Y; geo GB; node the sheet's labels.",
            f"{len(t)} rows; all companies' fuel used, MTOE as the sheet gives it: " + ", ".join(f"{y} {mtoe[y]:.3f}" for y in sorted(mtoe)
                                                                                          if y in (1996, 2010, 2020, 2025))
            + ". Every group's fuels checked against its total; nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/desnz_dukes_fuel_used.py; the workbook retrieved "
            f"{base['retrieved_at']}; last modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL} (listed on {PAGE})",
            "License: Open Government Licence v3.0. Credit: Department for Energy Security and Net Zero.",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="UK Department for Energy Security and Net Zero (DESNZ)",
                                    report="Digest of UK Energy Statistics (DUKES) 2026, table 5.3, Fuel used in generation",
                                    report_url=PAGE, document_list=PAGE, license="public", tables=[NAME])])  # OGL v3.0
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
