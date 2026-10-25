#!/usr/bin/env python3
"""The United Kingdom's electricity imports, exports and net imports over each interconnector a year, 1998 to 2025, their
utilisation, and the transfers between Scotland and England and between Scotland and Northern Ireland, as the Department for
Energy Security and Net Zero (DESNZ) gives them in the Digest of UK Energy Statistics (DUKES), table 5.13 (session 2715, chain
136, the session's second own choice; the companion of session 2714's power stations, from the same chapter; the ERW held the
UK's trade only as national totals).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/desnz_uk_interconnector_flows_annual.csv, from one
workbook (no key), its sheet "5.13":

    https://assets.publishing.service.gov.uk/media/6a6a3675bf4a7aca8f8976c3/DUKES_5.13.xlsx

    entity      desnz_ic:<pair>: france, ireland_northern_ireland (the North-South interconnector), netherlands, ireland_wales,
                belgium, norway, denmark, uk_total (the seven summed by DESNZ); scotland_england and scotland_northern_ireland
                (transfers within the UK, the second over Moyle)
    variable    imports, exports, net_imports (into the UK, or into Northern Ireland for ireland_northern_ireland; for the two
                transfers: transfers_to_england or transfers_to_northern_ireland, transfers_to_scotland and net_transfers, the
                first direction less the second, as DESNZ nets them) in MWh (the sheet's GWh times 1,000); utilisation (ratio:
                imports and exports over the capacity's possible flow in the year, DESNZ's note 3)
    ts_utc      the year's first day; freq P1Y; geo GB; node the sheet's column heading

Each heading is matched exactly to its pair and variable; a heading not in the map stops the run. "[x]" (not available) and a
blank cell have no row. Checks, within 0.01 GWh: each pair's imports less exports equal its net imports (and each transfer's
two directions its net); the seven interconnectors add to the UK's total imports, exports and net imports, where all are
given; a utilisation lies in [0, 1]. Nothing is filled. Table 5.13A (each interconnector's capacity) is not taken.

    python warehouse/connectors/desnz_dukes_interconnectors.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/desnz_dukes/ and reused unless --fresh; --offline asks nothing.
www.gov.uk's robots file closes only print and search paths; assets.publishing.service.gov.uk has no robots file (session 2678).
License: the Open Government Licence v3.0 (DESNZ's verdict on record, session 2167): public. DESNZ's notes: the flows are from
National Grid's demand data, except Ireland and Northern Ireland's, which EirGrid supplies.
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "desnz_uk_interconnector_flows_annual"
CONNECTOR = "desnz_dukes"
SOURCE = "desnz:dukes_5_13_interconnectors"
HOST = "assets.publishing.service.gov.uk"
URL = "https://assets.publishing.service.gov.uk/media/6a6a3675bf4a7aca8f8976c3/DUKES_5.13.xlsx"
PAGE = "https://www.gov.uk/government/statistics/electricity-chapter-5-digest-of-united-kingdom-energy-statistics-dukes"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = ("Department for Energy Security and Net Zero, Digest of UK Energy Statistics (DUKES) 2026, table 5.13, Net imports and "
            "utilisation of interconnectors")
PAIRS = [("France to UK", "UK to France", "France and UK", "france"),
         ("Ireland to Northern Ireland", "Northern Ireland to Ireland", "Ireland and Northern Ireland", "ireland_northern_ireland"),
         ("Netherlands to UK", "UK to Netherlands", "Netherlands and UK", "netherlands"),
         ("Ireland to Wales", "Wales to Ireland", "Ireland and Wales", "ireland_wales"),
         ("Belgium to UK", "UK to Belgium", "Belgium and UK", "belgium"),
         ("Norway to UK", "UK to Norway", "Norway and UK", "norway"),
         ("Denmark to UK", "UK to Denmark", "Denmark and UK", "denmark")]
COLUMNS = {"Year": None}
for imp, exp, use, key in PAIRS:
    COLUMNS[f"Imports ({imp})"] = (key, "imports")
    COLUMNS[f"Exports ({exp})"] = (key, "exports")
    COLUMNS[f"Net imports ({imp})"] = (key, "net_imports")
    COLUMNS[f"Utilisation ({use})"] = (key, "utilisation")
COLUMNS.update({"Total Imports (to UK)": ("uk_total", "imports"), "Total Exports (from UK)": ("uk_total", "exports"),
                "Net imports (to UK)": ("uk_total", "net_imports"),
                "Transfers to England (from Scotland)": ("scotland_england", "transfers_to_england"),
                "Transfers to Scotland (from England)": ("scotland_england", "transfers_to_scotland"),
                "Net transfers (Scotland to England)": ("scotland_england", "net_transfers"),
                "Transfers to NI (from Scotland)": ("scotland_northern_ireland", "transfers_to_northern_ireland"),
                "Transfers to Scotland (from NI)": ("scotland_northern_ireland", "transfers_to_scotland"),
                "Net transfers (to NI from Scotland)": ("scotland_northern_ireland", "net_transfers")})
NETS = [("imports", "exports", "net_imports"), ("transfers_to_england", "transfers_to_scotland", "net_transfers"),
        ("transfers_to_northern_ireland", "transfers_to_scotland", "net_transfers")]
X_COLS = []


def label(v):
    s = " ".join(str(v).replace(chr(0xA0), " ").split()) if v is not None else ""
    return re.sub(r"\s*\[note \d+\]$", "", s)


def parse_sheet(rows):
    h = next((i for i, r in enumerate(rows) if r and label(r[0]) == "Year"), None)
    if h is None:
        raise RuntimeError("no header row")
    cols = {}
    for j, x in enumerate(rows[h]):
        s = label(x)
        if not s:
            continue
        if s not in COLUMNS:
            raise RuntimeError(f"a heading not in the map: {s!r}")
        if COLUMNS[s]:
            cols[j] = COLUMNS[s] + (s,)
    if len(cols) != len(COLUMNS) - 1:
        raise RuntimeError(f"{len(COLUMNS) - 1 - len(cols)} headings of the map are not on the sheet")
    out = {}
    for r in rows[h + 1:]:
        if not r or r[0] is None or label(r[0]) == "":
            continue
        if not isinstance(r[0], int) or not 1990 <= r[0] <= 2100:
            raise RuntimeError(f"a year not a year: {r[0]!r}")
        for j, (pair, var, head) in cols.items():
            v = r[j] if j < len(r) else None
            if v is None or (isinstance(v, str) and v.strip() in ("", "[x]")):
                continue
            if not isinstance(v, (int, float)) or isinstance(v, bool):
                raise RuntimeError(f"{r[0]} {head}: not a number: {v!r}")
            if var == "utilisation" and not 0 <= v <= 1:
                raise RuntimeError(f"{r[0]} {head}: a utilisation outside [0, 1]: {v!r}")
            out[(r[0], pair, var)] = (float(v), head)
    g = lambda y, p, v: out.get((y, p, v), (None,))[0]
    for y in {k[0] for k in out}:
        for p in {k[1] for k in out if k[0] == y}:
            for a, b, n in NETS:
                x, z, w = g(y, p, a), g(y, p, b), g(y, p, n)
                if None not in (x, z, w) and abs(x - z - w) > 0.01:
                    raise RuntimeError(f"{y} {p}: {a} less {b} is {x - z:,.3f} GWh, {n} {w:,.3f}")
        for var in ("imports", "exports", "net_imports"):
            parts = [g(y, k, var) for *_, k in PAIRS]
            tot = g(y, "uk_total", var)
            if tot is not None and None not in [g(y, k, var) for *_, k in PAIRS if g(y, k, "imports") is not None]:
                s = sum(x for x in parts if x is not None)
                if abs(s - tot) > 0.01:
                    raise RuntimeError(f"{y}: the interconnectors' {var} add to {s:,.3f} GWh, the UK's total {tot:,.3f}")
    return out


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    if "5.13" not in wb.sheetnames:
        raise RuntimeError(f"the workbook has no sheet 5.13: {wb.sheetnames}")
    vals = parse_sheet([tuple(r) for r in wb["5.13"].iter_rows(values_only=True)])
    rows = [dict(entity=f"desnz_ic:{p}", variable=v, ts_utc=f"{y}-01-01T00:00:00Z",
                 value=round(x, 4) if v == "utilisation" else round(x * 1000, 3), unit="ratio" if v == "utilisation" else "MWh",
                 freq="P1Y", geo="GB", market="", node=head, **base) for (y, p, v), (x, head) in sorted(vals.items())]
    return pd.DataFrame(rows)


def main(argv=None):
    ap = argparse.ArgumentParser(description="DESNZ DUKES 5.13: the UK's interconnector flows a year")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t = parse(content, base)
        if t.duplicated(["entity", "variable", "ts_utc"]).any():
            raise RuntimeError("a key given twice")
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["entity", "variable", "ts_utc"]).reset_index(drop=True)
        net = t[(t.entity == "desnz_ic:uk_total") & (t.variable == "net_imports")].set_index("ts_utc").value
        head = [
            "Energy Research Warehouse (ERW): DESNZ, Digest of UK Energy Statistics (DUKES) 2026, table 5.13: the UK's electricity "
            f"imports, exports and net imports over each interconnector and their utilisation, {t.ts_utc.min()[:4]} to "
            f"{t.ts_utc.max()[:4]}, and the transfers between Scotland and England and Northern Ireland (session 2715)",
            "Shape: series (docs/datastandard.md v0). entity desnz_ic:<pair>; variable imports, exports, net_imports, utilisation, "
            "or for the two transfers within the UK transfers_to_<place> and net_transfers; unit MWh (the sheet's GWh times 1,000) "
            "or ratio (utilisation); ts_utc the year's first day; freq P1Y; geo GB; node the sheet's heading.",
            f"{len(t)} rows; the UK's net imports, TWh: " + ", ".join(f"{k[:4]} {v / 1e6:.1f}" for k, v in net.items()
                                                                       if k[:4] in ("1998", "2010", "2020", "2022", "2025"))
            + ". Each pair's net and the UK's totals checked against their parts; nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/desnz_dukes_interconnectors.py; the workbook retrieved "
            f"{base['retrieved_at']}; last modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL} (listed on {PAGE})",
            "License: Open Government Licence v3.0. Credit: Department for Energy Security and Net Zero (from National Grid's "
            "demand data, and EirGrid's for Ireland and Northern Ireland).",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="UK Department for Energy Security and Net Zero (DESNZ)",
                                    report="Digest of UK Energy Statistics (DUKES) 2026, table 5.13, Net imports and utilisation of interconnectors",
                                    report_url=PAGE, document_list=PAGE, license="public", tables=[NAME])])  # OGL v3.0
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
