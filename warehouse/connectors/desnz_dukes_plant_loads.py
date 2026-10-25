#!/usr/bin/env python3
"""The United Kingdom's maximum load, plant load factors and thermal efficiency a year, 1996 to 2025, as the Department for Energy
Security and Net Zero (DESNZ) gives them for the major power producers in the Digest of UK Energy Statistics (DUKES), table 5.10
(session 2728, chain 137, the session's own choice; the companion of chain 136's DUKES tables 5.11, 5.13 and 5.16; the ERW held no
UK peak load or load factor).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/desnz_uk_plant_loads_annual.csv, from one workbook (no
key), its sheets "5.10.A" and "5.10.B and 5.10.C":

    https://assets.publishing.service.gov.uk/media/6a6a364b0ddb7e4831c629f2/DUKES_5.10.xlsx

    entity      desnz_pl:<area or plant type>: uk, england_wales, scotland, great_britain, northern_ireland (the maximum load);
                ccgt, nuclear, pumped_storage, conventional_thermal, coal, all_plant (load factors); ccgt, coal, nuclear
                (thermal efficiency)
    variable    max_load (5.10.A, the simultaneous maximum load met, MW: load met by the transmission network, net of demand met
                by embedded generation, DESNZ's note 2); max_demand_share_of_capacity (5.10.A, maximum demand as a percentage of the
                major power producers' capacity, pct, on uk); load_factor (5.10.B, pct); system_load_factor (5.10.B, average
                electricity available as a percentage of maximum demand, DESNZ's note 9, pct, on uk); thermal_efficiency (5.10.C,
                electricity generated as a percentage of the fuel used, gross calorific value, pct)
    ts_utc      the first day of the year DESNZ labels; freq P1Y; geo GB, or GB-SCT and GB-NIR for those regions; node the
                sheet's row as written

DESNZ's note 3: the maximum load of a year is that of the 12 months ending March of the following year (2017 is the year ending
March 2018), so a winter's peak falls in the year it starts. Note 11: before 2005 Scotland's and Northern Ireland's demands are
those at the time of England and Wales's maximum; in 2005 to 2007 Northern Ireland's is that at Great Britain's. Note 6: the load
factors of conventional thermal stations and of all plant use the major power producers' supply and capacity, revised back to
2010. Note 8: all plant includes wind from 2008, solar from 2015 and natural flow hydro, at capacity not de-rated.

Each table is found by its title and its year row (each year from 1996 in turn); each row label (note marks read off) is matched
exactly to its entity and variable, and a label not in the map stops the run. "[x]" (not available) and a blank cell have no row;
a zero stays a zero (no coal after 2024). Checks, within 1 MW: England and Wales and Scotland add to Great Britain, and Great
Britain and Northern Ireland to the total, where all are given; every percentage lies in [0, 100]. Nothing is filled.

    python warehouse/connectors/desnz_dukes_plant_loads.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/desnz_dukes/ and reused unless --fresh; --offline asks nothing.
www.gov.uk's robots file closes only print and search paths; assets.publishing.service.gov.uk has no robots file (session 2678).
License: the Open Government Licence v3.0 (DESNZ's verdict on record, session 2167): public.
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "desnz_uk_plant_loads_annual"
CONNECTOR = "desnz_dukes"
SOURCE = "desnz:dukes_5_10_plant_loads"
HOST = "assets.publishing.service.gov.uk"
URL = "https://assets.publishing.service.gov.uk/media/6a6a364b0ddb7e4831c629f2/DUKES_5.10.xlsx"
PAGE = "https://www.gov.uk/government/statistics/electricity-chapter-5-digest-of-united-kingdom-energy-statistics-dukes"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = ("Department for Energy Security and Net Zero, Digest of UK Energy Statistics (DUKES) 2026, table 5.10, Plant loads, demand "
            "and efficiency of major power producers")
FIRST_YEAR = 1996
LOAD = "Simultaneous maximum load met, MW"
A_ROWS = {(LOAD, "Total"): ("uk", "max_load", "MW", "GB"), (LOAD, "England and Wales"): ("england_wales", "max_load", "MW", "GB"),
          (LOAD, "Scotland"): ("scotland", "max_load", "MW", "GB-SCT"), (LOAD, "Great Britain"): ("great_britain", "max_load", "MW", "GB"),
          (LOAD, "Northern Ireland"): ("northern_ireland", "max_load", "MW", "GB-NIR"),
          ("Maximum demand as a percentage of major power producers' capacity", "UK"): ("uk", "max_demand_share_of_capacity", "pct", "GB")}
B_TITLE = "Table 5.10.B Plant load factor, per cent"
B_ROWS = {"Combined cycle gas turbine stations": "ccgt", "Nuclear stations": "nuclear", "Pumped storage hydro": "pumped_storage",
          "Conventional thermal and other stations": "conventional_thermal", "of which coal-fired stations": "coal", "All plant": "all_plant",
          "System load factor": "uk"}
C_TITLE = "Table 5.10.C Thermal efficiency, gross calorific value basis, per cent"
C_ROWS = {"Combined cycle gas turbine stations": "ccgt", "Coal fired stations": "coal", "Nuclear stations": "nuclear"}
X_COLS = []


def label(v):
    s = " ".join(str(v).replace(chr(0xA0), " ").split()) if v is not None else ""
    return re.sub(r"\s*\[note \d+\]", "", s).strip()


def years_of(row, first):
    years = [label(x) for x in row[first:]]
    while years and not years[-1]:
        years.pop()
    if not years or years != [str(FIRST_YEAR + k) for k in range(len(years))]:
        raise RuntimeError(f"the years are not each year from {FIRST_YEAR}: {years[:3]} .. {years[-2:]}")
    return [int(y) for y in years]


def cell(v, where):
    if v is None or (isinstance(v, str) and v.strip() in ("", "[x]")):
        return None
    if not isinstance(v, (int, float)) or isinstance(v, bool):
        raise RuntimeError(f"{where}: not a number: {v!r}")
    return float(v)


def parse_a(rows):
    h = next((i for i, r in enumerate(rows) if r and label(r[0]) == "Indicator" and label(r[1]) == "Region"), None)
    if h is None:
        raise RuntimeError("5.10.A: no header row Indicator, Region")
    years = years_of(rows[h], 2)
    out = {}
    for r in rows[h + 1:]:
        if not r or not label(r[0]):
            continue
        key = (label(r[0]), label(r[1]))
        if key not in A_ROWS:
            raise RuntimeError(f"5.10.A: a row not in the map: {key}")
        ent, var, unit, geo = A_ROWS[key]
        for k, y in enumerate(years):
            v = cell(r[2 + k] if 2 + k < len(r) else None, f"5.10.A {key} {y}")
            if v is not None:
                out[(y, ent, var)] = (v, unit, geo, " ".join(key))
    return out


def parse_bc(rows):
    out = {}
    titles = {label(r[0]): i for i, r in enumerate(rows) if r and r[0] is not None}
    for title, rmap, var in ((B_TITLE, B_ROWS, "load_factor"), (C_TITLE, C_ROWS, "thermal_efficiency")):
        if title not in titles:
            raise RuntimeError(f"no table titled {title!r}")
        i = titles[title]
        if label(rows[i + 1][0]) != "Generator Type":
            raise RuntimeError(f"{title}: no header row")
        years = years_of(rows[i + 1], 1)
        for r in rows[i + 2:]:
            name = label(r[0]) if r else ""
            if not name or name.startswith("Table 5.10"):
                break
            if name not in rmap:
                raise RuntimeError(f"{title}: a row not in the map: {name!r}")
            ent, v_name = rmap[name], ("system_load_factor" if name == "System load factor" else var)
            for k, y in enumerate(years):
                v = cell(r[1 + k] if 1 + k < len(r) else None, f"{title} {name} {y}")
                if v is not None:
                    out[(y, ent, v_name)] = (v, "pct", "GB", f"{title[6:12]}: {name}")
    return out


def check(vals):
    g = lambda y, e: vals.get((y, e, "max_load"), (None,))[0]
    for y in {k[0] for k in vals}:
        ew, sc, gb, ni, uk = g(y, "england_wales"), g(y, "scotland"), g(y, "great_britain"), g(y, "northern_ireland"), g(y, "uk")
        if None not in (ew, sc, gb) and abs(ew + sc - gb) > 1:
            raise RuntimeError(f"{y}: England and Wales and Scotland make {ew + sc:,.1f} MW, Great Britain {gb:,.1f}")
        if None not in (gb, ni, uk) and abs(gb + ni - uk) > 1:
            raise RuntimeError(f"{y}: Great Britain and Northern Ireland make {gb + ni:,.1f} MW, the total {uk:,.1f}")
    for (y, e, v), (x, unit, *_r) in vals.items():
        if unit == "pct" and not 0 <= x <= 100:
            raise RuntimeError(f"{y} {e} {v}: a percentage outside [0, 100]: {x}")


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    for s in ("5.10.A", "5.10.B and 5.10.C"):
        if s not in wb.sheetnames:
            raise RuntimeError(f"the workbook has no sheet {s!r}: {wb.sheetnames}")
    vals = parse_a([tuple(r) for r in wb["5.10.A"].iter_rows(values_only=True)])
    bc = parse_bc([tuple(r) for r in wb["5.10.B and 5.10.C"].iter_rows(values_only=True)])
    if set(vals) & set(bc):
        raise RuntimeError("a key in both sheets")
    vals.update(bc)
    check(vals)
    rows = [dict(entity=f"desnz_pl:{e}", variable=v, ts_utc=f"{y}-01-01T00:00:00Z", value=round(x, 4), unit=unit, freq="P1Y",
                 geo=geo, market="", node=node, **base) for (y, e, v), (x, unit, geo, node) in sorted(vals.items())]
    return pd.DataFrame(rows)


def main(argv=None):
    ap = argparse.ArgumentParser(description="DESNZ DUKES 5.10: the UK's maximum load, load factors and efficiency a year")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t = parse(content, base)
        if t.duplicated(["entity", "variable", "ts_utc"]).any():
            raise RuntimeError("a key given twice")
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["entity", "variable", "ts_utc"]).reset_index(drop=True)
        peak = t[(t.entity == "desnz_pl:uk") & (t.variable == "max_load")].set_index("ts_utc").value
        head = [
            "Energy Research Warehouse (ERW): DESNZ, Digest of UK Energy Statistics (DUKES) 2026, table 5.10: the UK's maximum load, "
            f"plant load factors and thermal efficiency for the major power producers, {t.ts_utc.min()[:4]} to {t.ts_utc.max()[:4]} "
            "(session 2728)",
            "Shape: series (docs/datastandard.md v0). entity desnz_pl:<area or plant type>; variable max_load (MW), "
            "max_demand_share_of_capacity, load_factor, system_load_factor or thermal_efficiency (pct); ts_utc the first day of "
            "DESNZ's year (the maximum load's year runs to March of the next, note 3); freq P1Y; geo GB, GB-SCT or GB-NIR; node the "
            "sheet's row.",
            f"{len(t)} rows; the UK's maximum load, MW: " + ", ".join(f"{k[:4]} {v:,.0f}" for k, v in peak.items()
                                                                       if k[:4] in ("1996", "2005", "2010", "2020", "2025"))
            + ". Regions checked against their totals; nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/desnz_dukes_plant_loads.py; the workbook retrieved "
            f"{base['retrieved_at']}; last modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL} (listed on {PAGE})",
            "License: Open Government Licence v3.0. Credit: Department for Energy Security and Net Zero.",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="UK Department for Energy Security and Net Zero (DESNZ)",
                                    report="Digest of UK Energy Statistics (DUKES) 2026, table 5.10, Plant loads, demand and efficiency",
                                    report_url=PAGE, document_list=PAGE, license="public", tables=[NAME])])  # OGL v3.0
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
