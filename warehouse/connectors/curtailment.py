#!/usr/bin/env python3
"""Wind and solar curtailment by ISO, daily: CAISO and SPP as published, ERCOT output below HSL.

Energy Research Warehouse (ERW) connector, session 18 (platform tool 22, the curtailment
tracker). Method and the exact meaning of each ISO's figure: docs/methods/curtailment.md.
Writes three `series` tables, one row per day and variable, MWh, through the merge writer:

    caiso_curtailment_daily        entity caiso:ISO
    spp_curtailment_daily          entity spp:<BAA> (SPP, and SWPW since SPP's western expansion)
    ercot_wind_solar_hsl_daily     entity ercot:system

    python warehouse/connectors/curtailment.py              # the latest days (the daily run)
    python warehouse/connectors/curtailment.py --history    # every day each source still posts
    python warehouse/connectors/curtailment.py --iso caiso

Sources:
- CAISO, to 2025-12-31: "Production and curtailments data" (www.caiso.com/library/
  production-curtailments-data), one workbook per period, 5-minute MW: sheet Curtailments
  (wind and solar curtailment, only the intervals with one; since 2024 a Reason, Local or
  System) and sheet Production (load, solar, wind, ...; the June to December 2025 file has none, and repeats the 2025
  workbook's curtailment rows).
  MWh = MW x 5/60 per interval, summed over the local (Pacific) day of the Date column.
- CAISO, from 2026-01-01: the Daily Renewable Report (www.caiso.com/library/
  daily-renewable-reports), one HTML page per day. Its chart data are JavaScript arrays: the
  report day's hourly curtailment in MWh by fuel and category (economic, self-schedule cut,
  operator instruction; each local or system), and 5-minute ISO solar and wind output
  telemetry in MW (MWh = MW x 5/60). A day is written only from the report whose chart date
  is that day; its generation only if all 5-minute values are present (288, or 276 or 300 on
  a daylight saving change day).
- SPP: VER Curtailments (portal.spp.org, file browser ver-curtailments): 5-minute MW by
  balancing authority area of wind and solar curtailed by redispatch, manual (operator)
  curtailment, and "curtailed for energy" (economic dispatch). Annual zips to 2024, daily files
  since. The day is the Central (America/Chicago) date of each interval's start; a BAA's day is
  written only with every 5-minute interval of it.
- ERCOT publishes no curtailment figure. Its hourly reports NP4-732-CD (wind) and NP4-745-CD
  (solar), "Hourly Averaged Actual and Forecasted Values", give system-wide actual output (GEN)
  and the High Sustained Limit (HSL, the output the resources report they could sustain). The
  ERW writes, per operating day (Central): generation, HSL, and output below HSL = the sum over
  hours of max(0, HSL - GEN), an estimate of curtailment, not ERCOT's figure. Each report covers
  the 48 hours before it and ERCOT's list holds about a week of reports, so the history starts
  when the ERW began reading it; a day is written only with every hour's GEN and HSL.
MISO publishes no curtailment series (its market reports list hourly wind output only); the
site says so. Every run's requests are saved under warehouse/raw/curtailment/<run_id>/.
"""

import argparse
import datetime as dt
import io
import json
import os
import re
import sys
import time
import traceback
import zipfile

import pandas as pd
import requests

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-curtailment/0.1; +https://github.com/SamuelEnrique/erw)"}
CAISO = "https://www.caiso.com"
CAISO_LEGACY_PAGE = CAISO + "/library/production-curtailments-data"
CAISO_DRR_PAGE = CAISO + "/library/daily-renewable-reports"
CAISO_DRR_FROM = "2026-01-01"  # the legacy workbooks cover every day before this
SPP_BASE = "https://portal.spp.org/file-browser-api/download/ver-curtailments"
SPP_PAGE = "https://portal.spp.org/pages/ver-curtailments"
SPP_FIRST_DAILY_YEAR = 2025  # annual zips hold the years before
ERCOT_LIST = "https://www.ercot.com/misapp/servlets/IceDocListJsonWS?reportTypeId={}"
ERCOT_DOC = "https://www.ercot.com/misdownload/servlets/mirDownload?doclookupId={}"
ERCOT_REPORTS = {"wind": (13028, "NP4-732-CD", "https://www.ercot.com/mp/data-products/data-product-details?id=NP4-732-CD"),
                 "solar": (13483, "NP4-745-CD", "https://www.ercot.com/mp/data-products/data-product-details?id=NP4-745-CD")}
CAISO_CATS = ["econ_local", "econ_system", "ss_local", "ss_system", "oi_local", "oi_system"]
SPP_COLS = {"WindRedispatchCurtailments": "curtailed_wind_redispatch_mwh",
            "WindManualCurtailments": "curtailed_wind_manual_mwh",
            "WindCurtailedForEnergy": "curtailed_wind_economic_mwh",
            "SolarRedispatchCurtailments": "curtailed_solar_redispatch_mwh",
            "SolarManualCurtailments": "curtailed_solar_manual_mwh",
            "SolarCurtailedForEnergy": "curtailed_solar_economic_mwh"}
SOURCES = {
    "caiso:production_curtailments": ("California ISO (CAISO)", "Production and curtailments data (5-minute)",
                                      CAISO_LEGACY_PAGE, CAISO_LEGACY_PAGE),
    "caiso:daily_renewable_report": ("California ISO (CAISO)", "Daily Renewable Report (curtailment and output)",
                                     CAISO_DRR_PAGE, CAISO_DRR_PAGE),
    "spp:ver_curtailments": ("Southwest Power Pool (SPP)", "VER Curtailments (5-minute, by BAA)", SPP_PAGE,
                             SPP_BASE),
    "ercot:np4_732_745": ("Electric Reliability Council of Texas (ERCOT)",
                          "Wind and Solar Power Production, Hourly Averaged Actual and Forecasted Values "
                          "(NP4-732-CD, NP4-745-CD)", ERCOT_REPORTS["wind"][2], ERCOT_REPORTS["solar"][2]),
}


def get(url, log, what, ok404=False):
    def call():
        r = requests.get(url, headers=UA, timeout=300)
        if r.status_code == 404 and ok404:
            return r
        if r.status_code != 200:
            raise RuntimeError(f"{what}: HTTP {r.status_code} for {url}")
        return r
    return ip.with_retries(what, call, log)


def day_rows(entity, geo, daily, source, url_of, retrieved):
    """series rows from {(date, variable): value}; url_of(date) gives the row's source URL."""
    rows = [{"entity": entity, "variable": v, "ts_utc": f"{d}T00:00:00Z", "value": round(x, 3), "unit": "MWh",
             "freq": "P1D", "geo": geo, "market": "", "node": "", "source": source, "source_url": url_of(d),
             "retrieved_at": retrieved, "vintage": ""} for (d, v), x in sorted(daily.items())]
    return pd.DataFrame(rows, columns=ip.SERIES_COLS)


def expected_intervals(day, tz, minutes):
    d0 = pd.Timestamp(day).tz_localize(tz)
    d1 = (pd.Timestamp(day) + pd.Timedelta(days=1)).tz_localize(tz)
    return int((d1 - d0) / pd.Timedelta(minutes=minutes))


# ---------------------------------------------------------------- CAISO

def caiso_legacy(log, retrieved):
    """Every day of the legacy workbooks: (daily {(date, var): MWh}, url of each date)."""
    import openpyxl
    page = get(CAISO_LEGACY_PAGE, log, "CAISO production and curtailments page").text
    files = sorted(set(re.findall(r'"(/documents/[^"]*curtailments?-?data[^"]*\.xlsx)"', page, re.I)))
    if not files:
        raise RuntimeError(f"no workbook links on {CAISO_LEGACY_PAGE}")
    log(f"  CAISO legacy workbooks: {len(files)}: {files}")
    daily, url_of, gen_n = {}, {}, {}
    # the workbooks overlap (the "June to December 2025" file starts on 2025-01-01): each
    # 5-minute curtailment interval is counted once, from the first workbook (in name order) that lists it
    seen_c, dup_c = set(), 0
    reason_days = set()  # days of a workbook that gives each curtailment's reason (2024 on)
    for f in files:
        url = CAISO + f
        r = get(url, log, f"CAISO {f}")
        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=False)
        n_c = n_p = 0
        if "Curtailments" in wb.sheetnames:
            it = wb["Curtailments"].iter_rows(values_only=True)
            head = [str(h).strip() if h is not None else "" for h in next(it)]
            ci = {h: i for i, h in enumerate(head)}
            wb_days = []
            for row in it:
                when = row[ci["Date"]]
                if not isinstance(when, dt.datetime):
                    continue
                d = when.date().isoformat()
                wb_days.append(d)
                reason = str(row[ci["Reason"]]).strip().lower() if "Reason" in ci and row[ci["Reason"]] else ""
                key = (when, row[ci["Hour"]], row[ci["Interval"]], reason,
                       row[ci["Wind Curtailment"]], row[ci["Solar Curtailment"]])
                if key in seen_c:
                    dup_c += 1
                    continue
                seen_c.add(key)
                for fuel, col in (("wind", "Wind Curtailment"), ("solar", "Solar Curtailment")):
                    v = row[ci[col]]
                    if v in (None, "") or not isinstance(v, (int, float)):
                        continue
                    mwh = float(v) * 5 / 60
                    daily[(d, f"curtailed_{fuel}_mwh")] = daily.get((d, f"curtailed_{fuel}_mwh"), 0.0) + mwh
                    if reason in ("local", "system"):
                        k = (d, f"curtailed_{fuel}_{reason}_mwh")
                        daily[k] = daily.get(k, 0.0) + mwh
                    url_of[d] = url
                n_c += 1
            if "Reason" in ci and wb_days:
                reason_days |= set(pd.date_range(min(wb_days), max(wb_days), freq="D").strftime("%Y-%m-%d"))
        if "Production" in wb.sheetnames:
            it = wb["Production"].iter_rows(values_only=True)
            head = [str(h).strip() if h is not None else "" for h in next(it)]
            ci = {h: i for i, h in enumerate(head)}
            for row in it:
                when = row[ci["Date"]]
                if not isinstance(when, dt.datetime):
                    continue
                d = when.date().isoformat()
                # no dedup here: only one workbook per period has a Production sheet, and on a
                # fall-back day the repeated hour's intervals carry the same local timestamps
                s, w = row[ci["Solar"]], row[ci["Wind"]]
                if not isinstance(s, (int, float)) or not isinstance(w, (int, float)):
                    gen_n[d] = gen_n.get(d, 0) - 10_000  # a missing value spoils the day
                    continue
                daily[(d, "solar_generation_mwh")] = daily.get((d, "solar_generation_mwh"), 0.0) + float(s) * 5 / 60
                daily[(d, "wind_generation_mwh")] = daily.get((d, "wind_generation_mwh"), 0.0) + float(w) * 5 / 60
                gen_n[d] = gen_n.get(d, 0) + 1
                url_of.setdefault(d, url)
                n_p += 1
        wb.close()
        log(f"  {f}: {n_c} curtailment rows, {n_p} production rows")
    log(f"  CAISO legacy: curtailment intervals listed in two workbooks, counted once: {dup_c}")
    # a day's generation is kept only with every 5-minute interval of it
    bad = [d for d, n in gen_n.items() if n != expected_intervals(d, "America/Los_Angeles", 5)]
    for d in bad:
        daily.pop((d, "solar_generation_mwh"), None)
        daily.pop((d, "wind_generation_mwh"), None)
    log(f"  CAISO legacy: generation kept for {len(gen_n) - len(bad)} days, dropped for {len(bad)} incomplete days "
        f"({sorted(bad)[:5]})")
    # every day the curtailment sheets cover, with no row for a fuel, had none of it curtailed
    first_c = min(d for d, v in daily if v.startswith("curtailed_"))
    last_c = max(d for d, v in daily if v.startswith("curtailed_"))
    for d in pd.date_range(first_c, last_c, freq="D").strftime("%Y-%m-%d"):
        for fuel in ("wind", "solar"):
            daily.setdefault((d, f"curtailed_{fuel}_mwh"), 0.0)
    # and where the workbook gives reasons, a day with no row of a reason had none of it
    for d in reason_days:
        for fuel in ("wind", "solar"):
            for reason in ("local", "system"):
                daily.setdefault((d, f"curtailed_{fuel}_{reason}_mwh"), 0.0)
    return {k: v for k, v in daily.items() if k[0] < CAISO_DRR_FROM}, url_of


def js_array(html, name):
    m = re.search(r"\b" + re.escape(name) + r"\s*=\s*(?:JSON\.parse\(\[\")?\[([^\]]*)\]", html)
    if not m:
        return None
    out = []
    for x in m.group(1).split(","):
        x = x.strip().strip('"')
        out.append(None if x in ("", "NA", "null") else float(x))
    return out


def caiso_drr_days(log, since, until):
    """Daily Renewable Report URLs {date: url} from the monthly library pages, since..until."""
    index = get(CAISO_DRR_PAGE, log, "CAISO daily renewable reports index").text
    months = sorted(set(re.findall(r'"(/library/daily-renewable-reports-[a-z]{3}-\d{4})"', index)))
    out = {}
    for m in months:
        mon, year = m.rsplit("-", 2)[1:]
        first = pd.Timestamp(f"{mon} 1 {year}")
        if first + pd.offsets.MonthEnd(0) < pd.Timestamp(since) or first > pd.Timestamp(until):
            continue
        page = get(CAISO + m, log, f"CAISO {m}").text
        for u in set(re.findall(r'"(/documents/daily-renewable-report-([a-z]{3})-(\d{2})-(\d{4})\.html)"', page)):
            d = pd.Timestamp(f"{u[1]} {u[2]} {u[3]}").date().isoformat()
            if since <= d <= until:
                out[d] = CAISO + u[0]
        time.sleep(0.3)
    return out


def caiso_drr(log, since, until):
    daily, url_of = {}, {}
    urls = caiso_drr_days(log, since, until)
    log(f"  CAISO daily renewable reports {since}..{until}: {len(urls)} found")
    for d, url in sorted(urls.items()):
        html = get(url, log, f"CAISO report {d}").text
        m = re.search(r"var chart_date = new Date\(Date\.UTC\((\d{4}), (\d+), (\d+)\)\)", html)
        if not m or dt.date(int(m.group(1)), int(m.group(2)) + 1, int(m.group(3))).isoformat() != d:
            log(f"  CAISO report {d}: chart date {m.groups() if m else None} is not {d}; not used")
            continue
        ok = True
        for fuel in ("solar", "wind"):
            tot = {"local": 0.0, "system": 0.0}
            for cat in CAISO_CATS:
                arr = js_array(html, f"curt_hr_tot_{fuel}_{cat}_mwh")
                if arr is None or len(arr) not in (23, 24, 25) or any(v is None for v in arr):
                    log(f"  CAISO report {d}: curt_hr_tot_{fuel}_{cat}_mwh missing or incomplete; day not written")
                    ok = False
                    break
                tot[cat.split("_")[1]] += sum(arr)
            if not ok:
                break
            daily[(d, f"curtailed_{fuel}_local_mwh")] = tot["local"]
            daily[(d, f"curtailed_{fuel}_system_mwh")] = tot["system"]
            daily[(d, f"curtailed_{fuel}_mwh")] = tot["local"] + tot["system"]
            tel = js_array(html, f"tot_gen_{fuel}_iso_telem")
            if tel and all(v is not None for v in tel) and len(tel) == expected_intervals(d, "America/Los_Angeles", 5):
                daily[(d, f"{fuel}_generation_mwh")] = sum(tel) * 5 / 60
            else:
                log(f"  CAISO report {d}: {fuel} telemetry incomplete ({0 if not tel else sum(v is not None for v in tel)} "
                    f"values); generation not written")
        if not ok:
            for k in [k for k in daily if k[0] == d]:
                daily.pop(k)
            continue
        url_of[d] = url
        time.sleep(0.3)
    return daily, url_of


def run_caiso(history, log, run_id):
    retrieved = ip.utc_iso(pd.Timestamp.now(tz="UTC"))
    today = pd.Timestamp.now(tz="America/Los_Angeles").date()
    frames = []
    if history:
        legacy, lurl = caiso_legacy(log, retrieved)
        frames.append(day_rows("caiso:ISO", "US-CA", legacy, "caiso:production_curtailments", lurl.get, retrieved))
    since = CAISO_DRR_FROM if history else (today - dt.timedelta(days=14)).isoformat()
    drr, durl = caiso_drr(log, since, (today - dt.timedelta(days=1)).isoformat())
    if drr:
        frames.append(day_rows("caiso:ISO", "US-CA", drr, "caiso:daily_renewable_report", durl.get, retrieved))
    s = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=ip.SERIES_COLS)
    if s.empty:
        raise ip.SourceGap("no CAISO day could be read")
    header = [
        "Energy Research Warehouse (ERW): CAISO wind and solar curtailment and output, daily, MWh",
        "Shape: series (docs/datastandard.md v0). freq P1D: ts_utc is the local (Pacific) operating day at "
        "00:00:00Z (Decision 11). Method and meaning: docs/methods/curtailment.md.",
        f"Retrieved: {run_id} (UTC) by warehouse/connectors/curtailment.py",
        f"Run log: warehouse/output/logs/curtailment_{run_id}.log; raw files warehouse/raw/curtailment/{run_id}/",
        f"Source: caiso:production_curtailments {SOURCES['caiso:production_curtailments'][1]}, {CAISO_LEGACY_PAGE} "
        f"(days before {CAISO_DRR_FROM})",
        f"Source: caiso:daily_renewable_report {SOURCES['caiso:daily_renewable_report'][1]}, {CAISO_DRR_PAGE} "
        f"(days from {CAISO_DRR_FROM})",
        "Variables: curtailed_<fuel>_mwh (all curtailment), curtailed_<fuel>_local_mwh and _system_mwh (by "
        "CAISO's reason, where published: 2024 on), <fuel>_generation_mwh (ISO output). Curtailment is not in the output figures.",
    ]
    ip.write_csv(s, "caiso_curtailment_daily", header, log)
    return len(s)


# ---------------------------------------------------------------- SPP

def spp_frame(df):
    """5-minute rows to {(date, baa, var): MWh} and interval counts {(date, baa): n}. A row with
    no interval time (a blank line) is dropped; a time in any other format raises."""
    df = df[df["GMTIntervalEnding"].notna() & (df["GMTIntervalEnding"].astype(str).str.strip() != "")]
    end = pd.to_datetime(df["GMTIntervalEnding"], format="%m/%d/%Y %H:%M:%S.%f", utc=True)
    local_day = (end - pd.Timedelta(minutes=5)).dt.tz_convert("America/Chicago").dt.date.astype(str)
    df = df.assign(day=local_day.values, baa=df["BAA"].fillna("SPP") if "BAA" in df else "SPP")
    for c in SPP_COLS:
        if c not in df:
            df[c] = 0.0
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0)
    g = df.groupby(["day", "baa"])
    sums = g[list(SPP_COLS)].sum() * 5 / 60
    counts = g.size()
    return sums, counts


def run_spp(history, log, run_id):
    retrieved = ip.utc_iso(pd.Timestamp.now(tz="UTC"))
    today = pd.Timestamp.now(tz="America/Chicago").date()
    parts = []  # (sums, counts, url of each (day))
    if history:
        for y in range(2014, SPP_FIRST_DAILY_YEAR):
            url = f"{SPP_BASE}?path=/{y}/{y}.zip"
            r = get(url, log, f"SPP {y}.zip", ok404=True)
            if r.status_code == 404:
                log(f"  SPP {y}.zip: HTTP 404, no annual file")
                continue
            z = zipfile.ZipFile(io.BytesIO(r.content))
            # the zips also hold monthly files (VER-Curtailments-MONTHLY-<yyyymm>.csv) that repeat
            # the daily rows: only the daily files are read, and an interval is kept once per BAA
            names = sorted(n for n in z.namelist() if n.lower().endswith(".csv") and "monthly" not in n.lower())
            frames = [pd.read_csv(z.open(n)) for n in names]
            df = pd.concat(frames, ignore_index=True)
            key = ["GMTIntervalEnding"] + (["BAA"] if "BAA" in df else [])
            n0 = len(df)
            df = df.drop_duplicates(key, keep="last")
            if len(df) != n0:
                log(f"  SPP {y}.zip: {n0 - len(df)} rows repeat an interval of another daily file; kept the later file's")
            sums, counts = spp_frame(df)
            parts.append((sums, counts, url))
            log(f"  SPP {y}.zip: {len(frames)} daily files, {len(df)} rows")
        start = dt.date(SPP_FIRST_DAILY_YEAR, 1, 1)
    else:
        start = today - dt.timedelta(days=10)
    d = start
    while d < today:
        url = f"{SPP_BASE}?path=/{d:%Y}/{d:%m}/VER-Curtailments-{d:%Y%m%d}.csv"
        r = get(url, log, f"SPP {d}", ok404=True)
        if r.status_code == 404:
            log(f"  SPP {d}: HTTP 404, no daily file")
        else:
            sums, counts = spp_frame(pd.read_csv(io.BytesIO(r.content)))
            parts.append((sums, counts, url))
        d += dt.timedelta(days=1)
        time.sleep(0.3)
    if not parts:
        raise ip.SourceGap("no SPP file could be read")
    sums = pd.concat([p[0] for p in parts]).groupby(level=[0, 1]).sum()
    counts = pd.concat([p[1] for p in parts]).groupby(level=[0, 1]).sum()
    url_by_day = {}
    for s_, _, url in parts:
        for day in s_.index.get_level_values(0):
            url_by_day.setdefault(day, url)
    rows, gaps = [], []
    for (day, baa), row in sums.iterrows():
        want = expected_intervals(day, "America/Chicago", 5)
        if counts[(day, baa)] != want:
            gaps.append((day, baa, int(counts[(day, baa)]), want))
            continue
        vals = {v: float(row[c]) for c, v in SPP_COLS.items()}
        vals["curtailed_wind_mwh"] = sum(vals[k] for k in vals if k.startswith("curtailed_wind_"))
        vals["curtailed_solar_mwh"] = sum(vals[k] for k in vals if k.startswith("curtailed_solar_") and k != "curtailed_solar_mwh")
        for v, x in vals.items():
            rows.append({"entity": f"spp:{baa}", "variable": v, "ts_utc": f"{day}T00:00:00Z", "value": round(x, 3),
                         "unit": "MWh", "freq": "P1D", "geo": "", "market": "", "node": "",
                         "source": "spp:ver_curtailments", "source_url": url_by_day[day], "retrieved_at": retrieved,
                         "vintage": ""})
    for g in gaps[:20]:
        log(f"  SPP {g[0]} {g[1]}: {g[2]} of {g[3]} intervals; day not written")
    # the last local day may still be filling: never written incomplete (above)
    s = pd.DataFrame(rows, columns=ip.SERIES_COLS)
    if s.empty:
        raise ip.SourceGap("no complete SPP day")
    header = [
        "Energy Research Warehouse (ERW): SPP wind and solar curtailment by balancing authority area, daily, MWh",
        "Shape: series (docs/datastandard.md v0). freq P1D: ts_utc is the Central (America/Chicago) date of the "
        "intervals' starts, at 00:00:00Z (Decision 11). Method and meaning: docs/methods/curtailment.md.",
        f"Retrieved: {run_id} (UTC) by warehouse/connectors/curtailment.py",
        f"Run log: warehouse/output/logs/curtailment_{run_id}.log; raw files warehouse/raw/curtailment/{run_id}/",
        f"Source: spp:ver_curtailments {SOURCES['spp:ver_curtailments'][1]}, {SPP_PAGE} (files {SPP_BASE})",
        "Variables: curtailed_<fuel>_redispatch_mwh, _manual_mwh, _economic_mwh (SPP's 'curtailed for energy'), "
        "and curtailed_<fuel>_mwh, their sum. MWh = 5-minute MW x 5/60. A BAA's day is written only with every "
        f"5-minute interval; {len(gaps)} BAA-days were not complete in this run.",
    ]
    ip.write_csv(s, "spp_curtailment_daily", header, log)
    return len(s)


# ---------------------------------------------------------------- ERCOT

def run_ercot(history, log, run_id):
    retrieved = ip.utc_iso(pd.Timestamp.now(tz="UTC"))
    hourly = {}  # (fuel, delivery date, hour ending) -> (gen, hsl, url), newest report wins
    for fuel, (rtid, code, _) in ERCOT_REPORTS.items():
        lst = get(ERCOT_LIST.format(rtid), log, f"ERCOT {code} list").json()
        docs = [x["Document"] for x in lst["ListDocsByRptTypeRes"]["DocumentList"]]
        docs = [x for x in docs if x["FriendlyName"].endswith("_csv")]
        docs.sort(key=lambda x: x["PublishDate"])
        # the last report of each publish day covers that day's hours and the day before
        by_day = {}
        for x in docs:
            by_day[x["PublishDate"][:10]] = x
        pick = list(by_day.values()) if history else list(by_day.values())[-3:]
        log(f"  ERCOT {code}: {len(docs)} reports listed, reading {len(pick)}")
        for x in pick:
            url = ERCOT_DOC.format(x["DocID"])
            r = get(url, log, f"ERCOT {code} {x['PublishDate']}")
            z = zipfile.ZipFile(io.BytesIO(r.content))
            df = pd.read_csv(z.open(z.namelist()[0]))
            gen_col = "SYSTEM_WIDE_GEN" if "SYSTEM_WIDE_GEN" in df else "GEN_SYSTEM_WIDE"
            if gen_col not in df or "SYSTEM_WIDE_HSL" not in df:
                raise RuntimeError(f"ERCOT {code}: columns {list(df.columns)} lack system-wide GEN or HSL")
            for row in df.itertuples(index=False):
                g, h = getattr(row, gen_col), row.SYSTEM_WIDE_HSL
                if pd.isna(g) or pd.isna(h):
                    continue
                day = pd.Timestamp(row.DELIVERY_DATE).date().isoformat()
                key = (fuel, day, int(row.HOUR_ENDING), str(getattr(row, "DSTFlag", "N")))
                hourly[key] = (float(g), float(h), url)
            time.sleep(0.3)
    daily, url_of, gaps = {}, {}, []
    days = sorted({(f, d) for f, d, _, _ in hourly})
    for fuel, day in days:
        hrs = {(he, dst): v for (f, d, he, dst), v in hourly.items() if f == fuel and d == day}
        want = expected_intervals(day, "America/Chicago", 60)
        if len(hrs) != want:
            gaps.append((fuel, day, len(hrs), want))
            continue
        gen = sum(v[0] for v in hrs.values())
        hsl = sum(v[1] for v in hrs.values())
        below = sum(max(0.0, v[1] - v[0]) for v in hrs.values())
        daily[(day, f"{fuel}_generation_mwh")] = gen
        daily[(day, f"{fuel}_hsl_mwh")] = hsl
        daily[(day, f"{fuel}_below_hsl_mwh")] = below
        url_of[day] = next(iter(hrs.values()))[2]
    for g in gaps:
        log(f"  ERCOT {g[0]} {g[1]}: {g[2]} of {g[3]} hours with GEN and HSL; day not written")
    if not daily:
        raise ip.SourceGap("no complete ERCOT operating day in the reports listed")
    s = day_rows("ercot:system", "US-TX", daily, "ercot:np4_732_745", url_of.get, retrieved)
    header = [
        "Energy Research Warehouse (ERW): ERCOT wind and solar output and High Sustained Limit (HSL), daily, MWh",
        "Shape: series (docs/datastandard.md v0). freq P1D: ts_utc is the ERCOT operating day (Central) at "
        "00:00:00Z (Decision 11). Method and meaning: docs/methods/curtailment.md.",
        f"Retrieved: {run_id} (UTC) by warehouse/connectors/curtailment.py",
        f"Run log: warehouse/output/logs/curtailment_{run_id}.log; raw files warehouse/raw/curtailment/{run_id}/",
        f"Source: ercot:np4_732_745 {SOURCES['ercot:np4_732_745'][1]}, {ERCOT_REPORTS['wind'][2]}, "
        f"{ERCOT_REPORTS['solar'][2]}",
        "Variables: <fuel>_generation_mwh (sum of hourly system-wide GEN), <fuel>_hsl_mwh (sum of hourly "
        "SYSTEM_WIDE_HSL), <fuel>_below_hsl_mwh (sum over hours of max(0, HSL - GEN)). ERCOT publishes no "
        "curtailment figure: output below HSL is the ERW's estimate of it, not ERCOT's. A day is written only "
        f"with every hour's GEN and HSL; {len(gaps)} fuel-days were not complete in this run.",
    ]
    ip.write_csv(s, "ercot_wind_solar_hsl_daily", header, log)
    return len(s)


RUNNERS = {"caiso": ("caiso_curtailment_daily", run_caiso, ["caiso:production_curtailments", "caiso:daily_renewable_report"]),
           "spp": ("spp_curtailment_daily", run_spp, ["spp:ver_curtailments"]),
           "ercot": ("ercot_wind_solar_hsl_daily", run_ercot, ["ercot:np4_732_745"])}


def main(argv=None):
    ap = argparse.ArgumentParser(description="ERW wind and solar curtailment connector")
    ap.add_argument("--history", action="store_true", help="every day each source still posts")
    ap.add_argument("--iso", action="append", choices=sorted(RUNNERS))
    ap.add_argument("--out-dir")
    args = ap.parse_args(argv)
    if args.out_dir:
        ip.set_out_dir(args.out_dir)
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"curtailment_{run_id}.log"))
    ip.RAW.open("curtailment", run_id)
    log(f"ERW curtailment {run_id}: {'history' if args.history else 'latest days'}")
    results, done = [], []
    for iso in args.iso or sorted(RUNNERS):
        table, fn, srcs = RUNNERS[iso]
        log(f"{table}:")
        try:
            n = fn(args.history, log, run_id)
            results.append(dict(table=table, market=iso, status="ok", detail=f"{n} rows this run"))
            done.append((table, srcs))
        except Exception:
            tb = ip.redact(traceback.format_exc())
            last = tb.strip().splitlines()[-1]
            log(f"{table} FAILED, no output file written:\n{tb}")
            print(f"curtailment {table} FAILED, no output file written: {last}", file=sys.stderr)
            results.append(dict(table=table, market=iso, status="failed", detail=last[:300]))
    ip.update_sources([dict(source=s, publisher=SOURCES[s][0], report=SOURCES[s][1], report_url=SOURCES[s][2],
                            document_list=SOURCES[s][3], tables=[t]) for t, srcs in done for s in srcs])
    ip.write_status("curtailment", run_id, results)
    failures = sum(r["status"] == "failed" for r in results)
    log(f"done, failures={failures}")
    log.close()
    print(f"curtailment run log: {os.path.relpath(log.path, ip.ROOT)}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
