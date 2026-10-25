#!/usr/bin/env python3
"""Tajikistan's electricity a year, 1991 to 2025: output (hydro and thermal), imports, exports and internal consumption by
sector, with losses, as the Agency on Statistics under the President of the Republic of Tajikistan gives them in its analytical
table "Electricity output, trade and consumption, 1991-2025" (session 2659, chain 133; the ERW held Tajikistan only in EIA's
world tables).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/stat_tj_tajikistan_electricity_annual.csv, from
one workbook (no key):

    https://www.stat.tj/wp-content/uploads/2026/09/table_11-1.xlsx

    entity      tajstat:tj:<item>: output, hydro, thermal (variable generation); imports; exports; consumption (the total
                internal consumption) and industry, construction, agriculture, transport, other_sectors, households (variable
                consumption); losses (distribution_losses); discrepancy (statistical_difference)
    unit        MWh (the table's billion kWh times 1,000,000); ts_utc the year's first day; freq P1Y; geo TJ
    node        the row's label as the table gives it

The sheet is read by its title ("Electricity output, trade and consumption"), its unit line ("in billions of kilowatt hours")
and its header row of years, a run from 1991; each row is found by its label, and a label not in the map stops the run. A dash
or a blank cell has no row; nothing is filled. The table's parts do not always add to its totals: from 2015 the output exceeds
hydro and thermal power by up to 0.9 billion kWh (2017), and from 2016 the total internal consumption differs from its sectors
and losses. Every row is kept as published; the run stops only if a year's hydro and thermal power miss its output by more than
1 billion kWh, or if a value is negative where it cannot be.

    python warehouse/connectors/stat_tj_tajikistan_electricity.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/stat_tj/ and reused unless --fresh; --offline asks nothing. www.stat.tj's
robots file closes only the WordPress admin paths (stat.tj sends to www.stat.tj; session 2659). License: CC BY 4.0, the site's
footer ("All materials of the site are available under the license: Creative Commons Attribution 4.0 International"; "When using
materials published on this site, a link to www.stat.tj is required"): public.
"""

import argparse
import datetime as dt
import io
import os
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "stat_tj_tajikistan_electricity_annual"
CONNECTOR = "stat_tj"
SOURCE = "stat_tj:analytical_table_11"
HOST = "www.stat.tj"
URL = "https://www.stat.tj/wp-content/uploads/2026/09/table_11-1.xlsx"
PAGE = "https://www.stat.tj/en/analytical-tables/"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = ("Agency on Statistics under the President of the Republic of Tajikistan, analytical table: Electricity output, trade "
            "and consumption, 1991-2025 (www.stat.tj)")
TITLE = "Electricity output, trade and consumption"
UNIT = "in billions of kilowatt hours"
# the row's label: (ERW item, variable)
ROWS = {"Output": ("output", "generation"), "Hydropower": ("hydro", "generation"), "Thermal power": ("thermal", "generation"),
        "Imports": ("imports", "imports"), "Exports": ("exports", "exports"),
        "Total internal consumption (by sector)": ("consumption", "consumption"), "Industry": ("industry", "consumption"),
        "Construction": ("construction", "consumption"), "Agriculture": ("agriculture", "consumption"),
        "Transport": ("transport", "consumption"), "Other sectors": ("other_sectors", "consumption"),
        "Households": ("households", "consumption"), "Losses": ("losses", "distribution_losses"),
        "Discrepancy": ("discrepancy", "statistical_difference")}
SKIP = {"of which:"}
X_COLS = []


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def label(v):
    return " ".join(str(v).split()) if v is not None else ""


def parse_rows(rows):
    rows = [r for r in rows if any(x is not None and str(x).strip() != "" for x in r)]
    if not rows or not label(rows[0][0]).startswith(TITLE):
        raise RuntimeError(f"the title {TITLE!r} not found")
    if label(rows[1][0]) != UNIT:
        raise RuntimeError(f"the unit line is {label(rows[1][0])!r}")
    h = rows[2]
    years = []
    for x in h[1:]:
        if is_num(x) and x == int(x):
            years.append(int(x))
        elif x is None or str(x).strip() == "":
            continue
        else:
            raise RuntimeError(f"a header cell not a year: {x!r}")
    if not years or years[0] != 1991 or years != list(range(1991, 1991 + len(years))):
        raise RuntimeError(f"the header's years are not a run from 1991: {years[:3]} ... {years[-3:]}")
    out, seen = {}, set()
    for r in rows[3:]:
        lab = label(r[0])
        if lab in SKIP:
            continue
        if lab not in ROWS:
            raise RuntimeError(f"a row not in the map: {lab!r}")
        if lab in seen:
            raise RuntimeError(f"a row given twice: {lab!r}")
        seen.add(lab)
        for j, y in enumerate(years):
            v = r[1 + j] if 1 + j < len(r) else None
            if v is None or (isinstance(v, str) and v.strip() in ("", "-")):
                continue
            if not is_num(v):
                raise RuntimeError(f"{lab} {y}: not a number: {v!r}")
            if v < 0 and lab != "Discrepancy":
                raise RuntimeError(f"{lab} {y}: a negative value {v}")
            out[(ROWS[lab][0], y)] = float(v)
    missing = sorted(set(ROWS) - seen)
    if missing:
        raise RuntimeError(f"rows missing: {missing}")
    for y in years:
        if ("output", y) in out:
            s = out.get(("hydro", y), 0.0) + out.get(("thermal", y), 0.0)
            if abs(out[("output", y)] - s) > 1.0:
                raise RuntimeError(f"{y}: hydro and thermal make {s:.3f} billion kWh, the output {out[('output', y)]}")
    return out


def frame(vals, base):
    lab = {v[0]: k for k, v in ROWS.items()}
    var = {v[0]: v[1] for v in ROWS.values()}
    out = [dict(entity=f"tajstat:tj:{item}", variable=var[item], ts_utc=f"{y}-01-01T00:00:00Z", value=round(v * 1e6, 3), unit="MWh",
                freq="P1Y", geo="TJ", market="", node=lab[item], **base)
           for (item, y), v in sorted(vals.items())]
    return pd.DataFrame(out)


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    if len(wb.worksheets) != 1:
        raise RuntimeError(f"{len(wb.worksheets)} sheets, not one: {wb.sheetnames}")
    return frame(parse_rows([tuple(r) for r in wb.worksheets[0].iter_rows(values_only=True)]), base)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Agency on Statistics of Tajikistan: electricity output, trade and consumption a year")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t = parse(content, base)
        if t.duplicated(["entity", "variable", "ts_utc"]).any():
            raise RuntimeError("a key given twice")
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["entity", "ts_utc"]).reset_index(drop=True)
        out = t[t.entity == "tajstat:tj:output"].set_index("ts_utc").value
        head = [
            "Energy Research Warehouse (ERW): Agency on Statistics under the President of the Republic of Tajikistan, Tajikistan's "
            f"electricity output, trade and consumption, {t.ts_utc.min()[:4]} to {t.ts_utc.max()[:4]} (session 2659)",
            "Shape: series (docs/datastandard.md v0). entity tajstat:tj:<item>; variable generation, imports, exports, consumption, "
            "distribution_losses or statistical_difference; unit MWh (the table's billion kWh times 1,000,000); ts_utc the year's "
            "first day; freq P1Y; geo TJ; node the row's label.",
            f"{len(t)} rows; the output, billion kWh: " + ", ".join(f"{k[:4]} {v / 1e6:.1f}" for k, v in out.items()
                                                            if k[:4] in ("1991", "2000", "2010", "2020", "2025"))
            + ". Every row as published: the parts do not always add to the totals (from 2015 and 2016); nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/stat_tj_tajikistan_electricity.py; the workbook retrieved "
            f"{base['retrieved_at']}; last modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL} (listed on {PAGE})",
            "License: CC BY 4.0, the site's footer; a link to www.stat.tj is required. Credit: Agency on Statistics under the President "
            "of the Republic of Tajikistan, www.stat.tj.",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="Agency on Statistics under the President of the Republic of Tajikistan",
                                    report="Analytical table: Electricity output, trade and consumption, 1991-2025", report_url=URL,
                                    document_list=PAGE, license="public", tables=[NAME])])  # CC BY 4.0
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
