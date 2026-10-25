#!/usr/bin/env python3
"""Installed renewable electricity capacity at the end of each year by region and technology across the UK, 2003 to 2025:
England's nine regions, Wales, Scotland, Northern Ireland, England and the UK, as the Department for Energy Security and Net
Zero (DESNZ) publishes them in its Regional Renewable Statistics, the workbook "Regional installed capacity MW 2003-2025"
(session 2692, chain 135, the session's own choice; the companion of session 2684's generation by region, from the same page).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/desnz_regional_renewable_capacity_annual.csv, from
one workbook (no key), a sheet a year ("MW<year>"):

    https://assets.publishing.service.gov.uk/media/6aba6ee6997a4b2950ccefec/Regional_installed_capacity_MW_2003-2025.xlsx

    entity      desnz_rc:<region>:<technology>: region england, north_east, north_west, yorkshire_humber, east_midlands,
                west_midlands, east_of_england, london, south_east, south_west, wales, scotland, northern_ireland, other_sites
                (not attributed to a region), uk, co_firing (the non-fossil capacity used for co-firing, outside the UK
                total); technology as the year's columns give it (onshore_wind, offshore_wind, wind,
                wave_tidal, solar, hydro, landfill_gas, sewage_gas, anaerobic_digestion, biomass_waste, other_bioenergy,
                other_bioenergy_incl_sewage, total)
    variable    capacity
    unit        MW (at 31 December, DESNZ's note 2); ts_utc the year's first day; freq P1Y; geo GB (the UK, the other
                sites and co-firing), GB-SCT, GB-WLS, GB-NIR, GB-ENG (England and its regions: ISO 3166-2 has no code for them)
    node        the region and the column as DESNZ gives them (note marks dropped)

The columns change over the years (one "Wind" column in the early years, then onshore and offshore; "Other bioenergy" in the
early years, then sewage gas, anaerobic digestion and biomass and waste apart); each column is kept as its own technology, and
none is merged or split. Each sheet is checked for its title ("Installed capacity of sites generating electricity from renewable
sources, <year>"), its header and its fifteen region rows in order, as DESNZ names them ("Yorkshire and Humber" from 2015; and
"Eastern", "Other" and "Total" in 2022), then the co-firing row; the footnote rows after it are not read, and any other row
stops the run. "[x]" (too few sites to show), "[y]" (not available) and a blank cell have no row; nothing is filled. DESNZ's
note 1: components may not add exactly to totals because of rounding. Checks: a region's technologies add to its total within
1 MW where no cell is "[x]" (no more than it where one is); England's nine regions add to England within 1 MW, and the four
nations and the other sites to the UK (co-firing is outside the UK total).

    python warehouse/connectors/desnz_regional_renewable_capacity.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/desnz_regional/ and reused unless --fresh; --offline asks nothing.
www.gov.uk's robots file closes only print and search paths; assets.publishing.service.gov.uk has no robots file (session 2678).
License: the Open Government Licence v3.0 (DESNZ's verdict on record, session 2167): public.
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "desnz_regional_renewable_capacity_annual"
CONNECTOR = "desnz_regional"
SOURCE = "desnz:regional_renewable_capacity_2003_2025"
HOST = "assets.publishing.service.gov.uk"
URL = "https://assets.publishing.service.gov.uk/media/6aba6ee6997a4b2950ccefec/Regional_installed_capacity_MW_2003-2025.xlsx"
PAGE = "https://www.gov.uk/government/statistics/regional-renewable-statistics"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = "Department for Energy Security and Net Zero, Regional Renewable Statistics: Regional installed capacity MW 2003-2025"
REGIONS = {"England": ("england", "GB-ENG"), "North East": ("north_east", "GB-ENG"), "North West": ("north_west", "GB-ENG"),
           "Yorkshire and the Humber": ("yorkshire_humber", "GB-ENG"), "East Midlands": ("east_midlands", "GB-ENG"),
           "West Midlands": ("west_midlands", "GB-ENG"), "East of England": ("east_of_england", "GB-ENG"),
           "London": ("london", "GB-ENG"), "South East": ("south_east", "GB-ENG"), "South West": ("south_west", "GB-ENG"),
           "Wales": ("wales", "GB-WLS"), "Scotland": ("scotland", "GB-SCT"), "Northern Ireland": ("northern_ireland", "GB-NIR"),
           "Other Sites": ("other_sites", "GB"), "UK Total": ("uk", "GB")}
# the region as some sheets name it: "Yorkshire and Humber" from 2015; "Eastern", "Other" and "Total" in 2022
ALIASES = {"Yorkshire and Humber": "Yorkshire and the Humber", "Eastern": "East of England", "Other": "Other Sites",
           "Total": "UK Total"}
COFIRING = "Co-firing"
TECH = {"Onshore Wind": "onshore_wind", "Offshore Wind": "offshore_wind", "Wind": "wind", "Wave and tidal": "wave_tidal",
        "Solar PV": "solar", "Hydro": "hydro", "Landfill gas": "landfill_gas", "Sewage gas": "sewage_gas",
        "AD": "anaerobic_digestion", "Biomass and waste": "biomass_waste", "Other bioenergy": "other_bioenergy",
        "Other bioenergy (incl Sewage gas)": "other_bioenergy_incl_sewage", "Other bioenergy (incl Sewage gas": "other_bioenergy_incl_sewage",
        "Total": "total"}  # the 2003 sheet heads one column without its closing parenthesis
ENGLISH = ["north_east", "north_west", "yorkshire_humber", "east_midlands", "west_midlands", "east_of_england", "london",
           "south_east", "south_west"]
X_COLS = []


def label(v):
    return re.sub(r"\s*\[[^\]]*\]", "", " ".join(str(v).split())).strip() if v is not None else ""


def parse_sheet(rows, year):
    if label(rows[0][0]) != f"Installed capacity of sites generating electricity from renewable sources, {year}":
        raise RuntimeError(f"{year}: the title is {rows[0][0]!r}")
    h = next((i for i, r in enumerate(rows) if label(r[0]) == "Region"), None)
    if h is None:
        raise RuntimeError(f"{year}: no header row")
    cols = {}
    for j, x in enumerate(rows[h][1:], 1):
        s = label(x)
        if not s:
            continue
        if s not in TECH:
            raise RuntimeError(f"{year}: a column not in the map: {s!r}")
        cols[j] = (TECH[s], s)
    if "total" not in [c[0] for c in cols.values()]:
        raise RuntimeError(f"{year}: no Total column")
    lines = [r for r in rows[h + 1:] if label(r[0])]
    names = [ALIASES.get(label(r[0]), label(r[0])) for r in lines]
    if names[:len(REGIONS)] != list(REGIONS) or len(names) <= len(REGIONS) or names[len(REGIONS)] != COFIRING:
        raise RuntimeError(f"{year}: the regions are {names[:17]}")
    if not all("=" in n for n in names[len(REGIONS) + 1:]):
        raise RuntimeError(f"{year}: a row after co-firing not a footnote: {names[len(REGIONS) + 1:]}")
    body = lines[:len(REGIONS) + 1]
    out, hidden = {}, set()
    for r in body:
        name = ALIASES.get(label(r[0]), label(r[0]))
        reg = "co_firing" if name == COFIRING else REGIONS[name][0]
        for j, (tech, s) in cols.items():
            v = r[j] if j < len(r) else None
            if v is None or (isinstance(v, str) and v.strip() in ("", "[x]", "[y]")):
                if isinstance(v, str) and v.strip() == "[x]":
                    hidden.add(reg)
                continue
            if not isinstance(v, (int, float)) or v < 0:
                raise RuntimeError(f"{year} {reg} {s}: not a quantity: {v!r}")
            out[(reg, tech)] = (float(v), f"{name}: {s}")
    g = lambda reg, tech: out.get((reg, tech), (0.0, ""))[0]
    for reg in [x[0] for x in REGIONS.values()] + ["co_firing"]:
        if (reg, "total") not in out:
            continue
        s = sum(v for (rr, t), (v, _) in out.items() if rr == reg and t != "total")
        tot = g(reg, "total")
        if (reg not in hidden and abs(s - tot) > 1) or (reg in hidden and s > tot + 1):
            raise RuntimeError(f"{year} {reg}: the technologies make {s:,.1f} MW, the total {tot:,.1f}")
    for whole, parts in (("england", ENGLISH), ("uk", ["england", "wales", "scotland", "northern_ireland", "other_sites"])):
        s = sum(g(p, "total") for p in parts)
        if abs(s - g(whole, "total")) > 1:
            raise RuntimeError(f"{year}: the parts of {whole} make {s:,.1f} MW, its total {g(whole, 'total'):,.1f}")
    return out


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    sheet = {int(m.group(1)): s for s in wb.sheetnames for m in [re.fullmatch(r"MW ?(\d{4})", s)] if m}
    years = sorted(sheet)
    if not years or years[0] != 2003 or years != list(range(2003, years[-1] + 1)):
        raise RuntimeError(f"the yearly sheets are not a run from 2003: {wb.sheetnames}")
    geo = {v[0]: v[1] for v in REGIONS.values()}
    geo["co_firing"] = "GB"
    rows = []
    for y in years:
        vals = parse_sheet([tuple(r) for r in wb[sheet[y]].iter_rows(values_only=True)], y)
        rows += [dict(entity=f"desnz_rc:{reg}:{tech}", variable="capacity", ts_utc=f"{y}-01-01T00:00:00Z", value=round(v, 3),
                      unit="MW", freq="P1Y", geo=geo[reg], market="", node=lab, **base)
                 for (reg, tech), (v, lab) in sorted(vals.items())]
    return pd.DataFrame(rows)


def main(argv=None):
    ap = argparse.ArgumentParser(description="DESNZ: installed renewable capacity a year by UK region and technology")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t = parse(content, base)
        if t.duplicated(["entity", "variable", "ts_utc"]).any():
            raise RuntimeError("a key given twice")
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["entity", "ts_utc"]).reset_index(drop=True)
        uk = t[t.entity == "desnz_rc:uk:total"].set_index("ts_utc").value
        head = [
            "Energy Research Warehouse (ERW): DESNZ, Regional Renewable Statistics: installed renewable capacity at year end by UK "
            f"region and technology, {t.ts_utc.min()[:4]} to {t.ts_utc.max()[:4]} (session 2692)",
            "Shape: series (docs/datastandard.md v0). entity desnz_rc:<region>:<technology>; variable capacity; unit MW (at 31 "
            "December); ts_utc the year's first day; freq P1Y; geo GB, GB-SCT, GB-WLS, GB-NIR or GB-ENG (England "
            "and its regions); node the region and the column.",
            f"{len(t)} rows; the UK's renewable capacity, GW: " + ", ".join(f"{k[:4]} {v / 1e3:.1f}" for k, v in uk.items()
                                                                          if k[:4] in ("2003", "2010", "2020", "2025"))
            + ". Each year's columns kept as given; regions, England and the UK checked against their totals; nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/desnz_regional_renewable_capacity.py; the workbook retrieved "
            f"{base['retrieved_at']}; last modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL} (listed on {PAGE})",
            "License: Open Government Licence v3.0. Credit: Department for Energy Security and Net Zero.",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="UK Department for Energy Security and Net Zero (DESNZ)",
                                    report="Regional Renewable Statistics: Regional installed capacity MW 2003-2025", report_url=PAGE,
                                    document_list=PAGE, license="public", tables=[NAME])])  # Open Government Licence v3.0
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
