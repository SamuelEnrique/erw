#!/usr/bin/env python3
"""Anguilla's electricity generation a month, 1995 to 2025, as the Anguilla Statistics Department gives it in Table 2.4.2-1
("Electricity Generation, Month: 1995 - 2026") of its Energy Statistics workbook, from the Anguilla Electricity Company's
(ANGLEC) figures (session 2642, chain 131, the session's own choice; the ERW held Anguilla in no table but EIA's world one).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/statai_anguilla_electricity_monthly.csv, from one
workbook (no key):

    https://statistics.gov.ai/AllDocuments/2.4.2%20Energy%20Statistics.xlsx

    entity      statai:ai:anglec
    variable    generation
    unit        MWh (the table's kWh divided by 1,000); ts_utc the month's first day; freq P1M; geo AI
    node        the table's label

The sheet "Table 2.4.2-1" is read by its title cells ("Electricity Generation") and its header row ("Month", then the years
from 1995 without a gap, then "% Change ..."); the twelve month rows follow, then "Total". Only whole years are written: a year
whose twelve months are not all figures is left out. 2026 is so left out: it holds January to March only, and its February and
March repeat 2025's exactly (the table's % change 0). Check: each year's twelve months add to its Total within 1 kWh. The
annual % change and the other sheets (revenue in XCD, petroleum imports) are not taken. Nothing is filled.

    python warehouse/connectors/statai_anguilla_electricity.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/statai_ai/ and reused unless --fresh; --offline asks nothing.
statistics.gov.ai has no robots file (HTTP 404, session 2642). License: internal. The Department's site reads "(c) 2026 Anguilla
Statistics Department, GOA. All rights reserved." and links no terms of reuse; internal until a ruling.
"""

import argparse
import datetime as dt
import io
import os
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "statai_anguilla_electricity_monthly"
CONNECTOR = "statai_ai"
SOURCE = "statai_ai:energy_statistics_table_2_4_2_1"
HOST = "statistics.gov.ai"
URL = "https://statistics.gov.ai/AllDocuments/2.4.2%20Energy%20Statistics.xlsx"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = "Anguilla Statistics Department, Energy Statistics, Table 2.4.2-1: Electricity Generation by month (source: ANGLEC)"
SHEET = "Table 2.4.2-1"
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
X_COLS = []


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def parse_rows(rows):
    if not any(r and "Electricity Generation" in [str(x).strip() for x in r if x is not None] for r in rows[:8]):
        raise RuntimeError("the title Electricity Generation not found")
    hd = [i for i, r in enumerate(rows) if r and "Month" in [str(x).strip() if x is not None else "" for x in r]]
    if len(hd) != 1:
        raise RuntimeError(f"the header row found {len(hd)} times")
    h = list(rows[hd[0]])
    c = [str(x).strip() if x is not None else "" for x in h].index("Month")
    years = []
    for x in h[c + 1:]:
        if is_num(x) and x == int(x):
            years.append(int(x))
        else:
            break
    if not years or years[0] != 1995 or years != list(range(1995, 1995 + len(years))):
        raise RuntimeError(f"the header's years are not a run from 1995: {years[:3]} ... {years[-3:]}")
    body = rows[hd[0] + 1:hd[0] + 14]
    labels = [str(r[c]).strip() if r[c] is not None else "" for r in body]
    if labels != MONTHS + ["Total"]:
        raise RuntimeError(f"the rows are {labels}")
    out = {}
    for j, y in enumerate(years):
        vals = [body[m][c + 1 + j] for m in range(12)]
        if not all(is_num(v) for v in vals):
            continue  # not a whole year
        tot = body[12][c + 1 + j]
        if not is_num(tot) or abs(sum(vals) - tot) > 1:
            raise RuntimeError(f"{y}: the months make {sum(vals):,.2f} kWh, the Total {tot}")
        out[y] = [float(v) for v in vals]
    return out


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    if SHEET not in wb.sheetnames:
        raise RuntimeError(f"no sheet {SHEET}: {wb.sheetnames}")
    return frame(parse_rows([tuple(r) for r in wb[SHEET].iter_rows(values_only=True)]), base)


def frame(years, base):
    out = [dict(entity="statai:ai:anglec", variable="generation", ts_utc=f"{y}-{m + 1:02d}-01T00:00:00Z", value=round(v[m] / 1000, 3),
                unit="MWh", freq="P1M", geo="AI", market="", node="Electricity Generation", **base)
           for y, v in sorted(years.items()) for m in range(12)]
    return pd.DataFrame(out)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Anguilla Statistics Department: Anguilla's electricity generation a month")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t = parse(content, base)
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["entity", "ts_utc"]).reset_index(drop=True)
        yr = t.groupby(t.ts_utc.str[:4]).value.sum()
        head = [
            "Energy Research Warehouse (ERW): Anguilla Statistics Department, Energy Statistics, Table 2.4.2-1: Anguilla's electricity "
            "generation a month, whole years 1995 to 2025, from ANGLEC's figures (session 2642)",
            "Shape: series (docs/datastandard.md v0). entity statai:ai:anglec; variable generation; unit MWh (the table's kWh divided by "
            "1,000); ts_utc the month's first day; freq P1M; geo AI; node the table's label.",
            f"{len(t)} rows; {yr.iloc[0]:,.1f} MWh in {yr.index[0]}, {yr.iloc[-1]:,.1f} MWh in {yr.index[-1]}. Whole years only (2026's three "
            "months left out: February and March repeat 2025's); each year's months checked against its Total; nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/statai_anguilla_electricity.py; the workbook retrieved {base['retrieved_at']}; "
            f"last modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL}",
            f"License: internal. The Department's site states only its copyright line; internal until a ruling. Cite: {CITATION}.",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="Anguilla Statistics Department",
                                    report="Energy Statistics, Table 2.4.2-1: Electricity Generation by month", report_url=URL,
                                    document_list="https://statistics.gov.ai/StatisticsDept/Energy2_4_2", license="internal", tables=[NAME])])
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
