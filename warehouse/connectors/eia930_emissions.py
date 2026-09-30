#!/usr/bin/env python3
"""Hourly CO2 emissions estimates of EIA-930, from EIA's per-balancing-authority workbooks (session 32).

Energy Research Warehouse (ERW) connector. EIA publishes its hourly CO2 estimates for EIA-930 only in the
per-balancing-authority Excel workbooks of the Hourly Electric Grid Monitor (not in the API v2 and not in the six-month
bulk files: session 31), at

    https://www.eia.gov/electricity/gridmonitor/knownissues/xls/<BA>.xlsx     (CISO, ERCO, ISNE, MISO, NYIS, PJM, SWPP)
    https://www.eia.gov/electricity/gridmonitor/knownissues/xls/Region_US48.xlsx

Each workbook's sheet "Published Hourly Data" has one row per hour from 2015-07-01 (UTC time = the hour's END) and the
CO2 columns from 2018-07-01: by fuel (COL, NG, OIL, Other), Generated, Imported, Exported and Consumed, in metric tons,
with EIA's emission factors (lbs/kWh) and intensities. Its "Notes" sheet defines each column.

This connector writes one series table, warehouse/output/eia930_all_emissions.csv, partition column ba, tCO2:

    co2_emissions_generated   EIA's "CO2 Emissions Generated": CO2 from all sources generated in the BA
    co2_emissions_consumed    EIA's "CO2 Emissions Consumed": generated plus imported minus exported
    co2_emissions_coal        EIA's "CO2 Emissions: COL"      (session 34)
    co2_emissions_natural_gas EIA's "CO2 Emissions: NG"       (session 34)
    co2_emissions_oil         EIA's "CO2 Emissions: OIL"      (session 34)
    co2_emissions_other       EIA's "CO2 Emissions: Other"    (session 34)
    co2_emissions_imported    EIA's "CO2 Emissions Imported"  (session 34)
    co2_emissions_exported    EIA's "CO2 Emissions Exported"  (session 34)

Session 32 wrote the first two (its ceiling held two); session 34 added the other six from the same saved workbooks,
without a new download (--from-raw).

Time: ts_utc = EIA's UTC time minus one hour (the hour's start), as eia930.py. Values as EIA writes them (floats, not
rounded). Completeness per UTC day, EIA-930 generation's rule (session 6 ruling a; per day since sessions 13 and 16):
a day is written only when the core pair (generated, consumed) has all 24 hours; each of the six others is written for
that day only when it too has all 24 hours, and is otherwise left out for that day alone. Every day or variable-day
left out is recorded as a gap in run_status.csv. Adding the six changed no day of the pair.

    python warehouse/connectors/eia930_emissions.py                  # the whole series (session 32)
    python warehouse/connectors/eia930_emissions.py --days 3         # the daily run: the newest workbooks, last 3 days
    python warehouse/connectors/eia930_emissions.py --from-raw --variables co2_emissions_coal,...   # session 34

Raw files and resume: every workbook is downloaded through requests into warehouse/raw/eia930_emissions/<run_id>/ with
a manifest (URL, Last-Modified, sha256). Before a download, the connector asks for the workbook's Last-Modified (HEAD)
and reuses a saved copy with the same Last-Modified from any earlier run, so an interrupted pull resumes from the files
already saved and never downloads one twice. --from-raw asks nothing: it reads the newest saved copy of each workbook.

One pass (session 34): each workbook is read once, with openpyxl in read-only (streaming) mode, one at a time, and the
columns the ERW uses are written to an extract beside the raw files, warehouse/raw/eia930_emissions/<run_id>/
<ba>_hours.csv: the hour's start, EIA's Demand and Net generation (MWh), and the eight CO2 columns (tCO2), as EIA wrote
them. The table's rows are built from that extract, and warehouse/derived/carbon_intensity.py reads the workbooks' own
Demand and Net generation from it, so no workbook is opened twice. The table is written and merged in a stream.
"""

import argparse
import csv
import datetime as dt
import glob
import hashlib
import os
import sys
import traceback

import pandas as pd
import requests

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import eia930  # noqa: E402  (BAS: respondent codes and geo)
import iso_prices as ip  # noqa: E402

NAME = "eia930_all_emissions"
BASE = "https://www.eia.gov/electricity/gridmonitor/knownissues/xls/"
FILE = {"ciso": "CISO", "erco": "ERCO", "isne": "ISNE", "miso": "MISO", "nyis": "NYIS", "pjm": "PJM", "swpp": "SWPP",
        "us48": "Region_US48"}
SHEET = "Published Hourly Data"
# EIA's column -> the ERW variable. GROUPS[0] is the core: a day is written only when both have all 24 hours
COLUMNS = {"CO2 Emissions Generated": "co2_emissions_generated", "CO2 Emissions Consumed": "co2_emissions_consumed",
           "CO2 Emissions: COL": "co2_emissions_coal", "CO2 Emissions: NG": "co2_emissions_natural_gas",
           "CO2 Emissions: OIL": "co2_emissions_oil", "CO2 Emissions: Other": "co2_emissions_other",
           "CO2 Emissions Imported": "co2_emissions_imported", "CO2 Emissions Exported": "co2_emissions_exported"}
GROUPS = [["co2_emissions_generated", "co2_emissions_consumed"],
          ["co2_emissions_coal", "co2_emissions_natural_gas", "co2_emissions_oil", "co2_emissions_other",
           "co2_emissions_imported", "co2_emissions_exported"]]
# session 34: the workbooks' own demand and net generation, kept in the extract for carbon_intensity.py (not in the table)
DENOMS = {"Demand": "demand_mwh", "Net generation": "net_generation_mwh"}
EXTRACT_COLS = ["ts_utc", *DENOMS.values(), *COLUMNS.values()]
SOURCE = "eia:gridmonitor/knownissues/xls"
REPORT = ("Form EIA-930, Hourly Electric Grid Monitor: per-balancing-authority workbooks, sheet Published Hourly Data "
          "(estimated CO2 emissions)")
PAGE = "https://www.eia.gov/electricity/gridmonitor/about"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-eia930-emissions/0.1; +https://github.com/SamuelEnrique/erw)"}
COLS = ip.SERIES_COLS + ["ba"]
KEY = ["entity", "variable", "ts_utc"]
# session 34: the extract always holds every hour from here, whatever the table's window, so carbon_intensity.py has the
# workbooks' demand and net generation since 2018 from the same single read of each workbook
HISTORY_START = pd.Timestamp("2018-07-01", tz="UTC")


def saved_copy(url, last_modified):
    """A workbook saved by an earlier run with the same Last-Modified, or None."""
    for man in sorted(glob.glob(os.path.join(ip.RAW_DIR, "eia930_emissions", "*", "manifest.csv")), reverse=True):
        with open(man, encoding="utf-8") as f:
            for r in csv.DictReader(f):
                # the HEAD request is captured too, as a 0-byte 200: only a whole xlsx counts as a saved copy
                if r["url"] == url and r["last_modified"] == last_modified and r["status"] == "200" \
                        and int(r["bytes"]) > 0:
                    path = os.path.join(os.path.dirname(man), r["file"])
                    if os.path.exists(path) and os.path.getsize(path) == int(r["bytes"]):
                        with open(path, "rb") as f:
                            if f.read(2) == b"PK":
                                return path
    return None


def workbook(code, log):
    """(path, url, last_modified) of the newest workbook of a BA, downloaded unless a saved copy is current."""
    url = BASE + FILE[code] + ".xlsx"
    head = ip.with_retries(f"HEAD {url}", lambda: requests.head(url, headers=UA, timeout=60), log)
    lm = head.headers.get("Last-Modified", "")
    path = saved_copy(url, lm) if lm else None
    if path:
        log(f"  {code}: saved copy of {url} (Last-Modified {lm}): {os.path.relpath(path, ip.ROOT)}")
        return path, url, lm

    def get():
        r = requests.get(url, headers=UA, timeout=900)
        if r.status_code != 200 or not r.content.startswith(b"PK"):
            raise RuntimeError(f"{url}: HTTP {r.status_code}, {len(r.content)} bytes, not an xlsx")
        return r
    with ip.RAW.collect() as recs:
        r = ip.with_retries(f"GET {url}", get, log)
    rec = [x for x in recs if x["url"] == url and str(x["status"]) == "200"][-1]
    path = os.path.join(ip.RAW.dir, rec["file"])
    log(f"  {code}: downloaded {url}, {len(r.content):,} bytes, Last-Modified {r.headers.get('Last-Modified')}")
    return path, url, r.headers.get("Last-Modified", "")



def newest_saved(code, log):
    """Session 34: (path, url, last_modified, retrieved_at) of the newest workbook of a BA saved by any earlier run,
    from the manifests, checked against its sha256; no request is made."""
    url = BASE + FILE[code] + ".xlsx"
    best = None
    for man in glob.glob(os.path.join(ip.RAW_DIR, "eia930_emissions", "*", "manifest.csv")):
        with open(man, encoding="utf-8") as f:
            for r in csv.DictReader(f):
                if r["url"] == url and r["status"] == "200" and int(r["bytes"] or 0) > 0:
                    path = os.path.join(os.path.dirname(man), r["file"])
                    if os.path.exists(path) and (best is None or r["retrieved_at"] > best[3]):
                        best = (path, url, r["last_modified"], r["retrieved_at"], r["sha256"])
    if best is None:
        raise RuntimeError(f"{code}: no saved workbook of {url} under warehouse/raw/eia930_emissions/")
    with open(best[0], "rb") as f:
        if hashlib.sha256(f.read()).hexdigest() != best[4]:
            raise RuntimeError(f"{best[0]}: sha256 differs from its manifest")
    log(f"  {code}: saved workbook {os.path.relpath(best[0], ip.ROOT)} (Last-Modified {best[2]}, retrieved {best[3]}); "
        "no request")
    return best[:4]


def extract(path, code, start, end, log, url="", lm="", retrieved=""):
    """Session 34, the one pass: read the workbook once (streamed) and write the hours in [start, end) to
    <ba>_hours.csv in this run's raw folder: ts_utc (the hour's start), EIA's Demand and Net generation, and the eight
    CO2 columns, as EIA wrote them (empty where EIA's cell is empty). Returns the extract's path."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise RuntimeError(f"{path}: no sheet {SHEET!r} (sheets {wb.sheetnames})")
    ws = wb[SHEET]
    rows = ws.iter_rows(values_only=True)
    hdr = list(next(rows))
    # the BA workbooks name the first column "BA", the region workbook (US48) "Region"
    who = "BA" if "BA" in hdr else "Region"
    missing = [c for c in [who, "UTC time", *DENOMS, *COLUMNS] if c not in hdr]
    if missing:
        raise RuntimeError(f"{path}: columns {missing} not in {SHEET}; EIA changed the layout")
    ir, it = hdr.index(who), hdr.index("UTC time")
    iv = [hdr.index(c) for c in [*DENOMS, *COLUMNS]]
    want = eia930.BAS[code][0]
    out_path = os.path.join(ip.RAW.dir, f"{code}_hours.csv")
    n, other = 0, 0
    with open(out_path, "w", encoding="utf-8", newline="") as f:
        f.write(f"# ERW extract of {os.path.relpath(path, ip.ROOT)}, sheet {SHEET}, BA {want}; hours [{ip.utc_iso(start)}, "
                f"{ip.utc_iso(end)}); ts_utc is the hour's start; values as EIA wrote them (demand and net generation "
                "MWh, CO2 tCO2)\n")
        f.write(f"# workbook_url={url}; last_modified={lm}; retrieved_at={retrieved}\n")
        w = csv.writer(f, lineterminator="\n")
        w.writerow(EXTRACT_COLS)
        for row in rows:
            if row[it] is None:
                continue
            if row[ir] != want:
                other += 1
                continue
            ts = pd.Timestamp(row[it], tz="UTC") - pd.Timedelta(hours=1)  # EIA's UTC time is the hour's end
            if ts < start or ts >= end:
                continue
            w.writerow([ts.strftime("%Y-%m-%dT%H:%M:%SZ")]
                       + [repr(float(row[i])) if isinstance(row[i], (int, float)) else "" for i in iv])
            n += 1
    wb.close()
    if other:
        raise RuntimeError(f"{path}: {other} rows name a region other than {want}")
    log(f"  {code}: {n:,} hours in the window extracted to {os.path.relpath(out_path, ip.ROOT)}")
    return out_path


def latest_extract(code):
    """The newest extract of a BA under warehouse/raw/eia930_emissions/ (run folders sort by time), or None."""
    found = sorted(glob.glob(os.path.join(ip.RAW_DIR, "eia930_emissions", "*", f"{code}_hours.csv")))
    return found[-1] if found else None


def extract_meta(path):
    """(url, last_modified, retrieved_at) from an extract's second header line."""
    with open(path, encoding="utf-8") as f:
        f.readline()
        meta = dict(kv.split("=", 1) for kv in f.readline()[2:].strip().split("; "))
    return meta["workbook_url"], meta["last_modified"], meta["retrieved_at"]


def read_extract(path):
    with open(path, encoding="utf-8") as f:
        skip = sum(1 for ln in f if ln.startswith("#"))
    return pd.read_csv(path, skiprows=skip, dtype={"ts_utc": str})


def complete_rows(code, x, variables, url, retrieved, gaps):
    """Rows of the complete UTC days; the other days become gap records. x is the extract of one BA.

    The rule of EIA-930 generation (session 6 ruling a, per UTC day since sessions 13 and 16): the core pair
    (generated, consumed) must have all 24 hours for a day to be written at all; each of session 34's six (by fuel,
    imported, exported) is then written for that day only when it too has all 24 hours, and is otherwise left out for
    that day alone, as a per-fuel series is. Nothing is filled."""
    respondent, geo = eia930.BAS[code]
    entity = f"eia930:{respondent}"
    x = x.assign(day=x["ts_utc"].str[:10])
    first = x["day"].min() if len(x) else None
    core = GROUPS[0]

    def full_days(cols):
        by_day = x.assign(ok=x[cols].notna().all(axis=1)).groupby("day")["ok"].agg(["sum", "size"])
        return by_day, set(by_day.index[(by_day["sum"] == 24) & (by_day["size"] == 24)])

    found = {}  # (what, why) -> [(day, hours)]

    def gap(day, r, what, why):
        if day != first or r["size"] == 24:  # the partial first day of a window is not a gap of the source
            found.setdefault((what, why), []).append((day, int(r["sum"])))

    by_core, core_days = full_days(core)
    rows = []
    for var in [v for v in COLUMNS.values() if v in variables]:
        if var in core:
            days = core_days
            if var == [v for v in core if v in variables][0]:  # the core's gaps once, when the core is written
                for day, r in by_core.iterrows():
                    if day not in core_days:
                        gap(day, r, "generated+consumed", "with both CO2 values")
        elif x[var].isna().all():  # a CO2 column EIA leaves empty for this BA: not a variable of it (as eia930.py)
            continue
        else:
            by_var, var_days = full_days([var])
            days = var_days & core_days
            for day, r in by_var.iterrows():
                if day in core_days and day not in var_days:
                    gap(day, r, var.replace("co2_emissions_", ""), f"with {var}")
        y = x[x["day"].isin(days)]
        # tuples in COLS order: millions of rows as dicts would not fit in memory
        rows += [(entity, var, t, repr(float(v)), "tCO2", "PT1H", geo, "", "", SOURCE, url, retrieved, "", code)
                 for t, v in zip(y["ts_utc"], y[var])]
    for (what, why), days in found.items():
        if len(days) <= 3:  # a daily run: one gap row per day
            gaps += [dict(table=NAME, market=f"{code} {d} {what}", status="gap",
                          detail=f"{h} of 24 hours {why}; not written for this day") for d, h in days]
        else:  # a history run: one row per BA and variable, so run_status.csv stays readable
            gaps.append(dict(table=NAME, market=f"{code} {days[0][0]}..{days[-1][0]} {what}", status="gap",
                             detail=f"{len(days)} UTC days without 24 hours {why}, not written (first "
                                    f"{', '.join(d for d, _ in days[:3])}; last {days[-1][0]})"))
    return rows


def merge_write(chunks, header, log):
    """Stream the existing table, drop the rows the new ones replace, append the new, write through a temporary
    file: neither the table nor all the new rows are held whole. chunks: temporary CSV files of new rows (no header),
    one per BA. Returns (kept, replaced, added)."""
    path = os.path.join(ip.OUT_DIR, NAME + ".csv")
    tmp = path + ".tmp"
    newkeys = set()
    for c in chunks:
        with open(c, encoding="utf-8", newline="") as f:
            newkeys.update("|".join(r[:3]) for r in csv.reader(f))
    kept = replaced = n_new = 0
    with open(tmp, "w", encoding="utf-8", newline="") as out:
        for h in header:
            out.write("# " + h + "\n")
        w = csv.writer(out, lineterminator="\n")
        w.writerow(COLS)
        if os.path.exists(path):
            with open(path, encoding="utf-8", newline="") as f:
                rd = csv.reader(ln for ln in f if not ln.startswith("#"))
                if next(rd) != COLS:
                    raise RuntimeError(f"{path}: columns are not {COLS}; not merging into a file of another shape")
                for r in rd:
                    if "|".join(r[:3]) in newkeys:
                        replaced += 1
                        continue
                    w.writerow(r)
                    kept += 1
        for c in chunks:
            with open(c, encoding="utf-8", newline="") as f:
                for r in csv.reader(f):
                    w.writerow(r)
                    n_new += 1
    os.replace(tmp, path)
    return kept, replaced, n_new - replaced


def main(argv=None):
    ap = argparse.ArgumentParser(description="ERW EIA-930 CO2 emissions connector (session 32)")
    ap.add_argument("--days", type=int, help="keep only the last N complete UTC days (the daily run)")
    ap.add_argument("--ba", action="append", choices=sorted(FILE))
    ap.add_argument("--from-raw", action="store_true",
                    help="session 34: read the newest saved workbook of each BA, no request of any kind")
    ap.add_argument("--from-extract", action="store_true",
                    help="session 34: build the rows from the newest extract of each BA (<ba>_hours.csv), not a workbook")
    ap.add_argument("--dry-run", action="store_true", help="session 34: count the rows and gaps, write nothing to the table")
    ap.add_argument("--variables", help="session 34: only these variables (comma-separated; default all eight)")
    ap.add_argument("--max-new-rows", type=int, help="session 34: fail before writing if more new rows than this")
    ap.add_argument("--out-dir")
    args = ap.parse_args(argv)
    if args.out_dir:
        ip.set_out_dir(args.out_dir)
    variables = args.variables.split(",") if args.variables else list(COLUMNS.values())
    unknown = sorted(set(variables) - set(COLUMNS.values()))
    if unknown:
        ap.error(f"unknown variables {unknown}")
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"eia930_emissions_{run_id}.log"))
    ip.RAW.open("eia930_emissions", run_id)
    end = pd.Timestamp.now(tz="UTC").normalize()
    start = end - pd.Timedelta(days=args.days) if args.days else HISTORY_START
    codes = args.ba or list(FILE)
    results, gaps, chunks, used = [], [], [], []
    n_new = 0
    log(f"ERW eia930_emissions run {run_id}; window [{ip.utc_iso(start)}, {ip.utc_iso(end)}); BAs {codes}; "
        f"variables {variables}{'; from the saved workbooks, no request' if args.from_raw else ''}")
    try:
        now = ip.utc_iso(pd.Timestamp.now(tz="UTC"))
        for code in codes:
            if args.from_extract:
                ex = latest_extract(code)
                if ex is None:
                    raise RuntimeError(f"{code}: no extract under warehouse/raw/eia930_emissions/")
                url, lm, retrieved = extract_meta(ex)
                log(f"  {code}: extract {os.path.relpath(ex, ip.ROOT)} (no workbook read)")
            else:
                if args.from_raw:
                    path, url, lm, retrieved = newest_saved(code, log)  # the rows keep the workbook's retrieval time
                else:
                    (path, url, lm), retrieved = workbook(code, log), now
                ex = extract(path, code, HISTORY_START, end, log, url, lm, retrieved)
            used.append(f"{code}: {url} (Last-Modified {lm})")
            x = read_extract(ex)
            x = x[(x["ts_utc"] >= ip.utc_iso(start)) & (x["ts_utc"] < ip.utc_iso(end))]
            rows = complete_rows(code, x, variables, url, retrieved, gaps)
            rows.sort(key=lambda r: r[:3])
            log(f"  {code}: {len(rows):,} rows in complete UTC days; first {rows[0][2] if rows else None}")
            chunk = os.path.join(ip.RAW.dir, f"{code}_rows.tmp")
            with open(chunk, "w", encoding="utf-8", newline="") as f:
                csv.writer(f, lineterminator="\n").writerows(rows)
            chunks.append(chunk)
            n_new += len(rows)
            del rows, x
        if not n_new:
            raise RuntimeError("no complete day of CO2 values in the window")
        if args.max_new_rows is not None and n_new > args.max_new_rows:
            raise RuntimeError(f"{n_new:,} rows exceed the ceiling of {args.max_new_rows:,}; nothing written")
        if args.dry_run:
            for c in chunks:
                os.remove(c)
            msg = f"dry run: {n_new:,} rows in complete days, {len(gaps)} incomplete BA-day groups; nothing written"
            log(msg)
            print(msg)
            ip.write_status("eia930_emissions", run_id, [dict(table=NAME, market="all", status="ok", detail=msg)])
            log.close()
            return 0
        header = [
            "Energy Research Warehouse (ERW): EIA-930 hourly CO2 emissions estimates by balancing authority",
            "Shape: series (docs/datastandard.md v0), partition column ba. tCO2 = metric tons of CO2, as EIA states "
            "them. co2_emissions_generated: EIA's CO2 Emissions Generated (all sources in the BA); co2_emissions_consumed: "
            "EIA's CO2 Emissions Consumed (generated plus imported minus exported); since session 34 co2_emissions_coal, "
            "_natural_gas, _oil, _other (EIA's CO2 Emissions: COL, NG, OIL, Other) and co2_emissions_imported, "
            "_exported (EIA's CO2 Emissions Imported, Exported). ts_utc is the hour's start (EIA's UTC time, the hour's "
            "end, minus one hour).",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/eia930_emissions.py"
            + (" from the saved workbooks (--from-raw or --from-extract, no request); each row's retrieved_at is its "
               "workbook's download time" if args.from_raw or args.from_extract else ""),
            f"Run log: warehouse/output/logs/eia930_emissions_{run_id}.log",
            f"Raw files: warehouse/raw/eia930_emissions/<run_id>/ (not in git; manifest.csv lists each workbook, its URL, "
            f"Last-Modified and sha256; <ba>_hours.csv is the extract this run read). Workbooks read by this run: "
            + "; ".join(used),
            f"Source: {SOURCE} {REPORT}, {PAGE}",
            f"  document list: {BASE}",
            "EIA's method: estimated emissions = positive generation by fuel x EIA's CO2 emission factor per fuel and BA "
            "(lbs/kWh; https://www.eia.gov/tools/faqs/faq.php?id=74&t=11); imported and exported CO2 follow the "
            "interchange. See docs/methods/emissions.md.",
            "Completeness per UTC day, as EIA-930 generation (sessions 6, 13, 16): a day is written only when both "
            "generated and consumed have all 24 hours; each of the other six is written for that day only when it too "
            "has all 24 hours. Each day or variable-day left out is a gap row in warehouse/metadata/run_status.csv. "
            "Values are EIA's, not rounded.",
        ]
        kept, replaced, added = merge_write(chunks, header, log)
        for c in chunks:
            os.remove(c)
        summary = (f"{NAME}.csv: this run {n_new:,} rows ({added:,} new, {replaced:,} replacing earlier rows with "
                   f"the same key); {kept:,} rows kept from earlier runs; {len(gaps)} incomplete BA-day groups left out")
        log(summary)
        print(summary)
        ip.update_sources([{"source": SOURCE, "publisher": "U.S. Energy Information Administration (EIA)",
                            "report": REPORT, "report_url": PAGE, "document_list": BASE, "license": "public",
                            "tables": [NAME]}])
        results.append(dict(table=NAME, market="all", status="ok", detail=summary[:300]))
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"FAILED:\n{tb}")
        print(f"eia930_emissions {NAME} FAILED: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    ip.write_status("eia930_emissions", run_id, results + gaps)
    log.close()
    return 0 if all(r["status"] != "failed" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
