#!/usr/bin/env python3
"""The major power producers' de-rated generating capacity at each year's end in England and Wales, Scotland and Northern Ireland,
1999 to 2025, by fuel and by thermal plant type, and each country's share, as the Department for Energy Security and Net Zero
(DESNZ) gives them in the Digest of UK Energy Statistics (DUKES), table 5.8 (session 2743, chain 138, the session's second own
choice; the countries' split of chain 137's DUKES 5.7, whose major power producers' totals they make).

Energy Research Warehouse (ERW) connector. One series table, warehouse/output/desnz_uk_mpp_capacity_countries_annual.csv, from one
workbook (no key), its sheet "5.8", three tables one above another:

    https://assets.publishing.service.gov.uk/media/6a6a3631862aaf18d9c629ec/DUKES_5.8.xlsx

    entity      desnz_mppr:<country>:<item>: country england_wales, scotland, northern_ireland; item the row's fuel or plant type
                as a slug (coal_fired, gas_fired, nuclear_stations, onshore_wind, ccgt_stations, ...), or total (5.8.A's total
                major power producer capacity, carrying the share), total_thermal (5.8.C's total of thermal stations)
    variable    derated_capacity (5.8.A, by fuel, MW), thermal_capacity (5.8.C, by type of thermal plant, MW), capacity_share
                (5.8.B, the country's share of the major power producers' capacity, ratio: the sheet's heading says per cent, its
                cells hold fractions such as 0.7977)
    ts_utc      the year's first day; freq P1Y; geo GB-ENG,GB-WLS, GB-SCT or GB-NIR; node the sheet's two labels as written

DESNZ's notes: grid export capacity where available, else installed capacity, and no data for 1996 to 1998, which the sheet keeps
as "[x]" (note 1); the major power producers' capacity on both networks, small-scale hydro, wind and solar de-rated by 0.365,
0.43 and 0.17, and declared net capacity before 2006 (notes 3 and 5); major wind and solar companies counted as major power
producers from 2007 and 2015 (note 6); conventional steam stations leave out bioenergy (note 7).

Each table is found by its title and its year row (each year from 1996 in turn); each row's two labels (note marks read off) are
mapped to an entity, and a country not in the map or an item given twice stops the run. "[x]" and a blank cell have no row; a
zero stays a zero. Checks in every year: each country's fuels add to its total and its plant types to its thermal total, within
0.5 MW; the three shares add to 1, and each equals the country's total over the three, within 0.001. A check missed stops the
run. Nothing is filled.

    python warehouse/connectors/desnz_dukes_mpp_countries.py [--offline] [--fresh] [--out-dir DIR]

Ceiling: one request a run, saved in warehouse/raw/desnz_dukes/ and reused unless --fresh; --offline asks nothing.
www.gov.uk's robots file closes only print and search paths; assets.publishing.service.gov.uk has no robots file (session 2678).
License: the Open Government Licence v3.0 (DESNZ's verdict on record, session 2167): public.
"""

import argparse
import datetime as dt
import io
import os
import re
import sys
import traceback

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso_prices as ip  # noqa: E402

NAME = "desnz_uk_mpp_capacity_countries_annual"
CONNECTOR = "desnz_dukes"
SOURCE = "desnz:dukes_5_8_mpp_capacity_countries"
HOST = "assets.publishing.service.gov.uk"
URL = "https://assets.publishing.service.gov.uk/media/6a6a3631862aaf18d9c629ec/DUKES_5.8.xlsx"
PAGE = "https://www.gov.uk/government/statistics/electricity-chapter-5-digest-of-united-kingdom-energy-statistics-dukes"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ERW-connector/0.1; +https://github.com/SamuelEnrique/erw)"}
CITATION = ("Department for Energy Security and Net Zero, Digest of UK Energy Statistics (DUKES) 2026, table 5.8, Major power producers "
            "plant capacity, England and Wales, Scotland, and Northern Ireland")
COUNTRIES = {"England and Wales": ("england_wales", "GB-ENG,GB-WLS"), "Scotland": ("scotland", "GB-SCT"),
             "Northern Ireland": ("northern_ireland", "GB-NIR")}
TABLES = [("Table 5.8.A Major Power Producer Capacity by fuel, MW", "derated_capacity", ("Region", "Fuel")),
          ("Table 5.8.B Shares of major power producer capacity by country, per cent", "capacity_share", ("Generator Type", "Region")),
          ("Table 5.8.C Capacity by type of plant, MW", "thermal_capacity", ("Region", "Fuel"))]
TOTALS = {"Total MPP capacity": "total", "Total capacity of thermal stations": "total_thermal"}
FIRST_YEAR = 1996
X_COLS = []


def label(v):
    s = " ".join(str(v).replace(chr(0xA0), " ").split()) if v is not None else ""
    return re.sub(r"\s*\[note \d+\]", "", s).strip()


def slug(s):
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", s.lower())).strip("_")


def parse_sheet(rows):
    titles = {label(r[0]): i for i, r in enumerate(rows) if r and r[0] is not None}
    out, nodes = {}, {}
    for title, var, head_want in TABLES:
        if title not in titles:
            raise RuntimeError(f"no table titled {title!r}")
        i = titles[title]
        head = [label(x) for x in rows[i + 1]]
        if tuple(head[:2]) != head_want:
            raise RuntimeError(f"{title}: the header is {head[:2]}")
        years = head[2:]
        while years and not years[-1]:
            years.pop()
        if not years or years != [str(FIRST_YEAR + k) for k in range(len(years))]:
            raise RuntimeError(f"{title}: the years are not each year from {FIRST_YEAR}")
        for r in rows[i + 2:]:
            if not r or not any(label(x) for x in r[:2]):
                break
            if label(r[0]).startswith("Table 5.8"):
                break
            if var == "capacity_share":
                if label(r[0]) != "Major Power Producer":
                    raise RuntimeError(f"{title}: a row not of major power producers: {label(r[0])!r}")
                country, item = label(r[1]), "total"
            else:
                country, item = label(r[0]), label(r[1])
                item = TOTALS.get(item, slug(item))
            if country not in COUNTRIES:
                raise RuntimeError(f"{title}: a country not in the map: {country!r}")
            key, geo = COUNTRIES[country]
            ent = f"{key}:{item}"
            if (ent, var) in nodes:
                raise RuntimeError(f"{title}: {ent} given twice")
            nodes[(ent, var)] = (f"{label(r[0])}: {label(r[1])}", geo)
            for k, y in enumerate(years):
                v = r[2 + k] if 2 + k < len(r) else None
                if v is None or (isinstance(v, str) and v.strip() in ("", "[x]")):
                    continue
                if not isinstance(v, (int, float)) or isinstance(v, bool) or v < 0:
                    raise RuntimeError(f"{title} {country} {item} {y}: not a quantity: {v!r}")
                out[(int(y), ent, var)] = float(v)
    check(out)
    return out, nodes


def check(out):
    for y in sorted({k[0] for k in out}):
        g = {(k[1], k[2]): v for k, v in out.items() if k[0] == y}
        for c, _ in COUNTRIES.values():
            for var, tot_item in (("derated_capacity", "total"), ("thermal_capacity", "total_thermal")):
                tot = g.get((f"{c}:{tot_item}", var))
                parts = [v for (e, vv), v in g.items() if vv == var and e.startswith(c + ":") and not e.split(":")[1].startswith("total")]
                if tot is not None and parts and abs(sum(parts) - tot) > 0.5:
                    raise RuntimeError(f"{y} {c} {var}: the parts make {sum(parts):,.2f} MW, the total {tot:,.2f}")
        shares = {c: g.get((f"{c}:total", "capacity_share")) for c, _ in COUNTRIES.values()}
        totals = {c: g.get((f"{c}:total", "derated_capacity")) for c, _ in COUNTRIES.values()}
        if None not in shares.values():
            if abs(sum(shares.values()) - 1) > 0.001:
                raise RuntimeError(f"{y}: the shares add to {sum(shares.values()):.4f}")
            if None not in totals.values() and sum(totals.values()) > 0:
                for c in shares:
                    if abs(totals[c] / sum(totals.values()) - shares[c]) > 0.001:
                        raise RuntimeError(f"{y} {c}: the share {shares[c]:.4f} is not the total over the three")


def parse(content, base):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    if "5.8" not in wb.sheetnames:
        raise RuntimeError(f"the workbook has no sheet 5.8: {wb.sheetnames}")
    vals, nodes = parse_sheet([tuple(r) for r in wb["5.8"].iter_rows(values_only=True)])
    rows = [dict(entity=f"desnz_mppr:{e}", variable=v, ts_utc=f"{y}-01-01T00:00:00Z", value=round(x, 4),
                 unit="ratio" if v == "capacity_share" else "MW", freq="P1Y", geo=nodes[(e, v)][1], market="", node=nodes[(e, v)][0],
                 **base) for (y, e, v), x in sorted(vals.items())]
    return pd.DataFrame(rows)


def main(argv=None):
    ap = argparse.ArgumentParser(description="DESNZ DUKES 5.8: the major power producers' capacity by UK country a year")
    ap.add_argument("--offline", action="store_true", help="the saved workbook only, no request")
    ap.add_argument("--fresh", action="store_true", help="ask for the workbook again")
    ap.add_argument("--out-dir", help="a trial: the table under this folder, not warehouse/output")
    args = ap.parse_args(argv)
    if ip.paused(CONNECTOR) or ip.paused_host(HOST):
        print(f"{NAME}: the publisher is paused (warehouse/metadata/paused_sources.csv); nothing asked")
        return 0
    if args.out_dir:
        raw_dir = ip.RAW_DIR
        ip.set_out_dir(args.out_dir)
        ip.RAW_DIR = raw_dir
    os.makedirs(ip.LOG_DIR, exist_ok=True)
    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log = ip.Log(os.path.join(ip.LOG_DIR, f"{CONNECTOR}_{run_id}.log"))
    ip.RAW.open(CONNECTOR, run_id)
    results = []
    try:
        content, rec = ip.fetch_raw(CONNECTOR, URL, log, offline=args.offline, fresh=args.fresh, headers=UA, timeout=180, pause=1)
        base = dict(source=SOURCE, source_url=URL, retrieved_at=rec.get("retrieved_at") or "", vintage=ip.vintage_of(rec))
        t = parse(content, base)
        if t.duplicated(["entity", "variable", "ts_utc"]).any():
            raise RuntimeError("a key given twice")
        cols = ip.SERIES_COLS + X_COLS
        t = t[cols].sort_values(["entity", "variable", "ts_utc"]).reset_index(drop=True)
        last = t.ts_utc.max()
        tot = t[(t.variable == "derated_capacity") & t.entity.str.endswith(":total") & (t.ts_utc == last)].set_index("entity").value
        head = [
            "Energy Research Warehouse (ERW): DESNZ, Digest of UK Energy Statistics (DUKES) 2026, table 5.8: the major power producers' "
            "de-rated capacity in England and Wales, Scotland and Northern Ireland by fuel and thermal plant type, and each country's "
            f"share, {t.ts_utc.min()[:4]} to {last[:4]} (session 2743)",
            "Shape: series (docs/datastandard.md v0). entity desnz_mppr:<country>:<item>; variable derated_capacity or "
            "thermal_capacity (MW) or capacity_share (ratio); ts_utc the year's first day; freq P1Y; geo GB-ENG,GB-WLS, GB-SCT or "
            "GB-NIR; node the sheet's labels.",
            f"{len(t)} rows; {last[:4]} totals, MW: " + ", ".join(f"{k.split(':')[1]} {v:,.1f}" for k, v in tot.items())
            + ". 1996 to 1998 have no data (DESNZ's note 1); parts and shares checked; nothing is filled.",
            f"Retrieved: {run_id} (UTC) by warehouse/connectors/desnz_dukes_mpp_countries.py; the workbook retrieved "
            f"{base['retrieved_at']}; last modified {base['vintage']}",
            f"Run log: warehouse/output/logs/{CONNECTOR}_{run_id}.log",
            f"Raw files: warehouse/raw/{CONNECTOR}/ (not in git)",
            f"Source: {SOURCE} {CITATION}",
            f"  access: {URL} (listed on {PAGE})",
            "License: Open Government Licence v3.0. Credit: Department for Energy Security and Net Zero.",
        ]
        ip.write_snapshot(t, NAME, head, log, cols)
        results.append(dict(table=NAME, market="all", status="ok", detail=f"{len(t)} rows"))
        print(f"{NAME}: {len(t)} rows, {t.entity.nunique()} entities")
    except Exception:
        tb = ip.redact(traceback.format_exc())
        log(f"{NAME} FAILED, no output file written:\n{tb}")
        print(f"{NAME} FAILED, no output file written: {tb.strip().splitlines()[-1]}", file=sys.stderr)
        results.append(dict(table=NAME, market="all", status="failed", detail=tb.strip().splitlines()[-1][:300]))
    if not args.out_dir:
        if results and results[0]["status"] == "ok":
            ip.update_sources([dict(source=SOURCE, publisher="UK Department for Energy Security and Net Zero (DESNZ)",
                                    report="Digest of UK Energy Statistics (DUKES) 2026, table 5.8, Major power producers plant capacity by country",
                                    report_url=PAGE, document_list=PAGE, license="public", tables=[NAME])])  # OGL v3.0
        ip.write_status(CONNECTOR, run_id, results)
    log.close()
    return 0 if all(r["status"] == "ok" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
