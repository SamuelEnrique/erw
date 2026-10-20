#!/usr/bin/env python3
"""EIA-930's balance for the seven ISO grids, reported against adjusted (session 360).

Energy Research Warehouse (ERW). The Grid pack's item 3 (docs/reviews/packs/grid.md, section 6): for CAISO, PJM, MISO
and SPP the sum of a grid's ties does not equal its net generation minus its demand. This reruns that test, and the
same test on EIA's adjusted series, and splits the gap into its two parts:

    the tie gap      the sum of the grid's interchange with each neighbour (its DIBAs) less its total interchange (TI)
    the total gap    TI less (net generation NG minus demand D)

whose sum is the pack's gap. Read, no request:

    before   the live tables the pack used: eia930_all_interchange (each tie, EIA's API, interchange-data) and
             eia930_all_demand and eia930_all_generation (D and NG, EIA's API, region-data)
    EIA's    the BA workbooks EIA publishes (https://www.eia.gov/electricity/gridmonitor/knownissues/xls/<BA>.xlsx),
             as the emissions connector saved them (warehouse/raw/eia930_emissions/<run>/), sheet "Published Hourly
             Data": Demand, Net generation, Total interchange, their Imputed and Adjusted columns, and one column of
             interchange per neighbour; and the sheet "Imbalance Thresholds", EIA's own tolerances

EIA's API offers the region data as D, DF, NG and TI only (its facet "type", read 20 October 2026): the adjusted and
imputed series are in the workbooks alone. EIA's notes: an adjusted value is the reported value unless EIA imputed it
(a value missing, negative or zero, or 1.5 times past the BA's own extremes).

    python warehouse/analysis/eia930_balance.py [--run <workbook run folder>] [--from 2026-09-13] [--to 2026-10-17]

Writes docs/methods/eia930_balance.json (the figures the note quotes) and prints them.
"""

import argparse
import glob
import json
import os
import re
import sys

import numpy as np
import openpyxl
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, os.path.join(ROOT, "warehouse", "connectors"))
import iso_prices as ip  # noqa: E402

BAS = ["CISO", "ERCO", "ISNE", "MISO", "NYIS", "PJM", "SWPP"]
NAMES = {"CISO": "CAISO", "ERCO": "ERCOT", "ISNE": "ISO-NE", "MISO": "MISO", "NYIS": "NYISO", "PJM": "PJM", "SWPP": "SPP"}
RAW = os.path.join(ROOT, "warehouse", "raw", "eia930_emissions")
OUT = os.path.join(ROOT, "docs", "methods", "eia930_balance.json")
CODE = re.compile(r"^[A-Z][A-Z0-9]{1,5}$")


def read_live(first, last):
    """{BA: DataFrame(start, ties, ng, d)} from the live tables, as the pack's test read them."""
    def read(name, cols):
        p = os.path.join(ip.OUT_DIR, name + ".csv")
        out = []
        for ch in pd.read_csv(p, skiprows=ip.header_rows(p), usecols=cols, dtype=str, chunksize=500_000):
            out.append(ch[(ch["ts_utc"] >= first) & (ch["ts_utc"] < last)])
        return pd.concat(out, ignore_index=True)
    x = read("eia930_all_interchange", ["entity", "variable", "ts_utc", "value", "ba", "x_to_ba"])
    d = read("eia930_all_demand", ["entity", "variable", "ts_utc", "value"])
    g = read("eia930_all_generation", ["entity", "variable", "ts_utc", "value"])
    out = {}
    for ba in BAS:
        xi = x[x["entity"].str.startswith(f"eia930:{ba}-")]
        ties = xi.assign(v=xi["value"].astype(float)).groupby("ts_utc")["v"].sum()
        dd = d[(d["entity"] == f"eia930:{ba}") & (d["variable"] == "demand_mw")].set_index("ts_utc")["value"].astype(float)
        ng = g[(g["entity"] == f"eia930:{ba}") & (g["variable"] == "net_generation_mw")].set_index("ts_utc")["value"].astype(float)
        f = pd.concat([ties.rename("ties"), ng.rename("ng"), dd.rename("d")], axis=1).dropna()
        out[ba] = f
    return out


def workbook(path):
    """The hourly sheet's columns of one BA (start in UTC, D, NG, TI, imputed, adjusted, each tie) and its thresholds."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Published Hourly Data"]
    it = ws.iter_rows(values_only=True)
    head = list(next(it))
    k0 = head.index("Adjusted UNK Gen") + 1
    k1 = next(k for k in range(k0, len(head)) if str(head[k]).startswith(("Subregion", "CO2")))
    ties = [k for k in range(k0, k1) if CODE.match(str(head[k] or ""))]
    want = {c: head.index(c) for c in ("UTC time", "Demand", "Net generation", "Total interchange", "Imputed demand", "Imputed net generation",
                                       "Imputed total interchange", "Adjusted demand", "Adjusted net generation", "Adjusted total interchange")}
    rows = []
    for r in it:
        t = r[want["UTC time"]]
        if t is None:
            continue
        rec = {c: r[k] for c, k in want.items()}
        for k in ties:
            rec["tie:" + head[k]] = r[k]
        rows.append(rec)
    f = pd.DataFrame(rows)
    f["start"] = (pd.to_datetime(f["UTC time"]) - pd.Timedelta(hours=1)).dt.strftime("%Y-%m-%dT%H:%M:%SZ")  # EIA: the end of the hour
    th = {}
    for r in wb["Imbalance Thresholds"].iter_rows(values_only=True):
        if r and r[0] and r[0] != "BA":
            th[r[0]] = {"ng_d_ti": r[1], "ti_diba": r[2]}
    return f.set_index("start"), [head[k] for k in ties], th


def mean_abs(s):
    s = s.dropna()
    return round(float(s.abs().mean()), 1) if len(s) else None


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--run", help="the folder of the workbooks (default: the newest run that holds all seven)")
    ap.add_argument("--from", dest="first", default="2026-09-13T00:00:00Z")
    ap.add_argument("--to", dest="last", default="2026-10-18T00:00:00Z", help="exclusive")
    ap.add_argument("--pack-window", default="2026-09-13T00:00:00Z,2026-09-28T00:00:00Z", help="the pack's 360 hours, to reproduce its figures")
    a = ap.parse_args(argv)
    run = a.run
    if not run:
        runs = sorted(d for d in glob.glob(os.path.join(RAW, "*")) if os.path.isdir(d) and all(glob.glob(os.path.join(d, f"*_{b}.xlsx")) for b in BAS))
        run = runs[-1]
    live = read_live(a.first, a.last)
    p0, p1 = a.pack_window.split(",")
    res = {"window": [a.first, a.last], "pack_window": [p0, p1], "workbooks": os.path.relpath(run, ROOT).replace(os.sep, "/"), "bas": {}}
    for ba in BAS:
        path = sorted(glob.glob(os.path.join(run, f"*_{ba}.xlsx")))[-1]
        wbk, ties, th = workbook(path)
        w = wbk[(wbk.index >= a.first) & (wbk.index < a.last)].copy()
        num = lambda c: pd.to_numeric(w[c], errors="coerce")  # noqa: E731
        # the ties: the neighbour columns that hold a value in the window (EIA's sheet names every possible neighbour);
        # an hour's sum only when each of them is held
        tie_cols = [c for c in w.columns if c.startswith("tie:") and pd.to_numeric(w[c], errors="coerce").notna().any()]
        tv = w[tie_cols].apply(pd.to_numeric, errors="coerce")
        w["ties"] = tv.sum(axis=1).where(tv.notna().all(axis=1))
        ties = [c[4:] for c in tie_cols]
        d, ng, ti = num("Demand"), num("Net generation"), num("Total interchange")
        ad, ang, ati = num("Adjusted demand"), num("Adjusted net generation"), num("Adjusted total interchange")
        tie = pd.to_numeric(w["ties"], errors="coerce")
        lv = live[ba]
        lp = lv[(lv.index >= p0) & (lv.index < p1)]
        thr = th.get(ba, {})
        pct = lambda gap, dem: (100 * gap.abs() / dem.where(dem > 0)).dropna()  # noqa: E731
        adj_total = ati - (ang - ad)
        adj_tie = tie - ati
        r = {
            "name": NAMES[ba], "ties": ties, "hours": int(len(w)), "hours_all_ties": int(tie.notna().sum()),
            "pack_test_reproduced": {"hours": int(len(lp)), "mean_abs_gap_mw": mean_abs(lp["ties"] - (lp["ng"] - lp["d"]))},
            "live_window": {"hours": int(len(lv)), "mean_abs_gap_mw": mean_abs(lv["ties"] - (lv["ng"] - lv["d"]))},
            "reported": {"total_gap_mw": mean_abs(ti - (ng - d)), "tie_gap_mw": mean_abs(tie - ti), "pack_gap_mw": mean_abs(tie - (ng - d))},
            "adjusted": {"total_gap_mw": mean_abs(adj_total), "tie_gap_mw": mean_abs(adj_tie), "pack_gap_mw": mean_abs(tie - (ang - ad))},
            "imputed_hours": {"demand": int(pd.to_numeric(w["Imputed demand"], errors="coerce").notna().sum()),
                              "net_generation": int(pd.to_numeric(w["Imputed net generation"], errors="coerce").notna().sum()),
                              "total_interchange": int(pd.to_numeric(w["Imputed total interchange"], errors="coerce").notna().sum())},
            "eia_thresholds": thr,
            "adjusted_pct_of_demand": {"total_gap_mean": round(float(pct(adj_total, ad).mean()), 2) if len(pct(adj_total, ad)) else None,
                                       "tie_gap_mean": round(float(pct(adj_tie, ad).mean()), 2) if len(pct(adj_tie, ad)) else None,
                                       "tie_gap_hours_over_threshold": int((pct(adj_tie, ad) > float(thr.get("ti_diba") or np.inf)).sum()),
                                       "total_gap_hours_over_threshold": int((pct(adj_total, ad) > float(thr.get("ng_d_ti") or np.inf)).sum())},
            # EIA's sheet gives no unit for its thresholds: the hours over them read as MW too
            "adjusted_mw": {"tie_gap_hours_over_threshold": int((adj_tie.dropna().abs() > float(thr.get("ti_diba") or np.inf)).sum()),
                            "total_gap_hours_over_threshold": int((adj_total.dropna().abs() > float(thr.get("ng_d_ti") or np.inf)).sum()),
                            "hours_tie_gap_held": int(adj_tie.notna().sum()), "hours_total_gap_held": int(adj_total.notna().sum())},
        }
        res["bas"][ba] = r
        print(f"{ba}: pack {r['pack_test_reproduced']}, live {r['live_window']}, reported {r['reported']}, adjusted {r['adjusted']}, imputed {r['imputed_hours']}, "
              f"thresholds {thr}, pct {r['adjusted_pct_of_demand']}", flush=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(res, f, indent=1)
    print(f"written {os.path.relpath(OUT, ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
